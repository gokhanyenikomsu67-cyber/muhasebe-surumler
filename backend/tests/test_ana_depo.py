"""Ana Depo: alış faturalarının malzemeleri ana depoya girer, şantiyesi yazılan kalem o şantiyeye çıkar (sentetik veri)."""
import io
import uuid

from openpyxl import load_workbook

from tests.test_api import client, yeni_kullanici
from tests.test_malzeme import _fatura


def test_ana_depo_giris_cikis_ve_sevk():
    h = yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")
    s1 = client.post("/api/santiyeler", json={"ad": "Deneme Şantiye"}, headers=h).json()
    merkez = client.post("/api/santiyeler", json={"ad": "Merkez Depo"}, headers=h).json()
    c = client.post("/api/cariler", json={"unvan": "Örnek Yapı Market"}, headers=h).json()
    _fatura(h, c["id"], [{"aciklama": "ÇİMENTO 50 KG", "miktar": 10, "birim": "adet", "birim_fiyat": 100},
                         {"aciklama": "Kum", "miktar": 2, "birim": "ton", "birim_fiyat": 500, "santiye_id": s1["id"]}])
    _fatura(h, c["id"], [{"aciklama": "Çimento 50kg", "miktar": 5, "birim": "Adet", "birim_fiyat": 100}], santiye_id=s1["id"])
    _fatura(h, c["id"], [{"aciklama": "ÇİMENTO 50 KG", "miktar": 3, "birim": "adet", "birim_fiyat": 100}], santiye_id=merkez["id"])
    _fatura(h, c["id"], [{"aciklama": "ÇİMENTO 50 KG", "miktar": 1, "birim": "adet", "birim_fiyat": 100}], yon="giden")  # satış girmez

    v = client.get("/api/ana-depo", headers=h).json()
    ozet = {x["malzeme"]: x for x in v["ozet"]}
    cim = ozet["ÇİMENTO 50 KG"]
    # 18 adet girdi, 5'i şantiyeye çıktı; "Merkez Depo" ana deponun kendisi sayılır
    assert (cim["giren_miktar"], cim["cikan_miktar"], cim["kalan_miktar"], cim["kalan_tutar"]) == (18, 5, 13, 1300)
    assert {(x["santiye"], x["malzeme"], x["miktar"]) for x in v["cikislar"]} == {("Deneme Şantiye", "ÇİMENTO 50 KG", 5), ("Deneme Şantiye", "Kum", 2)}

    # Ana depoda bekleyen çimento (10 adet) şantiyeye çıkarılır, sonra geri alınır
    bekleyen = [x for x in v["satirlar"] if x["yer"] == "ana" and x["miktar"] == 10]
    r = client.post("/api/ana-depo/sevk", json={"kalem_idler": [bekleyen[0]["kalem_id"]], "santiye_id": s1["id"]}, headers=h).json()
    assert r["tasinan"] == 1
    cim = {x["malzeme"]: x for x in client.get("/api/ana-depo", headers=h).json()["ozet"]}["ÇİMENTO 50 KG"]
    assert (cim["cikan_miktar"], cim["kalan_miktar"]) == (15, 3)
    # Faturası şantiyeli kalem de ana depoya geri alınabilir (kalem ana depoya yazılır)
    satir = [x for x in client.get("/api/ana-depo", headers=h).json()["satirlar"] if x["yer"] == "santiye" and x["malzeme"] == "ÇİMENTO 50 KG"]
    client.post("/api/ana-depo/sevk", json={"kalem_idler": [x["kalem_id"] for x in satir], "santiye_id": None}, headers=h)
    cim = {x["malzeme"]: x for x in client.get("/api/ana-depo", headers=h).json()["ozet"]}["ÇİMENTO 50 KG"]
    assert (cim["cikan_miktar"], cim["kalan_miktar"]) == (0, 18)
    # Excel: 3 sayfa
    wb = load_workbook(io.BytesIO(client.get("/api/excel/ana-depo", headers=h).content))
    assert wb.sheetnames == ["Ana Depo Özeti", "Şantiyelere Çıkışlar", "Bütün Hareketler"]
