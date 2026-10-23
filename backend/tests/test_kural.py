"""Yapay zekasız okuma testleri."""
import io
import uuid

import pytest
from openpyxl import Workbook, load_workbook

from app import ai, kural
from tests.test_api import client, yeni_kullanici

UBL = """<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
 <cbc:ID>ABC2026000000123</cbc:ID>
 <cbc:IssueDate>2026-10-01</cbc:IssueDate>
 <cbc:Note>Sevk: Kartal Konut şantiyesi</cbc:Note>
 <cbc:DocumentCurrencyCode>TRY</cbc:DocumentCurrencyCode>
 <cac:AccountingSupplierParty><cac:Party>
  <cac:PartyIdentification><cbc:ID schemeID="MERSISNO">0123</cbc:ID></cac:PartyIdentification>
  <cac:PartyIdentification><cbc:ID schemeID="VKN">1234567890</cbc:ID></cac:PartyIdentification>
  <cac:PartyName><cbc:Name>Yılmaz Beton A.Ş.</cbc:Name></cac:PartyName>
  <cac:PostalAddress><cbc:StreetName>Sanayi Cad.</cbc:StreetName><cbc:CitySubdivisionName>Tuzla</cbc:CitySubdivisionName><cbc:CityName>İstanbul</cbc:CityName></cac:PostalAddress>
  <cac:PartyTaxScheme><cac:TaxScheme><cbc:Name>Tuzla</cbc:Name></cac:TaxScheme></cac:PartyTaxScheme>
 </cac:Party></cac:AccountingSupplierParty>
 <cac:AccountingCustomerParty><cac:Party>
  <cac:PartyIdentification><cbc:ID schemeID="VKN">9999999999</cbc:ID></cac:PartyIdentification>
  <cac:PartyName><cbc:Name>Test İnşaat</cbc:Name></cac:PartyName>
 </cac:Party></cac:AccountingCustomerParty>
 <cac:TaxTotal><cbc:TaxAmount currencyID="TRY">5600.00</cbc:TaxAmount></cac:TaxTotal>
 <cac:LegalMonetaryTotal>
  <cbc:LineExtensionAmount currencyID="TRY">28000.00</cbc:LineExtensionAmount>
  <cbc:TaxExclusiveAmount currencyID="TRY">28000.00</cbc:TaxExclusiveAmount>
  <cbc:TaxInclusiveAmount currencyID="TRY">33600.00</cbc:TaxInclusiveAmount>
  <cbc:PayableAmount currencyID="TRY">33600.00</cbc:PayableAmount>
 </cac:LegalMonetaryTotal>
 <cac:InvoiceLine><cbc:ID>1</cbc:ID><cbc:InvoicedQuantity unitCode="MTQ">10</cbc:InvoicedQuantity>
  <cbc:LineExtensionAmount currencyID="TRY">25000.00</cbc:LineExtensionAmount>
  <cac:TaxTotal><cbc:TaxAmount>5000</cbc:TaxAmount><cac:TaxSubtotal><cbc:Percent>20</cbc:Percent></cac:TaxSubtotal></cac:TaxTotal>
  <cac:Item><cbc:Name>C30 Hazır Beton</cbc:Name></cac:Item><cac:Price><cbc:PriceAmount>2500</cbc:PriceAmount></cac:Price>
 </cac:InvoiceLine>
 <cac:InvoiceLine><cbc:ID>2</cbc:ID><cbc:InvoicedQuantity unitCode="C62">1</cbc:InvoicedQuantity>
  <cbc:LineExtensionAmount currencyID="TRY">3000.00</cbc:LineExtensionAmount>
  <cac:TaxTotal><cbc:TaxAmount>600</cbc:TaxAmount><cac:TaxSubtotal><cbc:Percent>20</cbc:Percent></cac:TaxSubtotal></cac:TaxTotal>
  <cac:Item><cbc:Name>Beton Pompası</cbc:Name></cac:Item><cac:Price><cbc:PriceAmount>3000</cbc:PriceAmount></cac:Price>
 </cac:InvoiceLine>
</Invoice>""".encode()


def yeni():
    return yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")


def xlsx(satirlar):
    wb = Workbook()
    for s in satirlar:
        wb.active.append(s)
    t = io.BytesIO()
    wb.save(t)
    return t.getvalue()


@pytest.mark.parametrize("girdi,beklenen", [
    ("1.234,56", 1234.56), ("-1.234,56", -1234.56), ("1,234.56", 1234.56), ("12,5 TL", 12.5),
    ("1.234", 1234), ("250,00-", -250), (99, 99.0), ("", None), ("abc", None),
])
def test_para_coz(girdi, beklenen):
    assert kural.para_coz(girdi) == beklenen


def test_sade():
    assert kural.sade("İNŞAAT Şirketi") == "insaat sirketi"


def test_efatura_xml_uctan_uca():
    h = yeni()
    client.post("/api/santiyeler", json={"ad": "Kartal Konut"}, headers=h)
    r = client.post("/api/faturalar/oku", files={"dosya": ("fatura.xml", UBL, "text/xml")}, headers=h)
    assert r.status_code == 200, r.text
    f = r.json()[0]
    assert f["fatura_no"] == "ABC2026000000123" and f["tarih"] == "2026-10-01"
    assert f["genel_toplam"] == 33600 and f["kdv_toplam"] == 5600
    assert [k["birim"] for k in f["kalemler"]] == ["m³", "adet"]
    assert f["santiye_ad"] == "Kartal Konut" and f["gider_turu"] == "Beton"
    assert f["ai_okuma"]["satici"]["vkn"] == "1234567890"
    assert f["ai_okuma"]["satici"]["vergi_dairesi"] == "Tuzla"
    assert f["ai_okuma"]["uyarilar"] == [] and f["ai_okuma"]["bilgi"] == ["'Yılmaz Beton A.Ş.' carilerde yoktu; yeni cari açıldı."]
    r = client.post("/api/faturalar/toplu-onayla", json={"idler": [f["id"]]}, headers=h).json()
    assert r == {"onaylanan": 1, "hatalar": []}


def test_efatura_zip():
    import zipfile
    t = io.BytesIO()
    with zipfile.ZipFile(t, "w") as z:
        z.writestr("a.xml", UBL)
        z.writestr("b.xml", UBL.replace(b"ABC2026000000123", b"ABC2026000000124"))
        z.writestr("a.html", "<html></html>")
    r = client.post("/api/faturalar/oku", files={"dosya": ("faturalar.zip", t.getvalue(), "application/zip")}, headers=yeni())
    assert [f["fatura_no"] for f in r.json()] == ["ABC2026000000123", "ABC2026000000124"]


def test_bozuk_fotograf_nedeniyle_reddedilir(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setattr("app.ocr.tesseract", lambda: "tesseract")
    r = client.post("/api/faturalar/oku", files={"dosya": ("f.jpg", b"\xff\xd8", "image/jpeg")}, headers=yeni())
    assert r.status_code == 422 and "Resim açılamadı" in r.json()["detail"]


def pdf_olustur(satirlar):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    t = io.BytesIO()
    c = canvas.Canvas(t)
    y = 800
    for s in satirlar:
        c.drawString(40, y, s)
        y -= 16
    c.save()
    return t.getvalue()


def test_pdf_fatura(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur(["YILMAZ BETON A.S.", "VKN: 1234567890", "SAYIN TEST INSAAT", "VKN: 9999999999",
                       "Fatura No: ABC2026000000123", "Fatura Tarihi: 01-10-2026", "C30 Beton 10 m3",
                       "Mal Hizmet Toplam Tutari 28.000,00 TL", "Hesaplanan KDV (%20) 5.600,00 TL",
                       "Odenecek Tutar 33.600,00 TL"])
    r = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=yeni())
    assert r.status_code == 200, r.text
    f = r.json()[0]
    assert f["fatura_no"] == "ABC2026000000123" and f["tarih"] == "2026-10-01"
    # Kalem tablosu yok: geçici "KDV %20 matrahı" satırı açılmaz (kullanıcı kuralı); faturadaki toplamlar okunur
    o = f["ai_okuma"]
    assert f["kalemler"] == [] and o["kalemler_okunamadi"]
    assert (o["ara_toplam"], o["kdv_toplam"], o["genel_toplam"]) == (28000, 5600, 33600)
    assert o["satici"]["vkn"] == "1234567890"
    assert f["gider_turu"] == "Beton"


def test_pdf_efatura_gorunumu(monkeypatch):
    """e-Fatura görüntüsü: önce alıcı (SAYIN), sonra satıcı bloğu; kalem tablosu satır satır çözülür."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur([
        "SAYIN", "DENEME INSAAT LIMITED SIRKETI", "CUMHURIYET MAH. ATATURK CAD. NO:1", "34000 MERKEZ/ Istanbul",
        "Vergi Dairesi: KADIKOY", "VKN: 9999999999", "Senaryo: TICARIFATURA", "Fatura No: KRT2026000000045",
        "Fatura Tarihi: 05- 01- 2026", "Son Odeme Tarihi 20-01-2026",
        "KARTAL43 YAPI MARKET SAN. VE TIC. LTD. STI.", "SANAYI MAH. 10. SOKAK NO: 3", "43100 MERKEZ/ KUTAHYA",
        "Tel: 2740000000 Fax:", "E-Posta: info@ornek.com", "Vergi Dairesi: CINILI", "VKN: 1234567890",
        "e-FATURA", "Sira No Mal Hizmet Miktar Birim Fiyat Iskonto Orani KDV Orani KDV Tutari Mal Hizmet Tutari",
        "1 NERVURLU DEMIR 12MM 2,5 Ton 20.000 TL %0 %20,00 10.000,00 TL 50.000,00 TL",
        "2 BAGLAMA TELI 10 Kg 45,50 TL %0 %20,00 91,00 TL 455,00 TL",
        "Mal Hizmet Toplam Tutari 50.455,00 TL", "Toplam Iskonto 0,00 TL",
        "Hesaplanan KDV(%20) 10.091,00 TL", "Vergiler Dahil Toplam Tutar 60.546,00 TL", "Odenecek Tutar 60.546,00 TL"])
    h = yeni()
    r = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=h)
    assert r.status_code == 200, r.text
    f = r.json()[0]
    assert f["tarih"] == "2026-01-05" and f["vade"] == "2026-01-20" and f["genel_toplam"] == 60546
    assert f["ai_okuma"]["satici"] == {"unvan": "KARTAL43 YAPI MARKET SAN. VE TIC. LTD. STI.", "vkn": "1234567890",
                                       "vergi_dairesi": None, "adres": None, "telefon": None, "iban": None}
    assert f["ai_okuma"]["alici"]["vkn"] == "9999999999"
    assert [(k["aciklama"], k["miktar"], k["birim"], k["birim_fiyat"], k["tutar"]) for k in f["kalemler"]] == [
        ("NERVURLU DEMIR 12MM", 2.5, "Ton", 20000, 50000), ("BAGLAMA TELI", 10, "Kg", 45.5, 455)]
    assert f["gider_turu"] == "Demir"


def test_taranmis_pdf(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur([""])  # boş sayfa
    r = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=yeni())
    # Metinsiz PDF resim olarak okunmaya çalışılır; boş sayfada yazı yoktur
    from app import ocr
    assert r.status_code == 422 and ("okunabilir yazı" if ocr.var_mi() else "Tesseract") in r.json()["detail"]


def hazirlik():
    h = yeni()
    ahmet = client.post("/api/cariler", json={"unvan": "Ahmet Yapı San. Tic. Ltd. Şti.", "vkn": "5555555555"}, headers=h).json()
    demir = client.post("/api/cariler", json={"unvan": "Demir Ltd", "iban": "TR12 0001 0000 0000 0000 0000 01"}, headers=h).json()
    banka = client.post("/api/banka-hesaplari", json={"ad": "Ziraat"}, headers=h).json()
    return h, ahmet, demir, banka


def test_ekstre_excel_borc_alacak():
    h, ahmet, demir, banka = hazirlik()
    veri = xlsx([["ZİRAAT BANKASI HESAP HAREKETLERİ"], ["Hesap: 123"], [],
                 ["İşlem Tarihi", "Valör", "Açıklama", "Borç", "Alacak", "Bakiye"],
                 ["01.10.2026", "01.10.2026", "EFT AHMET YAPI SAN TIC", None, "5.000,00", "5.000,00"],
                 ["02.10.2026", "02.10.2026", "HAVALE TR120001000000000000000001", "2.000,00", None, "3.000,00"],
                 ["02.10.2026", "02.10.2026", "EFT MASRAFI BSMV", "8,50", None, "2.991,50"],
                 ["03.10.2026", "03.10.2026", "SHELL MOTORIN", "1.000,00", None, "1.991,50"],
                 ["", "", "TOPLAM", "3.008,50", "5.000,00", ""]])
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre?yontem=kural", files={"dosya": ("e.xlsx", veri)}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["eklenen"] == 4
    t = {x["aciklama"]: x for x in client.get("/api/hareketler?durum=taslak", headers=h).json()}
    assert t["EFT AHMET YAPI SAN TIC"]["cari_id"] == ahmet["id"] and t["EFT AHMET YAPI SAN TIC"]["tur"] == "tahsilat"
    assert t["HAVALE TR120001000000000000000001"]["cari_id"] == demir["id"] and t["HAVALE TR120001000000000000000001"]["tur"] == "odeme"
    assert t["EFT MASRAFI BSMV"]["tur"] == "banka_masrafi" and t["EFT MASRAFI BSMV"]["yon"] == "cikis"
    assert t["SHELL MOTORIN"]["tur"] == "masraf" and t["SHELL MOTORIN"]["gider_turu"] == "Akaryakıt"


def test_ekstre_csv_tek_tutar():
    h, ahmet, demir, banka = hazirlik()
    csv = "Tarih;Açıklama;Tutar;Bakiye\n05/10/2026;Gelen EFT 5555555555;1.500,00;1.500,00\n06/10/2026;Kira;-750,00;750,00\n"
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre", files={"dosya": ("e.csv", csv.encode("cp1254"), "text/csv")}, headers=h)
    assert r.json()["eklenen"] == 2
    t = {x["aciklama"]: x for x in client.get("/api/hareketler?durum=taslak", headers=h).json()}
    assert t["Gelen EFT 5555555555"]["cari_id"] == ahmet["id"]
    assert t["Kira"]["yon"] == "cikis" and t["Kira"]["tutar"] == 750


def test_ekstre_pdf():
    h, ahmet, demir, banka = hazirlik()
    pdf = pdf_olustur(["Tarih Aciklama Tutar Bakiye", "Devreden bakiye 10.000,00",
                       "01.10.2026 EFT AHMET YAPI SAN 5.000,00 15.000,00",
                       "02.10.2026 14:22 ODEME DEMIR LTD 2.000,00 13.000,00",
                       "03.10.2026 EFT UCRETI 8,50 12.991,50"])
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre?yontem=kural", files={"dosya": ("e.pdf", pdf, "application/pdf")}, headers=h)
    assert r.status_code == 200, r.text
    t = {x["aciklama"]: x for x in client.get("/api/hareketler?durum=taslak", headers=h).json()}
    assert t["EFT AHMET YAPI SAN"]["yon"] == "giris" and t["EFT AHMET YAPI SAN"]["cari_id"] == ahmet["id"]
    assert t["ODEME DEMIR LTD"]["yon"] == "cikis" and t["ODEME DEMIR LTD"]["cari_id"] == demir["id"]
    assert t["EFT UCRETI"]["tur"] == "banka_masrafi"


def test_ekstre_basliksiz():
    h, _, _, banka = hazirlik()
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre?yontem=kural", files={"dosya": ("e.xlsx", xlsx([["a", "b"], [1, 2]]))}, headers=h)
    assert r.status_code == 422 and "Tarih" in r.json()["detail"]


def test_cari_eslestirme_belirsiz():
    cariler = [{"id": 1, "unvan": "Yıldız Yapı Ltd"}, {"id": 2, "unvan": "Yıldız Yapı Market"}]
    assert kural.cari_eslestir("EFT YILDIZ YAPI", cariler) is None


def test_fatura_excel_yukleme():
    h = yeni()
    sablon = load_workbook(io.BytesIO(client.get("/api/excel/fatura-sablon", headers=h).content))
    ws = sablon.active
    ws.append(["S-1", "05.10.2026", None, "Satış", "Yeni Müşteri", "7777777777", None, None, "Daire satışı", 1, "adet", "100.000,00", 1, 0])
    t = io.BytesIO()
    sablon.save(t)
    r = client.post("/api/excel/faturalar/yukle", files={"dosya": ("f.xlsx", t.getvalue())}, headers=h)
    assert r.json() == {"eklenen": 2, "atlanan": 0, "yeni_cari": 2, "yeni_santiye": 0}
    tekrar = client.post("/api/excel/faturalar/yukle", files={"dosya": ("f.xlsx", t.getvalue())}, headers=h)
    assert tekrar.json()["eklenen"] == 0 and tekrar.json()["atlanan"] == 2
    alis = client.get("/api/faturalar?yon=gelen", headers=h).json()
    assert len(alis) == 1 and len(alis[0]["kalemler"]) == 2 and alis[0]["genel_toplam"] == 33600
    satis = client.get("/api/faturalar?yon=giden", headers=h).json()[0]
    assert satis["genel_toplam"] == 101000 and satis["durum"] == "taslak"


def test_fatura_excel_hatali_satir():
    h = yeni()
    veri = xlsx([["Fatura No", "Tarih", "Tür (Alış/Satış)", "Cari Ünvan", "Kalem Açıklaması", "Birim Fiyat", "Şantiye"],
                 ["X1", "tarih yok", "Alış", "A", "x", 10, None],
                 ["X2", "01.10.2026", "Alış", "A", "x", "fiyat yok", "Yeni Şantiye"]])
    r = client.post("/api/excel/faturalar/yukle", files={"dosya": ("f.xlsx", veri)}, headers=h)
    assert r.status_code == 400 and "2. satır" in r.json()["detail"] and "3. satır" in r.json()["detail"]
    assert client.get("/api/faturalar", headers=h).json() == []
    assert client.get("/api/cariler", headers=h).json() == []


def test_cek_excel_yukleme():
    h = yeni()
    sablon = load_workbook(io.BytesIO(client.get("/api/excel/cek-sablon", headers=h).content))
    sablon.active.append(["Verilen", "Senet", "S-9", None, None, None, None, "2.500,00", "15.12.2026", None, "Örnek Müşteri Ltd.", None, None, None])
    t = io.BytesIO()
    sablon.save(t)
    r = client.post("/api/excel/cekler/yukle", files={"dosya": ("c.xlsx", t.getvalue())}, headers=h)
    assert r.json() == {"eklenen": 2, "yeni_cari": 1, "yeni_santiye": 0}
    cekler = client.get("/api/cekler", headers=h).json()
    assert {(c["yon"], c["tur"], c["tutar"]) for c in cekler} == {("alinan", "cek", 4000), ("verilen", "senet", 2500)}
    bakiye = client.get("/api/raporlar/cari-bakiye", headers=h).json()[0]
    assert bakiye["bakiye"] == -4000 + 2500


def test_fatura_listesi_excel():
    """Şablon dışı fatura listesi: Tarihi / Ticari Unvanı / Stok Adı / Grubu / Tutarı / Tutarı (KDV Dahil)."""
    h = yeni()
    veri = xlsx([["Tarihi", "Fatura No", "Ticari Unvanı", "Stok Adı", "Grubu", "Birimi", "Miktarı", "Tutarı", "Tutarı (KDV Dahil)"],
                 ["01.08.2026", "ABC2026000000001", "Örnek Elektrik A.Ş.", "NYY KABLO 3X2,5", "Merkez Şantiye", "METRE", 100, 2500, 3000],
                 ["01.08.2026", "ABC2026000000001", "Örnek Elektrik A.Ş.", "Sigorta 16A", "Merkez Şantiye", "ADET", 3, 300, 330],
                 ["02.08.2026", "XYZ1", "Yemekçi", "Yemek", "Proje Ofis", "ADET", 1, 1000, 1000],
                 [None, None, None, None, "TOPLAM", None, 104, 3800, 4330]])
    r = client.post("/api/excel/faturalar/yukle", files={"dosya": ("l.xlsx", veri)}, headers=h)
    assert r.json() == {"eklenen": 2, "atlanan": 0, "yeni_cari": 2, "yeni_santiye": 2}, r.text
    f = {x["fatura_no"]: x for x in client.get("/api/faturalar?yon=gelen", headers=h).json()}
    a = f["ABC2026000000001"]
    assert a["santiye_ad"] == "Merkez Şantiye" and a["gider_turu"] == "Elektrik malzemesi"
    assert [(k["miktar"], k["birim_fiyat"], k["kdv_orani"], k["tutar"]) for k in a["kalemler"]] == [(100, 25, 20, 2500), (3, 100, 10, 300)]
    assert a["genel_toplam"] == 3330 and f["XYZ1"]["kdv_toplam"] == 0


def test_ekstre_halkbank_bicimi():
    """Üstte hesap bilgisi, B/A sütunu, eksi bakiye, kendi hesaplar arası transfer ve banka masrafı."""
    h, ahmet, _, banka = hazirlik()
    csv = ("Şube Kodu:   ;1546   ;;;;;\nIBAN No:   ;TR120001000000000000000099;;;;;\n;;;;;;\n"
           "İŞLEM TARİHİ;BİM REF   ;TUTAR   ;B/A;BAKİYE   ;AÇIKLAMA   ;VALOR\n"
           "01.10.2026  ;r1;   1.250.000,00 ;A  ;   1.250.000,00;Ahmet Yapı San,AHMET YAPI SAN 'DAN cari ödeme;01.10.2026\n"
           "02.10.2026  ;r2;     200.000,00 ;B  ;   1.050.000,00;HESAPLAR ARASI TRANSFER,TEST ŞİRKET,TR00;02.10.2026\n"
           "03.10.2026  ;r3;   1.060.000,00 ;B  ;-     10.000,00;cari ödeme,Yeni Tedarikçi Ltd,Akbank T.A.Ş.,TR00;03.10.2026\n"
           "03.10.2026  ;r4;          18,73 ;B  ;-     10.018,73;KMH FAİZİ;03.10.2026\n"
           ";;;;;;\nTicaret Unvanı: Banka A.Ş.;;;;;;\n")
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre?yontem=kural",
                    files={"dosya": ("e.csv", csv.encode("cp1254"), "text/csv")}, headers=h)
    assert r.json()["eklenen"] == 4, r.text
    t = {x["tutar"]: x for x in client.get("/api/hareketler?durum=taslak", headers=h).json()}
    assert (t[1250000]["yon"], t[1250000]["tur"], t[1250000]["cari_id"]) == ("giris", "tahsilat", ahmet["id"])
    assert (t[200000]["yon"], t[200000]["tur"]) == ("cikis", "virman")
    assert (t[1060000]["yon"], t[1060000]["tur"]) == ("cikis", "odeme")
    assert t[18.73]["tur"] == "banka_masrafi"
    hesap = next(x for x in client.get("/api/banka-hesaplari", headers=h).json() if x["id"] == banka["id"])
    assert hesap["iban"] == "TR120001000000000000000099"
    baska = client.post("/api/banka-hesaplari", json={"ad": "Diğer", "iban": "TR990000000000000000000001"}, headers=h).json()
    r = client.post(f"/api/banka-hesaplari/{baska['id']}/ekstre?yontem=kural",
                    files={"dosya": ("e.csv", csv.encode("cp1254"), "text/csv")}, headers=h)
    # Başka hesap seçilse de ekstre IBAN'ına göre kendi hesabına işlenir (satırlar zaten var: yeniden okunur, eklenmez)
    assert (r.json()["hesap"], r.json()["eklenen"], r.json()["yenilenen"]) == (banka["ad"], 0, 4)
    # Açılış bakiyesi ilk satırdan, ekstre bakiyesi son satırdan alınır; onaylanınca fark kalmaz
    client.post("/api/hareketler/onayla", json={"idler": [x["id"] for x in client.get("/api/hareketler?durum=taslak", headers=h).json()]}, headers=h)
    hesap = next(x for x in client.get("/api/banka-hesaplari", headers=h).json() if x["id"] == banka["id"])
    assert (hesap["acilis_bakiyesi"], hesap["ekstre_bakiye"], hesap["bakiye"], hesap["mutabakat_farki"]) == (0, -10018.73, -10018.73, 0)


def test_hesaplar_arasi_transfer_ve_kasa():
    """94 nolu hesaptan 53 nolu hesaba transfer: 94'ten çıkış, 53'e giriş; 53 otomatik açılır, iki ekstre çift kayıt yapmaz.
    Nakit çekim kasaya girer, kasadan cariye ödeme yapılır."""
    h = yeni()
    a94 = client.post("/api/banka-hesaplari", json={"ad": "Halkbank 94", "banka": "Halkbank"}, headers=h).json()
    bas = ";;;;;\nİŞLEM TARİHİ;BİM REF   ;TUTAR   ;B/A;BAKİYE   ;AÇIKLAMA\n"
    e94 = ("IBAN No:   ;TR980001200999900010101594;;;;;\n" + bas +
           "01.10.2026;r1;1.000,00;A;1.000,00;Gelen Firma Ltd,GELEN FİRMA LTD 'DAN ödeme\n"
           "02.10.2026;r2;300,00;B;700,00;HESAPLARIM ARASI TRANSFER,TEST ŞİRKET,TR41******************1553,TEST ŞİRKET,TR980001200999900010101594\n"
           "03.10.2026;r3;200,00;B;500,00;NAKİT ÇEKİLEN\n")
    yukle = lambda hid, metin: client.post(f"/api/banka-hesaplari/{hid}/ekstre?yontem=kural",  # noqa: E731
                                           files={"dosya": ("e.csv", metin.encode("cp1254"), "text/csv")}, headers=h).json()
    r = yukle(a94["id"], e94)
    assert r["eklenen"] == 3 and r["yeni_hesap"] == ["Halkbank 1553"], r
    hesaplar = {x["ad"]: x for x in client.get("/api/banka-hesaplari", headers=h).json()}
    a53 = hesaplar["Halkbank 1553"]
    taslak = client.get("/api/hareketler?durum=taslak", headers=h).json()
    assert {x["tur"] for x in taslak} == {"tahsilat", "virman", "nakit_cekim"}
    client.post("/api/hareketler/onayla", json={"idler": [x["id"] for x in taslak]}, headers=h)
    hesaplar = {x["id"]: x for x in client.get("/api/banka-hesaplari", headers=h).json()}
    assert hesaplar[a94["id"]]["bakiye"] == 500 and hesaplar[a94["id"]]["mutabakat_farki"] == 0
    assert hesaplar[a53["id"]]["bakiye"] == 300
    kasa = client.get("/api/kasa", headers=h).json()
    assert kasa["bakiye"] == 200

    # 53'ün ekstresi yüklenince aynı transfer ikinci kez yazılmaz, eşleşir; hesap IBAN'ı tamamlanır
    e53 = ("IBAN No:   ;TR410001200999900010101553;;;;;\n" + bas +
           "02.10.2026;r2;300,00;A;300,00;HESAPLARIM ARASI TRANSFER,TEST ŞİRKET,TR98******************1594,TEST ŞİRKET,TR410001200999900010101553\n")
    r = yukle(a94["id"], e53)
    assert (r["hesap"], r["eklenen"], r["eslesen"]) == ("Halkbank 1553", 0, 1), r
    hesaplar = {x["id"]: x for x in client.get("/api/banka-hesaplari", headers=h).json()}
    assert hesaplar[a53["id"]]["iban"] == "TR410001200999900010101553"
    assert hesaplar[a53["id"]]["bakiye"] == 300 and hesaplar[a53["id"]]["mutabakat_farki"] == 0

    # Kasadan cariye ödeme
    c = client.post("/api/cariler", json={"unvan": "Usta Ali"}, headers=h).json()
    client.post("/api/hareketler", json={"tarih": "2026-10-04", "yon": "cikis", "tur": "odeme", "tutar": 150, "cari_id": c["id"]}, headers=h)
    assert client.get("/api/kasa", headers=h).json()["bakiye"] == 50
    assert next(x for x in client.get("/api/raporlar/cari-bakiye", headers=h).json() if x["id"] == c["id"])["bakiye"] == 150

    # Bankadaki nakit çekim silinirse kasadaki karşılığı da silinir
    cekim = next(x for x in client.get("/api/hareketler?durum=onayli", headers=h).json() if x["tur"] == "nakit_cekim" and x["banka_hesap_id"])
    client.delete(f"/api/hareketler/{cekim['id']}", headers=h)
    assert client.get("/api/kasa", headers=h).json()["bakiye"] == -150


def test_toplu_fatura_islemleri():
    h = yeni()
    s = client.post("/api/santiyeler", json={"ad": "Merkez"}, headers=h).json()
    c = client.post("/api/cariler", json={"unvan": "Tedarikçi"}, headers=h).json()
    idler = [client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": f"F{i}", "tarih": "2026-10-01", "cari_id": c["id"],
             "kalemler": [{"aciklama": "x", "birim_fiyat": 100}]}, headers=h).json()["id"] for i in range(3)]
    r = client.post("/api/faturalar/toplu-guncelle", json={"idler": idler[:2], "santiye_id": s["id"], "gider_turu": "Beton"}, headers=h)
    assert r.json() == {"guncellenen": 2}
    f = {x["id"]: x for x in client.get("/api/faturalar", headers=h).json()}
    assert f[idler[0]]["santiye_ad"] == "Merkez" and f[idler[0]]["gider_turu"] == "Beton" and f[idler[2]]["santiye_id"] is None
    assert client.post(f"/api/faturalar/{idler[0]}/onayla", headers=h).status_code == 200
    r = client.post("/api/faturalar/toplu-sil", json={"idler": idler}, headers=h)
    assert r.json() == {"silinen": 2, "atlanan": 1}
    assert [x["id"] for x in client.get("/api/faturalar", headers=h).json()] == [idler[0]]


def test_pdf_kargo_ve_kodlu_kalemler(monkeypatch):
    """VKN 'V.D.' ile yazılmış, tutarlar TRY ile, kalemler malzeme koduyla başlıyor, vade not içinde; cari okurken açılır."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur([
        "Ornek Kargo Servisi A.S.", "Genel Mudurluk", "Ayazaga Mah. Kemerburgaz Cad. No:10", "Sariyer / Istanbul Tel: (0212) 365",
        "efatura@ornekkargo.com", "Buyuk Mukellefler V.D. 1111111111", "SAYIN", "TEST INSAAT LIMITED SIRKETI",
        "Yildirim Mah. No: 12", "Vergi No: 9999999999", "Fatura No: KRG2026004423895", "Fatura Tarihi: 15-08-2026",
        "Sira No Mal Hizmet Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani KDV Tutari Mal Hizmet Tutari",
        "1 HG.SR.01.01.005 Demontaj / Montaj Bedeli 1Adet 5,48TRY 0,00TRY %20,00 1,10TRY 5,48TRY",
        "2 Gonderici Odemeli Hizmet Geliri 1Adet 233,30TRY 0,00TRY %20,00 46,66TRY 233,30TRY",
        "Mal Hizmet Toplam Tutari 238,78TRY", "Toplam Iskonto 0,00TRY", "Hesaplanan KDV(%20) 47,76TRY",
        "Odenecek Tutar 286,54TRY", "Odeme Notu: Fatura vadesi 20.08.2026dir."])
    h = yeni()
    f = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    assert f["vade"] == "2026-08-20" and f["genel_toplam"] == 286.54
    assert [k["aciklama"] for k in f["kalemler"]] == ["Demontaj / Montaj Bedeli (HG.SR.01.01.005)", "Gonderici Odemeli Hizmet Geliri"]
    assert f["cari_unvan"] == "Ornek Kargo Servisi A.S."
    cariler = client.get("/api/cariler", headers=h).json()
    assert [(c["unvan"], c["vkn"]) for c in cariler] == [("Ornek Kargo Servisi A.S.", "1111111111")]
    # aynı firmanın ikinci faturası mevcut cariye işlenir
    client.post("/api/faturalar/oku", files={"dosya": ("f2.pdf", pdf, "application/pdf")}, headers=h)
    assert len(client.get("/api/cariler", headers=h).json()) == 1


def test_pdf_otel_konaklama_vergisi(monkeypatch):
    """Tutar sütunu KDV'den önce; konaklama vergisi 'diğer vergiler' olarak genel toplama eklenir."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur([
        "ORNEK OTELCILIK TURIZM LTD. STI.", "GURE MAH. TURGUT REIS SK. NO:1 EDREMIT", "Vergi Dairesi: Edremit", "VKN: 2222222222",
        "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999", "Fatura No: OTL2026000001660", "Tarih: 21-08-2026",
        "Sira No Mal Hizmet Miktar Birim Fiyat Mal Hizmet Tutari KDV Orani KDV Tutari Diger Vergiler",
        "1 Roommate 1 Adet 10.990,99 TL 10.990,99 TL %10,00 1.099,10 TL Konaklama Vergisi (%1,00)=109,91 TL",
        "2 Breakfast 1 Adet 720,72 TL 720,72 TL %10,00 72,07 TL",
        "Mal Hizmet Toplam Tutari: 11.711,71 TL", "Hesaplanan KDV (%10): 1.171,17 TL",
        "Hesaplanan Konaklama Vergisi (%1): 117,12 TL", "Vergiler Dahil Toplam Tutar: 13.000,00 TL", "Odenecek Tutar: 13.000,00 TL"])
    f = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=yeni()).json()[0]
    assert [(k["aciklama"], k["tutar"], k["kdv_orani"]) for k in f["kalemler"]] == [("Roommate", 10990.99, 10), ("Breakfast", 720.72, 10)]
    assert (f["diger_vergiler"], f["genel_toplam"]) == (117.12, 13000)
    assert f["gider_turu"] == "Konaklama"
    assert not any("tutmuyor" in u or "farklı" in u for u in f["ai_okuma"]["uyarilar"])


def test_eski_veritabani_guncellenir(tmp_path):
    from sqlalchemy import create_engine, inspect, text
    from app.db import db_guncelle
    motor = create_engine(f"sqlite:///{tmp_path / 'eski.db'}")
    with motor.begin() as b:
        b.execute(text("CREATE TABLE faturalar (id INTEGER PRIMARY KEY, tevkifat NUMERIC)"))
        b.execute(text("INSERT INTO faturalar (id, tevkifat) VALUES (1, 0)"))
    db_guncelle(motor)
    assert "diger_vergiler" in {c["name"] for c in inspect(motor).get_columns("faturalar")}
    with motor.connect() as b:
        assert b.execute(text("SELECT diger_vergiler FROM faturalar")).scalar() == 0


def test_ekstre_karsi_taraf_carisi():
    """Karşı taraf caride varsa (kısaltılmış yazılsa da) ona işlenir, yoksa yeni cari açılır; kendi firması virman kalır."""
    h, ahmet, _, banka = hazirlik()
    csv = ("Tarih;Tutar;B/A;Bakiye;Açıklama\n"
           "01.10.2026;1.000,00;B;-1.000,00;cari ödeme,AHMET YAPI SAN TIC LTD STI,Akbank T.A.Ş.,TR00\n"
           "02.10.2026;500,00;B;-1.500,00;kira ödemesi,mehmet yılmaz,Ziraat Bankası A.Ş.,TR00\n"
           "03.10.2026;700,00;B;-2.200,00;kira ödemesi,MEHMET YILMAZ,Ziraat Bankası A.Ş.,TR00\n"
           "04.10.2026;300,00;A;-1.900,00;ÖRNEK MÜHENDİSLİK LTD ŞTİ,ÖRNEK MÜHENDİSLİK LTD ŞTİ 'DAN avans,Garanti Bankası A.Ş.\n"
           "05.10.2026;250,00;B;-2.150,00;0000293907 |OEF2025113680 |TEST INSAAT|METEOR ENDÜSTRİ DBS\n"
           "06.10.2026;100,00;B;-2.250,00;geçiş,TEST İNŞAAT LTD ŞTİ,Yapı ve Kredi Bankası A.Ş.,TR00\n")
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre?yontem=kural",
                    files={"dosya": ("e.csv", csv.encode("cp1254"), "text/csv")}, headers=h).json()
    assert r["eklenen"] == 6 and r["yeni_cari"] == 3
    t = {x["tutar"]: x for x in client.get("/api/hareketler?durum=taslak", headers=h).json()}
    assert t[1000]["cari_id"] == ahmet["id"]
    assert t[500]["cari_unvan"] == "MEHMET YILMAZ" and t[700]["cari_id"] == t[500]["cari_id"]
    assert t[300]["cari_unvan"] == "ÖRNEK MÜHENDİSLİK LTD ŞTİ" and t[300]["tur"] == "tahsilat"
    assert t[250]["cari_unvan"] == "METEOR ENDÜSTRİ"
    assert t[100]["tur"] == "virman" and t[100]["cari_id"] is None


def test_kalem_bazinda_santiye():
    """Bir fatura iki şantiyeye ait: kalemlerde şantiye seçilir, boş kalem faturanın şantiyesine düşer."""
    h = yeni()
    a = client.post("/api/santiyeler", json={"ad": "A Şantiyesi"}, headers=h).json()
    b = client.post("/api/santiyeler", json={"ad": "B Şantiyesi"}, headers=h).json()
    c = client.post("/api/cariler", json={"unvan": "Beton AŞ"}, headers=h).json()
    f = client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": "F1", "tarih": "2026-10-05", "cari_id": c["id"],
                                            "santiye_id": a["id"], "gider_turu": "Beton", "kalemler": [
        {"aciklama": "Beton A", "miktar": 1, "birim_fiyat": 1000, "kdv_orani": 20},
        {"aciklama": "Beton B", "miktar": 1, "birim_fiyat": 3000, "kdv_orani": 20, "santiye_id": b["id"]}]}, headers=h).json()
    assert f["santiye_ad"] == "A Şantiyesi, B Şantiyesi" and f["kalemler"][1]["santiye_id"] == b["id"]
    client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    satis = client.post("/api/faturalar", json={"yon": "giden", "fatura_no": "S1", "tarih": "2026-09-30", "cari_id": c["id"],
                                                "santiye_id": b["id"], "kalemler": [{"aciklama": "Hakediş", "birim_fiyat": 10000, "kdv_orani": 20}]}, headers=h).json()
    client.post(f"/api/faturalar/{satis['id']}/onayla", headers=h)

    ozet = {s["ad"]: s for s in client.get("/api/raporlar/santiye-ozet", headers=h).json()}
    assert ozet["A Şantiyesi"]["gider"] == 1200 and ozet["B Şantiyesi"]["gider"] == 3600 and ozet["B Şantiyesi"]["gelir"] == 12000
    assert [x["fatura_no"] for x in client.get(f"/api/faturalar?santiye_id={b['id']}", headers=h).json()] == ["F1", "S1"]

    # Aylık alış listesi, yalnız B şantiyesi
    r = client.get(f"/api/raporlar/santiye-faturalari?yon=gelen&baslangic=2026-10-01&bitis=2026-10-31&santiye_ids={b['id']}", headers=h).json()
    assert [(x["santiye"], x["kalemler"], x["genel_toplam"]) for x in r] == [("B Şantiyesi", "Beton B", 3600)]
    assert [(d["aciklama"], d["tutar"], d["kdv"], d["toplam"]) for d in r[0]["kalem_detay"]] == [("Beton B", 3000, 600, 3600)]
    r = client.get(f"/api/excel/santiye-faturalari?yon=gelen&baslangic=2026-10-01&bitis=2026-10-31&santiye_ids={a['id']},{b['id']}", headers=h)
    ws = load_workbook(io.BytesIO(r.content)).active
    satirlar = [[c.value for c in s] for s in ws.iter_rows(min_row=2)]
    assert [s[0] for s in satirlar] == ["A Şantiyesi", "A Şantiyesi TOPLAMI", "B Şantiyesi", "B Şantiyesi TOPLAMI", "GENEL TOPLAMI"]
    assert [s[7] for s in satirlar] == ["Beton A", None, "Beton B", None, None]  # kalem kalem harcama
    assert satirlar[-1][10:13] == [4000, 800, 4800]


def test_turkce_harfli_kod_ve_vd_bicimi():
    assert kural._kodu_sona_al("K.Ö.HF.X.5X6 5x6mm KABLO") == "5x6mm KABLO (K.Ö.HF.X.5X6)"
    assert kural._VKN_SATIRI.search("Örnek Kurumlar V.D.  /1234567890").group("no") == "1234567890"
    kalemler = [{"aciklama": "Montaj bedeli", "tutar": 100}, {"aciklama": "Sigortalı hücre", "tutar": 5000}]
    assert kural.kalem_gider_turu(kalemler) == "Elektrik malzemesi"
    assert kural.kalem_gider_turu([{"aciklama": "Hizmet Geliri", "tutar": 10}], "Örnek Kargo A.Ş.") == "Nakliye"


def test_eski_veritabani_yeni_klasore_tasinir(tmp_path):
    """Paket her seferinde yeni klasöre açılsa da en çok kayıt içeren eski veritabanı sabit veri klasörüne taşınır."""
    import sqlite3

    from app import db as dbm
    eski_bos, eski_dolu = tmp_path / "Muhasebe2/backend/veri", tmp_path / "Muhasebe1/backend/veri"
    for klasor, kullanici in ((eski_bos, 0), (eski_dolu, 2)):
        klasor.mkdir(parents=True)
        with sqlite3.connect(klasor / "muhasebe.db") as b:
            for t in ("kullanicilar", "faturalar", "hareketler", "cariler", "cekler"):
                b.execute(f"CREATE TABLE {t} (id INTEGER)")
            b.executemany("INSERT INTO kullanicilar VALUES (?)", [(i,) for i in range(kullanici)])
    (eski_dolu / "belgeler").mkdir()
    (eski_dolu / "belgeler/f.pdf").write_bytes(b"%PDF")
    hedef = tmp_path / "Belgeler/Muhasebe Verileri"
    adaylar = dbm.eski_veritabanlari([tmp_path])
    assert len(adaylar) == 2
    assert dbm.eski_veriyi_tasi(hedef, adaylar) == eski_dolu / "muhasebe.db"
    assert (hedef / "muhasebe.db").exists() and (hedef / "belgeler/f.pdf").exists()
    assert dbm.eski_veriyi_tasi(hedef, adaylar) is None  # hedefte veri varsa dokunulmaz


def test_giris_hata_mesajlari():
    yeni()
    r = client.post("/api/giris", json={"eposta": "yok@ornek.com", "sifre": "x"})
    assert r.status_code == 401 and "kayıtlı hesap bulunamadı" in r.json()["detail"]
    e = f"k{uuid.uuid4().hex[:8]}@ornek.com"
    client.post("/api/kayit", json={"sirket_unvan": "Ş", "ad": "A", "eposta": e, "sifre": "gizli123"})
    r = client.post("/api/giris", json={"eposta": e, "sifre": "yanlis"})
    assert r.status_code == 401 and "Şifre yanlış" in r.json()["detail"]
    assert client.post("/api/giris", json={"eposta": f"  {e.upper()} ", "sifre": "gizli123"}).status_code == 200


def test_fatura_uyari_raporu():
    h = yeni()
    c = client.post("/api/cariler", json={"unvan": "Firma"}, headers=h).json()
    bos = client.post("/api/faturalar", json={"fatura_no": "B1", "cari_id": c["id"]}, headers=h).json()
    tamam = client.post("/api/faturalar", json={"fatura_no": "T1", "cari_id": c["id"], "kalemler": [{"aciklama": "x", "birim_fiyat": 5}]}, headers=h).json()
    r = client.get("/api/excel/fatura-uyarilari?yon=gelen", headers=h)
    satirlar = [[x.value for x in s] for s in load_workbook(io.BytesIO(r.content)).active.iter_rows(min_row=2)]
    assert [(s[0], s[2]) for s in satirlar] == [("İşlenemedi", "B1")] and tamam["id"] != bos["id"]


def test_ekstre_eft_ucreti_kredi_ve_toplu_atama():
    h = yeni()
    b = client.post("/api/banka-hesaplari", json={"ad": "Halk", "banka": "Halkbank"}, headers=h).json()
    csv = (";;;;;\nİŞLEM TARİHİ;BİM REF   ;TUTAR   ;B/A;BAKİYE   ;AÇIKLAMA\n"
           "02.10.2026;r1;200.000,00;B;800.000,00;Ahmet banko iade ödeme,Ahmet banko,Türkiye İş Bankası A.Ş.,TR930006400000100000000098\n"
           "02.10.2026;r2;16,76;B;799.983,24;Ahmet banko iade ödeme,Ahmet banko,Türkiye İş Bankası A.Ş.,TR930006400000100000000098\n"
           "02.10.2026;r3;22.329,82;B;777.653,42;01546 K0000001 NO'LU KREDİ HESABININ ANAPARA TAKSİT ÖDEMESİ\n"
           "02.10.2026;r4;280.000,00;B;497.653,42;1546-L0000002 KREDİ HESABI KAPANIŞI.\n")
    r = client.post(f"/api/banka-hesaplari/{b['id']}/ekstre?yontem=kural", files={"dosya": ("e.csv", csv.encode("cp1254"), "text/csv")}, headers=h).json()
    assert r.get("eklenen") == 4, r
    t = {x["tutar"]: x for x in client.get("/api/hareketler?durum=taslak", headers=h).json()}
    assert (t[200000]["tur"], t[200000]["cari_unvan"]) == ("odeme", "Ahmet banko")
    assert (t[16.76]["tur"], t[16.76]["cari_unvan"]) == ("banka_masrafi", "MUHTELİF BANKA İŞLEMLERİ")
    assert {t[22329.82]["tur"], t[280000]["tur"]} == {"kredi_odeme"} and t[280000]["cari_unvan"] == "MUHTELİF BANKA İŞLEMLERİ"
    assert not any(x["sorun"] for x in t.values())
    c = client.post("/api/cariler", json={"unvan": "Başka"}, headers=h).json()
    r = client.post("/api/hareketler/toplu-guncelle", json={"idler": [t[200000]["id"]], "cari_id": c["id"], "tur": "masraf"}, headers=h).json()
    assert r["guncellenen"] == 1
    x = next(x for x in client.get("/api/hareketler?durum=taslak", headers=h).json() if x["id"] == t[200000]["id"])
    assert (x["cari_id"], x["tur"]) == (c["id"], "masraf")


def _oku(satirlar, vkn="9999999999"):
    return kural.pdf_fatura_oku(pdf_olustur(satirlar), [], "DENEME INSAAT LTD. STI.", vkn, "gelen")


def test_pdf_abonelik_ozeti_kalemsiz():
    """Kalem tablosu olmayan internet/telefon faturası: KDV ve ÖİV matrah satırlarından kalem çıkar, genel toplam tutar."""
    r = _oku(["Vergi No: 9999999999", "Fatura ID A162026000000001", "Fatura Tarih 31-08-2026 23:59",
              "ORNEKNET ANONIM SIRKETI", "Ornek Mahallesi Deneme Sokak", "Vergi Numarasi: 1234567890",
              "SAYIN DENEME INSAAT LIMITED SIRKETI", "SON ODEME TARIHI : 16/09/2026", "ODENECEK TUTAR : 1.504,00 TL",
              "BIR SONRAKI FATURA DUZENLEME TARIHI : 30/09/2026", "Fiber DSL paket ucreti",
              "TOPLAM FATURA TUTARI 1.413,99", "Taksitler Toplami 90,00", "ODENECEK TOPLAM 1.504,00",
              "KDV %20 (Matrah 1.092.49 ) 218,50", "OIV %10 (Matrah 1.029.99 ) 103,00"])
    assert r["tarih"] == "2026-08-31" and r["vade"] == "2026-09-16"
    assert r["satici"]["vkn"] == "1234567890" and r["satici"]["unvan"] == "ORNEKNET ANONIM SIRKETI"
    assert (r["ara_toplam"], r["kdv_toplam"], r["diger_vergiler"], r["genel_toplam"]) == (1182.5, 218.5, 103, 1504)
    # Faturada ücret satırları yazmıyor: "Hizmet bedeli" gibi uydurma adla kalem açılmaz, kalemler elle girilir
    assert r["kalemler"] == [] and r["kalemler_okunamadi"]
    assert r["onerilen_gider_turu"] == "Haberleşme/İnternet"


def test_pdf_abonelik_ucret_satirlari_adlariyla():
    """Turkcell düzeni: vergisiz ücret satırları faturadaki adlarıyla kalem olur; aracılık edilen ödeme KDV'siz,
    %0'lık vergi satırı (telsiz taksidi) kendi adıyla; ödenecek tutar fatura tutarından farklıysa fatura tutarı esas."""
    r = _oku(["FATURA ID : 0012026000000001", "10/08/2026", "Vergi Dairesi/9999999999", "DENEME INSAAT",
              "LIMITED SIRKETI", "ODENECEK TUTAR : 1,000.00 TRY", "SON ODEME TARIHI : 21/08/2026", " ",
              "Tarife ve Paket Ucretleri 600.00", "KULLANIMLARINIZLA ILGILI UCRETLER", "Diger Ucretler 120.00",
              "Gecikme Bedeli 30.00", "Indirimler -50.00", "Duzeltmeler -0.40",
              "Tahsilatina Aracilik Edilen Odemeleriniz 100.00", "FATURADAKI DIGER KALEMLER", "FATURA TUTARI 1,060.40 TRY",
              "Katma Deger Vergisi %20 (Matrah:699.60 TRY) 139.92", "Ozel Iletisim Vergisi %10 (Matrah:699.60 TRY) 69.96",
              "Telsiz Kullanim Aylik Taksit %0 (Matrah:50.92 TRY) 50.92"], vkn="9999999999")
    assert r["satici"]["vkn"] is None and r["satici"]["unvan"] is None  # kendi adres bloğumuz satıcı olmaz
    assert [(k["aciklama"], k["tutar"], k["kdv_orani"]) for k in r["kalemler"]] == [
        ("Tarife ve Paket Ucretleri", 600, 20), ("Diger Ucretler", 120, 20), ("Gecikme Bedeli", 30, 20),
        ("Indirimler", -50, 20), ("Duzeltmeler", -0.4, 20), ("Tahsilatina Aracilik Edilen Odemeleriniz", 100, 0),
        ("Telsiz Kullanim Aylik Taksit", 50.92, 0)]
    assert round(sum(k.get("kdv_tutari") or 0 for k in r["kalemler"]), 2) == 139.92
    assert r["genel_toplam"] == 1060.4 and "fatura tutarıyla" in r["notlar"]


def test_pdf_abonelik_vergiler_dahil_satirlar():
    """TTNET düzeni: satırlar vergiler dahil; net tutarlar KDV matrahından dağıtılır, devir/taksit satırları KDV'siz."""
    r = _oku(["Vergi No: 9999999999", "Fatura ID A162026000000001", "Fatura Tarih 31-08-2026 23:59",
              "ORNEKNET ANONIM SIRKETI", "Vergi Numarasi: 1234567890", "SAYIN DENEME INSAAT LIMITED SIRKETI",
              "ODENECEK TUTAR : 1.430,00 TL", " ", "Tutar Indirim Toplam", "Internet 1.500,00 400,00 1.100,00",
              "Telefon 100,00 40,00 60,00", "Paket Ucreti 1.600,00 440,00 1.160,00", "Cihaz/ Baglanti Ucreti 80,00",
              "Gecikme Bedeli (Gunluk %0,13) 13,00", "TOPLAM FATURA TUTARI 1.253,00", "Taksitler Toplami 177,10",
              "Onceki Aydan Devir -0,10", "ODENECEK TOPLAM 1.430,00", "KDV %20 (Matrah 973.00 ) 194,60",
              "OIV %10 (Matrah 854.00 ) 85,40"], vkn="9999999999")
    assert [(k["aciklama"], k["kdv_orani"]) for k in r["kalemler"]] == [
        ("Internet", 20), ("Telefon", 20), ("Cihaz/ Baglanti Ucreti", 20), ("Gecikme Bedeli (Gunluk %0,13)", 20),
        ("Taksitler", 0), ("Onceki Aydan Devir", 0)]
    assert round(sum(k["tutar"] for k in r["kalemler"][:4]), 2) == 973
    assert round(sum(k["tutar"] for k in r["kalemler"]) + r["kdv_toplam"] + r["diger_vergiler"] + r["yuvarlama"], 2) == 1430


def test_pdf_abonelik_ozeti_ingilizce_sayi_ve_web_adresi():
    """Ünvanı yazmayan kurumsal GSM faturası: '4,231.30' biçimi okunur, satıcı web adresinden bulunur."""
    r = _oku(["FATURA ID : 0012026000000001", "10/08/2026", "Vergi Dairesi/9999999999", "www.orneknet.com.tr",
              "ODENECEK TUTAR : 6,700.97 TRY", "Katma Deger Vergisi %20 (Matrah:4,231.30 TRY) 846.26",
              "Ozel Iletisim Vergisi %10 (Matrah:4,231.30 TRY) 423.13",
              "Telsiz Kullanim Aylik Taksit %0 (Matrah:188.89 TRY) 188.89"])
    assert r["satici"]["unvan"] == "ORNEKNET"
    assert (r["kdv_toplam"], r["diger_vergiler"], r["genel_toplam"]) == (846.26, 423.13, 6700.97)
    assert round(r["ara_toplam"] + r["kdv_toplam"] + r["diger_vergiler"], 2) == 6700.97
    assert r["kalemler"] == []  # ücret satırları yazmıyor: uydurma "Hizmet bedeli" kalemi açılmaz


def test_pdf_sira_numarasiz_tablo_alici_once():
    """Önce alıcı, sonra satıcı yazılan ve kalemlerinde sıra no / KDV sütunu olmayan fatura."""
    r = _oku(["DENEME INSAAT LIMITED SIRKETI", "Vkn/Tckn : 9999999999", "Fatura No : SP12026000000001",
              "Fatura Tarihi : 11.08.2026", "Son Odeme", "Tarihi : 21.08.2026", "ETTN: 82f57c05-dc07-4b68-818c-",
              "70d4f3d1651c", "Ornek Iletisim A.S.", "Ornek Mah. Inonu Cad.", "Vergi Dairesi : Bogazici", "Vergi Kimlik",
              "No : 1234567890", "Urun/Hizmet Adi Miktar Birim Birim Fiyati Tutar",
              "Superbox Modem Kullanim Bedeli 1 Adet 111,67 TL 111,67 TL", "Toplam : 111,67 TL", "%20.0 KDV : 22,33 TL",
              "Indirim : 0 TL", "Genel Toplam : 134 TL", "Odenecek Tutar : 134 TL"])
    assert r["satici"] == {"unvan": "Ornek Iletisim A.S.", "vkn": "1234567890", "vergi_dairesi": None, "adres": None,
                           "telefon": None, "iban": None}
    assert r["vade"] == "2026-08-21"
    assert (r["ara_toplam"], r["kdv_toplam"], r["genel_toplam"]) == (111.67, 22.33, 134)
    assert [(k["aciklama"], k["birim_fiyat"], k["kdv_orani"]) for k in r["kalemler"]] == [
        ("Superbox Modem Kullanim Bedeli", 111.67, 20)]


def test_pdf_mal_kodlu_tablo_ve_ara_toplam():
    r = _oku(["Fatura No: REP2026000000001", "Fatura Tarihi: 31-08-2026", "SAYIN", "DENEME INSAAT LIMITED", "SIRKETI",
              "VKN: 9999999999", "ORNEK ENERJI DAGITIM", "A.S.", "Vergi Dairesi: Mecidiyekoy", "VKN: 1234567890",
              "MAL KODU MALIN/HIZMETIN CINSI MIKTAR BIRIM FIYAT ISKONTO ORANI KDV ORANI TUTAR",
              "01010020008  VADE FARKI  1,00 Adet  1.016,730000 TL     %20,00  1.016,73 TL",
              "Ara Toplam  1.016,73 TL", "Toplam Iskonto  0,00 TL", "KDV  203,35 TL", "Genel Toplam  1.220,08 TL"])
    assert (r["ara_toplam"], r["kdv_toplam"], r["genel_toplam"]) == (1016.73, 203.35, 1220.08)
    assert r["kalemler"][0]["aciklama"] == "VADE FARKI (01010020008)" and r["kalemler"][0]["kdv_orani"] == 20
    assert r["onerilen_gider_turu"] == "Vade farkı/Faiz"


def test_pdf_tevkifatli_diyez_baslikli_tablo():
    """'#' başlıklı tablo, satırda tevkifat kodu; ödenecek tutar etiketsiz, 'Vergi Dahil Toplam - Tevkifat' ile bulunur."""
    r = _oku(["Odenecek Tutar:", "ORNEK NAKLIYAT", "VKN/TCKN: 12345678901", "ALICI : DENEME INSAAT", "VKN/TCKN: 9999999999",
              "Fatura No: ALP2026000000039", "Tarih: 07.08.2026 20:18",
              "# Mal/Hizmet Miktar Birim Fiyat Tevkifat Kodu Tevkifat Tutari KDV KDV Tutari Tutar",
              "1 NAKLIYE BEDELI 1 Adet 18.000 TL 624 720,00 TL %20 3.600,00 TL 18.000,00 TL",
              "Mal Hizmet Toplam: 18.000,00 TL", "Vergi Haric Toplam: 18.000,00 TL", "KDV (%20) Matrah:",
              "Hesaplanan KDV (%20):", "18.000,00 TL", "3.600,00 TL", "Toplam Tevkifat: 720,00 TL",
              "Vergi Dahil Toplam: 21.600,00 TL", "20.880,00 TL"])
    assert (r["ara_toplam"], r["kdv_toplam"], r["tevkifat"], r["genel_toplam"]) == (18000, 3600, 720, 20880)
    assert r["kalemler"][0]["tutar"] == 18000 and r["onerilen_gider_turu"] == "Nakliye"


def test_pdf_no_baslikli_karisik_kdv_ve_alt_satirdaki_vkn():
    """'No Hizmet / Ürün' başlıklı tablo, %20 ve %10 karışık KDV; 'Vergi No' etiketi ile numara ayrı satırlarda."""
    r = _oku(["ORNEK OSGB IS SAGLIGI VE", "GUVENLIGI HIZMETLERI TICARET", "LIMITED SIRKETI", "Ornek Mah. Cesme Cad. No:33",
              "Telefon", "0 (543) 000 00 00", "Vergi No", "1234567890 (Ilyasbey)", "e-Fatura", "Sayin",
              "DENEME INSAAT LIMITED SIRKETI", "Vergi No", "9999999999 (CINILI)", "Fatura Numarasi", "NA12026000000001",
              "Fatura Tarihi", "31.08.2026 - 10:31:27",
              "No Hizmet / Urun Miktar Birim Birim Fiyat KDV KDV Tutari Toplam",
              "1 Is Guvenligi Uzmani", "Hizmeti 195 Saat 2.100,00 TL %20,00 81.900,00 TL 409.500,00 TL",
              "2 Is Yeri Hekimi", "Hizmeti 30 Saat 1.300,00 TL %10,00 3.900,00 TL 39.000,00 TL",
              "3 Diger Saglik", "Personeli Hizmeti 30 Saat 750,00 TL %10,00 2.250,00 TL 22.500,00 TL",
              "Mal Hizmet Toplam Tutari 471.000,00 TL", "Hesaplanan KDV GERCEK (%20.0) 81.900,00 TL",
              "Hesaplanan KDV GERCEK (%10.0) 6.150,00 TL", "Odenecek Tutar 559.050,00 TL"])
    assert r["satici"]["vkn"] == "1234567890"
    assert r["satici"]["unvan"] == "ORNEK OSGB IS SAGLIGI VE GUVENLIGI HIZMETLERI TICARET LIMITED SIRKETI"
    assert [(k["miktar"], k["birim"], k["kdv_orani"], k["tutar"]) for k in r["kalemler"]] == [
        (195, "Saat", 20, 409500), (30, "Saat", 10, 39000), (30, "Saat", 10, 22500)]
    assert (r["kdv_toplam"], r["genel_toplam"]) == (88050, 559050)
    assert r["onerilen_gider_turu"] == "Danışmanlık"


def test_pdf_birim_kodu_ve_istisna_satiri():
    """'1 C62' birim kodu ve KDV'siz satırda istisna sebebi yazısı (ör. 13/b) olan tablo."""
    r = _oku(["SAYIN", "DENEME INSAAT LTD STI", "VKN: 9999999999", "Fatura No: UW12026000237199", "Fatura Tarihi: 27-08-2026",
              "ORNEK SEYAHAT A.S.", "VKN: 1234567890",
              "Sira No Malzeme/Hizmet Aciklamasi Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani KDV Tutari "
              "Vergi Muafiyet Sebebi Malzeme/Hizmet Tutari",
              "1 28/08/2026 ESB GZT UCUS BILETI 1 C62 1.854,27 TL %0,00 0,00 TL %20 370,86 TL 1.854,27 TL",
              "2 Servis Ucreti : 1 C62 149,82 TL %0,00 0,00 TL %20 29,97 TL 149,82 TL",
              "3 Yakit Fiyat Farki : 1 C62 500,00 TL %0,00 0,00 TL %20 100,00 TL 500,00 TL",
              "4 Alan Vergisi: 1 C62 180,07 TL %0,00 0,00 TL %0 0,00 TL 13/b, Deniz ve hava tasima araclari icin",
              "liman ve hava meydanlarinda yapilan hizmetler 180,07 TL",
              "Malzeme/Hizmet Toplam Tutari 2.684,16 TL", "Hesaplanan GERCEK USULDE KATMA DEGER VERGISI(%20,00) 500,83 TL",
              "Vergiler Dahil Toplam Tutar 3.184,99 TL", "Odenecek Tutar 3.184,99 TL"])
    assert [(k["birim"], k["kdv_orani"], k["tutar"]) for k in r["kalemler"]] == [
        ("adet", 20, 1854.27), ("adet", 20, 149.82), ("adet", 20, 500), ("adet", 0, 180.07)]
    assert (r["kdv_toplam"], r["genel_toplam"]) == (500.83, 3184.99)
    assert r["onerilen_gider_turu"] == "Seyahat/Bilet"


def test_fatura_alti_iskonto_kalemlere_eksi_satir_olarak_islenir():
    r = _oku(["ORNEK AKARYAKIT A.S.", "VKN: 1234567890", "SAYIN", "DENEME INSAAT LTD STI", "VKN: 9999999999",
              "Fatura No: OEF2026000105379", "Fatura Tarihi: 31-08-2026",
              "Sira No Mal Hizmet Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani KDV Tutari Mal Hizmet Tutari",
              "1 KURSUNSUZ BENZIN 95 149,54 lt 60,68415 TL %20,00 1.814,94 TL 9.074,71 TL",
              "2 MOTORIN 172,29 lt 68,95946 TL %20,00 2.376,21 TL 11.881,03 TL",
              "Mal Hizmet Toplam Tutari 20.955,74 TL", "Toplam Iskonto 943,01 TL", "KDV Matrahi 20.012,73 TL",
              "Hesaplanan KDV(20.012,73 %20) 4.002,55 TL", "Vergiler Dahil Toplam Tutar 24.015,28 TL"])
    assert [(k["aciklama"], k["tutar"]) for k in r["kalemler"]][-1] == ("İskonto (%4,50)", -943.01)
    assert round(sum(k["tutar"] for k in r["kalemler"]), 2) == r["ara_toplam"] == 20012.73
    assert (r["kdv_toplam"], r["genel_toplam"]) == (4002.55, 24015.28)
    assert "iskonto var" not in r["notlar"]
    # İki ayrı KDV oranında iskonto her orana payı kadar bölünür
    k = kural._iskonto_kalemleri([{"tutar": 300, "kdv_orani": 20}, {"tutar": 100, "kdv_orani": 10}], 40)
    assert [(x["kdv_orani"], x["tutar"]) for x in k] == [(20, -30), (10, -10)]


def test_satirdaki_tevkifat_ve_parantezsiz_kdv_orani():
    r = _oku(["ORNEK METAL A.S.", "VKN: 1234567890", "SAYIN", "DENEME INSAAT LTD STI", "VKN: 9999999999",
              "Fatura No: FYS2026000004068", "Fatura Tarihi: 19-08-2026",
              "Sira No Malzeme / Hizmet Aciklamasi Miktar Birim Fiyat KDV Orani KDV Tutari Malzeme / Hizmet Tutari",
              "1 BAK-SOMB-001 Som Bakir 105 KG 840,0000 TL % 20,00 17.640,00 TL Tevkifat ( 619 - BAKIR TESLIMI )",
              "(% 70,00 )= 12.348,00 TL 88.200,00 TL", "Mal / Hizmet Toplam Tutari 88.200,00 TL",
              "Hesaplanan KDV (% 20,00 )", "Vergiler Dahil Toplam Tutar 105.840,00 TL", "Odenecek Tutar 93.492,00 TL"])
    assert (r["kdv_toplam"], r["tevkifat"], r["genel_toplam"]) == (17640, 12348, 93492)
    assert [(k["miktar"], k["kdv_orani"], k["tutar"]) for k in r["kalemler"]] == [(105, 20, 88200)]
    r = _oku(["Fatura No: HAZ2026000001211", "Fatura Tarihi: 30-08-2026", "ORNEK BETON LTD STI", "VKN: 1234567890",
              "Sira No Urun Adi Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
              "1 HAZIR BETON C25/30 12 M3 3.500,00 TL %20,00 8.400,00 TL 42.000,00 TL",
              "Mal Hizmet Toplam Tutari 42.000,00 TL", "Hesaplanan KDV %20 8.400,00 TL", "Hesaplanan OTV",
              "Odenecek Tutar 50.400,00 TL"])
    assert (r["kdv_toplam"], r["diger_vergiler"], r["genel_toplam"]) == (8400, 0, 50400)
    assert (r["kalemler"][0]["miktar"], r["kalemler"][0]["birim"]) == (12, "M3")


def test_bos_iskonto_sutunlari_ve_sifir_tutarli_kalemler():
    """Boş iskonto sütunları '|  |' olarak basılan tablo; 0 TL'lik satırlar da kalem olarak gelir."""
    r = _oku(["ORNEK ELEKTRONIK A.S.", "VKN: 1234567890", "SAYIN", "DENEME INSAAT LTD STI", "VKN: 9999999999",
              "Fatura No: NTS2026000000001", "Fatura Tarihi: 31-08-2026",
              "Sira No Mal Hizmet Satici Kodu Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani KDV Tutari "
              "Diger Vergiler Mal Hizmet Tutari",
              "1 ORNEK DIAFON PANELI M.U.AL", "G.00006 11 Adet 0,00 TL    |  | %20,00 0,00 TL 0,00 TL",
              "2 BAKIM ONARIM BKM00", "01 1 Adet 7.916,67 TL    |  | %20,00 1.583,33 TL 7.916,67 TL",
              "Mal Hizmet Toplam Tutari 7.916,67 TL", "Toplam Iskonto 0,00 TL",
              "Hesaplanan KATMA DEGER VERGISI(%20.00) 1.583,33 TL", "Vergiler Dahil Toplam Tutar 9.500,00 TL"])
    assert [(k["miktar"], k["tutar"]) for k in r["kalemler"]] == [(11, 0), (1, 7916.67)]
    assert not r["kalemler_okunamadi"] and r["genel_toplam"] == 9500


_UST = ["ORNEK TICARET A.S.", "VKN: 1234567890", "SAYIN", "DENEME INSAAT LTD STI", "VKN: 9999999999",
        "Fatura No: ORN2026000000001", "Fatura Tarihi: 31-08-2026"]
_BASLIK = ("Sira No Mal Hizmet Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani KDV Tutari Diger Vergiler "
           "Mal Hizmet Tutari")


def _kalemler(r):
    return [(k["miktar"], k["birim_fiyat"], k["kdv_orani"], k["tutar"]) for k in r["kalemler"]]


def test_satirda_kdv_tevkifati_ve_bitisik_birim():
    """'1Adet', boş iskonto tutarı ve iki satıra bölünen 'KDV TEVKIFAT (%20,00)=1.000,00 TL' (nakliye faturası)."""
    r = _oku(_UST + [_BASLIK, "1 Ankara Gebze", "Nakliye Bedeli 1Adet 25.000 TL %0 %20,00 5.000,00", "TL",
                     "KDV TEVKIFAT", "(%20,00)=1.000,00 TL 25.000,00 TL",
                     "2 Ankara Kutahya", "Nakliye Bedeli 1Adet 20.000TL %0 %20,00 4.000,00 TL KDV TEVKIFAT",
                     "(%20,00)=800,00 TL 20.000,00 TL",
                     "Mal Hizmet Toplam Tutari 45.000,00 TL", "Hesaplanan KDV (%20) 9.000,00 TL",
                     "KDV Tevkifat Tutari 1.800,00 TL", "Vergiler Dahil Toplam Tutar 54.000,00 TL",
                     "Odenecek Tutar 52.200,00 TL"])
    assert _kalemler(r) == [(1, 25000, 20, 25000), (1, 20000, 20, 20000)]
    assert r["kalemler"][0]["aciklama"] == "Ankara Gebze Nakliye Bedeli"
    assert (r["tevkifat"], r["genel_toplam"], r["kalemler_okunamadi"]) == (1800, 52200, False)


def test_satirda_konaklama_vergisi_ve_kdv_oncesi_aciklama():
    r = _oku(_UST + [_BASLIK, "1 KONAKLAMA BEDELI 1 Adet 22.500 TL %10,00 2.250,00 TL KONAKLAMA VERGISI (%1,00)=225,00 TL "
                     "22.500,00 TL", "Mal Hizmet Toplam Tutari 22.500,00 TL", "Hesaplanan KDV (%10) 2.250,00 TL",
                     "Konaklama Vergisi 225,00 TL", "Odenecek Tutar 24.975,00 TL"])
    assert _kalemler(r) == [(1, 22500, 10, 22500)]
    # Fiyat ile KDV arasında ürün açıklaması ya da iskonto nedeni
    r = _oku(_UST + ["Sira No Mal Hizmet Adi Miktar Birim Fiyat Aciklama KDV Orani KDV Tutari Mal Hizmet Tutari",
                     "1 ORNEK BETON KOSK 1 Adet 178.015,86 TL AS26-00553-14-K %20,00 35.603,17 TL 178.015,86 TL",
                     "2 KLIMA MONTAJ 1,0 Adet 9.000 TL %0,00 0,00 TL Iskonto - %20,00 1.800,00 TL 9.000,00 TL",
                     "Mal Hizmet Toplam Tutari 187.015,86 TL", "Hesaplanan KDV (%20) 37.403,17 TL",
                     "Odenecek Tutar 224.419,03 TL"])
    assert _kalemler(r) == [(1, 178015.86, 20, 178015.86), (1, 9000, 20, 9000)]


def test_satir_iskontosu_net_tutarla_ve_bolunmus_yazimlar():
    """Tutar sütunu brüt, KDV net üzerinden: kalem net işlenir. '%20,0 0' ve '152.915,00T L' bölünmeleri."""
    r = _oku(_UST + [_BASLIK, "1 GORUNTULU DIAFON 342 Adet 4.300,000000TL %70,00 1.029.420,00 TL %20,00 88.236,00TL "
                     "1.470.600,00TL", "2 ZIL PANELI 17 Adet 12.850,000000TL %70,00 152.915,00T L %20,00 13.107,00TL 218.450,00TL",
                     "3 KABLO KBL.N 641 Mt 138,07 TL %20,0 0 17.700,57 TL 88.502,87 TL",
                     "Mal Hizmet Toplam Tutari 1.777.552,87 TL", "Toplam Iskonto 1.182.335,00 TL",
                     "Hesaplanan KDV (%20) 119.043,57 TL", "Odenecek Tutar 714.261,44 TL"])
    assert _kalemler(r) == [(342, 4300, 20, 441180), (17, 12850, 20, 65535), (641, 138.07, 20, 88502.87)]
    assert [(k["iskonto_orani"], k["iskonto_tutari"]) for k in r["kalemler"]] == [(70, 1029420), (70, 152915), (0, 0)]
    assert r["kalemler"][0]["aciklama"] == "GORUNTULU DIAFON"
    assert r["onerilen_gider_turu"] == "Elektrik malzemesi"


def test_fiyatsiz_akaryakit_ve_kdv_tutarsiz_satirlar():
    r = _oku(_UST + ["Sira No Satici Urun Kodu Mal Hizmet Miktar Birim Mal Hizmet Tutari",
                     "1 VIS 500003761 Motorin - Diesel 449,19 LT 30.486,23 TL 2 VIS 500006552 Kursunsuz 95 331,95 LT 20.398,52 TL",
                     "Mal Hizmet Toplam Tutari 50.884,75 TL", "Hesaplanan KDV (%20) 10.176,95 TL",
                     "Odenecek Tutar 61.061,70 TL"])
    assert [(k["miktar"], k["birim"], k["kdv_orani"], k["tutar"]) for k in r["kalemler"]] == [
        (449.19, "LT", 20, 30486.23), (331.95, "LT", 20, 20398.52)]
    assert r["onerilen_gider_turu"] == "Akaryakıt"
    r = _oku(_UST + ["No Hizmet / Urun Miktar Birim Fiyat KDV Toplam", "1 galvanizli tel 1,5 100,00 TL %20,00 150,00 TL",
                     "Senaryo TICARIFATURA", "Mal Hizmet Toplam Tutari 150,00 TL", "Hesaplanan KDV (%20) 30,00 TL",
                     "Odenecek Tutar 180,00 TL"])
    assert _kalemler(r) == [(1.5, 100, 20, 150)]


def test_kdv_tutari_etiketli_satir_ve_toplama_girmeyen_bilgi_satiri():
    """Muayene istasyonu: tutar sütunu yok; 'Kartlı ödemeye konu toplam tutar' satırı fatura toplamına girmez."""
    r = _oku(_UST + ["Sira No Mal Hizmet Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani Diger Vergiler",
                     "1 1010-Kamyonet PM 1 2.740TL %20,00 KDV Tutari 548,00TL",
                     "2 Kartli Odeme Hizmetine Konu Toplam Tutar 1 3.288,00TL %20,00 KDV Tutari 0,00TL",
                     "3 Odeme Kurulusu Hizmet Bedeli 1 139,22TL %20,00 KDV Tutari 27,84TL",
                     "Mal Hizmet Toplam Tutari 2.879,22TL", "Hesaplanan KDV(%20.00) 575,84TL",
                     "Odenecek Tutar 3.455,06TL"])
    assert [(k["aciklama"], k["tutar"]) for k in r["kalemler"]] == [
        ("1010-Kamyonet PM", 2740), ("Odeme Kurulusu Hizmet Bedeli", 139.22)]


def test_yanlis_baslik_atlanir_ve_kdv_dahil_tutar():
    """'# Yalniz ...' başlık sanılmaz; asıl tablo 'SIRA' ile sonra gelir ve tutar sütunu KDV dahildir."""
    r = _oku(_UST + ["# Yalniz BinUcYuzDoksanDokuz TL #", "ARA TUTAR 1.165,83TL", "HESAPLANAN KDV(%20.0) 233,17TL",
                     "ODENECEK TUTAR 1.399,00TL", "SIRA ACIKLAMA MIKTAR BIRIM FIYAT KDV ORANI KDV TUTARI TUTAR",
                     "1 SSI-6632 Buharli Utu 1 Adet 1.165,8334TL %20,00 233,17TL 1.399,00TL", "ALICI BILGILERI"])
    assert _kalemler(r) == [(1, 1165.8334, 20, 1165.83)]


def test_elektrik_faturasi_kademeler_ve_btv():
    r = _oku(["ORNEK ELEKTRIK PERAKENDE SATIS A.S.", "VKN: 1234567890", "Fatura No: EAH2026000000001",
              "FATURA TARIHI: 12-08-2026", "FATURA BILGILERI", "ACIKLAMA MIKTAR BIRIM FIYAT NET TUTAR",
              "Dusuk Kademe ENERJI TUKETIM BEDELI 960KWH 5,352448 TL 5.138,35 TL",
              "Yuksek Kademe ENERJI TUKETIM BEDELI 301,778KWH 5,934064 TL 1.790,77 TL", "TALIMATLI BANKA VAR TRT PAYI",
              "BTV 190.04 TL", "KDV MATRAHI (%20) 7119,16 TL", "KDV (%20) 1423,83 TL", "FATURA TUTARI 8.542,99 TL",
              "ODENECEK TUTAR 8.542,99 TL"])
    assert [(k["aciklama"][:12], k["miktar"], k["tutar"]) for k in r["kalemler"]] == [
        ("Dusuk Kademe", 960, 5138.35), ("Yuksek Kadem", 301.778, 1790.77), ("BTV", 1, 190.04)]
    assert r["onerilen_gider_turu"] == "Elektrik/Su/Doğalgaz"


def test_tarihle_baslayan_tasima_satiri():
    r = _oku(_UST + ["ATF Tarihi ATF No Gonderen / Alici Adet Birim Kg Desi Birim Fiyat Tutar",
                     "7.08.2026 275817 GELIS ATF222774 ORNEK KOMPOZIT LTD.STI. 1 PALET", "PALET VB 0,000 TL 6.000,00 TL",
                     "Malzeme / Hizmet Toplam Tutari 6.000,00 TL", "Hesaplanan KDV (%20) 1.200,00 TL",
                     "Odenecek Tutar 7.200,00 TL"])
    assert _kalemler(r) == [(1, 6000, 20, 6000)] and r["kalemler"][0]["aciklama"].startswith("7.08.2026")


def test_kalemleri_okunamayan_fatura_kirmizi_isaretlenir():
    r = _oku(_UST + ["Bu faturada kalem tablosu yok", "Mal Hizmet Toplam Tutari 1.000,00 TL",
                     "Hesaplanan KDV (%20) 200,00 TL", "Odenecek Tutar 1.200,00 TL"])
    assert r["kalemler_okunamadi"] and "tek tek okunamadı" in r["notlar"]
    assert kural.gider_turu_tahmin("Okuma donemi") is None  # 'okuma' içinde 'kum' aranmaz


def test_unvan_bloku_adres_ve_etiket_satirlarindan_ayrilir():
    """Ünvan üç satır + iki satır adres (Nova); 'TEL ÇİT' ünvanın parçası, 'Ticaret Sicil No' etiket (şahıs firması)."""
    r = _oku(["Fatura No: NVC2026000000001", "Fatura Tarihi: 03-08-2026", "ORNEK GRUP MUHENDISLIK PROJE ELEKTRIK",
              "ELEKTRONIK ILETISIM SISTEMLERI INSAAT VE", "TEKS.SAN.TIC.LTD.STI.", "Asagi Ovecler Mah. 1332.Cad. No:8 / Cankaya /",
              "ANKARA Ovecler", "Tel: 312311 02 00", "Vergi Dairesi: Baskent", "VKN: 1234567890", "SAYIN",
              "DENEME INSAAT LTD STI", "VKN: 9999999999", "Mal Hizmet Toplam Tutari 100,00 TL", "Hesaplanan KDV (%20) 20,00 TL",
              "Odenecek Tutar 120,00 TL"])
    assert r["satici"]["unvan"] == "ORNEK GRUP MUHENDISLIK PROJE ELEKTRIK ELEKTRONIK ILETISIM SISTEMLERI INSAAT VE TEKS.SAN.TIC.LTD.STI."
    r = _oku(["AHMET YILMAZ GARANTI TEL CIT", "17. Sk 7 / Yeni Sanayi Sitesi Kutahya", "Yeni sanayi sitesi, KUTAHYA - Turkiye",
              "Vergi No", "12345678901 (CINILI)", "Ticaret Sicil No", "-", "Mersis No", "-", "Sayin", "DENEME INSAAT LTD STI",
              "Vergi No", "9999999999 (30 AGUSTOS)", "Fatura Numarasi", "GR02026000000001", "Fatura Tarihi", "17.08.2026",
              "No Hizmet / Urun Miktar Birim Fiyat Indirim Tutari KDV Toplam",
              "1 75MM DATA BORU 300 37,50 TL 0,00 TL %20,00 11.250,00 TL", "Senaryo TICARIFATURA",
              "Mal Hizmet Toplam Tutari 11.250,00 TL", "Hesaplanan KDV (%20) 2.250,00 TL", "Odenecek Tutar 13.500,00 TL"])
    assert r["satici"] == {**r["satici"], "unvan": "AHMET YILMAZ GARANTI TEL CIT", "vkn": "12345678901"}
    assert _kalemler(r) == [(300, 37.5, 20, 11250)]
    assert [kural.kotu_unvan(x) for x in ["Ticaret Sicil No", "Mersis No", "Vergi No", "AHMET YILMAZ GARANTI TEL CIT"]] == [
        True, True, True, False]


def test_kalem_ozdenetimi_miktar_fiyat_tutar():
    k = kural._kalemleri_denetle([
        {"aciklama": "001 MOTORIN 1.000,21 lt", "miktar": 1, "birim": None, "birim_fiyat": 63.53358, "tutar": 63546.92},
        {"aciklama": "KUTU KOLA", "miktar": 2, "birim": "Adet", "birim_fiyat": 45455, "tutar": 90.91},
        {"aciklama": "EK MUFU", "miktar": 2, "birim": "Adet", "birim_fiyat": 875, "tutar": 1137.5},
        {"aciklama": "Netelsan panel", "miktar": 11, "birim": "Adet", "birim_fiyat": 0, "tutar": 0}])
    assert [(x["aciklama"], x["miktar"], x["birim_fiyat"]) for x in k] == [
        ("001 MOTORIN", 1000.21, 63.53358), ("KUTU KOLA", 2, 45.455),
        ("EK MUFU (İsk. %35, liste fiyatı 875)", 2, 568.75), ("Netelsan panel", 11, 0)]


@pytest.mark.parametrize("dip", [
    ["Genel Toplam 71.600,00 TL", "Toplam Iskonto 42.960,00 TL", "Hesaplanan KDV(%20) 5.728,00 TL"],
    ["Genel Toplam 71.600,00 TL", "Toplam Iskonto 42.960,00 TL", "Net Toplam 28.640,00 TL", "Hesaplanan KDV(%20) 5.728,00 TL",
     "Vergiler Dahil Toplam 34.368,00 TL"],
])
def test_pdf_genel_toplam_iskonto_oncesi_brut(monkeypatch, dip):
    """'Genel Toplam' iskonto öncesi brüt olabilir; ödenecek tutar sanılmaz, fark KDV'siz satırla kapatılmaz.
    Tabloda KDV oranı sütunu yok: İskonto1, İskonto2, Net Birim Fiyat, Net Tutar, KDV Tutarı, KDV'li Net Tutar."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur([
        "ORNEK ELEKTRIK INSAAT A.S.", "VKN: 3333333333", "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999",
        "Fatura No: OEL2026000017962", "Fatura Tarihi: 06-08-2026",
        "Sira No Mal Hizmet Kodu Mal Hizmet Adi Miktar Birim Fiyat Iskonto1 Iskonto2 Net Birim Fiyat Net Tutar KDV Tutari",
        "KDV'li Net Tutar",
        "1 NET-", "MU.VYU.00004 NETELSAN EQMAX/XL VIDEO", "YUKSELTICI(1GIRIS/4 CIKIS) 16 Adet 1.600 TL %60,00 640,00 10.240,00 TL 2.048,00 TL 12.288,00 TL",
        "2 NET-MU.PEQ- NETELSAN SMALL KAMERALI", "00004 DIJITAL TUSLU KARTLI GECISLI ZIL PANELI 2 Adet 12.600 TL %60,00 5.040,00 10.080,00 TL 2.016,00 TL 12.096,00 TL",
        "3 NET- EQMAX ADAPTOR-SMPS 24V", "MU.SMP.00004 3.2A 8 Adet 2.600 TL %60,00 1.040,00 8.320,00 TL 1.664,00 TL 9.984,00 TL",
        *dip])
    f = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=yeni()).json()[0]
    assert (f["ara_toplam"], f["kdv_toplam"], f["genel_toplam"]) == (28640, 5728, 34368)
    assert [(k["miktar"], k["birim_fiyat"], k["iskonto_orani"], k["iskonto_tutari"], k["kdv_orani"], k["tutar"])
            for k in f["kalemler"]] == [(16, 1600, 60, 15360, 20, 10240), (2, 12600, 60, 15120, 20, 10080), (8, 2600, 60, 12480, 20, 8320)]
    assert "VIDEO" in f["kalemler"][0]["aciklama"] and "ZIL PANELI" in f["kalemler"][1]["aciklama"]
    assert not f["ai_okuma"]["kalemler_okunamadi"]


def test_pdf_satir_tevkifati_oran_ve_kod(monkeypatch):
    """Dört kalemin ikisi tevkifatlı nakliye (2/10, 624): oran ve kod satıra yazılır, fatura tevkifatı satırların toplamıdır."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur([
        "ORNEK MADENCILIK NAKLIYE LTD. STI.", "VKN: 4444444444", "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999",
        "Fatura No: ONM2026000000764", "Fatura Tarihi: 10-08-2026", "Fatura Tipi : TEVKIFAT",
        "Sira No Mal Hizmet Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani KDV Tutari Diger Vergiler Mal Hizmet Tutari",
        "1 KUM 0-5 MM 100 Ton 300 TL %0,00 0,00 TL %20,00 6.000,00 TL 30.000,00 TL",
        "2 NAKLIYE BEDELI (TEVKIFATLI 2/10) 87,73 Ton 264 TL %0,00 0,00 TL %20,00 4.632,14 TL",
        "KDV TEVKIFAT (%20,00)=926,43 TL 23.160,72 TL",
        "3 MICIR 50 Ton 400 TL %0,00 0,00 TL %20,00 4.000,00 TL 20.000,00 TL",
        "4 NAKLIYE BEDELI (TEVKIFATLI 2/10) 34,55 Ton 264 TL %0,00 0,00 TL %20,00 1.824,24 TL",
        "KDV TEVKIFAT (%20,00)=364,85 TL 9.121,20 TL",
        "Tevkifat Sebebi: 624-2/10 YUK TASIMACILIGI HIZMETI",
        "Mal Hizmet Toplam Tutari 82.281,92 TL", "Hesaplanan KDV(%20) 16.456,38 TL", "Hesaplanan KDV Tevkifat(%20) : 1.291,28 TL",
        "Vergiler Dahil Toplam Tutar 98.738,30 TL", "Odenecek Tutar 97.447,02 TL"])
    h = yeni()
    f = client.post("/api/faturalar/oku", files={"dosya": ("t.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    assert [(k["tevkifat_orani"], k["tevkifat_kodu"]) for k in f["kalemler"]] == [(0, None), (2, "624"), (0, None), (2, "624")]
    assert (f["tevkifat"], f["genel_toplam"]) == (1291.28, 97447.02)
    assert not [u for u in f["ai_okuma"].get("uyarilar", []) if "farklı" in u]


def test_kalem_tevkifati_satis_faturasi_rapor_ve_wolvox_aktarimi():
    """Satış faturasında kalem bazında tevkifat seçilir; fatura tevkifatı elle girilen değil satırların toplamıdır.
    Tevkifat raporu koda göre gruplar; Wolvox kalem aktarımında oran, kod ve tutar sütunları vardır."""
    h = yeni()
    cari = client.post("/api/cariler", json={"unvan": "Ornek Belediye", "vkn": "5555555555"}, headers=h).json()
    kalemler = [{"aciklama": "Yapım işi hakedişi", "miktar": 1, "birim_fiyat": 100000, "kdv_orani": 20,
                 "tevkifat_orani": 4, "tevkifat_kodu": "601"},
                {"aciklama": "Malzeme", "miktar": 1, "birim_fiyat": 10000, "kdv_orani": 20}]
    f = client.post("/api/faturalar", json={"yon": "giden", "fatura_no": "AFS2026000000001", "tarih": "2026-09-05",
                                            "cari_id": cari["id"], "tevkifat": 999, "kalemler": kalemler}, headers=h).json()
    assert (f["kdv_toplam"], f["tevkifat"], f["genel_toplam"]) == (22000, 8000, 124000)
    assert client.post(f"/api/faturalar/{f['id']}/onayla", headers=h).status_code == 200
    gelen = client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": "NAK2026000000001", "tarih": "2026-09-10",
                                                "cari_id": cari["id"], "kalemler": [{"aciklama": "Nakliye", "birim_fiyat": 5000,
                                                "kdv_orani": 20, "tevkifat_orani": 2, "tevkifat_kodu": "624"}]}, headers=h).json()
    client.post(f"/api/faturalar/{gelen['id']}/onayla", headers=h)
    r = client.get("/api/raporlar/tevkifat?yil=2026&ay=9", headers=h).json()
    assert [(g["kod"], g["oran"], g["tevkifat"]) for g in r["satis"]] == [("601", "4/10", 8000)]
    assert [(g["kod"], g["oran"], g["matrah"], g["tevkifat"]) for g in r["alis"]] == [("624", "2/10", 5000, 200)]
    assert (r["kdv2"], r["satis_tevkifat"]) == (200, 8000)
    assert client.get("/api/raporlar/tevkifat?yil=2026&ay=8", headers=h).json()["faturalar"] == []
    assert client.get("/api/excel/tevkifat?yil=2026&ay=9", headers=h).status_code == 200
    ws = load_workbook(io.BytesIO(client.get("/api/excel/fatura-kalemleri", headers=h).content)).active
    satirlar = [[c.value for c in r] for r in ws.iter_rows()]
    i = satirlar[0].index("Tevkifat Oranı")
    assert satirlar[0][i:i + 3] == ["Tevkifat Oranı", "Tevkifat Kodu", "Tevkifat Tutarı"]
    assert [x[i:i + 3] for x in satirlar[1:]] == [["4/10", "601", 8000], [None, None, None], ["2/10", "624", 200]]
    # Kalemlerde oran yoksa (eski kayıt) faturaya yazılan tevkifat kalır
    eski = client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": "ESK1", "cari_id": cari["id"], "tevkifat": 100,
                                               "kalemler": [{"aciklama": "x", "birim_fiyat": 1000, "kdv_orani": 20}]}, headers=h).json()
    assert (eski["tevkifat"], eski["genel_toplam"]) == (100, 1100)
    assert "TEVKİFAT" in client.get(f"/api/yazdir/fatura/{f['id']}", headers=h).text


def test_pdf_satirda_yazan_kdv_tutari_esas_alinir(monkeypatch):
    """Öztaş: 1. satırda oran %20 ama KDV 833,00 yazılmış; faturanın KDV toplamı da buna göre. Yazan tutar esas alınır."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = pdf_olustur([
        "ORNEK PARKE BETON LTD. STI.", "VKN: 4444444444", "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999",
        "Fatura No: OPB2026000000764", "Fatura Tarihi: 21-08-2026",
        "Sira No Mal Hizmet Miktar Birim Fiyat Iskonto Orani Iskonto Tutari KDV Orani KDV Tutari Diger Vergiler Mal Hizmet Tutari",
        "1 KUM 100 Ton 300 TL %0,00 0,00 TL %20,00 833,00 TL 30.000,00 TL",
        "2 NAKLIYE BEDELI (TEVKIFATLI 2/10) 87,73 Ton 264 TL %0,00 0,00 TL %20,00 4.632,14 TL",
        "KDV TEVKIFAT (%20,00)=926,43 TL 23.160,72 TL",
        "3 MICIR 50 Ton 400 TL %0,00 0,00 TL %20,00 4.000,00 TL 20.000,00 TL",
        "4 NAKLIYE BEDELI (TEVKIFATLI 2/10) 34,55 Ton 264 TL %0,00 0,00 TL %20,00 1.824,24 TL",
        "KDV TEVKIFAT (%20,00)=364,85 TL 9.121,20 TL", "Tevkifat Sebebi: 624-2/10 YUK TASIMACILIGI HIZMETI",
        "Mal Hizmet Toplam Tutari 82.281,92 TL", "Hesaplanan KDV ( % 20.0 ) 11.289,38 TL", "Hesaplanan KDV Tevkifat(%20) : 1.291,28 TL",
        "Vergiler Dahil Toplam Tutar 93.571,30 TL", "Odenecek Tutar 92.280,02 TL"])
    h = yeni()
    f = client.post("/api/faturalar/oku", files={"dosya": ("o.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    assert [k["kdv_tutari"] for k in f["kalemler"]] == [833, None, None, None]
    assert (f["kdv_toplam"], f["tevkifat"], f["genel_toplam"]) == (11289.38, 1291.28, 92280.02)
    assert not [u for u in f["ai_okuma"]["uyarilar"] if "farklı" in u]
    # Kullanıcı satırın KDV tutarını düzeltirse toplamlar ona göre hesaplanır
    kalemler = f["kalemler"]
    kalemler[0]["kdv_tutari"] = 6000
    g = client.put(f"/api/faturalar/{f['id']}", json={**{a: f[a] for a in ("yon", "fatura_no", "tarih", "cari_id")},
                                                      "kalemler": kalemler}, headers=h).json()
    assert (g["kdv_toplam"], g["genel_toplam"]) == (16456.38, 97447.02)
    # Eski sürüm kaydı (KDV oran x tutar): açılıştaki onarım faturadaki KDV tutarını yazar
    from app import onarim
    from app.db import SessionLocal
    from app.models import Fatura, Sirket
    client.put("/api/sirket", json={"unvan": "Test Insaat Ltd.", "vkn": "9999999999"}, headers=h)
    db = SessionLocal()
    try:
        fa = db.get(Fatura, f["id"])
        fa.ai_okuma = {**fa.ai_okuma}
        rapor = onarim.onar(db, db.get(Sirket, fa.sirket_id))
    finally:
        db.close()
    assert rapor["kdv_tutari_duzeltilen"] == ["OPB2026000000764"]
    g = client.get(f"/api/faturalar/{f['id']}", headers=h).json()
    assert (g["kdv_toplam"], g["genel_toplam"]) == (11289.38, 92280.02)


def tablolu_pdf(ust, baslik, satirlar, alt):
    """Çizgili kalem tablosu olan PDF (e-Fatura görüntüsü gibi); hücre içinde alt satıra bölünmüş metin olabilir."""
    platypus = pytest.importorskip("reportlab.platypus")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    t = io.BytesIO()
    kucuk = ParagraphStyle("k", fontSize=6, leading=7)
    hucre = lambda x: platypus.Paragraph(x.replace("\n", "<br/>"), kucuk) if "\n" in x else x  # noqa: E731
    tablo = platypus.Table([[hucre(x) for x in baslik]] + [[hucre(x) for x in s] for s in satirlar])
    tablo.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 6)])
    platypus.SimpleDocTemplate(t, pagesize=landscape(A4)).build(
        [platypus.Paragraph(x) for x in ust] + [tablo] + [platypus.Paragraph(x) for x in alt])
    return t.getvalue()


def test_kalem_tablosu_basliklara_gore_okunur(monkeypatch):
    """Firmaya özel düzen aranmaz: sütunlar başlık adından bulunur (4M Altın Elektrik düzeni: KDV oranı sütunu yok,
    İskonto1/İskonto2, Net Birim Fiyat, Net Tutar, KDV Tutarı, KDV'li Net Tutar; kod hücresi iki satıra bölünmüş)."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = tablolu_pdf(
        ["ORNEK ELEKTRIK INSAAT A.S.", "VKN: 3333333333", "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999",
         "Fatura No: OEL2026000017962", "Fatura Tarihi: 06-08-2026"],
        ["Sira No", "Mal Hizmet Kodu", "Mal Hizmet Adi", "Miktar", "Birim Fiyat", "Iskonto1", "Iskonto2", "Net Birim Fiyat",
         "Net Tutar", "KDV Tutari", "KDV'li Net Tutar"],
        [["1", "NET-\nMU.VYU.00004", "NETELSAN EQMAX/XL VIDEO\nYUKSELTICI", "16 Adet", "1.600 TL", "%60,00", "", "640,00",
          "10.240,00 TL", "2.048,00 TL", "12.288,00 TL"],
         ["2", "NET-MU.PEQ-\n00004", "NETELSAN SMALL KAMERALI\nZIL PANELI", "2 Adet", "12.600 TL", "%60,00", "", "5.040,00",
          "10.080,00 TL", "2.016,00 TL", "12.096,00 TL"],
         ["3", "NET-\nMU.SMP.00004", "EQMAX ADAPTOR-SMPS 24V\n3.2A", "8 Adet", "2.600 TL", "%60,00", "", "1.040,00",
          "8.320,00 TL", "1.664,00 TL", "9.984,00 TL"]],
        ["Genel Toplam 71.600,00 TL", "Toplam Iskonto 42.960,00 TL", "Hesaplanan KDV(%20) 5.728,00 TL"])
    f = client.post("/api/faturalar/oku", files={"dosya": ("t.pdf", pdf, "application/pdf")}, headers=yeni()).json()[0]
    assert [k["aciklama"] for k in f["kalemler"]] == [
        "NETELSAN EQMAX/XL VIDEO YUKSELTICI (NET-MU.VYU.00004)",
        "NETELSAN SMALL KAMERALI ZIL PANELI (NET-MU.PEQ-00004)",
        "EQMAX ADAPTOR-SMPS 24V 3.2A (NET-MU.SMP.00004)"]
    assert [(k["miktar"], k["birim"], k["birim_fiyat"], k["iskonto_orani"], k["iskonto_tutari"], k["kdv_orani"], k["tutar"])
            for k in f["kalemler"]] == [(16, "Adet", 1600, 60, 15360, 20, 10240), (2, "Adet", 12600, 60, 15120, 20, 10080),
                                        (8, "Adet", 2600, 60, 12480, 20, 8320)]
    assert (f["ara_toplam"], f["kdv_toplam"], f["genel_toplam"]) == (28640, 5728, 34368)


def test_okuma_raporu_genel_kurallari_etiketli_taraf_noktali_kdv_oran_matrahi():
    """Gebze İş Güvenliği düzeni: 'Satıcı (Şube):' etiketli ünvan, 'Vergi No / Dairesi: VKN / VD', K.D.V. yazımı,
    'Ürün / Hizmet Toplam Tutarı', oran başına matrah; alıcı ünvanında satır sonuna kalan tek harf ('SIRKET I')."""
    r = _oku(["Fatura No: GEB2026000002976", "Fatura Tarihi: 08/07/2026", "Urun / Hizmet Toplam Tutari: 6.600,00 TL",
              "K.D.V. Matrahi (%20): 800,00 TL", "Hesaplanan K.D.V. (%20): 160,00 TL", "K.D.V. Matrahi (%10): 5.800,00 TL",
              "Hesaplanan K.D.V. (%10): 580,00 TL", "Vergiler Dahil Toplam Tutar: 7.340,00 TL", "ODENECEK TUTAR: 7.340,00 TL",
              "Satici (Sube): ORNEK IS GUVENLIGI INS TEKSTIL", "SAN ve TIC LTD STI", "Adres: Guzeller Mah. Ibrahimaga Cad.",
              "No:36/A", "GEBZE, KOCAELI, TURKIYE", "Telefon / Faks: 2620000000", "Vergi No / Dairesi: 3800000002 / ILYASBEY",
              "Alici (Sube): DENEME INSAAT LIMITED", "SIRKET I", "Adres: YILDIRIM BEYAZIT MAH.",
              "Vergi No / Dairesi: 9999999999 / 30 AGUSTOS", "Sira No Urun / Hizmet Adi Miktar Birim Birim Fiyat K.D.V. Orani "
              "K.D.V. Tutari Diger Vergiler Urun / Hizmet Tutari",
              "1 FOSFORLU IKAZ YELEGI 50 Adet 70,00 TL %10 350,00 TL 3.500,00 TL",
              "2 IS AYAKKABISI 1 Adet 2.300,00 TL %10 230,00 TL 2.300,00 TL",
              "3 BARET TIRNAK AYARLI 10 Adet 80,00 TL %20 160,00 TL 800,00 TL"])
    assert (r["satici"]["unvan"], r["satici"]["vkn"]) == ("ORNEK IS GUVENLIGI INS TEKSTIL SAN ve TIC LTD STI", "3800000002")
    assert (r["ara_toplam"], r["kdv_toplam"], r["genel_toplam"]) == (6600, 740, 7340)
    assert [k["aciklama"] for k in r["kalemler"]] == ["FOSFORLU IKAZ YELEGI", "IS AYAKKABISI", "BARET TIRNAK AYARLI"]
    assert [k["kdv_orani"] for k in r["kalemler"]] == [10, 10, 20]


def test_okuma_raporu_genel_kurallari_ek_tutar_yuvarlama_ve_sahis_unvani():
    """Kalem tablosu dışında matraha eklenen 'Kargo Ücreti' (Toplam Ek Tutar) faturadaki adıyla kalem olur, KDV oranı
    o oranın matrahından bulunur. 'Yuvarlama 0,05 TL' KDV'siz kalem olur. Şahıs ünvanındaki '(00)' kodu atılır;
    ünvanın üstündeki 'Bu fatura ... ile üretilmiştir' satırı ünvana katılmaz."""
    r = _oku(["SAYIN", "DENEME INSAAT LTD STI", "VKN: 9999999999", "Fatura No: AVA2026000285213", "Fatura Tarihi: 18-07-2026",
              "Bu Fatura E-Donusum Merkezi ile Uretilmistir", "MEHMET ORNEK(00)", "KURTULUS MAH.127.SK NO:12/2", "AHLAT/ Bitlis",
              "TCKN: 10000000146", _BASLIK, "1 FOTOKOPI KAGIDI A4 2 Adet 663,55 TL %10,00 132,71 TL 1.327,10 TL",
              "Mal Hizmet Toplam Tutari 1.327,10 TL", "Kargo Ucreti 124,92 TL", "Toplam Ek Tutar 124,92 TL",
              "KDV Matrahi (%10) 1.327,10 TL", "Hesaplanan KDV (%10) 132,71 TL", "KDV Matrahi (%20) 124,92 TL",
              "Hesaplanan KDV (%20) 24,98 TL", "Vergiler Dahil Toplam Tutar 1.609,71 TL", "Yuvarlama 0,05 TL",
              "Odenecek Tutar 1.609,76 TL"])
    assert (r["satici"]["unvan"], r["satici"]["vkn"]) == ("MEHMET ORNEK", "10000000146")
    assert [(k["aciklama"], k["kdv_orani"], k["tutar"]) for k in r["kalemler"]] == [
        ("FOTOKOPI KAGIDI A4", 10, 1327.10), ("Kargo Ucreti", 20, 124.92), ("Yuvarlama", 0, 0.05)]
    assert (r["ara_toplam"], r["kdv_toplam"], r["genel_toplam"]) == (1452.07, 157.69, 1609.76)


def test_yuzde_yuz_iskontolu_bedelsiz_kalem(monkeypatch):
    """Nova: '120 CM L AYAK 6 Adet 566,99 İskonto %100' tutar sütununda brüt yazar, KDV 0: Mal Hizmet Tutarı 0,
    iskonto sütunlarında %100 ve tutarı; miktar x fiyat - iskonto = tutar."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = tablolu_pdf(_UST, ["Sira No", "Mal Hizmet", "Miktar", "Birim Fiyat", "Iskonto Orani", "Iskonto Tutari", "KDV Orani",
                             "KDV Tutari", "Mal Hizmet Tutari"],
                      [["1", "CANAK ANTEN", "28 Adet", "1.889,98 TL", "", "", "%20,00", "10.583,89 TL", "52.919,44 TL"],
                       ["2", "120 CM L AYAK", "6 Adet", "566,99 TL", "%100,00", "3.401,94 TL", "%20,00", "0,00 TL", "3.401,94 TL"]],
                      ["Mal Hizmet Toplam Tutari 56.321,38 TL", "Toplam Iskonto 3.401,94 TL", "Hesaplanan KDV(%20) 10.583,89 TL",
                       "Odenecek Tutar 63.503,33 TL"])
    r = kural.pdf_fatura_oku(pdf, [], "DENEME INSAAT LTD. STI.", "9999999999", "gelen")
    k = r["kalemler"][1]
    assert (k["miktar"], k["birim_fiyat"], k["iskonto_orani"], k["iskonto_tutari"], k["tutar"]) == (6, 566.99, 100, 3401.94, 0)
    assert (r["ara_toplam"], r["kdv_toplam"], r["genel_toplam"]) == (52919.44, 10583.89, 63503.33)


def test_alt_satira_tasan_mal_hizmet_adi():
    """'1 ELEKTRIK 1Adet 625TL % 20 125,00TL 625,00TL' ve alt satırda 'MALZEMESI': ad tam okunur; alttaki ETTN ya da
    'Senaryo ...' bilgisi ada katılmaz. Parçalı harfli metin: kural._BOLUNMUS_HARF ('ELEKTR İ K') tanınır."""
    r = _oku(["CAPRAZ ELEKTRIK", "Vergi No 20000000046", "SAYIN", "DENEME INSAAT LTD STI", "Vergi Kimlik No 9999999999",
              "Fatura No 4DH2026000000062", "Fatura Tarihi 28-08-2026",
              "# Mal/Hizmet Miktar Birim Fiyat Iskonto / Arttirim Orani KDV Orani KDV Tutari Diger Vergiler Toplam",
              "1 ELEKTRIK 1Adet 625TL % 20 125,00TL 625,00TL", "MALZEMESI",
              "ETTN: be08a3c7-42be-4b6f-a246-424ba3a1d959 Mal Hizmet Toplam Tutari 625,00TL",
              "Hesaplanan KDV (%20,00) 125,00TL", "Odenecek Tutar 750,00TL"])
    assert [(k["aciklama"], k["tutar"]) for k in r["kalemler"]] == [("ELEKTRIK MALZEMESI", 625)]
    assert len(kural._BOLUNMUS_HARF.findall("ELEKTR İ K MALZEMES İ Y İĞİ THAN")) == 3


def test_nakliye_faturasi_tablosu():
    """Nakliye düzeni: taşınan iş "Gönderen / Alıcı" sütununda, miktar "Adet" sütununda; KDV oranı tabloda yok,
    altta "Hesaplanan KDV (%20)" yazar."""
    platypus = pytest.importorskip("reportlab.platypus")
    from reportlab.lib import colors
    t = io.BytesIO()
    belge = platypus.SimpleDocTemplate(t)
    tablo = platypus.Table([
        ["ATF Tarih", "ATF No", "Musteri Irs.No", "Gonderen / Alici", "Adet", "Birim", "Mamul Cinsi", "Kg Desi",
         "Birim Fiyat", "Tutar"],
        ["", "", "", "BALIKESIR-SINDIRGI ARA NAKLIYE", "1", "ADET", "", "", "6.000,000 TL", "6.000,00 TL"],
        ["", "", "", "", "", "", "", "", "", ""]],
        style=[("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 6)])
    belge.build([platypus.Paragraph("NAKLIYE TASIMACILIK LTD STI"), platypus.Paragraph("VKN: 1234567890"),
                 platypus.Paragraph("Fatura No: EF02026000001091"), platypus.Paragraph("Fatura Tarihi: 14-09-2026"), tablo,
                 platypus.Paragraph("Malzeme / Hizmet Toplam Tutari 6.000,00 TL"),
                 platypus.Paragraph("Hesaplanan KDV (%20) 1.200,00 TL"),
                 platypus.Paragraph("Vergiler Dahil Toplam Tutar 7.200,00 TL"), platypus.Paragraph("Odenecek Tutar 7.200,00 TL")])
    o = kural.pdf_fatura_oku(t.getvalue())
    assert [(k["aciklama"], k["miktar"], k["birim_fiyat"], k["kdv_orani"], k["tutar"]) for k in o["kalemler"]] == [
        ("BALIKESIR-SINDIRGI ARA NAKLIYE", 1, 6000, 20, 6000)]


def konumlu_pdf(satirlar):
    """Çizgisiz PDF: her satır [(x, yazı), ...]; tablo çizgisi yok, sütunlar yalnız konumla belli."""
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    t = io.BytesIO()
    c = canvas.Canvas(t)
    c.setFont("Helvetica", 9)
    y = 800
    for s in satirlar:
        for x, yazi in ([(40, s)] if isinstance(s, str) else s):
            c.drawString(x, y, yazi)
        y -= 14
    c.save()
    return t.getvalue()


def test_cizgisiz_tablo_kelime_konumlarindan_okunur():
    pdf = konumlu_pdf(["ORNEK E TICARET LTD STI", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD STI", "VKN: 9999999999",
                       "Fatura No: TEK2026000006237", "Fatura Tarihi: 21.06.2026",
                       [(40, "Kod"), (100, "Mal / Hizmet"), (330, "Miktar"), (390, "Birim Fiyat"), (460, "Indirim %"), (530, "Tutar")],
                       [(40, "20.05.298"), (100, "KABLOSUZ SUPURGE XW7263"), (330, "1 Adet"), (390, "11.665,83"), (520, "11.665,83 TRY")],
                       "Toplam Tutar 11.665,83 TRY", "KDV (%20) 2.333,17 TRY", "Vergiler Dahil Toplam 13.999,00 TRY",
                       "Odenecek Tutar 13.999,00 TRY"])
    o = kural.pdf_fatura_oku(pdf)
    assert [(k["aciklama"], k["miktar"], k["birim_fiyat"], k["tutar"]) for k in o["kalemler"]] == [
        ("KABLOSUZ SUPURGE XW7263 (20.05.298)", 1, 11665.83, 11665.83)]
    assert (o["ara_toplam"], o["kdv_toplam"], o["genel_toplam"]) == (11665.83, 2333.17, 13999)


def test_kdv_dahil_satirlar_net_tutara_cevrilir():
    """Market faturası: satır tutarları KDV dahil, "Ara Toplam" da KDV dahil; matrah oran başına yazılır."""
    pdf = tablolu_pdf(["MARKET TICARET A.S.", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999",
                       "Fatura No: M022026000184613", "Fatura Tarihi: 24-01-2026"],
                      ["Urun Kodu", "Urun Adi", "Miktar", "Birim", "Birim Fiyat", "Brut Tutar", "Iskonto Tutar",
                       "Isk. Sonrasi Brut Tutar", "KDV"],
                      [["125", "SUZME YOGURT 400 G", "1", "Adet", "101,00 TRY", "101.00", "", "101.00", "1"],
                       ["320", "PLASTIK POSET", "4", "Adet", "1,50 TRY", "6.00", "", "6.00", "20"]],
                      ["ARA TOPLAM : 107.00", "FATURA TOPLAMI : 107.00", "TOPLAM KDV : 2.00",
                       "K.D.V. MATRAHI % 1 : 100.00", "K.D.V. % 1 : 1.00", "K.D.V. MATRAHI % 20 : 5.00", "K.D.V. % 20 : 1.00"])
    o = kural.pdf_fatura_oku(pdf)
    assert (o["ara_toplam"], o["kdv_toplam"], o["genel_toplam"]) == (105, 2, 107)
    assert [(k["aciklama"], k["miktar"], k["birim_fiyat"], k["kdv_orani"], k["tutar"]) for k in o["kalemler"]] == [
        ("SUZME YOGURT 400 G (125)", 1, 100, 1, 100), ("PLASTIK POSET (320)", 4, 1.25, 20, 5)]


def test_otel_konaklama_vergisi_satirda_ve_tutar_vergiler_dahil():
    pdf = tablolu_pdf(["ORNEK KONAKLAMA LTD. STI.", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999",
                       "Fatura No: TKH2026000000413", "Fatura Tarihi: 18-04-2026"],
                      ["Sira No", "Mal Hizmet", "Miktar", "Birim Fiyat", "KDV Orani", "KDV Tutari", "Diger Vergiler",
                       "Mal Hizmet Tutari"],
                      [["1", "KON.BEDELI (Oda No: 402)", "1 Adet", "2.800 TL", "%10,00", "250,00 TL",
                        "KONAKLAMA VERGISI (%2,00)=50,00 TL", "2.800,00 TL"]],
                      ["Mal Hizmet Toplam Tutari 2.500,00 TL", "Hesaplanan KDV (%10) 250,00 TL",
                       "Vergiler Dahil Toplam Tutar 2.800,00 TL", "Odenecek Tutar 2.800,00 TL"])
    o = kural.pdf_fatura_oku(pdf)
    assert [(k["aciklama"], k["birim_fiyat"], k["kdv_orani"], k["tutar"]) for k in o["kalemler"]] == [
        ("KON.BEDELI (Oda No: 402)", 2500, 10, 2500)]
    assert (o["ara_toplam"], o["kdv_toplam"], o["diger_vergiler"], o["genel_toplam"]) == (2500, 250, 50, 2800)


def test_kasa_fisi_duzeni_mersisten_vkn_kendi_vkn_satici_olmaz():
    """Mağaza fişi düzeninde e-Fatura: satıcı VKN'si etiketsiz, "TCKN/VKN" bizim; tarih "Tarih: 29. 03. 2026";
    kalemler "% 20 11.999,00" (KDV dahil) satırlarıyla."""
    pdf = pdf_olustur(["TEST INSAAT LIMITED SIRKETI", "ORNEK MAGAZA TURKEY TICARET LIMITED", "SIRKETI", "  6100000006",
                       "Mersis No: 0610000000600256", "e-FATURA", "TCKN/VKN: 9999999999", "Tarih: 29. 03. 2026",
                       "Fatura No: 5I82026000000197", "Kasiyer: 4243 Magaza: T026",
                       "40594645.1248994 ANDROID TABLET", "LENOVO IDEA TAB 11", "Seri No: HA2BN697", "% 20 11.999,00",
                       "67.1195257 SERVIS POSET", "% 20 1,00", "KDV Toplam: 2.000,00", "Toplam: 12.000,00",
                       "Net Mal Degeri: 10.000,00"])
    o = kural.pdf_fatura_oku(pdf, None, "TEST INSAAT LIMITED SIRKETI", "9999999999")
    assert o["satici"]["vkn"] == "6100000006" and o["satici"]["unvan"] == "ORNEK MAGAZA TURKEY TICARET LIMITED SIRKETI"
    assert o["tarih"] == "2026-03-29"
    assert (o["ara_toplam"], o["kdv_toplam"], o["genel_toplam"]) == (10000, 2000, 12000)
    assert [(k["aciklama"], k["kdv_orani"], k["tutar"]) for k in o["kalemler"]] == [
        ("ANDROID TABLET LENOVO IDEA TAB 11", 20, 9999.17), ("SERVIS POSET", 20, 0.83)]


def test_kendi_vknmiz_alicida_iki_kez_yazsa_da_satici_olmaz():
    pdf = pdf_olustur(["SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999", "Vergi Dairesi: MERKEZ", "VKN: 9999999999",
                       "Fatura No: RMD2026000001848", "Fatura Tarihi: 28-02-2026", "Mal Hizmet Toplam Tutari 100,00 TL",
                       "Hesaplanan KDV %20 20,00 TL", "Odenecek Tutar 120,00 TL", "ORNEK TURIZM A.S.", "Vergi Dairesi: SELCUK",
                       "VKN: 1100000001"])
    o = kural.pdf_fatura_oku(pdf, None, "TEST INSAAT LTD. STI.", "9999999999")
    assert o["satici"]["vkn"] == "1100000001"


def test_kurus_farki_olan_kalemler_alinir_fark_uyari_olur(monkeypatch):
    """Satırlar kendi içinde tutarlı (19 x 496,67) ama toplamı faturadakinden 6 kuruş fazla (satıcının yuvarlaması):
    kalemler alınır, fark kalem uydurulmadan "yuvarlama" alanına yazılır, kayıt faturadaki toplamla yapılır (sarı not)."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    pdf = tablolu_pdf(["ORNEK YAPI MALZEME", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD. STI.", "VKN: 9999999999",
                       "Fatura No: GEF2026000000240", "Fatura Tarihi: 22-09-2026"],
                      ["Sira No", "Mal Hizmet", "Miktar", "Birim Fiyat", "KDV Orani", "KDV Tutari", "Mal Hizmet Tutari"],
                      [["1", "CIMENTO 50 KG", "2Adet", "165TL", "%20,00", "66,00TL", "330,00TL"],
                       ["2", "KORGE BORU 6 MT", "19Adet", "496,67TL", "%20,00", "1.887,33TL", "9.436,73TL"]],
                      ["Mal Hizmet Toplam Tutari 9.766,67TL", "Toplam Iskonto 0,00TL", "Hesaplanan KDV (%20) 1.953,33TL",
                       "Odenecek Tutar 11.720,00TL"])
    f = client.post("/api/faturalar/oku", files={"dosya": ("g.pdf", pdf, "application/pdf")}, headers=yeni()).json()[0]
    assert [k["aciklama"] for k in f["kalemler"]] == ["CIMENTO 50 KG", "KORGE BORU 6 MT"]
    assert (f["yuvarlama"], f["kdv_toplam"], f["genel_toplam"]) == (-0.06, 1953.33, 11720)
    assert any("0,06 TL yuvarlama" in u for u in f["ai_okuma"]["bilgi"]) and not f["ai_okuma"]["uyarilar"]


def test_oran_basina_tevkifat_satirlari_toplanir():
    pdf = pdf_olustur(["ORNEK METAL LTD STI", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD STI", "VKN: 9999999999",
                       "Fatura No: OZM2026000000115", "Fatura Tarihi: 11-02-2026", "Mal Hizmet Toplam Tutari 10.834,00 TL",
                       "Hesaplanan KDV(%20.0) 2.166,80 TL", "Hesaplanan KDV Tevkifat(%50) 1.000,16 TL",
                       "Hesaplanan KDV Tevkifat(%20) 33,36 TL", "Odenecek Tutar 11.967,28 TL"])
    assert kural.pdf_fatura_oku(pdf)["tevkifat"] == 1033.52


def test_doviz_faturasi_doviz_cinsinden_okunur():
    pdf = pdf_olustur(["ORNEK ELEKTRIK SIST. A.S.", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD STI", "VKN: 9999999999",
                       "Fatura No: VEK2026000010362", "Fatura Tarihi: 30-06-2026",
                       "Malzeme/Hizmet Toplam Tutari 34.500,00 USD", "Hesaplanan KDV(%20,00) 6.900,00 USD",
                       "Odenecek Tutar 41.400,00 USD", "Doviz Kuru(USD) 46,4935 TL",
                       "Malzeme/Hizmet Toplam Tutari 1.604.025,75 TL", "Hesaplanan KDV(%20,00) 320.805,15 TL",
                       "Odenecek Tutar 1.924.830,90 TL"])
    o = kural.pdf_fatura_oku(pdf)
    assert (o["ara_toplam"], o["kdv_toplam"], o["genel_toplam"], o["para_birimi"]) == (34500, 6900, 41400, "USD")


def test_doviz_kuru_basta_yazsa_da_toplamlar_okunur():
    """'Döviz Kuru: 55,28 TL' başlıkta (toplamlardan önce) yazıyor: toplamlar kesilmez, kur alana gelir."""
    pdf = pdf_olustur(["Fatura No: EB12026000000001", "Fatura Tarihi: 07-10-2026", "Doviz Kuru: 55.284700 TL",
                       "ORNEK ITHALAT LTD STI", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD STI", "VKN: 9999999999",
                       "Mal Hizmet Toplam Tutari 694,40 EUR", "Hesaplanan KDV(%20) 138,88 EUR",
                       "Vergiler Dahil Toplam Tutar 833,28 EUR", "Odenecek Tutar 833,28 EUR",
                       "Mal Hizmet Toplam Tutari(TL) 38.389,70 TL", "Odenecek Tutar(TL) 46.067,63 TL"])
    o = kural.pdf_fatura_oku(pdf)
    assert (o["ara_toplam"], o["kdv_toplam"], o["genel_toplam"], o["para_birimi"], o["kur"], o["tl_toplam"]) == (
        694.4, 138.88, 833.28, "EUR", 55.2847, 46067.63)


def test_doviz_kuru_fatura_kuru_ve_tl_toplamindan():
    satirlar = ["ORNEK ENERJI A.S.", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD STI", "VKN: 9999999999",
                "Fatura No: E002026000000001", "Fatura Tarihi: 13-07-2026", "Mal Hizmet Toplam Tutari: 4.210,90 USD",
                "Hesaplanan KDV (%0): 0,00 USD", "Odenecek Tutar: 4.210,90 USD"]
    o = kural.pdf_fatura_oku(pdf_olustur(satirlar + ["Fatura Kuru: 45,9911 TL", "Mal ve Hizmet Toplam Tutari (TL): 193.663,92 TL"]))
    assert (o["genel_toplam"], o["kur"]) == (4210.9, 45.9911)
    # Kur yazmıyor, yalnız TL karşılığı yazıyor: kur ondan bulunur
    o = kural.pdf_fatura_oku(pdf_olustur(satirlar + ["Mal ve Hizmet Toplam Tutari (TL): 193.663,92 TL"]))
    assert o["kur"] == 45.9911


def test_kalem_adindaki_sus_isaretleri_atilir():
    assert kural.susu_temizle({"aciklama": "## ULS 12U RACK KABINET (SOHO) ## (OEM-110)"})["aciklama"] == \
        "ULS 12U RACK KABINET (SOHO) (OEM-110)"
    assert kural.susu_temizle({"aciklama": "C# kursu"})["aciklama"] == "C# kursu"


def test_tcmb_kuru_listeden(monkeypatch):
    from datetime import date

    from app import tcmb
    xml = (b'<?xml version="1.0"?><Tarih_Date><Currency Kod="USD" CurrencyCode="USD"><Unit>1</Unit>'
           b'<ForexBuying>45.9911</ForexBuying></Currency><Currency Kod="JPY"><Unit>100</Unit>'
           b'<ForexBuying>30.50</ForexBuying></Currency></Tarih_Date>')
    istenen = []
    monkeypatch.setattr(tcmb, "_liste", lambda gun: istenen.append(gun) or (xml if gun == date(2026, 7, 10) else None))
    # 13.07.2026 pazartesi: bir önceki iş günü (cuma 10.07) yayımlanan liste
    assert tcmb.kur("USD", date(2026, 7, 13)) == (45.9911, date(2026, 7, 10))
    assert istenen == [date(2026, 7, 10)]
    assert tcmb.kur_bul(xml, "JPY") == 0.305


def test_karekod_fatura_tutari_ve_kdvsiz_kalemler():
    """Ödenecek tutardan önceki ödeme düşülmüşse fatura tutarı; KDV'siz kalemli faturada ara toplam mal/hizmet toplamı."""
    from app import karekod
    t = karekod.toplamlar({"malhizmettoplam": "7910.25", "kdvmatrah(20)": "6371.51", "hesaplanankdv(20)": "1274.30",
                           "vergidahil": "9819.60", "odenecek": "9540.00"})
    assert (t["ara_toplam"], t["genel_toplam"]) == (7910.25, 9819.6)
    t = karekod.toplamlar({"malhizmettoplam": "1000", "kdvmatrah(20)": "900", "hesaplanankdv(20)": "180",
                           "vergidahil": "1080", "odenecek": "1080"})
    assert (t["ara_toplam"], t["genel_toplam"]) == (900, 1080)  # iskonto düşülmüş


def test_karekoddan_fatura_bilgileri():
    from app import karekod
    kod = {"vkntckn": "20000000046", "avkntckn": "9999999999", "no": "UPT2026000000009", "tarih": "2026-02-05",
           "parabirimi": "TRY", "malhizmettoplam": "1166.67", "kdvmatrah(20)": "1166.67", "hesaplanankdv(20)": "233.33",
           "vergidahil": "1400", "odenecek": "1400"}
    o = karekod.fatura(kod, "9999999999")
    assert (o["fatura_no"], o["tarih"], o["satici"]["vkn"]) == ("UPT2026000000009", "2026-02-05", "20000000046")
    assert (o["ara_toplam"], o["kdv_toplam"], o["genel_toplam"]) == (1166.67, 233.33, 1400)
    assert o["kalemler"] == [] and o["kalemler_okunamadi"]  # kalemler uydurulmaz
    yazidan = {"fatura_no": "UPT2026000000003", "tarih": None, "satici": {"unvan": "ORNEK", "vkn": None},
               "ara_toplam": 1166.0, "kdv_toplam": 233.33, "genel_toplam": 1400, "notlar": ""}
    u = karekod.uygula(yazidan, kod, "9999999999")
    assert (u["fatura_no"], u["tarih"], u["satici"], u["ara_toplam"]) == (
        "UPT2026000000009", "2026-02-05", {"unvan": "ORNEK", "vkn": "20000000046"}, 1166.67)


def test_hareketli_banka_hesabi_hareketleriyle_birlikte_silinir():
    """Ekstresi yüklenmiş hesap: onay bekleyen hareketler de sayılır; hareketlerle=true ile yedek alınıp silinir,
    nakit çekimin kasadaki karşılığı da gider."""
    h = yeni()
    a = client.post("/api/banka-hesaplari", json={"ad": "Halkbank 94", "banka": "Halkbank"}, headers=h).json()
    e = ("IBAN No:   ;TR980001200999900010101594;;;;;\n;;;;;\nİŞLEM TARİHİ;BİM REF   ;TUTAR   ;B/A;BAKİYE   ;AÇIKLAMA\n"
         "01.10.2026;r1;1.000,00;A;1.000,00;Gelen Firma Ltd,GELEN FİRMA LTD 'DAN ödeme\n"
         "03.10.2026;r3;200,00;B;800,00;NAKİT ÇEKİLEN\n")
    client.post(f"/api/banka-hesaplari/{a['id']}/ekstre?yontem=kural", files={"dosya": ("e.csv", e.encode("cp1254"), "text/csv")}, headers=h)
    hesap = next(x for x in client.get("/api/banka-hesaplari", headers=h).json() if x["id"] == a["id"])
    assert (hesap["hareket_sayisi"], hesap["onay_bekleyen"]) == (2, 2)
    assert len(client.get(f"/api/hareketler?banka_hesap_id={a['id']}", headers=h).json()) == 2
    taslak = client.get("/api/hareketler?durum=taslak", headers=h).json()
    client.post("/api/hareketler/onayla", json={"idler": [x["id"] for x in taslak]}, headers=h)
    assert client.get("/api/kasa", headers=h).json()["bakiye"] == 200
    r = client.delete(f"/api/banka-hesaplari/{a['id']}", headers=h)
    assert r.status_code == 409 and "hareketleriyle" in r.json()["detail"]
    o = client.get(f"/api/banka-hesaplari/{a['id']}/silme-ozeti", headers=h).json()
    assert (o["hareket"], o["giris"], o["cikis"], o["cari"]) == (2, 1000, 200, 1)
    r = client.delete(f"/api/banka-hesaplari/{a['id']}?hareketlerle=true", headers=h)
    assert r.status_code == 200 and r.json()["silinen_hareket"] == 2
    assert client.get("/api/kasa", headers=h).json()["bakiye"] == 0
    assert all(x["id"] != a["id"] for x in client.get("/api/banka-hesaplari", headers=h).json())
    assert client.get("/api/yedekler", headers=h).json()


_TT = ["TTNET ANONIM SIRKETI", "Vergi Numarasi: 1234567890", "SAYIN DENEME INSAAT LIMITED SIRKETI", "Vergi No: 9999999999",
       "Fatura ID A162026000000001", "Fatura Tarih 30-09-2026 23:59", "ODENECEK TUTAR : 1.454,00 TL",
       "Tutar Indirim Toplam", "Internet 1.900,00 790,01 1.109,99", "Telefon 75,00 36,00 39,00",
       "Paket Ucreti 1.975,00 826,01 1.148,99", "Katma Degerli Servis Ucretler Toplami 229,99", "Cihaz/ Baglanti Ucreti 75,00",
       "Yuvarlama 0,01", "TOPLAM FATURA TUTARI 1.453,99", "Onceki Aydan Devir 0,24", "Gelecek Aya Devir -0,21",
       "ODENECEK TOPLAM 1.454,00", "DEVLETE ODENEN VERGILER TOPLAMI 330,73", "KDV %20 (Matrah 1.123.26 ) 224,65",
       "OIV %10 (Matrah 1.060.76 ) 106,08"]


def test_telekom_faturasi_hizmet_tablosundan_adlariyla_okunur():
    """Vergiler dahil hizmet tablosu: kalemler adları ve indirimleriyle, net tutara çevrilerek gelir; ara toplam satırı
    (Paket Ücreti) kalem olmaz; ÖİV diğer vergiye, devir/yuvarlama KDV'siz kaleme; geçici "Hizmet bedeli" satırı yok."""
    r = kural.pdf_fatura_oku(pdf_olustur(_TT), [], "DENEME INSAAT", "9999999999")
    assert [k["aciklama"] for k in r["kalemler"]] == ["Internet", "Telefon", "Katma Degerli Servis Ucretler",
                                                       "Cihaz/ Baglanti Ucreti", "Yuvarlama", "Onceki Aydan Devir", "Gelecek Aya Devir"]
    internet = r["kalemler"][0]
    assert (internet["birim_fiyat"], internet["tutar"]) == (1461.54, 853.84) and internet["iskonto_tutari"] == 607.7
    assert (r["kdv_toplam"], r["diger_vergiler"], r["genel_toplam"]) == (224.65, 106.08, 1454)
    assert abs(sum(k["tutar"] for k in r["kalemler"]) + r["yuvarlama"] + r["kdv_toplam"] + r["diger_vergiler"] - 1454) < 0.005
    assert not ai.tutarlilik_uyarilari(r)


def test_kendi_firmamiz_hicbir_yoldan_cari_olmaz():
    from app import onarim
    from app.db import SessionLocal
    from app.models import Cari, Fatura, Sirket
    h = yeni()
    client.put("/api/sirket", json={"unvan": "DENEME INSAAT LIMITED SIRKETI", "vkn": "9999999999"}, headers=h)
    assert client.post("/api/cariler", json={"unvan": "Deneme", "vkn": "9999999999"}, headers=h).status_code == 400
    # Satıcı VKN'si okunamayan faturada alıcı (biz) adıyla cari seçilmez
    pdf = pdf_olustur(["SAYIN", "DENEME INSAAT LIMITED SIRKETI", "VKN: 9999999999", "Fatura No: KND2026000000001",
                       "Fatura Tarihi: 01-09-2026", "Mal Hizmet Toplam Tutari 100,00 TL", "Hesaplanan KDV (%20) 20,00 TL",
                       "Odenecek Tutar 120,00 TL"])
    f = client.post("/api/faturalar/oku", files={"dosya": ("k.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    assert f["cari_id"] is None
    # Eski sürümde kendi adımızla açılmış cari: fatura ondan çözülür, cari silinir; elle de seçilemez
    db = SessionLocal()
    try:
        fa = db.get(Fatura, f["id"])
        kendi = Cari(sirket_id=fa.sirket_id, unvan="DENEME INSAAT LTD. STI.")
        db.add(kendi)
        db.flush()
        fa.cari_id, kendi_id = kendi.id, kendi.id
        db.commit()
        r = client.put(f"/api/faturalar/{f['id']}", json={"yon": "gelen", "cari_id": kendi_id, "kalemler": []}, headers=h)
        assert r.status_code == 400
        onarim.onar(db, db.get(Sirket, fa.sirket_id))
        assert db.get(Fatura, f["id"]).cari_id is None and db.get(Cari, kendi_id) is None
    finally:
        db.close()


def test_not_satiri_aciklamaya_faaliyet_aciklamasi_unvandan_ayrilir():
    r = kural.pdf_fatura_oku(tablolu_pdf(
        ["AHMET YILMAZ Her Türlü İnşaat Malzemeleri Satışı-", "VKN: 1234567890", "SAYIN", "TEST INSAAT LTD. STI.",
         "VKN: 9999999999", "Fatura No: AYM2026000000001", "Fatura Tarihi: 22-09-2026"],
        ["Sira No", "Mal Hizmet", "Miktar", "Birim Fiyat", "KDV Orani", "KDV Tutari", "Mal Hizmet Tutari"],
        [["1", "CIMENTO 50 KG", "2Adet", "165TL", "%20,00", "66,00TL", "330,00TL"],
         ["2", "ORNEK INSAAT FIRMASI ALDI", "1Adet", "0TL", "%20,00", "0,00TL", "0,00TL"]],
        ["Mal Hizmet Toplam Tutari 330,00TL", "Toplam Iskonto 0,00TL", "Hesaplanan KDV (%20) 66,00TL",
         "Odenecek Tutar 396,00TL", "Not: BU FATURADAKI MALZEMELERI ORNEK INSAAT FIRMASI ALDI"]), [], "TEST", "9999999999")
    assert [k["aciklama"] for k in r["kalemler"]] == ["CIMENTO 50 KG"] and r["aciklama"] == "ORNEK INSAAT FIRMASI ALDI"
    assert r["satici"]["unvan"] == "AHMET YILMAZ" and "iskonto" not in r["notlar"]


def test_hareketleri_silinen_hesapta_eski_ekstre_bakiyesi_kalmaz():
    """Hareketler silinince ekstre bakiyesi/tarihi ve ekstreden gelen açılış bakiyesi temizlenir (sahte mutabakat
    farkı çıkmaz); ekstre yeniden yüklenince bakiye ekstreyle tutar. Açılıştaki onarım hareketleri silmez."""
    from app import onarim
    from app.db import SessionLocal
    from app.models import BankaHesabi, Sirket
    h = yeni()
    a = client.post("/api/banka-hesaplari", json={"ad": "Halkbank 53", "banka": "Halkbank"}, headers=h).json()
    e = ("IBAN No:   ;TR410001200999900010101553;;;;;\n;;;;;\nİŞLEM TARİHİ;BİM REF   ;TUTAR   ;B/A;BAKİYE   ;AÇIKLAMA\n"
         "01.10.2026;r1;1.000,00;A;2.854,70;Gelen Firma Ltd,GELEN FİRMA LTD 'DAN ödeme\n"
         "02.10.2026;r2;500,00;A;3.354,70;GELEN FİRMA LTD.,GELEN FİRMA LTD. 'DAN ödeme\n"
         "09.10.2026;r3;300,00;B;3.054,70;EFT MASRAFI\n")
    yukle = lambda: client.post(f"/api/banka-hesaplari/{a['id']}/ekstre?yontem=kural",
                                files={"dosya": ("e.csv", e.encode("cp1254"), "text/csv")}, headers=h)
    hesap = lambda: next(x for x in client.get("/api/banka-hesaplari", headers=h).json() if x["id"] == a["id"])
    yukle()
    onayla = lambda: client.post("/api/hareketler/onayla", json={"idler": [x["id"] for x in client.get(
        "/api/hareketler?durum=taslak", headers=h).json()]}, headers=h)
    onayla()
    b = hesap()
    assert (b["acilis_bakiyesi"], b["ekstre_bakiye"], b["bakiye"], b["mutabakat_farki"]) == (1854.70, 3054.70, 3054.70, 0)
    # Açılıştaki onarım (cari/şantiye birleştirme dahil) hareketlere dokunmaz
    with SessionLocal() as db:
        sirket_id = db.get(BankaHesabi, a["id"]).sirket_id
        onarim.onar(db, db.get(Sirket, sirket_id))
    assert hesap()["hareket_sayisi"] == 3 and hesap()["mutabakat_farki"] == 0
    # İlk satır silinir: açılış kalan ilk satırdan yeniden hesaplanır, bakiye ekstreyle tutmaya devam eder
    ilk = client.get(f"/api/hareketler?banka_hesap_id={a['id']}", headers=h).json()[-1]
    assert client.delete(f"/api/hareketler/{ilk['id']}", headers=h).status_code == 200
    b = hesap()
    assert (b["acilis_bakiyesi"], b["bakiye"], b["mutabakat_farki"]) == (2854.70, 3054.70, 0)
    # Hepsi silinir: eski ekstre bakiyesi, tarihi ve açılış kalmaz
    idler = [x["id"] for x in client.get(f"/api/hareketler?banka_hesap_id={a['id']}", headers=h).json()]
    assert client.post("/api/hareketler/toplu-sil", json={"idler": idler}, headers=h).json()["silinen"] == 2
    b = hesap()
    assert (b["hareket_sayisi"], b["acilis_bakiyesi"], b["bakiye"], b["ekstre_bakiye"], b["ekstre_tarihi"],
            b["mutabakat_farki"]) == (0, 0, 0, None, None, None)
    # Yeniden yükleme: bakiye ekstreyle tutar
    assert yukle().json()["eklenen"] == 3
    onayla()
    b = hesap()
    assert (b["hareket_sayisi"], b["acilis_bakiyesi"], b["bakiye"], b["mutabakat_farki"]) == (3, 1854.70, 3054.70, 0)


def test_eski_surumde_hareketleri_silinmis_hesabin_ekstre_bilgisi_onarimda_temizlenir():
    from app import onarim
    from datetime import date
    from app.db import SessionLocal
    from app.models import BankaHesabi, Sirket
    h = yeni()
    a = client.post("/api/banka-hesaplari", json={"ad": "Eski Hesap", "acilis_bakiyesi": 1854.70}, headers=h).json()
    with SessionLocal() as db:
        b = db.get(BankaHesabi, a["id"])
        b.ekstre_bakiye, b.ekstre_tarihi = 502627, date(2026, 10, 9)
        db.commit()
        rapor = onarim.onar(db, db.get(Sirket, b.sirket_id))
    assert "Eski Hesap" in rapor["ekstre_bilgisi_temizlenen_hesaplar"]
    b = next(x for x in client.get("/api/banka-hesaplari", headers=h).json() if x["id"] == a["id"])
    assert (b["acilis_bakiyesi"], b["ekstre_bakiye"], b["mutabakat_farki"]) == (0, None, None)


def test_benzer_cari_uyarisi_ve_ayni_vkn():
    assert kural.benzer_unvan("Gülmez", "Gulmez İnşaat Ltd. Şti.") and kural.benzer_unvan("Yüksel Yalçın", "YUKSEL YALCIN")
    assert not kural.benzer_unvan("Ahmet Usta", "Mehmet Usta") and not kural.benzer_unvan("Akın", "Akınsoft")
    h = yeni()
    g = client.post("/api/cariler", json={"unvan": "Gulmez İnşaat Ltd. Şti.", "vkn": "1234567890"}, headers=h).json()
    b = client.get("/api/cariler/benzer?unvan=Gülmez", headers=h).json()
    assert [x["id"] for x in b] == [g["id"]] and not b[0]["ayni_vkn"]
    assert client.get("/api/cariler/benzer?unvan=Başka Firma&vkn=1234567890", headers=h).json()[0]["ayni_vkn"]
    assert client.get("/api/cariler/benzer?unvan=Gülmez&vkn=9999999999", headers=h).json() == []  # VKN'si farklı firma
    # Aynı VKN ile ikinci cari açılmaz, var olan döner; yalnız ek farkı olan ad da aynı cari sayılır
    assert client.post("/api/cariler", json={"unvan": "X", "vkn": "1234567890"}, headers=h).json()["id"] == g["id"]
    assert client.post("/api/cariler", json={"unvan": "GULMEZ İNŞAAT LTD"}, headers=h).json()["id"] == g["id"]
    assert len(client.get("/api/cariler", headers=h).json()) == 1


def test_toplu_santiye_atamasi_tum_kalemlere_gecer():
    """Toplu şantiye ataması onaylı faturada da çalışır; kalemlerdeki ayrı şantiye kalkar, hepsi yeni şantiyeye düşer."""
    h = yeni()
    eski = client.post("/api/santiyeler", json={"ad": "Eski"}, headers=h).json()
    yeni_s = client.post("/api/santiyeler", json={"ad": "Yeni"}, headers=h).json()
    c = client.post("/api/cariler", json={"unvan": "Demir AŞ"}, headers=h).json()
    f = client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": "T1", "tarih": "2026-10-05", "cari_id": c["id"], "kalemler": [
        {"aciklama": "Demir", "birim_fiyat": 1000, "kdv_orani": 20, "santiye_id": eski["id"]},
        {"aciklama": "Nakliye", "birim_fiyat": 100, "kdv_orani": 20}]}, headers=h).json()
    client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    r = client.post("/api/faturalar/toplu-guncelle", json={"idler": [f["id"]], "santiye_id": yeni_s["id"]}, headers=h).json()
    assert r["guncellenen"] == 1
    f = client.get(f"/api/faturalar/{f['id']}", headers=h).json()
    assert f["santiye_ad"] == "Yeni" and [k["santiye_id"] for k in f["kalemler"]] == [None, None]
    ozet = {s["ad"]: s["gider"] for s in client.get("/api/raporlar/santiye-ozet", headers=h).json()}
    assert ozet == {"Eski": 0, "Yeni": 1320}


def test_sema_surumlu_guncelleme_once_yedek_alir(tmp_path, monkeypatch):
    from sqlalchemy import create_engine, text
    from app import db as dbm
    monkeypatch.setattr(dbm, "SEMA_YEDEK_KLASORU", tmp_path / "yedekler")
    motor = create_engine(f"sqlite:///{tmp_path / 'eski.db'}")
    with motor.begin() as b:
        b.execute(text("CREATE TABLE sirketler (id INTEGER PRIMARY KEY, unvan TEXT)"))
        b.execute(text("INSERT INTO sirketler (unvan) VALUES ('ARMA')"))
    dbm.db_guncelle(motor)
    assert dbm.sema_surumu(motor) == dbm.SEMA_ADIMLARI[-1][0]
    assert len(list((tmp_path / "yedekler").glob("sema-*-oncesi-*.db"))) == 2  # son 2 kopya kalır
    dbm.db_guncelle(motor)  # ikinci açılışta adım tekrar çalışmaz, yedek alınmaz
    assert len(list((tmp_path / "yedekler").glob("*.db"))) == min(2, len(dbm.SEMA_ADIMLARI))  # son 2 kopya kalır

    # Yarıda kalan adım geri alınır, sürüm ilerlemez; yedekte eski veri durur
    def bozuk(baglanti, _motor):
        baglanti.execute(text("UPDATE sirketler SET unvan = 'BOZUK'"))
        raise ValueError("adım hatası")
    monkeypatch.setattr(dbm, "SEMA_ADIMLARI", [*dbm.SEMA_ADIMLARI, (99, "Deneme", bozuk)])
    import pytest
    with pytest.raises(ValueError):
        dbm.db_guncelle(motor)
    with motor.connect() as b:
        assert b.execute(text("SELECT unvan FROM sirketler")).scalar() == "ARMA"
    assert dbm.sema_surumu(motor) == dbm.SEMA_ADIMLARI[-2][0]  # bozuk adımdan önceki son sürüm
    import sqlite3
    yedek = next((tmp_path / "yedekler").glob("sema-99-oncesi-*.db"))
    assert sqlite3.connect(yedek).execute("SELECT unvan FROM sirketler").fetchone()[0] == "ARMA"


def test_pdf_yonu_ust_vkn_bizim_karsi_firma_unvanli_alis():
    """Bizim VKN'miz adımız yazılmadan en üstte (müşteri bilgisi), karşı firma ünvanı ve VKN'siyle: alış faturası."""
    v = pdf_olustur(["Musteri No: 12345", "Vergi No: 9999999999", "ORNEKNET ANONIM SIRKETI", "Vergi Numarasi: 1234567890",
                     "Fatura No: ORN2026000000001", "Fatura Tarihi: 31-08-2026", "ODENECEK TUTAR : 120,00 TL"])
    assert kural.pdf_yonu(v, "9999999999") == "gelen"
    # Karekod varsa karar ondan: satıcı VKN'si bizimse satış
    assert kural.pdf_yonu(v, "9999999999", lambda: "giden") == "giden"


def test_bosluklu_ve_bitisik_vkn_okunur():
    r = _oku(["ORNEK TICARET A.S.", "VKN: 123 456 7890", "SAYIN", "DENEME INSAAT LTD STI", "VKN: 9999999999TICARETSICILNO 55",
              "Fatura No: ORN2026000000001", "Fatura Tarihi: 31-08-2026", "Mal Hizmet Toplam Tutari 100,00",
              "Hesaplanan KDV (%20) 20,00", "Odenecek Tutar 120,00"])
    assert r["satici"]["vkn"] == "1234567890"


def test_alt_alta_kdv_etiketleri_tutarlariyla_eslesir():
    metin = kural._alt_alta_etiketler("Hesaplanan KDV %1\nHesaplanan KDV %10\n0,20 TRY\n80,00 TRY\nOdenecek Tutar 900,20")
    assert metin.splitlines()[:2] == ["Hesaplanan KDV %1 0,20 TRY", "Hesaplanan KDV %10 80,00 TRY"]
