import io
import os
import tempfile

os.environ["MUHASEBE_VERI"] = tempfile.mkdtemp()

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import ai
from app.main import app

client = TestClient(app)

OKUNAN = {
    "belge_fatura_mi": True, "fatura_no": "ABC2026000000123", "tarih": "2026-10-01", "vade": "2026-11-01",
    "satici": {"unvan": "Yılmaz Beton A.Ş.", "vkn": "1234567890", "vergi_dairesi": "Kadıköy",
               "adres": "İstanbul", "telefon": None, "iban": None},
    "alici": {"unvan": "Test İnşaat", "vkn": "9999999999"},
    "kalemler": [
        {"aciklama": "C30 Hazır Beton", "miktar": 10, "birim": "m3", "birim_fiyat": 2500, "kdv_orani": 20, "tutar": 25000},
        {"aciklama": "Pompa", "miktar": 1, "birim": "adet", "birim_fiyat": 3000, "kdv_orani": 20, "tutar": 3000},
    ],
    "ara_toplam": 28000, "kdv_toplam": 5600, "tevkifat": 0, "genel_toplam": 33600, "para_birimi": "TRY",
    "onerilen_santiye": "Kartal Konut", "onerilen_gider_turu": "Beton", "notlar": "",
}


def yeni_kullanici(eposta):
    r = client.post("/api/kayit", json={"sirket_unvan": "Test İnşaat", "ad": "Ali", "eposta": eposta, "sifre": "gizli123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


@pytest.fixture
def h():
    import uuid
    return yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")


def test_giris_gerekli():
    assert client.get("/api/cariler").status_code == 401


def test_giris(h):
    yeni_kullanici("giris@ornek.com")
    assert client.post("/api/giris", json={"eposta": "giris@ornek.com", "sifre": "yanlis"}).status_code == 401
    assert client.post("/api/giris", json={"eposta": "GIRIS@ornek.com", "sifre": "gizli123"}).status_code == 200


def test_ai_fatura_akisi(h, monkeypatch):
    monkeypatch.setattr(ai, "fatura_oku", lambda *a, **kw: OKUNAN)
    s = client.post("/api/santiyeler", json={"ad": "Kartal Konut", "butce": 100000}, headers=h).json()

    r = client.post("/api/faturalar/oku?yontem=ai", files={"dosya": ("f.pdf", b"%PDF-1.4", "application/pdf")}, headers=h)
    assert r.status_code == 200, r.text
    f = r.json()[0]
    assert f["durum"] == "taslak"
    assert f["santiye_id"] == s["id"]
    assert f["gider_turu"] == "Beton"
    assert f["genel_toplam"] == 33600
    assert f["cari_unvan"] == "Yılmaz Beton A.Ş."  # cari okurken açılır
    assert any("yeni cari açıldı" in u for u in f["ai_okuma"]["bilgi"]) and f["ai_okuma"]["uyarilar"] == []

    r = client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["cari_unvan"] == "Yılmaz Beton A.Ş."

    # Aynı fatura ikinci kez yüklenince yeni taslak açılmaz, kayıtlı fatura "zaten var" olarak döner.
    f2 = client.post("/api/faturalar/oku?yontem=ai", files={"dosya": ("f.pdf", b"%PDF-1.4", "application/pdf")}, headers=h).json()[0]
    assert f2["id"] == f["id"] and f2["zaten_var"] is True
    # Aynı dosya yeniden okunmaz (hızlı atlanır): okuyucu hiç çağrılmaz
    monkeypatch.setattr(ai, "fatura_oku", lambda *a, **kw: (_ for _ in ()).throw(AssertionError("yeniden okundu")))
    f2 = client.post("/api/faturalar/oku?yontem=ai", files={"dosya": ("g.pdf", b"%PDF-1.4", "application/pdf")}, headers=h).json()[0]
    assert f2["id"] == f["id"] and f2["zaten_var"] is True
    assert client.get("/api/faturalar?durum=taslak", headers=h).json() == []
    # Elle aynı numarayla girilen fatura onaylanamaz
    v = {k: f[k] for k in ("yon", "fatura_no", "tarih", "cari_id")} | {"kalemler": [{"aciklama": "x", "birim_fiyat": 1}]}
    f3 = client.post("/api/faturalar", json=v, headers=h).json()
    assert client.post(f"/api/faturalar/{f3['id']}/onayla", headers=h).status_code == 409
    client.delete(f"/api/faturalar/{f3['id']}", headers=h)

    ozet = client.get("/api/raporlar/santiye-ozet", headers=h).json()[0]
    assert ozet["gider"] == 33600 and ozet["butce_kalan"] == 66400 and ozet["gider_turleri"] == {"Beton": 33600}

    bakiye = client.get("/api/raporlar/cari-bakiye", headers=h).json()
    assert bakiye[0]["bakiye"] == -33600

    ekstre = client.get(f"/api/raporlar/cari-ekstre/{r.json()['cari_id']}", headers=h).json()
    assert ekstre["bakiye"] == -33600 and len(ekstre["satirlar"]) == 1

    assert client.get(f"/api/faturalar/{f['id']}/belge", headers=h).content == b"%PDF-1.4"


def test_elle_fatura_ve_toplamlar(h):
    c = client.post("/api/cariler", json={"unvan": "Demir Ltd", "vkn": "1111111111"}, headers=h).json()
    r = client.post("/api/faturalar", json={
        "yon": "gelen", "fatura_no": "D1", "tarih": "2026-10-05", "cari_id": c["id"], "tevkifat": 100,
        "kalemler": [{"aciklama": "Nervürlü demir", "miktar": 2, "birim_fiyat": 1000, "kdv_orani": 20}],
    }, headers=h).json()
    assert (r["ara_toplam"], r["kdv_toplam"], r["genel_toplam"]) == (2000, 400, 2300)


def test_sirketler_birbirini_gormez(h):
    client.post("/api/cariler", json={"unvan": "Gizli Cari"}, headers=h)
    baska = yeni_kullanici("baska@ornek.com")
    assert client.get("/api/cariler", headers=baska).json() == []


def test_excel(h):
    client.post("/api/cariler", json={"unvan": "Mevcut", "vkn": "2222222222"}, headers=h)
    sablon = client.get("/api/excel/cari-sablon", headers=h).content
    wb = load_workbook(io.BytesIO(sablon))
    ws = wb.active
    ws.append(["Yeni Cari", "musteri", "3333333333"])
    ws.append(["Tekrar", "musteri", "2222222222"])
    tampon = io.BytesIO()
    wb.save(tampon)
    r = client.post("/api/excel/cariler/yukle", files={"dosya": ("c.xlsx", tampon.getvalue())}, headers=h).json()
    assert r == {"eklenen": 1, "atlanan": 1}

    ws = load_workbook(io.BytesIO(client.get("/api/excel/cariler", headers=h).content)).active
    assert [row[0] for row in ws.iter_rows(min_row=2, values_only=True)] == ["Mevcut", "Yeni Cari"]
    for yol in ["faturalar", "fatura-kalemleri", "santiye-ozet", "cari-bakiye"]:
        assert client.get(f"/api/excel/{yol}", headers=h).status_code == 200


def test_tutarlilik_uyarilari():
    bozuk = {**OKUNAN, "ara_toplam": 27000}
    assert len(ai.tutarlilik_uyarilari(bozuk)) == 2
    assert ai.tutarlilik_uyarilari(OKUNAN) == []


def test_yedek(h):
    import zipfile
    from app import yedek
    r = client.post("/api/yedekler", headers=h)
    assert r.status_code == 200, r.text
    hedef = r.json()["hedefler"][0]
    assert hedef["ok"]
    with zipfile.ZipFile(f"{hedef['klasor']}/{r.json()['dosya']}") as z:
        assert "muhasebe.db" in z.namelist()
    assert client.get("/api/yedekler", headers=h).json()


def test_ekran_dosyalari_onbellekte_kalmaz():
    """Yeni paket kurulunca tarayıcı eski app.js'i çalıştırmasın."""
    r = client.get("/")
    assert r.headers["cache-control"].startswith("no-cache") and "/web/app.js?v=" in r.text
    assert client.get("/web/app.js").headers["cache-control"].startswith("no-cache")


def test_okuma_ozeti_sorunlu_faturalari_listeler(h, monkeypatch):
    monkeypatch.setattr(ai, "fatura_oku", lambda *a, **kw: {**OKUNAN, "fatura_no": "ABC2026000000999", "kalemler": []})
    f = client.post("/api/faturalar/oku?yontem=ai", files={"dosya": ("s.pdf", b"%PDF-1.4 s", "application/pdf")}, headers=h).json()[0]
    o = client.get("/api/faturalar/okuma-raporu/ozet", headers=h).json()
    assert any(x["kategori"] == "kalem_okunamadi" for x in o["sorunlar"][str(f["id"])])


def test_bilgi_notlari_sorun_sayilmaz():
    from types import SimpleNamespace

    from app.tani import bilgi_notu_mu
    kurlu, kursuz = SimpleNamespace(kur=45.1), SimpleNamespace(kur=None)
    assert bilgi_notu_mu("Fatura USD cinsinden (kur 45,1 TL).", kurlu) and not bilgi_notu_mu("Fatura USD cinsinden.", kursuz)
    assert bilgi_notu_mu("Faturada 0,01 TL yuvarlama farkı var; kayıt faturadaki toplamla yapıldı.", kursuz)
    assert not bilgi_notu_mu("Faturada 3,61 TL yuvarlama farkı var; kayıt faturadaki toplamla yapıldı.", kursuz)
    assert bilgi_notu_mu("Ödenecek tutar (9.540 TL) fatura tutarından (9.819,6 TL) farklı; fatura, fatura tutarıyla kaydedildi.", kursuz)
