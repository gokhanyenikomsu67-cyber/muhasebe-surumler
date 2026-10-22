"""Malzeme kartları: biraz farklı yazılmış aynı malzeme sorulup tek kartta toplanır; Malzeme Bazlı Alışlar raporu
(sentetik veri)."""
import io
import uuid

from openpyxl import load_workbook

from app import malzeme
from tests.test_api import client, yeni_kullanici


def _fatura(h, cari_id, kalemler, yon="gelen", tarih="2026-09-10", santiye_id=None, tipi=None):
    veri = {"yon": yon, "fatura_no": f"DNM{uuid.uuid4().hex[:6]}", "tarih": tarih, "cari_id": cari_id,
            "santiye_id": santiye_id, "kalemler": [{"kdv_orani": 20, **k} for k in kalemler]}
    if tipi:
        veri["fatura_tipi"] = tipi
    f = client.post("/api/faturalar", json=veri, headers=h).json()
    assert client.post(f"/api/faturalar/{f['id']}/onayla", headers=h).status_code == 200
    return f


def test_anahtar_ve_benzerlik():
    a = malzeme.anahtar
    assert malzeme.ayni_mi(a("ÇİMENTO 50KG"), a("Cimento 50 kg."))
    assert malzeme.ayni_mi(a("NYY 3 X 2,5 mm²"), a("nyy 3x2.5 mm2"))
    assert malzeme.benzerlik(a("NYY 3x2,5 KABLO"), a("NYY 3x4 KABLO")) == 0  # ölçü farklı = ayrı malzeme
    assert malzeme.benzerlik(a("İNCE KUM"), a("KUM İNCE")) > 0.9
    assert malzeme.benzerlik(a("PORTLAND ÇİMENTOSU 50 KG"), a("PORTLAND CIMENTO 50 KG")) >= malzeme.BENZERLIK
    assert malzeme.benzerlik(a("Çimento"), a("Kireç")) < malzeme.BENZERLIK


def test_benzer_adlar_sorulur_karar_hatirlanir_rapor_toplar():
    h = yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")
    s = client.post("/api/santiyeler", json={"ad": "Deneme Şantiye"}, headers=h).json()
    c1 = client.post("/api/cariler", json={"unvan": "Örnek Yapı Market"}, headers=h).json()
    c2 = client.post("/api/cariler", json={"unvan": "Deneme Hırdavat"}, headers=h).json()
    _fatura(h, c1["id"], [{"aciklama": "PORTLAND ÇİMENTOSU 50 KG", "miktar": 10, "birim": "adet", "birim_fiyat": 100},
                          {"aciklama": "NYY 3x2,5 KABLO", "miktar": 100, "birim": "m", "birim_fiyat": 20}], santiye_id=s["id"])
    _fatura(h, c1["id"], [{"aciklama": "Portland Çimentosu 50kg", "miktar": 5, "birim": "Adet", "birim_fiyat": 110}])
    _fatura(h, c2["id"], [{"aciklama": "PORTLAND CIMENTO 50 KG", "miktar": 20, "birim": "adet", "birim_fiyat": 95},
                          {"aciklama": "NYY 3x4 KABLO", "miktar": 50, "birim": "m", "birim_fiyat": 30}])
    # Yazılışı aynı olan kendiliğinden aynı kart; benzer olan sorulur; ölçüsü farklı kablo ayrı kart
    bekleyen = client.get("/api/malzemeler/bekleyen", headers=h).json()
    assert [(b["ad"], b["oneriler"][0]["ad"]) for b in bekleyen] == [("PORTLAND CIMENTO 50 KG", "PORTLAND ÇİMENTOSU 50 KG")]
    kartlar = {m["ad"] for m in client.get("/api/malzemeler", headers=h).json()}
    assert kartlar == {"PORTLAND ÇİMENTOSU 50 KG", "NYY 3x2,5 KABLO", "NYY 3x4 KABLO"}
    rapor = client.get("/api/raporlar/malzeme", headers=h).json()["ozet"]
    assert [x["eslesme_bekliyor"] for x in rapor if x["malzeme"] == "PORTLAND CIMENTO 50 KG"] == [True]

    # Toplu onay: benzeri o karta işlenir, karar hatırlanır (bir daha sorulmaz)
    hedef = bekleyen[0]["oneriler"][0]["id"]
    assert client.post("/api/malzemeler/esle", json={"kararlar": [{"ad": "PORTLAND CIMENTO 50 KG", "malzeme_id": hedef}]},
                       headers=h).json() == {"eslenen": 1, "ayri": 0}
    _fatura(h, c2["id"], [{"aciklama": "portland cimento 50 kg.", "miktar": 1, "birim": "adet", "birim_fiyat": 100}])
    assert client.get("/api/malzemeler/bekleyen", headers=h).json() == []

    ozet = {x["malzeme"]: x for x in client.get("/api/raporlar/malzeme", headers=h).json()["ozet"]}
    c = ozet["PORTLAND ÇİMENTOSU 50 KG"]
    assert (c["miktar"], c["birim"], c["tutar"], c["fatura_sayisi"], c["cari_sayisi"]) == (36, "adet", 3550, 4, 2)
    assert c["ortalama_fiyat"] == round(3550 / 36, 4) and len(c["yazilislar"]) == 4
    # Firma ve şantiye süzgeci
    yalniz = client.get(f"/api/raporlar/malzeme?cari_id={c1['id']}", headers=h).json()["ozet"]
    assert {x["malzeme"]: x["miktar"] for x in yalniz} == {"PORTLAND ÇİMENTOSU 50 KG": 15, "NYY 3x2,5 KABLO": 100}
    santiyede = client.get(f"/api/raporlar/malzeme?santiye_id={s['id']}", headers=h).json()["ozet"]
    assert {x["malzeme"] for x in santiyede} == {"PORTLAND ÇİMENTOSU 50 KG", "NYY 3x2,5 KABLO"}
    # Ayrıntı: kalem kendi adını korur
    ayr = client.get(f"/api/raporlar/malzeme?malzeme_id={hedef}", headers=h).json()["ayrinti"]
    assert "PORTLAND CIMENTO 50 KG" in [x["kalem"] for x in ayr]
    # Excel: toplamlar + kalemler
    wb = load_workbook(io.BytesIO(client.get("/api/excel/malzeme", headers=h).content))
    assert wb.sheetnames == ["Malzeme Toplamları", "Kalemler"]
    assert wb["Kalemler"].max_row == 7


def test_ayri_malzeme_karari_ad_degistirme_ve_birlestirme():
    h = yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")
    c = client.post("/api/cariler", json={"unvan": "Örnek Yapı Market"}, headers=h).json()
    _fatura(h, c["id"], [{"aciklama": "İNCE KUM", "miktar": 2, "birim": "ton", "birim_fiyat": 500},
                         {"aciklama": "Demir Ø12", "miktar": 1, "birim": "ton", "birim_fiyat": 20000}])
    _fatura(h, c["id"], [{"aciklama": "KUM İNCE", "miktar": 3, "birim": "ton", "birim_fiyat": 500},
                         {"aciklama": "NERVÜRLÜ DEMİR 12", "miktar": 2, "birim": "ton", "birim_fiyat": 21000}])
    b = client.get("/api/malzemeler/bekleyen", headers=h).json()
    assert [x["ad"] for x in b] == ["KUM İNCE"]
    # "Ayrı malzeme" kararı: ayrı kart açılır, bir daha sorulmaz
    client.post("/api/malzemeler/esle", json={"kararlar": [{"ad": "KUM İNCE", "malzeme_id": None}]}, headers=h)
    assert client.get("/api/malzemeler/bekleyen", headers=h).json() == []
    kartlar = {m["ad"]: m["id"] for m in client.get("/api/malzemeler", headers=h).json()}
    # Kartları sonradan birleştirme ve ad değiştirme (eski ad kartta kalır)
    client.post(f"/api/malzemeler/{kartlar['İNCE KUM']}/birlestir", json={"idler": [kartlar["KUM İNCE"]]}, headers=h)
    client.post(f"/api/malzemeler/{kartlar['Demir Ø12']}/birlestir", json={"idler": [kartlar["NERVÜRLÜ DEMİR 12"]]}, headers=h)
    m = client.put(f"/api/malzemeler/{kartlar['Demir Ø12']}", json={"ad": "Nervürlü Demir Ø12"}, headers=h).json()
    assert m["eski_adlar"] == ["Demir Ø12"]  # "NERVÜRLÜ DEMİR 12" yeni adla aynı yazılış
    ozet = {x["malzeme"]: x["miktar"] for x in client.get("/api/raporlar/malzeme", headers=h).json()["ozet"]}
    assert ozet == {"İNCE KUM": 5, "Nervürlü Demir Ø12": 3}
    # Başka kartın adı verilemez
    assert client.put(f"/api/malzemeler/{kartlar['İNCE KUM']}", json={"ad": "demir ø12"}, headers=h).status_code == 409
    # Yanlış eşleşen yazılış karttan çıkarılınca yeniden sorulur
    client.post(f"/api/malzemeler/{kartlar['İNCE KUM']}/ad-cikar", json={"ad": "KUM İNCE"}, headers=h)
    assert [x["ad"] for x in client.get("/api/malzemeler/bekleyen", headers=h).json()] == ["KUM İNCE"]


def test_satis_tarafi_ve_iade():
    h = yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")
    c = client.post("/api/cariler", json={"unvan": "Deneme Müşteri"}, headers=h).json()
    _fatura(h, c["id"], [{"aciklama": "Beton C30", "miktar": 10, "birim": "m3", "birim_fiyat": 2000}], yon="giden")
    _fatura(h, c["id"], [{"aciklama": "Beton C30", "miktar": 4, "birim": "m3", "birim_fiyat": 1500}])
    _fatura(h, c["id"], [{"aciklama": "Beton C30", "miktar": 1, "birim": "m³", "birim_fiyat": 1500}], yon="giden", tipi="IADE")
    alis = client.get("/api/raporlar/malzeme", headers=h).json()["ozet"]
    assert [(x["miktar"], x["birim"], x["tutar"]) for x in alis] == [(3, "m³", 4500)]  # tedarikçiye iade düşer
    satis = client.get("/api/raporlar/malzeme?yon=giden", headers=h).json()["ozet"]
    assert [(x["miktar"], x["tutar"]) for x in satis] == [(10, 20000)]
