import io
import uuid

from openpyxl import Workbook

from app import ai
from tests.test_api import client, yeni_kullanici


def kur():
    h = yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")
    musteri = client.post("/api/cariler", json={"unvan": "Ahmet Yapı", "tur": "musteri", "vkn": "5555555555"}, headers=h).json()
    tedarikci = client.post("/api/cariler", json={"unvan": "Demir Ltd", "vkn": "6666666666", "iban": "TR000000000000000000000001"}, headers=h).json()
    banka = client.post("/api/banka-hesaplari", json={"ad": "Ziraat TL", "acilis_bakiyesi": 1000}, headers=h).json()
    return h, musteri, tedarikci, banka


def bakiye(h, cari_id):
    return next(c["bakiye"] for c in client.get("/api/raporlar/cari-bakiye", headers=h).json() if c["id"] == cari_id)


def banka_bakiye(h, hesap_id):
    return next(b["bakiye"] for b in client.get("/api/banka-hesaplari", headers=h).json() if b["id"] == hesap_id)


def test_alinan_cek_akisi():
    h, musteri, tedarikci, banka = kur()
    # Müşteriye 10.000 satış faturası
    client.post("/api/faturalar", json={"yon": "giden", "fatura_no": "S1", "cari_id": musteri["id"],
                                        "kalemler": [{"aciklama": "Daire", "birim_fiyat": 10000, "kdv_orani": 0}]}, headers=h)
    f = client.get("/api/faturalar?yon=giden", headers=h).json()[0]
    client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    assert bakiye(h, musteri["id"]) == 10000

    c = client.post("/api/cekler", json={"yon": "alinan", "tutar": 4000, "vade": "2026-12-01", "cari_id": musteri["id"],
                                         "cek_no": "123", "banka": "Garanti"}, headers=h).json()
    assert c["durum"] == "portfoyde" and c["durum_adi"] == "Portföyde"
    assert bakiye(h, musteri["id"]) == 6000

    # Portföydeki çek ödenemez, sadece alınan işlemleri geçerli
    assert client.post(f"/api/cekler/{c['id']}/islem", json={"islem": "ode", "tarih": "2026-12-01"}, headers=h).status_code == 400
    r = client.post(f"/api/cekler/{c['id']}/islem", json={"islem": "tahsile_ver", "tarih": "2026-11-25", "banka_hesap_id": banka["id"]}, headers=h)
    assert r.json()["durum"] == "tahsilde"
    r = client.post(f"/api/cekler/{c['id']}/islem", json={"islem": "tahsil", "tarih": "2026-12-01"}, headers=h)
    assert r.json()["durum"] == "tahsil_edildi"
    assert banka_bakiye(h, banka["id"]) == 5000
    assert bakiye(h, musteri["id"]) == 6000
    # Tekrar tahsil edilemez
    assert client.post(f"/api/cekler/{c['id']}/islem", json={"islem": "tahsil", "tarih": "2026-12-01"}, headers=h).status_code == 400


def test_ciro_ve_karsiliksiz():
    h, musteri, tedarikci, banka = kur()
    c = client.post("/api/cekler", json={"yon": "alinan", "tutar": 3000, "vade": "2026-12-01", "cari_id": musteri["id"]}, headers=h).json()
    client.post(f"/api/cekler/{c['id']}/islem", json={"islem": "ciro", "tarih": "2026-11-01", "cari_id": tedarikci["id"]}, headers=h)
    assert bakiye(h, musteri["id"]) == -3000   # müşteriden peşin çek aldık
    assert bakiye(h, tedarikci["id"]) == 3000  # tedarikçiye çek verdik
    r = client.post(f"/api/cekler/{c['id']}/islem", json={"islem": "karsiliksiz", "tarih": "2026-12-02"}, headers=h).json()
    assert r["durum"] == "karsiliksiz"
    assert bakiye(h, musteri["id"]) == 0
    assert bakiye(h, tedarikci["id"]) == 0
    ekstre = client.get(f"/api/raporlar/cari-ekstre/{musteri['id']}", headers=h).json()
    assert len(ekstre["satirlar"]) == 2 and ekstre["bakiye"] == 0


def test_verilen_cek_ve_silme():
    h, musteri, tedarikci, banka = kur()
    c = client.post("/api/cekler", json={"yon": "verilen", "tutar": 2500, "vade": "2026-12-15", "cari_id": tedarikci["id"]}, headers=h).json()
    assert c["durum"] == "verildi" and bakiye(h, tedarikci["id"]) == 2500
    client.post(f"/api/cekler/{c['id']}/islem", json={"islem": "ode", "tarih": "2026-12-15", "banka_hesap_id": banka["id"]}, headers=h)
    assert banka_bakiye(h, banka["id"]) == -1500
    ozet = client.get("/api/raporlar/ozet", headers=h).json()
    assert ozet["odenecek_cek"] == 0 and ozet["banka_toplam"] == -1500
    client.delete(f"/api/cekler/{c['id']}", headers=h)
    assert bakiye(h, tedarikci["id"]) == 0 and banka_bakiye(h, banka["id"]) == 1000


def test_kaynak_dosya_dogrulama():
    h, musteri, *_ = kur()
    r = client.post("/api/cekler", json={"yon": "alinan", "tutar": 1, "vade": "2026-12-01", "kaynak_dosya": "../../etc/passwd"}, headers=h)
    assert r.status_code == 400


def test_ekstre_okuma(monkeypatch):
    h, musteri, tedarikci, banka = kur()
    alinan = {}

    def sahte(dosya, mime, cariler, gider_turleri):
        alinan["mime"], alinan["cariler"] = mime, list(cariler)
        alinan["metin"] = ai.tablo_metni(dosya, mime)
        return {"notlar": "", "satirlar": [
            {"tarih": "2026-10-01", "aciklama": "EFT AHMET YAPI", "tutar": 5000, "bakiye": 6000,
             "onerilen_cari_id": musteri["id"], "onerilen_tur": "tahsilat", "onerilen_gider_turu": None},
            {"tarih": "2026-10-02", "aciklama": "HAVALE DEMIR LTD", "tutar": -2000, "bakiye": 4000,
             "onerilen_cari_id": tedarikci["id"], "onerilen_tur": "odeme", "onerilen_gider_turu": None},
            {"tarih": "2026-10-02", "aciklama": "EFT MASRAFI", "tutar": -12.5, "bakiye": 3987.5,
             "onerilen_cari_id": None, "onerilen_tur": "banka_masrafi", "onerilen_gider_turu": None},
            {"tarih": "2026-10-03", "aciklama": "Yabancı", "tutar": -100, "bakiye": None,
             "onerilen_cari_id": 999999, "onerilen_tur": "tahsilat", "onerilen_gider_turu": "Uydurma"},
        ]}
    monkeypatch.setattr(ai, "ekstre_oku", sahte)

    wb = Workbook()
    wb.active.append(["Tarih", "Açıklama", "Tutar"])
    wb.active.append(["01.10.2026", "EFT AHMET YAPI", 5000])
    tampon = io.BytesIO()
    wb.save(tampon)
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre?yontem=ai", files={"dosya": ("ekstre.xlsx", tampon.getvalue(), "application/octet-stream")}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["eklenen"] == 4
    assert "EFT AHMET YAPI" in alinan["metin"] and len(alinan["cariler"]) == 2

    taslaklar = client.get("/api/hareketler?durum=taslak", headers=h).json()
    assert len(taslaklar) == 4
    yabanci = next(t for t in taslaklar if t["aciklama"] == "Yabancı")
    assert yabanci["cari_id"] is None and yabanci["tur"] == "odeme" and yabanci["gider_turu"] is None
    # Taslaklar bakiyeyi etkilemez
    assert banka_bakiye(h, banka["id"]) == 1000

    # Aynı ekstre tekrar yüklenince satırlar atlanır
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre?yontem=ai", files={"dosya": ("ekstre.xlsx", tampon.getvalue())}, headers=h)
    j = r.json()
    assert (j["eklenen"], j["atlanan"], j["yeni_cari"], j["islenemeyen"]) == (0, 4, 0, [])

    client.post("/api/hareketler/onayla", json={"idler": [t["id"] for t in taslaklar]}, headers=h)
    assert banka_bakiye(h, banka["id"]) == round(1000 + 5000 - 2000 - 12.5 - 100, 2)
    assert bakiye(h, musteri["id"]) == -5000
    assert bakiye(h, tedarikci["id"]) == 2000
    # Banka kesintisi MUHTELİF BANKA İŞLEMLERİ carisine gider ve ödeme olarak işlenir (bakiye sıfır)
    masraf = next(t for t in taslaklar if t["aciklama"] == "EFT MASRAFI")
    assert masraf["cari_unvan"] == "MUHTELİF BANKA İŞLEMLERİ" and bakiye(h, masraf["cari_id"]) == 0


def test_eski_xls_reddedilir():
    h, _, _, banka = kur()
    r = client.post(f"/api/banka-hesaplari/{banka['id']}/ekstre", files={"dosya": ("e.xls", b"x")}, headers=h)
    assert r.status_code == 422 and "xlsx" in r.json()["detail"]


def test_masraf_santiyeye_yansir():
    h, *_ , banka = kur()
    s = client.post("/api/santiyeler", json={"ad": "Pendik"}, headers=h).json()
    client.post("/api/hareketler", json={"tarih": "2026-10-05", "yon": "cikis", "tur": "masraf", "tutar": 750,
                                         "banka_hesap_id": banka["id"], "santiye_id": s["id"], "gider_turu": "Akaryakıt"}, headers=h)
    ozet = client.get("/api/raporlar/santiye-ozet", headers=h).json()[0]
    assert ozet["gider"] == 750 and ozet["gider_turleri"] == {"Akaryakıt": 750}
    for yol in ["cekler", "hareketler", "cari-bakiye"]:
        assert client.get(f"/api/excel/{yol}", headers=h).status_code == 200


def test_onayli_faturalar_toplu_silinir_bakiyeden_duser():
    h, musteri, tedarikci, banka = kur()
    for no, tutar in [("A1", 1000), ("A2", 2000), ("A3", 500)]:
        client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": no, "cari_id": tedarikci["id"],
                                            "kalemler": [{"aciklama": "Demir", "birim_fiyat": tutar, "kdv_orani": 20,
                                                          "tevkifat_orani": 2}]}, headers=h)
    faturalar = client.get("/api/faturalar?yon=gelen", headers=h).json()
    for f in faturalar:
        client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    once = bakiye(h, tedarikci["id"])
    sil = [f["id"] for f in faturalar if f["fatura_no"] in ("A1", "A2")]
    # Ekranda ayrıca onay alınmadan onaylı fatura silinmez
    r = client.post("/api/faturalar/toplu-sil", json={"idler": sil}, headers=h).json()
    assert (r["silinen"], r["atlanan"]) == (0, 2)
    r = client.post("/api/faturalar/toplu-sil", json={"idler": sil, "onaylilar_dahil": True}, headers=h).json()
    assert r["silinen"] == 2
    kalan = client.get("/api/faturalar?yon=gelen", headers=h).json()
    assert [f["fatura_no"] for f in kalan] == ["A3"]
    assert bakiye(h, tedarikci["id"]) == round(once * kalan[0]["genel_toplam"] / sum(f["genel_toplam"] for f in faturalar), 2)


def test_bagli_kaydi_olan_cari_ve_santiye_silinmez_nedeni_yazilir():
    h = yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")
    dolu = client.post("/api/cariler", json={"unvan": "Dolu Ltd"}, headers=h).json()
    bos = client.post("/api/cariler", json={"unvan": "Bos Ltd"}, headers=h).json()
    st = client.post("/api/santiyeler", json={"ad": "Ofis"}, headers=h).json()
    st2 = client.post("/api/santiyeler", json={"ad": "Depo"}, headers=h).json()
    for no in ("D1", "D2"):
        client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": no, "cari_id": dolu["id"],
                                            "kalemler": [{"aciklama": "Kira", "birim_fiyat": 100, "kdv_orani": 20, "santiye_id": st["id"]}]},
                    headers=h)
    r = client.delete(f"/api/cariler/{dolu['id']}", headers=h)
    assert r.status_code == 409 and r.json()["detail"] == "2 faturası var, silinemez."
    r = client.delete(f"/api/santiyeler/{st['id']}", headers=h)
    assert r.status_code == 409 and r.json()["detail"] == "2 faturası var, silinemez."
    assert client.delete(f"/api/cariler/{bos['id']}", headers=h).status_code == 200
    assert client.delete(f"/api/santiyeler/{st2['id']}", headers=h).status_code == 200


def test_cek_metninden_alanlar_ve_yaziyla_tutar():
    from app import cek_okuma
    assert cek_okuma.yaziyla_tutar("Yalnız: #ONİKİBİNBEŞYÜZ TÜRK LİRASI ELLİ KURUŞ#") == 12500.5
    assert cek_okuma.yaziyla_tutar("YÜZELLİBİN ÜÇYÜZ YİRMİBEŞ TL") == 150325
    metin = ("T.C. ZİRAAT BANKASI A.Ş. Seri No: B 7654321\nTAVŞANLI ŞUBESİ Hesap No: 98765432\n"
             "Keşide Yeri: ANKARA Keşide Tarihi: 20.01.2027\nYalnız: DÖRTBİNİKİYÜZELLİ TÜRK LİRASI ELLİ KURUŞ #4.250,50#\n"
             "Keşideci: AHMET YILMAZ\nVKN: 12345678901")
    c = cek_okuma.cek_metni_coz(metin)
    assert (c["banka"], c["sube"], c["cek_no"], c["hesap_no"], c["vade"], c["tutar"], c["kesideci"], c["vkn"], c["belirsiz"]) == \
        ("Ziraat Bankası", "TAVŞANLI", "B7654321", "98765432", "2027-01-20", 4250.5, "AHMET YILMAZ", "12345678901", [])
    # Rakamla ve yazıyla tutar tutmazsa tutar kırmızı
    assert "tutar" in cek_okuma.cek_metni_coz(metin.replace("#4.250,50#", "#4.520,50#"))["belirsiz"]
    assert cek_okuma.banka_bul("ZİRAAT KATILIM BANKASI A.Ş.") == "Ziraat Katılım"


def test_ayni_cek_ikinci_kez_girilmez_ve_toplu_giris():
    h, musteri, tedarikci, banka = kur()
    ilk = client.post("/api/cekler", json={"yon": "alinan", "cek_no": "0012345", "banka": "T.C. Ziraat Bankası A.Ş.",
                                           "tutar": 1000, "vade": "2026-12-01"}, headers=h)
    assert ilk.status_code == 200
    r = client.post("/api/cekler", json={"yon": "alinan", "cek_no": "12345", "banka": "ZİRAAT BANKASI", "tutar": 1000,
                                         "vade": "2026-12-01"}, headers=h)
    assert r.status_code == 409 and "zaten kayıtlı" in r.json()["detail"]
    # Aynı numara başka bankada: ayrı çek
    assert client.post("/api/cekler", json={"yon": "alinan", "cek_no": "12345", "banka": "Akbank", "tutar": 1000,
                                            "vade": "2026-12-01"}, headers=h).status_code == 200
    r = client.post("/api/cekler/toplu", json={"satirlar": [
        {"yon": "alinan", "cek_no": "A1", "banka": "Garanti", "tutar": 500, "vade": "2026-11-10", "cari_unvan": "Yeni Müşteri Ltd. Şti.", "cari_vkn": "7777777777"},
        {"yon": "alinan", "cek_no": "A2", "banka": "Garanti", "tutar": 700, "vade": "2026-11-20", "cari_unvan": "YENİ MÜŞTERİ LTD"},
        {"yon": "alinan", "cek_no": "12345", "banka": "Ziraat", "tutar": 1000, "vade": "2026-12-01", "cari_id": musteri["id"]},
        {"yon": "alinan", "cek_no": "A3", "banka": "Garanti", "tutar": None, "vade": "2026-11-20", "cari_unvan": "Ahmet Yapı"},
        {"yon": "alinan", "cek_no": "A4", "banka": "Garanti", "tutar": 300, "vade": "2026-11-25", "cari_unvan": "Başka", "cari_vkn": "5555555555"},
    ]}, headers=h).json()
    assert [x["ok"] for x in r["sonuc"]] == [True, True, False, False, True]
    assert "zaten kayıtlı" in r["sonuc"][2]["hata"] and r["sonuc"][3]["hata"] == "Tutar yok." and r["yeni_cari"] == 1
    cekler = {c["cek_no"]: c for c in client.get("/api/cekler?yon=alinan", headers=h).json()}
    assert cekler["A1"]["cari_id"] == cekler["A2"]["cari_id"]  # ek farkı olan ad aynı cari
    assert cekler["A4"]["cari_id"] == musteri["id"]  # aynı VKN'li cari seçildi
    assert bakiye(h, musteri["id"]) == -300


def test_cek_excel_sablonu_onizleme():
    h, musteri, tedarikci, banka = kur()
    sablon = client.get("/api/excel/cek-sablon", headers=h).content
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(sablon))
    ws = wb.active
    ws.append(["Verilen", "Çek", "9", "Akbank", "", "", "", 100, "01.12.2026", "", "Demir Ltd", "", "", ""])
    ws.append(["Alınan", "Senet", "S5", "", "", "", "", "", "05.12.2026", "", "", "", "", ""])
    ws.append(["", "", "", "", "", "", "TOPLAM", 4100, "", "", "", "", "", ""])  # toplam satırı çek değildir
    tampon = io.BytesIO(); wb.save(tampon)
    r = client.post("/api/excel/cekler/onizle?yon=alinan", files={"dosya": ("c.xlsx", tampon.getvalue())}, headers=h).json()
    assert r["atlanan"] == 1 and len(r["satirlar"]) == 2 and r["toplam_satiri"] == 1 and r["uyari"] is None
    ws.append(["", "", "", "", "", "", "GENEL TOPLAM", 5000, "", "", "", "", "", ""])
    tampon = io.BytesIO(); wb.save(tampon)
    r2 = client.post("/api/excel/cekler/onizle?yon=alinan", files={"dosya": ("c.xlsx", tampon.getvalue())}, headers=h).json()
    assert len(r2["satirlar"]) == 2 and "900,00 fark" in r2["uyari"]
    ornek, senet = r["satirlar"]
    assert (ornek["cek_no"], ornek["tutar"], ornek["vade"], ornek["cari_unvan"], ornek["cari_id"]) == ("0012345", 4000, "2026-12-22", "Örnek Müşteri Ltd.", None)
    assert senet["tur"] == "senet" and senet["belirsiz"] == ["tutar"]
    assert client.get("/api/cekler", headers=h).json() == []  # önizleme kaydetmez


def test_cek_fotografindan_toplu_okuma():
    import pytest
    from app import ocr
    if not ocr.var_mi():
        pytest.skip("Tesseract kurulu değil")
    from PIL import Image, ImageDraw, ImageFont
    yazi = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    try:
        f, b = ImageFont.truetype(yazi, 30), ImageFont.truetype(yazi, 38)
    except OSError:
        pytest.skip("Yazı tipi yok")
    img = Image.new("L", (1700, 2300), 255)
    d = ImageDraw.Draw(img)
    for y0, banka, no, tutar, yazi_tutar, tarih, kesideci in (
            (60, "TÜRKİYE İŞ BANKASI A.Ş.", "A 1234567", "12.500,00", "ONİKİBİNBEŞYÜZ TÜRK LİRASI", "15.12.2026", "AHMET YAPI"),
            (1000, "T.C. ZİRAAT BANKASI A.Ş.", "B 7654321", "4.250,50", "DÖRTBİNİKİYÜZELLİ TÜRK LİRASI ELLİ KURUŞ", "20.01.2027", "MEHMET USTA")):
        d.rectangle((40, y0, 1660, y0 + 720), outline=0, width=3)
        d.text((80, y0 + 30), banka, font=b, fill=0)
        d.text((80, y0 + 90), "KÜTAHYA ŞUBESİ", font=f, fill=0)
        d.text((1100, y0 + 30), f"Seri No: {no}", font=f, fill=0)
        d.text((80, y0 + 180), f"Keşide Yeri: ANKARA      Keşide Tarihi: {tarih}", font=f, fill=0)
        d.text((80, y0 + 320), f"Yalnız: {yazi_tutar}", font=f, fill=0)
        d.text((1150, y0 + 320), f"#{tutar}#", font=b, fill=0)
        d.text((80, y0 + 420), f"Keşideci: {kesideci}", font=f, fill=0)
    tampon = io.BytesIO(); img.save(tampon, "PNG")
    h, musteri, tedarikci, banka = kur()
    r = client.post("/api/cekler/oku-toplu", files=[("dosyalar", ("cekler.png", tampon.getvalue(), "image/png"))], headers=h).json()
    s = r["satirlar"]
    assert [(x["banka"], x["cek_no"], x["tutar"], x["vade"]) for x in s] == [
        ("İş Bankası", "A1234567", 12500, "2026-12-15"), ("Ziraat Bankası", "B7654321", 4250.5, "2027-01-20")]
    # Cari sütunu yok: okumada boş gelir, kırmızı sayılmaz; kayıtta çeki yazan firmadan kurulur
    assert [(x["kesideci"], x["cari_id"], x["cari_unvan"]) for x in s] == [("AHMET YAPI", None, None), ("MEHMET USTA", None, None)]
    assert s[0]["belirsiz"] == []


def test_karekodlu_cek_ve_taranmis_pdf_kurallari():
    """Banka karekodu (KKB) kesin bilgidir; anlamsız okunan yazı alanı doldurulmaz; aynı keşidecinin çeklerinde ad eşitlenir."""
    from app import cek_okuma
    k = cek_okuma.kkb_karekodu("KKB 00 0001234567 0012 01546 0000000012345678 01234567890 0000000000000000 abc")
    assert (k["cek_no"], k["banka"], k["sube_kodu"], k["hesap_no"], k["vkn"]) == ("1234567", "Halkbank", "1546", "12345678", "1234567890")
    assert cek_okuma.kkb_karekodu("KKB 00 0001234567 0012 01546 0000000012345678 12345678901 0")["vkn"] == "12345678901"
    metin = "HALKBANK\nSERİ NO. 7654321 ÖRNEK YAPI LİMİTED\nŞİRKETİ\nKeşidecinin Adı, Soyadı ve İmzası\nV.K.N. 9999999999"
    c = cek_okuma.cek_metni_coz(metin, "KKB 00 0001234567 0012 01546 0000000012345678 01234567890 0000000000000000 x")
    assert (c["cek_no"], c["hesap_no"], c["vkn"], c["kesideci"]) == ("1234567", "12345678", "1234567890", "ÖRNEK YAPI LİMİTED ŞİRKETİ")
    assert not cek_okuma.anlamli_mi("in dı") and cek_okuma.anlamli_mi("ZAFER MEYDANI")
    c = cek_okuma.cek_metni_coz("Keşideci: in dı\nTarih 20.01.2027\n12,50 9,08")
    assert c["kesideci"] is None and c["tutar"] is None and {"kesideci", "tutar"} <= set(c["belirsiz"])
    assert cek_okuma.cek_metni_coz("Tarih 20.01.2027 TL #350.000#")["tutar"] == 350000
    cekler = [{"vkn": "1", "kesideci": "ÖRNEK YAPI LİMİTED", "sube": "MERKEZ", "belirsiz": []},
              {"vkn": "1", "kesideci": "ÖRNEK YAPI LİMİTED", "sube": "1546", "belirsiz": []},
              {"vkn": "1", "kesideci": "ŞİRKETİ X", "sube": None, "belirsiz": ["kesideci"]}]
    sonuc = cek_okuma._ayni_kesideciyi_esitle(cekler)
    assert [(c["kesideci"], c["sube"]) for c in sonuc] == [("ÖRNEK YAPI LİMİTED", "MERKEZ")] * 3 and sonuc[2]["belirsiz"] == []


def kasa_bakiye(h):
    return client.get("/api/kasa", headers=h).json()["bakiye"]


def test_cek_akibeti_teminat_kasa_toplu_islem_ve_vade():
    from datetime import date, timedelta
    h, musteri, tedarikci, banka = kur()
    bugun = date.today()
    gecen = (bugun - timedelta(days=3)).isoformat()
    ileri = (bugun + timedelta(days=40)).isoformat()
    a1 = client.post("/api/cekler", json={"yon": "alinan", "tutar": 1000, "vade": gecen, "cari_id": musteri["id"], "cek_no": "A1",
                                          "alis_tarihi": (bugun - timedelta(days=33)).isoformat()}, headers=h).json()
    a2 = client.post("/api/cekler", json={"yon": "alinan", "tutar": 2000, "vade": ileri, "cari_id": musteri["id"], "cek_no": "A2"}, headers=h).json()
    v1 = client.post("/api/cekler", json={"yon": "verilen", "tutar": 500, "vade": bugun.isoformat(), "cari_id": tedarikci["id"], "cek_no": "V1"}, headers=h).json()
    # Wolvox'taki vade günü / kalan vade
    assert a1["vade_gunu"] == 30 and a1["kalan_gun"] == -3
    assert v1["durum_adi"] == "Ödenmedi"
    # Vadesi gelenler: geçmiş ve bugün; 40 gün sonraki ancak önceden haber istenirse
    assert {c["cek_no"] for c in client.get("/api/cekler/vadesi-gelen", headers=h).json()} == {"A1", "V1"}
    assert {c["cek_no"] for c in client.get("/api/cekler/vadesi-gelen?gun=45", headers=h).json()} == {"A1", "A2", "V1"}

    # Teminata verilen çek portföy toplamında kalır, bankası yazılır; sonra portföye geri alınır
    r = client.post(f"/api/cekler/{a2['id']}/islem", json={"islem": "teminata_ver", "tarih": bugun.isoformat(), "banka_hesap_id": banka["id"]}, headers=h).json()
    assert r["durum"] == "teminatta" and r["banka_hesap_ad"] == "Ziraat TL" and r["islem_tarihi"] == bugun.isoformat()
    r = client.post(f"/api/cekler/{a2['id']}/islem", json={"islem": "portfoye_al", "tarih": bugun.isoformat()}, headers=h).json()
    assert r["durum"] == "portfoyde" and r["banka_hesap_id"] is None

    # Kasadan ödendi: Nakit Kasa'dan çıkar, banka değişmez; vadesi gelenlerden düşer
    k0, b0 = kasa_bakiye(h), banka_bakiye(h, banka["id"])
    r = client.post(f"/api/cekler/{v1['id']}/islem", json={"islem": "ode", "tarih": bugun.isoformat(), "kasa": True}, headers=h).json()
    assert r["durum"] == "odendi" and r["durum_adi"] == "Ödendi (kasadan)" and r["kalan_gun"] is None
    assert kasa_bakiye(h) == k0 - 500 and banka_bakiye(h, banka["id"]) == b0
    assert {c["cek_no"] for c in client.get("/api/cekler/vadesi-gelen", headers=h).json()} == {"A1"}

    # Toplu işlem: ikisi de bankaya tahsil edilir; yapılamayanın nedeni döner
    r = client.post("/api/cekler/toplu-islem", json={"idler": [a1["id"], a2["id"], v1["id"]], "islem": "tahsil", "tarih": bugun.isoformat(),
                                                     "banka_hesap_id": banka["id"]}, headers=h).json()
    assert r["yapilan"] == 2 and len(r["hatalar"]) == 1 and "V1" in r["hatalar"][0]
    assert banka_bakiye(h, banka["id"]) == b0 + 3000
    assert client.get("/api/cekler/vadesi-gelen", headers=h).json() == []

    # Çek/Senet Raporları: süzgeçler ve Excel
    rap = client.get("/api/raporlar/cekler", headers=h).json()
    assert len(rap["satirlar"]) == 3 and rap["ozet"]["alinan"]["tutar"] == 3000 and rap["ozet"]["verilen"]["acik"] == 0
    assert [x["cek_no"] for x in client.get(f"/api/raporlar/cekler?yon=alinan&banka_hesap_id={banka['id']}", headers=h).json()["satirlar"]] == ["A1", "A2"]
    assert client.get(f"/api/raporlar/cekler?vade_bas={ileri}", headers=h).json()["satirlar"][0]["cek_no"] == "A2"
    assert client.get("/api/excel/cekler?yon=alinan", headers=h).status_code == 200


def test_toplu_giriste_akibet():
    h, musteri, tedarikci, banka = kur()
    satir = {"yon": "alinan", "tutar": 700, "vade": "2026-12-01", "cari_id": musteri["id"], "alis_tarihi": "2026-10-01"}
    r = client.post("/api/cekler/toplu", json={"satirlar": [
        {**satir, "cek_no": "T1", "akibet": {"islem": "ciro", "tarih": "2026-10-05", "cari_id": tedarikci["id"]}},
        {**satir, "cek_no": "T2", "akibet": {"islem": "tahsile_ver", "tarih": "2026-10-05", "banka_hesap_id": banka["id"]}},
        {**satir, "cek_no": "T3", "akibet": {"islem": "ciro", "tarih": "2026-10-05"}},  # cari seçilmemiş: kaydedilmez
        {**satir, "cek_no": "T4"},
    ]}, headers=h).json()
    assert [s["ok"] for s in r["sonuc"]] == [True, True, False, True] and "cari" in r["sonuc"][2]["hata"]
    cekler = {c["cek_no"]: c for c in client.get("/api/cekler?yon=alinan", headers=h).json()}
    assert cekler["T1"]["durum"] == "ciro_edildi" and cekler["T1"]["ciro_cari_unvan"] == "Demir Ltd"
    assert cekler["T2"]["durum"] == "tahsilde" and cekler["T4"]["durum"] == "portfoyde" and "T3" not in cekler
    assert bakiye(h, tedarikci["id"]) == 700


def test_verileri_sifirla_once_yedek_alir_firma_kalir():
    h, musteri, tedarikci, banka = kur()
    client.post("/api/santiyeler", json={"ad": "Deneme Şantiye"}, headers=h)
    client.post("/api/cekler", json={"yon": "alinan", "tutar": 100, "vade": "2026-12-01", "cari_id": musteri["id"]}, headers=h)
    client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": "D1", "cari_id": tedarikci["id"],
                                        "kalemler": [{"aciklama": "Demir", "birim_fiyat": 100, "kdv_orani": 20}]}, headers=h)
    # Başka firmanın verisine dokunulmaz
    h2, m2, _, _ = kur()
    assert client.post("/api/sistem/sifirla", json={"onay": "evet"}, headers=h).status_code == 400
    assert client.get("/api/cariler", headers=h).json()
    r = client.post("/api/sistem/sifirla", json={"onay": "sıfırla"}, headers=h).json()
    assert r["ok"] and r["yedek"] and r["fatura"] == 1 and r["cek"] == 1
    for yol in ("/api/cariler", "/api/santiyeler", "/api/cekler", "/api/faturalar", "/api/banka-hesaplari", "/api/hareketler"):
        assert client.get(yol, headers=h).json() == [], yol
    assert client.get("/api/ben", headers=h).status_code == 200
    assert any(c["id"] == m2["id"] for c in client.get("/api/cariler", headers=h2).json())


def test_yapi_kredi_tipi_cek_kurallari():
    """Ayraçsız tarih, ':' okunan binlik ayraç, IBAN'dan banka, 'SER! NO', basım tarihi kırıntısı, önceki çekten keşideci."""
    from datetime import date
    from app import cek_okuma, el_yazisi
    assert el_yazisi.vade_bul(["KUTAHYA", "Tah.05122026"], bugun=date(2026, 10, 21)) == (date(2026, 12, 5), True)
    assert el_yazisi.vade_bul(["KUTAA Tarh:05 01 2024"], bugun=date(2026, 10, 21)) == (date(2027, 1, 5), False)
    assert el_yazisi.rakamdan_tutar(["TL #300:000#"])[0] == 300000
    c = cek_okuma.cek_metni_coz("ÇEK SER! NO, 1000005\nIBAN. TR17 0006 7010 0000 0011 2233 44\nBAS.TRH. 17 OCAK 2025 ÖRNEK ENERJİ LİMİTED")
    assert c["cek_no"] == "1000005" and c["banka"] == "Yapı Kredi" and c["kesideci"] == "ÖRNEK ENERJİ LİMİTED"
    h, musteri, _, _ = kur()
    client.post("/api/cekler", json={"yon": "alinan", "tutar": 10, "vade": "2026-12-01", "cek_no": "1", "banka": "Yapı Kredi",
                                     "sube": "KÜTAHYA", "hesap_no": "11223344", "kesideci": "ÖRNEK ENERJİ LİMİTED ŞİRKETİ"}, headers=h)
    from app.db import SessionLocal
    from app.models import Cek, Kullanici
    from app.routers.cekler import _okunanlari_hazirla
    db = SessionLocal()
    sid = db.query(Cek).filter(Cek.hesap_no == "11223344", Cek.cek_no == "1").order_by(Cek.id.desc()).first().sirket_id
    k = db.query(Kullanici).filter(Kullanici.sirket_id == sid).first()
    r = _okunanlari_hazirla(db, k, [{"cek_no": "2", "hesap_no": "TR170006701000000011223344", "kesideci": "P LTD",
                                      "vkn": "1234567890", "banka": None, "belirsiz": ["kesideci", "banka"]}], "x.pdf", "x.pdf")[0]
    # VKN çeki yazan firmanın; cari okumada boş (kayıtta çeki yazan firmadan kurulur)
    assert r["cari_id"] is None and r["cari_unvan"] is None and r["kesideci_vkn"] == "1234567890" and "cari_unvan" not in r["belirsiz"]
    assert (r["kesideci"], r["banka"], r["sube"], r["hesap_no"]) == ("ÖRNEK ENERJİ LİMİTED ŞİRKETİ", "Yapı Kredi", "KÜTAHYA", "11223344")
    assert "kesideci" not in r["belirsiz"]
    db.close()


def test_lehtar_bizsek_firma_adi_yazilir():
    from app import kural
    u = "ARMA GRUP ENERJİ MÜH. İNŞ. LTD. ŞTİ."
    assert kural.firma_adi_benzer_mi("ARMA CRDP ENEPSIMUH.", u) and not kural.firma_adi_benzer_mi("ORNEK ENERJI MUHENDISLIK", u)
    from app import cek_okuma
    assert cek_okuma.cek_metni_coz("Bu çek karşılığında ÖRNEK YAPI LTD. ŞTİ. emrine\nTarih 01.12.2026")["lehtar"] == "ÖRNEK YAPI LTD. ŞTİ"


def test_lehtar_bozuk_okunursa_bos_kisa_unvanla_tanınir():
    from app import cek_okuma
    from app.db import SessionLocal
    from app.models import Kullanici, Sirket
    from app.routers.cekler import _okunanlari_hazirla
    # Bozuk okunan lehtar uydurma yazı olarak gösterilmez
    assert cek_okuma.cek_metni_coz("Bu çek karşılığında NES SA Ae 2. Mid. Ihe. LTO. emrine\nTarih 01.12.2026")["lehtar"] is None
    db = SessionLocal()
    k = db.query(Kullanici).first()
    s = db.get(Sirket, k.sirket_id)
    eski = s.unvan
    s.unvan = "ARMA"  # firma bilgilerinde kısa ünvan
    db.commit()
    try:
        r = _okunanlari_hazirla(db, k, [{"cek_no": "71", "lehtar": None, "kesideci": "ÖRNEK ENERJİ", "belirsiz": [],
                                          "_metinler": ["0765 KUTAHYA SUBESI", "ARMA CRUP ENEPG MUH."]}], None, "x.pdf")[0]
        assert r["lehtar"] == "ARMA"
        # Çeki biz yazdıysak basılı adımız lehtar sayılmaz
        r = _okunanlari_hazirla(db, k, [{"cek_no": "72", "lehtar": None, "kesideci": "ARMA GRUP", "belirsiz": [],
                                          "_metinler": ["ARMA GRUP ENERJİ"]}], None, "x.pdf")[0]
        assert r["lehtar"] is None and "lehtar" in r["belirsiz"]
    finally:
        s.unvan = eski
        db.commit()
        db.close()


def test_toplu_cekte_cari_cekin_yazanindan_kurulur():
    # Alınan çekte cari = çeki yazan firma (VKN ile caride varsa ona), verilen çekte cari = lehtar; Excel'de yazılan cari önce gelir
    h, musteri, tedarikci, banka = kur()
    vkn = str(uuid.uuid4().int)[:10]
    satirlar = [
        {"yon": "alinan", "cek_no": "OT1", "banka": "Garanti", "tutar": 100, "vade": "2026-12-10", "kesideci": "Önem Enerji Ltd", "kesideci_vkn": vkn},
        {"yon": "alinan", "cek_no": "OT2", "banka": "Garanti", "tutar": 200, "vade": "2026-12-11", "kesideci": "ÖRNEK ENERJİ LTD. (yazım farkı)", "kesideci_vkn": vkn},
        {"yon": "alinan", "cek_no": "OT3", "banka": "Garanti", "tutar": 300, "vade": "2026-12-12", "kesideci": "Önem Enerji Ltd", "cari_id": tedarikci["id"]},
        {"yon": "verilen", "cek_no": "OT4", "banka": "Garanti", "tutar": 400, "vade": "2026-12-13", "lehtar": "Demirci Yapı " + vkn[:4]},
    ]
    r = client.post("/api/cekler/toplu", json={"satirlar": satirlar}, headers=h).json()
    assert r["kaydedilen"] == 4 and r["yeni_cari"] == 2
    cekler = {c["cek_no"]: c for c in client.get("/api/cekler", headers=h).json()}
    assert cekler["OT1"]["cari_id"] == cekler["OT2"]["cari_id"] != tedarikci["id"]
    assert cekler["OT3"]["cari_id"] == tedarikci["id"]
    assert cekler["OT4"]["cari_unvan"] == "Demirci Yapı " + vkn[:4]


def test_yapay_zekayla_cek_okuma_bos_ve_kirmizi_alanlari_doldurur(monkeypatch):
    from app import cek_okuma
    kural_cekleri = [
        {"cek_no": "1000001", "vade": None, "tutar": 750000.0, "lehtar": None, "belirsiz": ["vade"]},
        {"cek_no": "1000002", "vade": "2027-01-05", "tutar": 750000.0, "belirsiz": ["vade", "tutar"]},
        {"cek_no": "1000003", "vade": None, "tutar": None, "belirsiz": ["vade", "tutar"]},
    ]
    yz = [
        {"cek_no": "1000003", "banka": "Yapı Kredi", "sube": None, "hesap_no": None, "cekin_yazani": "ÖRNEK ENERJİ",
         "cekin_yazani_vkn": "1234567890", "lehtar": "ARMA GRUP ENERJİ", "vade": "2026-12-20", "tutar_rakamla": 400000,
         "tutar_yaziyla": "dörtyüzbin", "tutar_yaziyla_sayi": 400000, "para_birimi": "TL"},
        {"cek_no": "1000001", "banka": None, "sube": None, "hesap_no": None, "cekin_yazani": None, "cekin_yazani_vkn": None,
         "lehtar": "ARMA GRUP", "vade": "2026-12-10", "tutar_rakamla": 700000, "tutar_yaziyla": "yediyüzellibin",
         "tutar_yaziyla_sayi": 750000, "para_birimi": "TRY"},
        {"cek_no": "1000002", "banka": None, "sube": None, "hesap_no": None, "cekin_yazani": None, "cekin_yazani_vkn": None,
         "lehtar": None, "vade": "2027-01-05", "tutar_rakamla": 750000, "tutar_yaziyla": None, "tutar_yaziyla_sayi": None,
         "para_birimi": "TRY"},
        {"cek_no": "1000009", "banka": "Yapı Kredi", "sube": "Kütahya", "hesap_no": "11223344", "cekin_yazani": "ÖRNEK ENERJİ",
         "cekin_yazani_vkn": None, "lehtar": None, "vade": "2027-02-20", "tutar_rakamla": 300000, "tutar_yaziyla": "üçyüzbin",
         "tutar_yaziyla_sayi": 300000, "para_birimi": "TRY"},
    ]
    c = {x["cek_no"]: x for x in cek_okuma.yapay_zekayla_tamamla([dict(x, belirsiz=list(x["belirsiz"])) for x in kural_cekleri], yz)}
    # Rakam ve yazı tutuyor: tutar emin; vade yapay zekadan ama kural okumasıyla doğrulanmadığı için kırmızı
    assert (c["1000003"]["tutar"], c["1000003"]["vade"]) == (400000, "2026-12-20") and c["1000003"]["belirsiz"] == ["vade", "banka", "kesideci"]  # boştu, yapay zekadan: kontrol edin
    # Kural okumasının emin tutarına dokunulmaz; vade doldurulur ama kırmızı kalır
    assert c["1000001"]["tutar"] == 750000.0 and c["1000001"]["vade"] == "2026-12-10" and c["1000001"]["belirsiz"] == ["vade"]
    # Vade iki okumada aynı: yeşil; yazı okunamadı: tutar kırmızı kalır
    assert c["1000002"]["belirsiz"] == ["tutar"]
    # Kural okumasının bulamadığı çek eklenir: tutar dışında her şey kırmızı, cari hiç doldurulmaz
    assert c["1000009"]["tutar"] == 300000 and "tutar" not in c["1000009"]["belirsiz"] and "vade" in c["1000009"]["belirsiz"]
    assert all("cari_unvan" not in x and "cari_id" not in x for x in c.values())

    # Uç: anahtar yoksa kısa açıklama; varsa yapay zeka sonucu lehtar kuralından geçer (firmamıza benziyorsa firma ünvanı)
    h, *_ = kur()
    png = io.BytesIO()
    from PIL import Image
    Image.new("RGB", (200, 100), "white").save(png, "PNG")
    dosya = [("dosyalar", ("cek.png", png.getvalue(), "image/png"))]
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    r = client.post("/api/cekler/oku-toplu?yapay_zeka=true", files=dosya, headers=h)
    assert r.status_code == 400 and "ANTHROPIC_API_KEY" in r.json()["detail"]
    monkeypatch.setenv("ANTHROPIC_API_KEY", "deneme")
    monkeypatch.setattr(cek_okuma, "dosyadan_oku", lambda *a: [])
    unvan = client.get("/api/ben", headers=h).json()["sirket"]
    monkeypatch.setattr(ai, "cek_sayfasi_oku", lambda jpeg: [dict(yz[0], lehtar=unvan.upper())])
    s = client.post("/api/cekler/oku-toplu?yapay_zeka=true", files=dosya, headers=h).json()["satirlar"]
    assert len(s) == 1 and s[0]["lehtar"] == unvan and s[0]["tutar"] == 400000 and s[0]["cari_unvan"] is None
    assert "lehtar" not in s[0]["belirsiz"] and "tutar" not in s[0]["belirsiz"] and s[0]["kesideci"] == "ÖRNEK ENERJİ"


def test_yapay_zekaya_yalniz_eksik_cekin_goruntusu_gider(monkeypatch):
    """Kural okuması çekleri bulduysa yapay zekaya yalnız boş/kırmızı alanı olan çekin bandı gönderilir."""
    from PIL import Image
    from app import ai, cek_okuma, ocr
    monkeypatch.setattr(ocr, "_goruntuler", lambda veri, pdf: [Image.new("RGB", (1000, 2000), "white")])
    giden = []

    def sahte(jpeg):
        import io
        giden.append(Image.open(io.BytesIO(jpeg)).size)
        return [{"cek_no": "1000002", "banka": None, "sube": None, "hesap_no": None, "cekin_yazani": None,
                 "cekin_yazani_vkn": None, "lehtar": None, "vade": "2027-01-05", "tutar_rakamla": 750000,
                 "tutar_yaziyla": "yediyüzellibin", "tutar_yaziyla_sayi": 750000, "para_birimi": "TRY"}]
    monkeypatch.setattr(ai, "cek_sayfasi_oku", sahte)
    tamam = {"cek_no": "1000001", "banka": "Yapı Kredi", "kesideci": "ÖRNEK", "tutar": 1.0, "vade": "2027-01-01",
             "belirsiz": [], "_bant": (1, 0, 900)}
    eksik = {"cek_no": "1000002", "banka": "Yapı Kredi", "kesideci": "ÖRNEK", "tutar": None, "vade": "2027-01-05",
             "belirsiz": ["tutar"], "_bant": (1, 1000, 1900)}
    sonuc = cek_okuma.yapay_zekayla_eksikleri_oku(b"%PDF-", "x.pdf", [tamam, eksik])
    assert giden == [(1000, 920)]  # tek çekin bandı (alt kenar payıyla), bütün sayfa değil
    assert sonuc[1]["tutar"] == 750000 and "tutar" not in sonuc[1]["belirsiz"] and sonuc[0]["tutar"] == 1.0


def test_okuma_raporu_zip_indirilir(monkeypatch):
    import io
    import json
    import zipfile
    from app import cek_okuma
    h = yeni_kullanici(f"rapor{uuid.uuid4().hex[:6]}@ornek.com")
    sahte = [{"cek_no": "1000001", "tutar": None, "vade": None, "belirsiz": ["tutar", "vade"],
              "nedenler": {"tutar": "Tutar okunamadı."}, "_tani": {"bolme": "karekod", "el_yazisi": [{"m": "x", "g": 0.5}]}}]
    monkeypatch.setattr(cek_okuma, "dosyadan_oku", lambda *a, **k: sahte)
    r = client.post("/api/cekler/okuma-raporu", files=[("dosyalar", ("cek.png", b"x", "image/png"))], headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    rapor = json.loads(z.read("001_cek.png/rapor.json"))
    assert rapor[0]["nedenler"]["tutar"] == "Tutar okunamadı." and rapor[0]["okuma"]["bolme"] == "karekod"
