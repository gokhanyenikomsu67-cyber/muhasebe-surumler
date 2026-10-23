"""Şantiye içindeki depolar (taşeron): kalem depoya ayrılır, depo carisine borç yazılır, satıcı carisi değişmez;
taşeron sayfası, elden harcama, bordrodan personel/SGK (sentetik veri)."""
import io
import uuid

from openpyxl import Workbook, load_workbook

from tests.test_api import client, yeni_kullanici
from tests.test_cek_banka import kasa_bakiye


def _kur():
    h = yeni_kullanici(f"{uuid.uuid4().hex}@ornek.com")
    s = client.post("/api/santiyeler", json={"ad": "Deneme 1.Grup"}, headers=h).json()
    satici = client.post("/api/cariler", json={"unvan": "Örnek Yapı Market", "vkn": "1111111111"}, headers=h).json()
    taseron = client.post("/api/cariler", json={"unvan": "Taşeron Deneme"}, headers=h).json()
    d = client.post("/api/depolar", json={"ad": "Taşeron Deneme", "santiye_id": s["id"], "cari_id": taseron["id"]}, headers=h).json()
    return h, s, satici, taseron, d


def _fatura(h, s, satici, depo_id=None):
    f = client.post("/api/faturalar", json={
        "yon": "gelen", "fatura_no": f"DNM{uuid.uuid4().hex[:6]}", "tarih": "2026-09-10", "cari_id": satici["id"], "santiye_id": s["id"],
        "kalemler": [{"aciklama": "Çimento", "miktar": 10, "birim_fiyat": 100, "kdv_orani": 20, "depo_id": depo_id},
                     {"aciklama": "Kum", "miktar": 1, "birim_fiyat": 500, "kdv_orani": 20}]}, headers=h).json()
    assert client.post(f"/api/faturalar/{f['id']}/onayla", headers=h).status_code == 200
    return f


def _bakiye(h, cari_id):
    return client.get(f"/api/raporlar/cari-ekstre/{cari_id}", headers=h).json()["bakiye"]


def test_depoya_ayrilan_kalem_taseron_carisine_borc_satici_degismez():
    h, s, satici, taseron, d = _kur()
    assert d["santiye_ad"] == "Deneme 1.Grup" and d["cari_unvan"] == "Taşeron Deneme"
    # Aynı şantiyede aynı adla ikinci depo açılmaz
    assert client.post("/api/depolar", json={"ad": "TAŞERON  deneme", "santiye_id": s["id"]}, headers=h).json()["id"] == d["id"]
    f = _fatura(h, s, satici, d["id"])
    assert _bakiye(h, satici["id"]) == -1800  # satıcıya borcumuz faturanın tamamı
    assert _bakiye(h, taseron["id"]) == 1200  # yalnız depoya ayrılan çimento (KDV dahil) taşerona borç
    sayfa = client.get(f"/api/taseron/{taseron['id']}", headers=h).json()
    assert [(x["kalem"], x["toplam"], x["satici"], x["santiye"]) for x in sayfa["kalemler"]] == \
        [("Çimento", 1200, "Örnek Yapı Market", "Deneme 1.Grup")]
    assert sayfa["ozet"]["kalemler"] == 1200 and sayfa["ozet"]["bakiye"] == 1200
    # Depo raporu ve şantiye Excel'inde depo görünür; depoya ayrılmayan kum şantiyenin kendisinde
    rapor = client.get(f"/api/raporlar/depo-harcamalari?santiye_id={s['id']}", headers=h).json()
    assert [(x["depo"], x["kalem"]) for x in rapor] == [("Taşeron Deneme", "Çimento")]
    ws = load_workbook(io.BytesIO(client.get("/api/excel/depo-harcamalari", headers=h).content)).active
    assert ws.cell(2, 2).value == "Taşeron Deneme" and ws.cell(3, 7).value == "TOPLAM"
    ws = load_workbook(io.BytesIO(client.get("/api/excel/santiye-faturalari", headers=h).content)).active
    assert [(r[1], r[7]) for r in ws.iter_rows(min_row=2, max_row=3, values_only=True)] == [("Taşeron Deneme", "Çimento"), (None, "Kum")]
    # Depo kalemi olan depo silinemez; fatura iptal edilince taşeron borcu kalkar
    assert client.delete(f"/api/depolar/{d['id']}", headers=h).status_code == 409
    assert client.post(f"/api/faturalar/{f['id']}/iptal", json={"neden": "deneme"}, headers=h).status_code == 200
    assert _bakiye(h, taseron["id"]) == 0


def test_toplu_depo_atamasi_ve_kaldirma():
    h, s, satici, taseron, d = _kur()
    f = _fatura(h, s, satici)
    assert _bakiye(h, taseron["id"]) == 0
    assert client.post("/api/faturalar/toplu-guncelle", json={"idler": [f["id"]], "depo_id": d["id"]}, headers=h).status_code == 200
    assert _bakiye(h, taseron["id"]) == 1800
    client.post("/api/faturalar/toplu-guncelle", json={"idler": [f["id"]], "depo_id": 0}, headers=h)
    assert _bakiye(h, taseron["id"]) == 0


def test_elden_harcama_kasadan_duser_ve_taserona_yazilir():
    h, s, satici, taseron, d = _kur()
    k0 = kasa_bakiye(h)
    g = client.post("/api/taseron-giderleri", json={"cari_id": taseron["id"], "tur": "elden", "tarih": "2026-09-12",
                                                    "tutar": 750, "aciklama": "Yemek", "santiye_id": s["id"], "kasadan": True}, headers=h)
    assert g.status_code == 200, g.text
    assert kasa_bakiye(h) == k0 - 750 and _bakiye(h, taseron["id"]) == 750
    # Verilen para da taşeron sayfasında ayrı sekmede
    client.post("/api/hareketler", json={"tarih": "2026-09-15", "yon": "cikis", "tur": "odeme", "tutar": 2000,
                                         "cari_id": taseron["id"]}, headers=h)
    sayfa = client.get(f"/api/taseron/{taseron['id']}", headers=h).json()
    assert sayfa["ozet"]["elden"] == 750 and sayfa["ozet"]["paralar"] == 2000 and sayfa["ozet"]["bakiye"] == 2750
    assert client.post("/api/taseron-giderleri/toplu-sil", json={"idler": [g.json()["id"]]}, headers=h).status_code == 200
    assert kasa_bakiye(h) == k0 - 2000 and _bakiye(h, taseron["id"]) == 2000


def _bordro() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "EYLÜL AYI MAAŞ BORDROSU"
    ws.append(["Kurum Adı", None, "DENEME LTD"])
    ws.append(["Ay / Yıl ", None, "EYLÜL / 2026"])
    ws.append(["#", "Adı Soyadı", "TCKN", "Giriş T", "Çıkış T", "T.Gün", "Nor.Kazanç", "SSK M.", "SSK İşçi", "İşs.P.İşçi",
               "G.V.M", "Gel.Ver.", "Kalan GV", "Damga V", "N.Ödenen", "icra"])
    ws.append([1, "PERSONEL BİR", "10000000146", None, None, 30, 1000, 1000, 140, 10, 850, 50, 20, 5, 825, None])
    ws.append([2, "PERSONEL İKİ", "", None, None, 15, 500, 500, 70, 5, 425, 0, 0, 0, 425, None])
    ws.append([None, "TOPLAM", None, None, None, 45, 1500, 1500, 210, 15, 1275, 50, 20, 5, 1250, None])
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def test_bordro_yalniz_depo_secilerek():
    """Bordro Yükle'de yalnız taşeron (depo) seçilir: carisi depodan alınır, depoda cari yoksa depo adıyla açılıp bağlanır."""
    h, s, _, taseron, d = _kur()
    v = client.post("/api/taseron-giderleri/bordro-onizle", files={"dosya": ("bordro.xlsx", _bordro(), "application/octet-stream")}, headers=h).json()
    kayit = {"donem": v["donem"], "depo_id": d["id"], "kisiler": v["kisiler"]}
    assert client.post("/api/taseron-giderleri/bordro", json=kayit, headers=h).json() == {"eklenen": 2, "guncellenen": 0}
    assert _bakiye(h, taseron["id"]) == 1841.25
    bos = client.post("/api/depolar", json={"ad": "Deneme Taşeron İki", "santiye_id": s["id"]}, headers=h).json()
    assert client.post("/api/taseron-giderleri/bordro", json={**kayit, "depo_id": bos["id"]}, headers=h).json()["eklenen"] == 2
    bagli = next(x for x in client.get("/api/depolar", headers=h).json() if x["id"] == bos["id"])
    assert bagli["cari_unvan"] == "Deneme Taşeron İki" and _bakiye(h, bagli["cari_id"]) == 1841.25
    assert client.post("/api/taseron-giderleri/bordro", json={**kayit, "depo_id": None}, headers=h).status_code == 422


def test_bordro_excelinden_personel_sgk():
    h, s, satici, taseron, d = _kur()
    r = client.post("/api/taseron-giderleri/bordro-onizle", files={"dosya": ("bordro.xlsx", _bordro(), "application/octet-stream")}, headers=h)
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["donem"] == "2026-09" and [x["personel"] for x in v["kisiler"]] == ["PERSONEL BİR", "PERSONEL İKİ"]  # TOPLAM alınmaz
    bir = v["kisiler"][0]
    assert (bir["brut"], bir["net"], bir["sgk_isci"], bir["vergi"], bir["gun"]) == (1000, 825, 150, 25, 30)
    kayit = {"cari_id": taseron["id"], "donem": v["donem"], "santiye_id": s["id"], "kisiler": v["kisiler"]}  # SGK oranı varsayılan 37,75
    assert client.post("/api/taseron-giderleri/bordro", json=kayit, headers=h).json() == {"eklenen": 2, "guncellenen": 0}
    # SGK masrafı = normal kazanç x %37,75; taşerona net + vergi + SGK masrafı: (825+25+377,5) + (425+188,75)
    assert _bakiye(h, taseron["id"]) == 1841.25
    # Aynı ay ikinci kez yüklenince çift yazılmaz
    assert client.post("/api/taseron-giderleri/bordro", json=kayit, headers=h).json() == {"eklenen": 0, "guncellenen": 2}
    sayfa = client.get(f"/api/taseron/{taseron['id']}", headers=h).json()
    assert sayfa["ozet"]["personel"] == 1841.25 and len(sayfa["personel"]) == 2
    assert [x["sgk_isci"] for x in sayfa["personel"]] == [377.5, 188.75]
    ekstre = client.get(f"/api/raporlar/cari-ekstre/{taseron['id']}", headers=h).json()["satirlar"]
    assert [(x["belge"], x["borc"]) for x in ekstre] == [("2026-09", 1841.25)]  # ekstrede ay tek satır
    wb = load_workbook(io.BytesIO(client.get(f"/api/excel/taseron/{taseron['id']}", headers=h).content))
    assert wb.sheetnames == ["Özet", "Verilen Paralar", "Fatura Çıkışları", "Elden Harcamalar", "Personel SGK"]
    # Taşeronu olan cari silinemez
    assert client.delete(f"/api/cariler/{taseron['id']}", headers=h).status_code == 409
