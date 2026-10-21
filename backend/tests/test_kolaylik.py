"""Hızlı arama, ödeme takvimi / nakit akışı, düzenli ödemeler, fatura kopyalama ve cari vade günü (sentetik veri)."""
import io
import os
import tempfile
from datetime import date, timedelta

os.environ.setdefault("MUHASEBE_VERI", tempfile.mkdtemp())

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests.test_mali import cari, fatura, yeni

client = TestClient(app)
BUGUN = date.today()


def test_hizli_arama_ad_vkn_ve_tutar():
    h = yeni()
    c = cari(h, "Çağlar İnşaat Ltd.", "3333333333")
    fatura(h, "gelen", c, "2026-03-01", 1000, no="ABC2026000000001")  # genel toplam 1.200
    client.post("/api/santiyeler", json={"ad": "Şişli Konut"}, headers=h)
    r = client.get("/api/ara", params={"q": "caglar"}, headers=h).json()
    assert [x["unvan"] for x in r["cariler"]] == ["Çağlar İnşaat Ltd."]
    assert r["faturalar"][0]["fatura_no"] == "ABC2026000000001"
    assert client.get("/api/ara", params={"q": "sisli"}, headers=h).json()["santiyeler"][0]["ad"] == "Şişli Konut"
    assert client.get("/api/ara", params={"q": "1.200,00"}, headers=h).json()["faturalar"][0]["genel_toplam"] == 1200
    assert client.get("/api/ara", params={"q": "3333"}, headers=h).json()["cariler"]
    assert client.get("/api/ara", params={"q": "x"}, headers=h).json()["cariler"] == []


def test_cari_vade_gunu_ve_odeme_takvimi():
    h = yeni()
    t = client.post("/api/cariler", json={"unvan": "Beton A.Ş.", "vkn": "4444444444", "vade_gun": 30}, headers=h).json()
    assert t["vade_gun"] == 30
    # Vade günü gönderilmeyen güncelleme (eski ekran) onu silmez
    t = client.put(f"/api/cariler/{t['id']}", json={"unvan": "Beton A.Ş.", "vkn": "4444444444"}, headers=h).json()
    assert t["vade_gun"] == 30
    m = cari(h, "Müşteri Ltd.", "5555555555")
    fatura(h, "gelen", t["id"], str(BUGUN - timedelta(days=10)), 1000)          # vade = +20 gün, 1.200 ödenecek
    fatura(h, "giden", m, str(BUGUN - timedelta(days=40)), 500, vade=str(BUGUN - timedelta(days=5)))  # 600 gecikmiş alacak
    v = client.get("/api/raporlar/nakit-akisi?gun=60", headers=h).json()
    fatura_kalemleri = {x["cari"]: x for x in v["kalemler"] if x["tur"] == "fatura"}
    assert fatura_kalemleri["Beton A.Ş."]["tarih"] == str(BUGUN + timedelta(days=20))
    assert fatura_kalemleri["Müşteri Ltd."]["gecikmis"] and fatura_kalemleri["Müşteri Ltd."]["gecikme_gun"] == 5
    assert v["ozet"]["gecikmis_giris"] == 600 and v["ozet"]["ay_cikis"] == 1200
    assert v["donemler"][0]["ad"] == "Vadesi geçmiş" and v["donemler"][-1]["bakiye"] == round(v["mevcut"] + 600 - 1200, 2)
    ws = load_workbook(io.BytesIO(client.get("/api/excel/nakit-akisi?gun=60", headers=h).content)).active
    assert ws["A1"].value == "Vade" and ws.max_row == 3


def test_duzenli_odeme_takvim_ve_odendi():
    h = yeni()
    kiraci = cari(h, "Ev Sahibi", "6666666666")
    hesap = client.post("/api/banka-hesaplari", json={"ad": "Halkbank", "acilis_bakiyesi": 50000}, headers=h).json()
    bas = (BUGUN.replace(day=1) - timedelta(days=1)).replace(day=1)  # geçen ayın başı
    d = client.post("/api/duzenli-odemeler", json={"ad": "Ofis kirası", "tutar": 10000, "gun": 31, "baslangic": str(bas),
                                                   "cari_id": kiraci}, headers=h).json()
    assert d["geciken"] == 1  # geçen ayın son günü geçti; bu ayın son günü henüz gelmedi
    v = client.get("/api/raporlar/nakit-akisi?gun=40", headers=h).json()
    kiralar = [x for x in v["kalemler"] if x["tur"] == "duzenli"]
    assert kiralar and kiralar[0]["gecikmis"]
    # En eski vade ödendi: banka hareketi de kaydedilir, cari bakiyesi kapanır
    r = client.post(f"/api/duzenli-odemeler/{d['id']}/odendi", json={"hareket": True, "banka_hesap_id": hesap["id"]}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["son_donem"] == f"{bas.year:04d}-{bas.month:02d}"
    hareketler = client.get("/api/hareketler?durum=onayli", headers=h).json()
    assert len(hareketler) == 1 and hareketler[0]["tutar"] == 10000 and hareketler[0]["tur"] == "odeme"
    # Yalnız işaretle: hareket kaydedilmez
    client.post(f"/api/duzenli-odemeler/{d['id']}/odendi", json={}, headers=h)
    assert len(client.get("/api/hareketler?durum=onayli", headers=h).json()) == 1
    assert client.get("/api/duzenli-odemeler", headers=h).json()[0]["son_donem"] == f"{BUGUN.year:04d}-{BUGUN.month:02d}"
    assert client.delete(f"/api/duzenli-odemeler/{d['id']}", headers=h).status_code == 200


def test_fatura_kopyala():
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Deneme İnşaat", "fatura_seri": "AFS"}, headers=h)
    m = cari(h, "Kiracı A.Ş.", "7777777777")
    f = fatura(h, "giden", m, "2026-01-05", 2500, no=f"AFS{BUGUN.year}000000004", vade="2026-01-20")
    k = client.post(f"/api/faturalar/{f['id']}/kopyala", headers=h)
    assert k.status_code == 200, k.text
    k = k.json()
    assert k["durum"] == "taslak" and k["tarih"] == str(BUGUN) and k["vade"] == str(BUGUN + timedelta(days=15))
    assert k["fatura_no"] == f"AFS{BUGUN.year}000000005" and k["cari_id"] == m
    assert [(x["aciklama"], x["tutar"]) for x in k["kalemler"]] == [("Kalem", 2500)] and k["genel_toplam"] == 3000
