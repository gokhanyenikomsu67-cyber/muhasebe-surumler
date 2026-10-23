"""e-Fatura XML (UBL-TR) okuma: iskonto, karışık KDV, tevkifat, diğer vergiler, görüntü, mükerrer kayıt (sentetik veri)."""
import base64
import io
import zipfile

import pytest

from app import kural
from tests.test_api import client
from tests.test_kural import pdf_olustur, yeni

XSLT = b"""<xsl:stylesheet version="1.0" xmlns:xsl="http://www.w3.org/1999/XSL/Transform"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
 <xsl:template match="/"><html><body><h1>FATURA <xsl:value-of select="//cbc:ID[1]"/></h1></body></html></xsl:template>
</xsl:stylesheet>"""


def vergi(tutar, oran, kod="0015"):
    return (f"<cac:TaxSubtotal><cbc:TaxAmount>{tutar}</cbc:TaxAmount><cbc:Percent>{oran}</cbc:Percent>"
            f"<cac:TaxCategory><cac:TaxScheme><cbc:TaxTypeCode>{kod}</cbc:TaxTypeCode></cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal>")


def satir(no, ad, miktar, birim, fiyat, net, kdv, oran, isk="", ek=""):
    return (f"<cac:InvoiceLine><cbc:ID>{no}</cbc:ID><cbc:InvoicedQuantity unitCode=\"{birim}\">{miktar}</cbc:InvoicedQuantity>"
            f"<cbc:LineExtensionAmount>{net}</cbc:LineExtensionAmount>{isk}"
            f"<cac:TaxTotal><cbc:TaxAmount>{kdv}</cbc:TaxAmount>{vergi(kdv, oran)}{ek}</cac:TaxTotal>"
            f"<cac:Item><cbc:Name>{ad}</cbc:Name></cac:Item><cac:Price><cbc:PriceAmount>{fiyat}</cbc:PriceAmount></cac:Price>"
            "</cac:InvoiceLine>")


def ubl(no="ORN2026000000001", satici_vkn="1234567890", alici_vkn="9999999999", xslt=False, tip="SATIS"):
    """%20 kalemde %10 satır iskontosu, %10 tevkifatlı hizmet, %0 istisna satırı, %10 fatura altı iskonto, ÖİV."""
    ek = (f"<cac:AdditionalDocumentReference><cbc:ID>x</cbc:ID><cac:Attachment><cbc:EmbeddedDocumentBinaryObject "
          f"mimeCode=\"application/xml\" filename=\"gorunum.xslt\">{base64.b64encode(XSLT).decode()}"
          f"</cbc:EmbeddedDocumentBinaryObject></cac:Attachment></cac:AdditionalDocumentReference>") if xslt else ""
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
 <cbc:ID>{no}</cbc:ID><cbc:UUID>3f1c0a2e-0000-4000-8000-{no[-12:].lower()}</cbc:UUID>
 <cbc:IssueDate>2026-09-15</cbc:IssueDate><cbc:InvoiceTypeCode>{tip}</cbc:InvoiceTypeCode>
 <cbc:DocumentCurrencyCode>TRY</cbc:DocumentCurrencyCode>{ek}
 <cac:AccountingSupplierParty><cac:Party>
  <cac:PartyIdentification><cbc:ID schemeID="VKN">{satici_vkn}</cbc:ID></cac:PartyIdentification>
  <cac:PartyName><cbc:Name>Örnek Elektrik Malzemeleri A.Ş.</cbc:Name></cac:PartyName>
 </cac:Party></cac:AccountingSupplierParty>
 <cac:AccountingCustomerParty><cac:Party>
  <cac:PartyIdentification><cbc:ID schemeID="VKN">{alici_vkn}</cbc:ID></cac:PartyIdentification>
  <cac:PartyName><cbc:Name>Deneme İnşaat Ltd.</cbc:Name></cac:PartyName>
 </cac:Party></cac:AccountingCustomerParty>
 <cac:PaymentMeans><cbc:PaymentDueDate>2026-10-15</cbc:PaymentDueDate>
  <cac:PayeeFinancialAccount><cbc:ID>TR00 0000 0000 0000 0000 0000 01</cbc:ID></cac:PayeeFinancialAccount></cac:PaymentMeans>
 <cac:AllowanceCharge><cbc:ChargeIndicator>false</cbc:ChargeIndicator><cbc:AllowanceChargeReason>Kampanya</cbc:AllowanceChargeReason>
  <cbc:Amount>750</cbc:Amount></cac:AllowanceCharge>
 <cac:TaxTotal><cbc:TaxAmount>1040</cbc:TaxAmount>{vergi(810, 20)}{vergi(180, 10)}{vergi(0, 0)}{vergi(50, 10, "4080")}</cac:TaxTotal>
 <cac:WithholdingTaxTotal><cbc:TaxAmount>126</cbc:TaxAmount>{vergi(126, 70, "9015")}</cac:WithholdingTaxTotal>
 <cac:LegalMonetaryTotal><cbc:LineExtensionAmount>7500</cbc:LineExtensionAmount>
  <cbc:TaxExclusiveAmount>6750</cbc:TaxExclusiveAmount><cbc:TaxInclusiveAmount>7790</cbc:TaxInclusiveAmount>
  <cbc:AllowanceTotalAmount>750</cbc:AllowanceTotalAmount><cbc:PayableAmount>7664</cbc:PayableAmount></cac:LegalMonetaryTotal>
 {satir(1, "NYY 3x2,5 Kablo", 100, "MTR", 50, 4500, 900, 20,
        isk="<cac:AllowanceCharge><cbc:ChargeIndicator>false</cbc:ChargeIndicator>"
            "<cbc:MultiplierFactorNumeric>0.10</cbc:MultiplierFactorNumeric><cbc:Amount>500</cbc:Amount></cac:AllowanceCharge>",
        ek=vergi(50, 10, "4080"))}
 {satir(2, "Montaj hizmeti", 1, "C62", 2000, 2000, 200, 10)}
 {satir(3, "Yurt dışı teslim (istisna)", 1, "C62", 1000, 1000, 0, 0)}
</Invoice>""".encode()


def test_xml_iskonto_karisik_kdv_tevkifat_ve_diger_vergi():
    r = kural.efatura_xml_oku(ubl())
    assert [(k["kdv_orani"], k["tutar"]) for k in r["kalemler"]] == [
        (20, 4500), (10, 2000), (0, 1000), (20, -450), (10, -200), (0, -100)]
    k = r["kalemler"][0]  # standart sütunlar: birim fiyat liste fiyatı, iskonto ayrı
    assert (k["aciklama"], k["birim_fiyat"], k["iskonto_orani"], k["iskonto_tutari"]) == ("NYY 3x2,5 Kablo", 50, 10, 500)
    assert r["kalemler"][0]["birim"] == "m"
    assert (r["ara_toplam"], r["kdv_toplam"], r["tevkifat"], r["diger_vergiler"], r["genel_toplam"]) == (6750, 990, 126, 50, 7664)
    assert r["satici"]["iban"] == "TR000000000000000000000001" and r["vade"] == "2026-10-15"
    assert r["ettn"].startswith("3f1c0a2e") and r["notlar"] == ""
    assert "İADE" in kural.efatura_xml_oku(ubl(tip="IADE"))["notlar"]


def test_xml_yukleme_toplamlar_uyarisiz_ve_xml_indirilir():
    h = yeni()
    f = client.post("/api/faturalar/oku", files={"dosya": ("a.xml", ubl(), "text/xml")}, headers=h).json()[0]
    assert (f["ara_toplam"], f["kdv_toplam"], f["tevkifat"], f["diger_vergiler"], f["genel_toplam"]) == (6750, 990, 126, 50, 7664)
    assert f["ai_okuma"]["uyarilar"] == []
    assert f["xml_dosya"] and client.get(f"/api/faturalar/{f['id']}/belge?tur=xml", headers=h).content == ubl()


def test_xml_icindeki_xslt_gorunumu_korumali_html_olur():
    pytest.importorskip("lxml")
    h = yeni()
    f = client.post("/api/faturalar/oku", files={"dosya": ("a.xml", ubl(xslt=True), "text/xml")}, headers=h).json()[0]
    r = client.get(f"/api/faturalar/{f['id']}/belge", headers=h)
    assert b"FATURA ORN2026000000001" in r.content and r.headers["content-security-policy"] == "sandbox"


def test_zip_icinde_coklu_xml_ve_ayni_adli_pdf():
    t = io.BytesIO()
    with zipfile.ZipFile(t, "w") as z:
        z.writestr("ORN1.xml", ubl("ORN2026000000001"))
        z.writestr("ORN1.pdf", b"%PDF-1.4 ornek")
        z.writestr("ORN2.xml", ubl("ORN2026000000002"))
        z.writestr("ic.zip", (lambda b: (zipfile.ZipFile(b, "w").writestr("ORN3.xml", ubl("ORN2026000000003")), b)[1])(io.BytesIO()).getvalue())
    h = yeni()
    r = client.post("/api/faturalar/oku", files={"dosya": ("toplu.zip", t.getvalue(), "application/zip")}, headers=h).json()
    assert sorted(f["fatura_no"] for f in r) == ["ORN2026000000001", "ORN2026000000002", "ORN2026000000003"]
    birinci = next(f for f in r if f["fatura_no"].endswith("1"))
    assert birinci["kaynak_dosya"].endswith(".pdf")
    assert client.get(f"/api/faturalar/{birinci['id']}/belge", headers=h).content == b"%PDF-1.4 ornek"
    assert len({f["ai_okuma"]["satici"]["vkn"] for f in r}) == 1  # aynı cari, bir kez açıldı


def test_once_pdf_sonra_xml_ayni_taslak_duzeltilir():
    h = yeni()
    pdf = pdf_olustur(["ORNEK ELEKTRIK MALZEMELERI A.S.", "VKN: 1234567890", "SAYIN DENEME INSAAT", "VKN: 9999999999",
                       "Fatura No: ORN2026000000001", "Fatura Tarihi: 15-09-2026", "Mal Hizmet Toplam Tutari 6.750,00 TL",
                       "Hesaplanan KDV (%20) 990,00 TL", "Odenecek Tutar 7.740,00 TL"])
    p = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    santiye = client.post("/api/santiyeler", json={"ad": "Kütahya Trafo"}, headers=h).json()["id"]
    client.post("/api/faturalar/toplu-guncelle", json={"idler": [p["id"]], "santiye_id": santiye}, headers=h)
    x = client.post("/api/faturalar/oku", files={"dosya": ("f.xml", ubl(), "text/xml")}, headers=h).json()[0]
    assert x["id"] == p["id"] and not x["zaten_var"]
    assert (x["genel_toplam"], x["tevkifat"], len(x["kalemler"])) == (7664, 126, 6)
    assert x["kaynak_dosya"].endswith(".pdf") and x["xml_dosya"] and x["santiye_id"] == santiye
    assert "XML'iyle güncellendi" in x["ai_okuma"]["bilgi"][0]
    taslaklar = client.get("/api/faturalar?yon=gelen&durum=taslak", headers=h).json()
    assert len(taslaklar) == 1
    # Aynı XML ve PDF tekrar gelirse yeni kayıt açılmaz
    assert client.post("/api/faturalar/oku", files={"dosya": ("f.xml", ubl(), "text/xml")}, headers=h).json()[0]["zaten_var"]
    assert client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=h).json()[0]["zaten_var"]


def test_once_xml_sonra_pdf_goruntu_olarak_eklenir():
    h = yeni()
    x = client.post("/api/faturalar/oku", files={"dosya": ("f.xml", ubl(), "text/xml")}, headers=h).json()[0]
    assert x["kaynak_dosya"] == x["xml_dosya"]
    pdf = pdf_olustur(["ORNEK ELEKTRIK MALZEMELERI A.S.", "VKN: 1234567890", "Fatura No: ORN2026000000001",
                       "Fatura Tarihi: 15-09-2026", "Odenecek Tutar 7.664,00 TL"])
    p = client.post("/api/faturalar/oku", files={"dosya": ("f.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    assert p["id"] == x["id"] and p["zaten_var"] and p["kaynak_dosya"].endswith(".pdf")
    assert p["genel_toplam"] == 7664  # XML'deki rakamlar korunur


def test_xml_yonu_vkn_den_bulunur():
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    # Satıcısı biz olan XML "alış" seçiliyken yüklense de satış faturası olur
    f = client.post("/api/faturalar/oku?yon=gelen", files={"dosya": ("s.xml", ubl(satici_vkn="9999999999", alici_vkn="1234567890"),
                                                                   "text/xml")}, headers=h).json()[0]
    assert f["yon"] == "giden" and f["cari_unvan"] == "Deneme İnşaat Ltd."
    f = client.post("/api/faturalar/oku?yon=giden", files={"dosya": ("a.xml", ubl("ORN2026000000009"), "text/xml")}, headers=h).json()[0]
    assert f["yon"] == "gelen"


def test_kendi_vkn_cariye_yazilmaz_ve_kotu_ad_duzelir():
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    # Eski sürüm "Ticaret Sicil No" adıyla cari açmış olsun; aynı VKN'li fatura gelince ad düzelir
    c = client.post("/api/cariler", json={"unvan": "Ticaret Sicil No", "vkn": "1234567890"}, headers=h).json()
    f = client.post("/api/faturalar/oku", files={"dosya": ("a.xml", ubl(), "text/xml")}, headers=h).json()[0]
    assert f["cari_id"] == c["id"] and f["cari_unvan"] == "Örnek Elektrik Malzemeleri A.Ş."
    # Kullanıcının verdiği kısa ad korunur
    client.put(f"/api/cariler/{c['id']}", json={"unvan": "Örnek Elektrik", "vkn": "1234567890"}, headers=h)
    f = client.post("/api/faturalar/oku", files={"dosya": ("b.xml", ubl("ORN2026000000002"), "text/xml")}, headers=h).json()[0]
    assert f["cari_unvan"] == "Örnek Elektrik"
    # Satıcı VKN'si yerine bizim VKN okunmuşsa cariye yazılmaz
    r = client.post("/api/faturalar/oku", files={"dosya": ("c.xml", ubl("ORN2026000000003", satici_vkn="9999999999", alici_vkn="9999999999"),
                                                             "text/xml")}, headers=h).json()[0]
    if r["cari_id"]:
        cari = next(x for x in client.get("/api/cariler", headers=h).json() if x["id"] == r["cari_id"])
        assert cari["vkn"] != "9999999999"
    assert client.get("/api/cariler/kontrol", headers=h).json()["sorunlu"] == []


def test_gider_turu_satici_bazinda_hatirlanir_ve_hatali_zip():
    h = yeni()
    f = client.post("/api/faturalar/oku", files={"dosya": ("a.xml", ubl(), "text/xml")}, headers=h).json()[0]
    client.put(f"/api/faturalar/{f['id']}", json={**{k: f[k] for k in ("yon", "cari_id", "fatura_no", "tarih", "tevkifat",
                                                                         "diger_vergiler", "para_birimi")},
                                                  "gider_turu": "Tesisat malzemesi",
                                                  "kalemler": f["kalemler"]}, headers=h)
    client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    f2 = client.post("/api/faturalar/oku", files={"dosya": ("b.xml", ubl("ORN2026000000002"), "text/xml")}, headers=h).json()[0]
    assert f2["gider_turu"] == "Tesisat malzemesi"
    assert client.get("/api/faturalar/hatali-disa-aktar", headers=h).status_code == 404  # uyarısız taslak dışa aktarılmaz
    pdf = pdf_olustur(["ORNEK A.S.", "VKN: 1111111111", "Fatura No: ORN2026000000077", "Bu faturada kalem tablosu yok",
                       "Mal Hizmet Toplam Tutari 1.000,00 TL", "Hesaplanan KDV (%20) 200,00 TL", "Odenecek Tutar 1.200,00 TL"])
    client.post("/api/faturalar/oku", files={"dosya": ("k.pdf", pdf, "application/pdf")}, headers=h)
    r = client.get("/api/faturalar/hatali-disa-aktar", headers=h)
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert sorted(z.namelist()) == ["ORN2026000000077.pdf", "nedenler.csv"]
    assert "Kalemler tek tek okunamadı" in z.read("nedenler.csv").decode("utf-8-sig")


def test_acilista_eski_yanlis_cari_adi_ve_tek_satir_kalemler_duzelir():
    from app import onarim
    from app.db import SessionLocal
    from app.models import Fatura, FaturaKalem, Kullanici, Sirket
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    c = client.post("/api/cariler", json={"unvan": "ANKARA Ovecler", "vkn": "1234567890"}, headers=h).json()
    pdf = pdf_olustur(["ORNEK GRUP ELEKTRIK", "TEKS.SAN.TIC.LTD.STI.", "Asagi Ovecler Mah. No:8 / Cankaya /", "ANKARA Ovecler",
                       "VKN: 1234567890", "SAYIN", "DENEME INSAAT LTD", "VKN: 9999999999", "Fatura No: NVC2026000000009",
                       "Fatura Tarihi: 03-08-2026", "Sira No Mal Hizmet Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
                       "1 PIGTAIL 15MT 25 Adet 95,000000TL %20,00 475,00TL 2.375,00TL", "Mal Hizmet Toplam Tutari 2.375,00TL",
                       "Hesaplanan KDV(%20) 475,00TL", "Odenecek Tutar 2.850,00TL"])
    f = client.post("/api/faturalar/oku", files={"dosya": ("n.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    assert f["cari_id"] == c["id"] and f["cari_unvan"] == "ANKARA Ovecler"  # elle açılmış ad okumayla değişmez
    db = SessionLocal()
    try:
        # Eski sürüm kalemleri tek satıra toplamış olsun
        fa = db.get(Fatura, f["id"])
        fa.kalemler.clear()
        fa.kalemler.append(FaturaKalem(aciklama="KDV %20 matrahı (kalemler PDF'ten ayrıştırılamadı)", miktar=1, birim_fiyat=2375,
                                       kdv_orani=20, tutar=2375))
        db.commit()
        sirket = db.get(Sirket, db.get(Kullanici, db.query(Kullanici.id).filter(Kullanici.sirket_id == fa.sirket_id).scalar()).sirket_id)
        rapor = onarim.onar(db, sirket)
    finally:
        db.close()
    assert rapor["duzelen_cariler"] == ["ANKARA Ovecler · 1234567890 → ORNEK GRUP ELEKTRIK TEKS.SAN.TIC.LTD.STI. · 1234567890"]
    assert rapor["kalemleri_okunan_faturalar"] == ["NVC2026000000009"]
    f = client.get(f"/api/faturalar/{f['id']}", headers=h).json()
    assert [(k["miktar"], k["birim_fiyat"], k["tutar"]) for k in f["kalemler"]] == [(25, 95, 2375)]


def test_acilista_uydurma_kdvsiz_satirli_faturalar_duzelir():
    """Eski sürüm tutmayan farkı "KDV'siz / istisna tutar" satırıyla kapatıyordu: taslak da onaylı da belgeden yeniden okunur."""
    from app import onarim
    from app.db import SessionLocal
    from app.models import Fatura, FaturaKalem, Sirket
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    idler = []
    for no in ("UYD2026000000001", "UYD2026000000002"):
        pdf = pdf_olustur(["ORNEK ELEKTRIK A.S.", "VKN: 3333333333", "SAYIN", "DENEME INSAAT LTD", "VKN: 9999999999",
                           f"Fatura No: {no}", "Fatura Tarihi: 06-08-2026",
                           "Sira No Mal Hizmet Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
                           "1 KABLO 10 Adet 100,00 TL %20,00 200,00 TL 1.000,00 TL",
                           "Genel Toplam 2.000,00 TL", "Toplam Iskonto 1.000,00 TL", "Hesaplanan KDV(%20) 200,00 TL"])
        idler.append(client.post("/api/faturalar/oku", files={"dosya": (f"{no}.pdf", pdf, "application/pdf")}, headers=h).json()[0]["id"])
    db = SessionLocal()
    try:
        for i in idler:
            fa = db.get(Fatura, i)
            fa.kalemler.clear()
            fa.kalemler.append(FaturaKalem(aciklama="KDV %20 matrahı (kalemler PDF'ten ayrıştırılamadı)", miktar=1,
                                           birim_fiyat=1000, kdv_orani=20, tutar=1000))
            fa.kalemler.append(FaturaKalem(aciklama="KDV'siz / istisna tutar", miktar=1, birim_fiyat=800, kdv_orani=0, tutar=800))
            fa.genel_toplam = 2000
        db.get(Fatura, idler[1]).durum = "onayli"
        db.commit()
        rapor = onarim.onar(db, db.get(Sirket, db.get(Fatura, idler[0]).sirket_id))
    finally:
        db.close()
    assert rapor["kalemleri_okunan_faturalar"] == ["UYD2026000000001"]
    assert rapor["yeniden_okunan_onaylilar"] == ["UYD2026000000002"]
    for i in idler:
        f = client.get(f"/api/faturalar/{i}", headers=h).json()
        assert [(k["aciklama"], k["tutar"]) for k in f["kalemler"]] == [("KABLO", 1000)] and f["genel_toplam"] == 1200
    assert f["durum"] == "onayli"
    assert client.get("/api/cariler/kontrol", headers=h).json()["sorunlu_faturalar"] == []


def test_onayli_faturada_eski_okuma_yenilenir_okunamayan_sade_listelenir():
    """Eski okumada KDV 0 yazılmış (kayıt doğru): yalnız okuma yenilenir, kullanıcının kalem adı kalır. Belgesi olmayan
    geçici satırlı onaylı fatura Genel Durum'da tek kısa cümleyle listelenir."""
    from app import onarim
    from app.db import SessionLocal
    from app.models import Fatura, FaturaKalem, Sirket
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    pdf = pdf_olustur(["ORNEK BETON LTD. STI.", "VKN: 4444444444", "SAYIN", "DENEME INSAAT LTD", "VKN: 9999999999",
                       "Fatura No: HZB2026000000001", "Fatura Tarihi: 06-08-2026",
                       "Sira No Mal Hizmet Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
                       "1 HAZIR BETON 12 m3 3.500,00 TL %20,00 8.400,00 TL 42.000,00 TL",
                       "Mal Hizmet Toplam Tutari 42.000,00 TL", "Hesaplanan KDV(%20) 8.400,00 TL", "Odenecek Tutar 50.400,00 TL"])
    i = client.post("/api/faturalar/oku", files={"dosya": ("b.pdf", pdf, "application/pdf")}, headers=h).json()[0]["id"]
    db = SessionLocal()
    try:
        fa = db.get(Fatura, i)
        fa.durum, fa.kalemler[0].aciklama = "onayli", "HAZIR BETON C25"
        fa.ai_okuma = {**fa.ai_okuma, "kdv_toplam": 0.0}
        elle = Fatura(sirket_id=fa.sirket_id, yon="gelen", durum="onayli", fatura_no="ELL2026000000001", cari_id=fa.cari_id,
                      ara_toplam=1000, kdv_toplam=200, genel_toplam=1200)
        elle.kalemler.append(FaturaKalem(aciklama="KDV %20 matrahı (kalemler PDF'ten ayrıştırılamadı)", miktar=1,
                                         birim_fiyat=1000, kdv_orani=20, tutar=1000))
        db.add(elle)
        db.commit()
        rapor = onarim.onar(db, db.get(Sirket, fa.sirket_id))
    finally:
        db.close()
    assert rapor["yeniden_okunan_onaylilar"] == ["HZB2026000000001"] and rapor["tutari_degisen_onaylilar"] == []
    f = client.get(f"/api/faturalar/{i}", headers=h).json()
    assert [k["aciklama"] for k in f["kalemler"]] == ["HAZIR BETON C25"] and f["kdv_toplam"] == 8400
    k = client.get("/api/cariler/kontrol", headers=h).json()
    assert [(x["fatura_no"], x["sorun"]) for x in k["sorunlu_faturalar"]] == [
        ("ELL2026000000001", "Kalemleri faturadan okunamadı; faturayı açıp kalemleri girin.")]


def test_kamu_kurumu_adi_cari_uyarisi_vermez():
    assert not kural.kotu_unvan("T.C. KARAYOLLARI 14. BURSA BÖLGE MÜDÜRLÜĞÜ")
    assert not kural.kotu_unvan("Bursa Büyükşehir Belediyesi")
    assert kural.kotu_unvan("Atatürk Mah. Gazi Cad. No 12") and kural.kotu_unvan("Ticaret Sicil No")


def test_duzenleyen_basligi_ve_koy_adresi_unvana_karismaz():
    metin = "\n".join(["DÜZENLEYEN", "AHMET YILMAZ", "ÇAMLICA KÖYÜ", "MERKEZ/ KÜTAHYA", "Tel: 05000000000", "Web Sitesi:",
                        "Vergi Dairesi: TAVŞANLI VERGİ DAİRESİ", "TCKN: 11111111110", "SAYIN", "DENEME INSAAT LTD STI",
                        "VKN: 9999999999"])
    satici, alici = kural.pdf_taraflari(metin)
    assert satici == {"unvan": "AHMET YILMAZ", "vkn": "11111111110"} and alici["vkn"] == "9999999999"
    assert kural.kotu_unvan("DÜZENLEYEN AHMET YILMAZ ÇAMLICA KÖYÜ")
    assert not kural.kotu_unvan("ÇAMLICA KÖYÜ TARIM KALKINMA KOOPERATİFİ")


def test_ayni_adli_santiye_ve_cari_ikinci_kez_acilmaz_eskiler_birlesir():
    from app import onarim
    from app.db import SessionLocal
    from app.models import Cari, Fatura, Santiye, Sirket
    h = yeni()
    s1 = client.post("/api/santiyeler", json={"ad": "OFİS"}, headers=h).json()
    s2 = client.post("/api/santiyeler", json={"ad": " ofis"}, headers=h).json()
    assert s1["id"] == s2["id"] and len(client.get("/api/santiyeler", headers=h).json()) == 1
    c1 = client.post("/api/cariler", json={"unvan": "Demir Ltd. Şti."}, headers=h).json()
    c2 = client.post("/api/cariler", json={"unvan": "DEMİR LTD ŞTİ", "vkn": "5555555555"}, headers=h).json()
    assert c1["id"] == c2["id"] and c2["vkn"] == "5555555555"
    baska = client.post("/api/cariler", json={"unvan": "DEMİR LTD ŞTİ", "vkn": "6666666666"}, headers=h).json()
    assert baska["id"] != c1["id"]  # VKN'si farklı: başka firma
    st = client.post("/api/santiyeler", json={"ad": "Depo"}, headers=h).json()
    assert client.put(f"/api/santiyeler/{st['id']}", json={"ad": "Ofis"}, headers=h).status_code == 409
    # Eski sürümde açılmış tekrarlar: açılışta birleşir, fatura ve kalemler tek kayda taşınır
    db = SessionLocal()
    try:
        sirket_id = db.get(Santiye, s1["id"]).sirket_id
        eski = Santiye(sirket_id=sirket_id, ad="Ofis")
        eski_cari = Cari(sirket_id=sirket_id, unvan="Demir Ltd Şti")
        db.add_all([eski, eski_cari])
        db.commit()
        fid = client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": "OFS1", "cari_id": eski_cari.id, "santiye_id": eski.id,
                                                  "kalemler": [{"aciklama": "Kira", "birim_fiyat": 100, "kdv_orani": 20,
                                                                "santiye_id": eski.id}]}, headers=h).json()["id"]
        rapor = onarim.onar(db, db.get(Sirket, sirket_id))
    finally:
        db.close()
    assert rapor["birlesen_santiyeler"] == ["OFİS"] and rapor["birlesen_cariler"] == ["Demir Ltd. Şti."]
    assert [x["ad"] for x in client.get("/api/santiyeler", headers=h).json()] == ["Depo", "OFİS"]
    f = client.get(f"/api/faturalar/{fid}", headers=h).json()
    assert f["santiye_id"] == s1["id"] and f["kalemler"][0]["santiye_id"] == s1["id"] and f["cari_id"] == c1["id"]


def test_xml_satir_tevkifati_oran_ve_kod():
    x = ubl().decode()
    x = x.replace("<cbc:TaxAmount>126</cbc:TaxAmount>", "<cbc:TaxAmount>140</cbc:TaxAmount>")
    x = x.replace("<cac:Item><cbc:Name>Montaj hizmeti", f"<cac:WithholdingTaxTotal><cbc:TaxAmount>140</cbc:TaxAmount>"
                  f"{vergi(140, 70, '603')}</cac:WithholdingTaxTotal><cac:Item><cbc:Name>Montaj hizmeti")
    r = kural.efatura_xml_oku(x.encode())
    assert [(k["tevkifat_orani"], k["tevkifat_kodu"]) for k in r["kalemler"] if k.get("tevkifat_orani")] == [(7, "603")]
    assert r["tevkifat"] == 140 and "farklı" not in r["notlar"]
    # Yalnız belge düzeyinde yazılmışsa (Percent 70) oranı tutan satırlara dağıtılır
    r = kural.efatura_xml_oku(ubl())
    assert {k["aciklama"][:14] for k in r["kalemler"] if k.get("tevkifat_orani") == 7} == {"Montaj hizmeti", "İskonto (%10,0"}


def test_acilista_eski_tevkifatli_faturaya_satir_orani_yazilir():
    from app import onarim
    from app.db import SessionLocal
    from app.models import Fatura, Sirket
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    pdf = pdf_olustur(["ORNEK NAKLIYE LTD. STI.", "VKN: 4444444444", "SAYIN", "DENEME INSAAT LTD", "VKN: 9999999999",
                       "Fatura No: NAK2026000000002", "Fatura Tarihi: 10-08-2026",
                       "Sira No Mal Hizmet Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
                       "1 NAKLIYE 1 Sefer 10.000,00 TL %20,00 2.000,00 TL 10.000,00 TL",
                       "Tevkifat Sebebi: 624-YUK TASIMACILIGI HIZMETI",
                       "Mal Hizmet Toplam Tutari 10.000,00 TL", "Hesaplanan KDV(%20) 2.000,00 TL",
                       "Hesaplanan KDV Tevkifat(%20) 400,00 TL", "Odenecek Tutar 11.600,00 TL"])
    f = client.post("/api/faturalar/oku", files={"dosya": ("n.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    assert [(k["tevkifat_orani"], k["tevkifat_kodu"]) for k in f["kalemler"]] == [(2, "624")] and f["genel_toplam"] == 11600
    db = SessionLocal()
    try:
        fa = db.get(Fatura, f["id"])
        fa.kalemler[0].tevkifat_orani, fa.kalemler[0].tevkifat_kodu = 0, None  # eski sürüm kaydı
        db.commit()
        rapor = onarim.onar(db, db.get(Sirket, fa.sirket_id))
    finally:
        db.close()
    assert rapor["tevkifati_kalemlere_yazilan"] == ["NAK2026000000002"]
    k = client.get(f"/api/faturalar/{f['id']}", headers=h).json()["kalemler"][0]
    assert (k["tevkifat_orani"], k["tevkifat_kodu"]) == (2, "624")


def test_standart_kalem_sutunlari_iskonto_elle_excel_yazdir_ve_eski_kayit():
    """Kalem: birim fiyat liste fiyatı, iskonto oranı/tutarı ayrı, Mal Hizmet Tutarı = miktar x fiyat - iskonto.
    Eski kayıtta açıklamaya yazılmış '(İsk. %60, liste fiyatı 1.600)' notu açılışta sütunlara taşınır, tutar değişmez."""
    from openpyxl import load_workbook
    from app import onarim
    from app.db import SessionLocal
    from app.models import Fatura, FaturaKalem, Sirket
    h = yeni()
    cari = client.post("/api/cariler", json={"unvan": "Ornek Elektrik A.S.", "vkn": "3333333333"}, headers=h).json()
    f = client.post("/api/faturalar", json={"yon": "gelen", "fatura_no": "ISK2026000000001", "tarih": "2026-09-01", "cari_id": cari["id"],
                                            "kalemler": [{"aciklama": "Kablo", "miktar": 10, "birim_fiyat": 100, "iskonto_orani": 25,
                                                          "kdv_orani": 20},
                                                         {"aciklama": "Pano", "miktar": 2, "birim_fiyat": 500, "iskonto_tutari": 100,
                                                          "kdv_orani": 20}]}, headers=h).json()
    assert [(k["iskonto_orani"], k["iskonto_tutari"], k["tutar"]) for k in f["kalemler"]] == [(25, 250, 750), (10, 100, 900)]
    assert (f["ara_toplam"], f["kdv_toplam"], f["genel_toplam"]) == (1650, 330, 1980)
    client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    ws = load_workbook(io.BytesIO(client.get("/api/excel/fatura-kalemleri", headers=h).content)).active
    satirlar = [[c.value for c in r] for r in ws.iter_rows()]
    i = satirlar[0].index("Sıra No")
    assert satirlar[0][i:i + 10] == ["Sıra No", "Mal Hizmet Adı", "Miktar", "Birim", "Birim Fiyat", "İskonto Oranı", "İskonto Tutarı",
                                     "KDV Oranı", "KDV Tutarı", "Mal Hizmet Tutarı"]
    assert satirlar[1][i:i + 10] == [1, "Kablo", 10, None, 100, 25, 250, 20, 150, 750]
    yazdir = client.get(f"/api/yazdir/fatura/{f['id']}", headers=h).text
    assert "İskonto Tutarı" in yazdir and "Mal Hizmet Tutarı" in yazdir
    db = SessionLocal()
    try:
        fa = db.get(Fatura, f["id"])
        fa.kalemler[0].aciklama, fa.kalemler[0].birim_fiyat = "Kablo (NET-1) (İsk. %25, liste fiyatı 100)", 75
        fa.kalemler[0].iskonto_orani = fa.kalemler[0].iskonto_tutari = 0
        db.commit()
        rapor = onarim.onar(db, db.get(Sirket, fa.sirket_id))
        kl = db.get(Fatura, f["id"]).kalemler[0]
        assert (kl.aciklama, float(kl.birim_fiyat), float(kl.iskonto_orani), float(kl.iskonto_tutari), float(kl.tutar)) == \
            ("Kablo (NET-1)", 100, 25, 250, 750)
        assert rapor["iskontosu_sutuna_tasinan_kalem"] == 1
    finally:
        db.close()


def test_okuma_hata_raporu_ozet_zip_ve_okunamayan_dosya():
    """Toplu yüklemeden sonra: sorunlu taslaklar kategoriyle, okunamayan dosyalar nedeniyle tek ZIP'te."""
    from openpyxl import load_workbook
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    iyi = pdf_olustur(["ORNEK ELEKTRIK A.S.", "VKN: 3333333333", "SAYIN", "DENEME INSAAT LTD", "VKN: 9999999999",
                       "Fatura No: IYI2026000000001", "Fatura Tarihi: 06-08-2026",
                       "Sira No Mal Hizmet Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
                       "1 KABLO 10 Adet 100,00 TL %20,00 200,00 TL 1.000,00 TL",
                       "Mal Hizmet Toplam Tutari 1.000,00 TL", "Hesaplanan KDV(%20) 200,00 TL", "Odenecek Tutar 1.200,00 TL"])
    kotu = pdf_olustur(["ORNEK YAPI A.S.", "VKN: 4444444444", "SAYIN", "DENEME INSAAT LTD", "VKN: 9999999999",
                        "Fatura No: KOT2026000000001", "Fatura Tarihi: 07-08-2026",
                        "Sira No Mal Hizmet Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
                        "1 KUM 10 Adet 100,00 TL %20,00 200,00 TL 1.000,00 TL",
                        "Mal Hizmet Toplam Tutari 1.000,00 TL", "Hesaplanan KDV(%20) 260,00 TL", "Odenecek Tutar 1.260,00 TL"])
    for ad, pdf in (("iyi.pdf", iyi), ("kotu.pdf", kotu)):
        assert client.post("/api/faturalar/oku?yontem=kural", files={"dosya": (ad, pdf, "application/pdf")}, headers=h).status_code == 200
    r = client.post("/api/faturalar/oku?yontem=kural", files={"dosya": ("bos.pdf", b"%PDF-1.4 bozuk", "application/pdf")}, headers=h)
    assert r.status_code == 422
    o = client.get("/api/faturalar/okuma-raporu/ozet", headers=h).json()
    assert (o["toplam"], o["sorunsuz"], o["hatali"], o["okunamayan_dosya"]) == (2, 1, 1, 1)
    assert "kdv" in [x["kategori"] for x in o["kategoriler"]]
    z = zipfile.ZipFile(io.BytesIO(client.get("/api/faturalar/okuma-raporu", headers=h).content))
    adlar = z.namelist()
    assert any(a.endswith("KOT2026000000001/belge.pdf") for a in adlar) and not any("IYI2026" in a for a in adlar)
    assert any(a.endswith("KOT2026000000001/metin.txt") for a in adlar) and any(a.endswith("/tablolar.json") for a in adlar)
    assert any(a.startswith("okunamayan/") and a.endswith("belge.pdf") for a in adlar)
    wb = load_workbook(io.BytesIO(z.read("okuma-hata-raporu.xlsx")))
    hatalar = [[c.value for c in r] for r in wb["Hatalı Faturalar"].iter_rows()]
    kdv = next(x for x in hatalar if x[7] == "KDV farkı")
    assert (kdv[3], kdv[9], kdv[10]) == ("KOT2026000000001", 260, 200)
    assert [x[1] for x in wb["Okunamayan Dosyalar"].iter_rows(min_row=2, values_only=True)] == ["bos.pdf"]
    assert client.delete("/api/faturalar/okuma-raporu/okunamayanlar", headers=h).json()["silinen"] == 1
    assert client.get("/api/faturalar/okuma-raporu/ozet", headers=h).json()["okunamayan_dosya"] == 0


def test_acilista_rapordaki_sorunlu_taslaklar_yeniden_okunur():
    """Eski sürüm kalemleri okuyamamış ve satıcı yerine alıcının bölünmüş adıyla cari açmışsa: açılışta taslak yeniden
    okunur, doğru cariye bağlanır, artık kullanılmayan yanlış otomatik cari silinir."""
    from app import onarim
    from app.db import SessionLocal
    from app.models import Cari, Fatura, Sirket
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Başka Firma Ltd.", "vkn": "9999999999"}, headers=h)
    pdf = pdf_olustur(["ORNEK ELEKTRIK A.S.", "VKN: 3333333333", "SAYIN", "DENEME INSAAT LTD", "VKN: 9999999999",
                       "Fatura No: TKR2026000000001", "Fatura Tarihi: 06-08-2026",
                       "Sira No Mal Hizmet Miktar Birim Fiyat KDV Orani KDV Tutari Mal Hizmet Tutari",
                       "1 KABLO 10 Adet 100,00 TL %20,00 200,00 TL 1.000,00 TL",
                       "Mal Hizmet Toplam Tutari 1.000,00 TL", "Hesaplanan KDV(%20) 200,00 TL", "Odenecek Tutar 1.200,00 TL"])
    f = client.post("/api/faturalar/oku", files={"dosya": ("t.pdf", pdf, "application/pdf")}, headers=h).json()[0]
    db = SessionLocal()
    try:
        fa = db.get(Fatura, f["id"])
        dogru = fa.cari
        yanlis = Cari(sirket_id=fa.sirket_id, tur="tedarikci", unvan="INSAAT LIMITED SIRKET I", ad_otomatik=True)
        db.add(yanlis)
        db.flush()
        fa.cari, fa.ara_toplam, fa.kdv_toplam, fa.genel_toplam = yanlis, 0, 0, 0
        fa.kalemler.clear()
        fa.ai_okuma = {**fa.ai_okuma, "kalemler_okunamadi": True}
        db.commit()
        yanlis_id = yanlis.id
        rapor = onarim.onar(db, db.get(Sirket, fa.sirket_id))
        fa = db.get(Fatura, f["id"])
        assert rapor["yeniden_okunan_taslaklar"] == ["TKR2026000000001"]
        assert fa.cari_id == dogru.id and float(fa.genel_toplam) == 1200 and [k.aciklama for k in fa.kalemler] == ["KABLO"]
        assert db.get(Cari, yanlis_id) is None
    finally:
        db.close()


def test_acilista_kendi_vkn_ogrenilir_kur_yazilir_sus_temizlenir():
    """Firma bilgisinde VKN'miz yoksa faturalardan öğrenilir; kendi VKN'mizle açılmış cari düzelir; dövizli taslağın
    kuru okunan bilgiden alana yazılır; kalem adlarındaki "## ... ##" süsü temizlenir."""
    from app import onarim
    from app.db import SessionLocal
    from sqlalchemy import select

    from app.models import Cari, Fatura, FaturaKalem, Kullanici, Sirket
    h = yeni()
    BEN = client.get("/api/ben", headers=h).json()
    client.put("/api/sirket", json={"unvan": "Deneme İnşaat Ltd.", "vkn": ""}, headers=h)
    db = SessionLocal()
    try:
        sid = db.scalar(select(Kullanici.sirket_id).where(Kullanici.eposta == BEN["eposta"]))
        sirket = db.get(Sirket, sid)
        sirket.vkn = None
        kendi = Cari(sirket_id=sid, tur="tedarikci", unvan="DENEME INSAAT LTD", vkn="9999999999", ad_otomatik=True)
        db.add(kendi)
        db.flush()
        for i in range(3):
            db.add(Fatura(sirket_id=sid, yon="gelen", fatura_no=f"ABC202600000000{i}", ara_toplam=100, genel_toplam=100,
                          ai_okuma={"alici": {"vkn": "9999999999"}, "karekod": {"avkntckn": "9999999999"}}))
        doviz = Fatura(sirket_id=sid, yon="gelen", fatura_no="ORG2026000000001", para_birimi="USD", kur=None, ara_toplam=10,
                       genel_toplam=12, ai_okuma={"kur": 45.1191, "alici": {"vkn": "9999999999"},
                                                  "uyarilar": ["Fatura USD cinsinden; kuru girin."]})
        doviz.kalemler.append(FaturaKalem(aciklama="## RACK KABINET ## (OEM-110)", miktar=1, birim_fiyat=10, kdv_orani=20, tutar=10))
        db.add(doviz)
        db.commit()
        rapor = onarim.onar(db, sirket, 12)
        assert rapor["ogrenilen_kendi_vkn"] == "9999999999" and db.get(Sirket, sid).vkn == "9999999999"
        assert rapor["kuru_yazilan_faturalar"] == ["ORG2026000000001"] and float(db.get(Fatura, doviz.id).kur) == 45.1191
        assert db.get(Fatura, doviz.id).kalemler[0].aciklama == "RACK KABINET (OEM-110)"
        assert db.get(Cari, kendi.id) is None  # kendi firmamız adına açılmış, hiçbir kayda bağlı olmayan cari silindi
    finally:
        db.close()
