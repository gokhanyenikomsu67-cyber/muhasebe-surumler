"""Gelir/kurumlar vergisi stopajı: okunur, kaleme oranıyla yazılır, ödenecekten düşülür (sentetik veri)."""
from app import kural
from tests.test_kural import _UST, tablolu_pdf

_BASLIK = ["Sira No", "Mal Hizmet", "Miktar", "Birim Fiyat", "Iskonto Orani", "Iskonto Tutari", "KDV Orani", "KDV Tutari",
           "Diger Vergiler", "Mal Hizmet Tutari"]


def stopajli_pdf(alt_stopaj=True):
    """e-Dönüşüm düzeni: Diğer Vergiler sütununda stopaj ve KDV tevkifatı, altta 'Hesaplanan STPJ(%5)'."""
    return tablolu_pdf(_UST, _BASLIK, [[
        "1", "KONUT INSAATI ELEKTRIK\nHAKEDIS BEDELI", "1 Adet", "100.000,00 TL", "", "", "%20,00", "20.000,00 TL",
        "GELIR VERGISI STOPAJI\n(%5,00)=5.000,00 TL KDV\nTEVKIFAT (%40,00)=8.000,00 TL", "100.000,00 TL"]],
        ["Mal Hizmet Toplam Tutari 100.000,00 TL", "Toplam Iskonto 0,00 TL"]
        + (["Hesaplanan STPJ(%5) 5.000,00 TL"] if alt_stopaj else [])
        + ["Hesaplanan KDV GERCEK(%20) 20.000,00 TL", "Hesaplanan KDV Tevkifat(%40) 8.000,00 TL",
           "Vergiler Dahil Toplam Tutar 120.000,00 TL", "Odenecek Tutar 107.000,00 TL"])


def test_pdf_stopaj_okunur_ve_odenecekten_duser(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    for alt in (True, False):  # stopaj yalnız kalem satırında yazsa da okunur
        r = kural.pdf_fatura_oku(stopajli_pdf(alt), [], "DENEME INSAAT LTD. STI.", "9999999999", "gelen")
        assert (r["ara_toplam"], r["kdv_toplam"], r["tevkifat"], r["stopaj"], r["diger_vergiler"], r["genel_toplam"]) == \
            (100000, 20000, 8000, 5000, 0, 107000), alt
        k = r["kalemler"][0]
        assert (k["tutar"], k.get("stopaj_orani"), k.get("tevkifat_orani")) == (100000, 5, 4)
        assert "farklı" not in r["notlar"]


def test_xml_stopaj_ayri_okunur_diger_vergiye_katilmaz():
    from tests.test_efatura import ubl
    x = ubl().decode()
    # ÖİV (4080) yerine gelir vergisi stopajı (0003, %5): ödenecek 7664 - 50 (ÖİV yok) - 337,50 stopaj
    x = x.replace('<cbc:TaxTypeCode>4080</cbc:TaxTypeCode>', '<cbc:TaxTypeCode>0003</cbc:TaxTypeCode>')
    x = x.replace("<cbc:TaxAmount>50</cbc:TaxAmount><cbc:Percent>10</cbc:Percent>",
                  "<cbc:TaxAmount>337.5</cbc:TaxAmount><cbc:Percent>5</cbc:Percent>")
    x = x.replace("<cbc:PayableAmount>7664</cbc:PayableAmount>", "<cbc:PayableAmount>7276.5</cbc:PayableAmount>")
    r = kural.efatura_xml_oku(x.encode())
    assert (r["stopaj"], r["diger_vergiler"], r["genel_toplam"]) == (337.5, 0, 7276.5)
    assert r["kalemler"][0]["stopaj_orani"] == 5


def test_stopajli_fatura_kaydi_cari_rapor_ve_wolvox(monkeypatch):
    import io
    from openpyxl import load_workbook
    from tests.test_api import client
    from tests.test_kural import yeni
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    h = yeni()
    f = client.post("/api/faturalar/oku", files={"dosya": ("s.pdf", stopajli_pdf(), "application/pdf")}, headers=h).json()[0]
    assert (f["stopaj"], f["tevkifat"], f["genel_toplam"]) == (5000, 8000, 107000)
    assert f["kalemler"][0]["stopaj_orani"] == 5
    assert not [u for u in f["ai_okuma"]["uyarilar"] if "tutmuyor" in u or "farklı" in u]
    # Elle düzeltme: oran %10 olunca stopaj ve ödenecek yeniden hesaplanır
    kalem = {**f["kalemler"][0], "stopaj_orani": 10}
    g = client.put(f"/api/faturalar/{f['id']}", json={**f, "kalemler": [kalem]}, headers=h).json()
    assert (g["stopaj"], g["genel_toplam"]) == (10000, 102000)
    client.put(f"/api/faturalar/{f['id']}", json={**f, "kalemler": [f["kalemler"][0]]}, headers=h)
    assert client.post(f"/api/faturalar/{f['id']}/onayla", headers=h).status_code == 200
    r = client.get("/api/raporlar/stopaj?yil=2026&ay=8", headers=h).json()
    assert r["alis_stopaj"] == 5000 and r["alis"][0]["oran"] == "%5" and r["faturalar"][0]["odenecek"] == 107000
    assert client.get("/api/excel/stopaj?yil=2026&ay=8", headers=h).status_code == 200
    ws = load_workbook(io.BytesIO(client.get("/api/excel/fatura-kalemleri", headers=h).content)).active
    basliklar = [c.value for c in ws[1]]
    satir = dict(zip(basliklar, [c.value for c in ws[2]]))
    assert (satir["Stopaj Oranı"], satir["Stopaj Tutarı"]) == ("%5", 5000)
