"""Mali raporlar, yazdırma sayfaları, firma bilgileri ve kullanıcılar (sentetik veri)."""
import io
import os
import tempfile
import uuid
from datetime import date, timedelta

os.environ.setdefault("MUHASEBE_VERI", tempfile.mkdtemp())

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)


def yeni():
    r = client.post("/api/kayit", json={"sirket_unvan": "Deneme İnşaat", "ad": "Ali",
                                        "eposta": f"{uuid.uuid4().hex}@ornek.com", "sifre": "gizli123"})
    return {"Authorization": f"Bearer {r.json()['token']}"}


def cari(h, unvan, vkn):
    return client.post("/api/cariler", json={"unvan": unvan, "vkn": vkn}, headers=h).json()["id"]


def fatura(h, yon, cari_id, tarih, tutar, kdv=20, tevkifat=0, no=None, vade=None, gider=None):
    f = client.post("/api/faturalar", json={
        "yon": yon, "cari_id": cari_id, "tarih": tarih, "vade": vade, "fatura_no": no, "tevkifat": tevkifat,
        "gider_turu": gider, "kalemler": [{"aciklama": "Kalem", "miktar": 1, "birim_fiyat": tutar, "kdv_orani": kdv}]},
        headers=h).json()
    r = client.post(f"/api/faturalar/{f['id']}/onayla", headers=h)
    assert r.status_code == 200, r.text
    return f


def test_kdv_ozeti_devir_ve_tevkifat():
    h = yeni()
    m, t = cari(h, "Müşteri A.Ş.", "1111111111"), cari(h, "Tedarikçi Ltd.", "2222222222")
    fatura(h, "gelen", t, "2026-01-10", 10000)                  # indirilecek 2.000
    fatura(h, "giden", m, "2026-01-20", 5000)                   # hesaplanan 1.000 -> 1.000 devreder
    fatura(h, "giden", m, "2026-02-05", 20000, tevkifat=1600)   # hesaplanan 4.000, alıcı 1.600'ünü öder
    fatura(h, "gelen", t, "2026-02-07", 1000, tevkifat=40)      # indirilecek 200, KDV-2 40
    v = client.get("/api/raporlar/kdv?yil=2026", headers=h).json()["aylar"]
    assert (v[0]["hesaplanan"], v[0]["indirilecek"], v[0]["odenecek"], v[0]["devreden"]) == (1000, 2000, 0, 1000)
    assert (v[1]["onceki_devreden"], v[1]["satis_tevkifat"], v[1]["odenecek"], v[1]["kdv2"]) == (1000, 1600, 1200, 40)
    assert v[1]["oranlar"] == [{"oran": 20, "matrah": 20000, "kdv": 4000}]
    ws = load_workbook(io.BytesIO(client.get("/api/excel/kdv?yil=2026", headers=h).content)).active
    assert ws["A3"].value == "Şubat" and ws["H3"].value == 1200


def test_ba_bs_limit_ve_toplama():
    h = yeni()
    a, b = cari(h, "Büyük Tedarik A.Ş.", "3333333333"), cari(h, "Küçük Esnaf", "4444444444")
    fatura(h, "gelen", a, "2026-03-01", 3000)
    fatura(h, "gelen", a, "2026-03-15", 2500)  # ay içinde toplam 5.500 -> Ba'ya girer
    fatura(h, "gelen", b, "2026-03-20", 4999)  # limit altı
    fatura(h, "gelen", a, "2026-04-01", 9000)  # başka ay
    fatura(h, "giden", b, "2026-03-25", 7000)
    v = client.get("/api/raporlar/ba-bs?yil=2026&ay=3", headers=h).json()
    assert [(x["unvan"], x["belge_sayisi"], x["tutar"]) for x in v["ba"]] == [("Büyük Tedarik A.Ş.", 2, 5500)]
    assert [(x["vkn"], x["tutar"]) for x in v["bs"]] == [("4444444444", 7000)]
    assert client.get("/api/excel/ba-bs?yil=2026&ay=3", headers=h).status_code == 200


def test_gelir_gider_tablosu():
    h = yeni()
    m, t = cari(h, "Müşteri", "5555555555"), cari(h, "Beton A.Ş.", "6666666666")
    fatura(h, "giden", m, "2026-05-02", 50000)
    fatura(h, "gelen", t, "2026-05-03", 30000, gider="Beton")
    client.post("/api/hareketler", json={"tarih": "2026-05-04", "yon": "cikis", "tur": "masraf", "tutar": 750,
                                         "gider_turu": "Yemek"}, headers=h)
    v = client.get("/api/raporlar/gelir-gider?yil=2026", headers=h).json()
    assert v["gelir"][4] == 50000
    assert {g["tur"]: g["aylar"][4] for g in v["gider"]} == {"Beton": 30000, "Yemek": 750}
    assert v["net"][4] == 19250
    assert client.get("/api/excel/gelir-gider?yil=2026", headers=h).status_code == 200


def test_yaslandirma_odemeler_en_eski_faturayi_kapatir():
    h = yeni()
    m = cari(h, "Geciken Müşteri", "7777777777")
    bugun = date.today()
    g = lambda n: (bugun - timedelta(days=n)).isoformat()  # noqa: E731
    fatura(h, "giden", m, g(120), 1000, kdv=0, vade=g(100))   # 100 gün geçmiş
    fatura(h, "giden", m, g(50), 2000, kdv=0, vade=g(45))     # 45 gün geçmiş
    fatura(h, "giden", m, g(5), 3000, kdv=0, vade=g(-25))     # vadesi gelmemiş
    client.post("/api/hareketler", json={"tarih": g(10), "yon": "giris", "tur": "tahsilat", "tutar": 1500, "cari_id": m},
                headers=h)  # en eski faturayı (1000) ve ikincinin 500'ünü kapatır
    v = client.get("/api/raporlar/yaslandirma", headers=h).json()
    r = v["alacak"][0]
    assert (r["bakiye"], r["vadesi_gelmemis"], r["g31_60"], r["g90"]) == (4500, 3000, 1500, 0)
    assert [x["acik_tutar"] for x in v["acik_faturalar"]] == [1500, 3000]
    assert client.get("/api/excel/yaslandirma", headers=h).status_code == 200


def test_firma_bilgileri_seri_ve_sonraki_fatura_no():
    h = yeni()
    assert client.put("/api/sirket", json={"unvan": "Deneme İnşaat", "fatura_seri": "abcd"}, headers=h).status_code == 422
    r = client.put("/api/sirket", json={"unvan": "Deneme İnşaat Ltd.", "vkn": "9999999999", "fatura_seri": "dnm",
                                        "iban": "TR000000000000000000000000"}, headers=h).json()
    assert r["fatura_seri"] == "DNM"
    yil = date.today().year
    assert client.get("/api/faturalar/sonraki-no", headers=h).json()["fatura_no"] == f"DNM{yil}000000001"
    m = cari(h, "Müşteri", "8888888888")
    fatura(h, "giden", m, date.today().isoformat(), 100, no=f"DNM{yil}000000007")
    assert client.get("/api/faturalar/sonraki-no", headers=h).json()["fatura_no"] == f"DNM{yil}000000008"


def test_yazdirma_sayfalari():
    h = yeni()
    client.put("/api/sirket", json={"unvan": "Deneme <İnşaat>", "adres": "Kütahya"}, headers=h)
    m = cari(h, "Müşteri & Ortak", "1212121212")
    f = fatura(h, "giden", m, "2026-06-01", 1000, no="DNM2026000000001")
    client.post("/api/hareketler", json={"tarih": "2026-06-10", "yon": "giris", "tur": "tahsilat", "tutar": 200, "cari_id": m},
                headers=h)
    r = client.get(f"/api/yazdir/fatura/{f['id']}", headers=h)
    assert r.status_code == 200 and "SATIŞ FATURASI" in r.text and "1.200,00" in r.text
    assert "Deneme &lt;İnşaat&gt;" in r.text and "Müşteri &amp; Ortak" in r.text  # HTML kaçışı
    e = client.get(f"/api/yazdir/cari-ekstre/{m}?baslangic=2026-06-05", headers=h).text
    assert "Devreden bakiye" in e and "1.000,00" in e  # devir 1.200, sonra 200 tahsilat
    mt = client.get(f"/api/yazdir/mutabakat/{m}", headers=h).text
    assert "1.000,00 TL borçlu" in mt
    assert client.get(f"/api/yazdir/fatura/{f['id']}", headers=yeni()).status_code == 404  # başka şirket göremez


def test_kullanici_yetkileri():
    h = yeni()
    r = client.post("/api/kullanicilar", json={"ad": "Ayşe", "eposta": f"{uuid.uuid4().hex}@ornek.com", "sifre": "123456",
                                               "rol": "izleyici"}, headers=h)
    assert r.status_code == 200, r.text
    liste = client.get("/api/kullanicilar", headers=h).json()["kullanicilar"]
    izleyici = next(x for x in liste if not x["ben"])
    t = client.post("/api/giris", json={"eposta": izleyici["eposta"], "sifre": "123456"}).json()["token"]
    hi = {"Authorization": f"Bearer {t}"}
    assert client.get("/api/cariler", headers=hi).status_code == 200
    assert client.post("/api/cariler", json={"unvan": "X"}, headers=hi).status_code == 403
    assert client.post("/api/kullanicilar", json={"ad": "Y", "eposta": "y@y.y", "sifre": "123456"}, headers=hi).status_code == 403
    assert client.put(f"/api/kullanicilar/{izleyici['id']}", json={"rol": "muhasebeci"}, headers=h).status_code == 200
    assert client.post("/api/cariler", json={"unvan": "X"}, headers=hi).status_code == 200
    ben = next(x for x in liste if x["ben"])
    assert client.delete(f"/api/kullanicilar/{ben['id']}", headers=h).status_code == 422
    assert client.delete(f"/api/kullanicilar/{izleyici['id']}", headers=h).status_code == 200
    assert client.get("/api/cariler", headers=hi).status_code == 401


def test_cek_bordrosu():
    h = yeni()
    m, t = cari(h, "Çek Veren Müşteri", "1313131313"), cari(h, "Tedarikçi", "1414141414")
    for no, tutar in (("A-1", 4000), ("A-2", 6000)):
        assert client.post("/api/cekler", json={"yon": "alinan", "tutar": tutar, "vade": "2026-12-01", "cari_id": m,
                                                "cek_no": no}, headers=h).status_code == 200
    client.post("/api/cekler", json={"yon": "verilen", "tutar": 2500, "vade": "2026-12-05", "cari_id": t, "cek_no": "V-1"}, headers=h)
    r = client.get("/api/yazdir/cek-bordro?yon=alinan&durum=portfoyde", headers=h).text
    assert "A-1" in r and "A-2" in r and "V-1" not in r and "10.000,00" in r and "Toplam (2 adet)" in r
    r = client.get(f"/api/yazdir/cek-bordro?cari_id={t}", headers=h).text
    assert "V-1" in r and "A-1" not in r


def test_mali_raporlar_beyan_degil_notu_ve_mevzuat_uyarilari():
    h = yeni()
    t = cari(h, "Nakliyeci Ltd.", "5555555551")
    for no, tutar in (("K1", 9000), ("K2", 11000)):  # KDV dahil 10.800 ve 13.200; 2026 sınırı 12.000
        f = client.post("/api/faturalar", json={"yon": "gelen", "cari_id": t, "tarih": "2026-05-10", "fatura_no": no,
                                                "kalemler": [{"aciklama": "Yük taşıma", "miktar": 1, "birim_fiyat": tutar,
                                                              "kdv_orani": 20, "tevkifat_orani": 2, "tevkifat_kodu": "624"}]},
                        headers=h).json()
        assert client.post(f"/api/faturalar/{f['id']}/onayla", headers=h).status_code == 200
    v = client.get("/api/raporlar/tevkifat?yil=2026&ay=5", headers=h).json()
    uyari = {x["fatura_no"]: x["uyari"] for x in v["faturalar"]}
    assert "12.000 TL" in uyari["K1"] and uyari["K2"] is None
    # Ba-Bs Eylül 2024'ten beri verilmiyor
    assert client.get("/api/raporlar/ba-bs?yil=2026&ay=5", headers=h).json()["kaldirildi"] is True
    assert client.get("/api/raporlar/ba-bs?yil=2024&ay=8", headers=h).json()["kaldirildi"] is False
    for yol in ("kdv?yil=2026", "tevkifat?yil=2026&ay=5", "ba-bs?yil=2026&ay=5", "gelir-gider?yil=2026"):
        ws = load_workbook(io.BytesIO(client.get(f"/api/excel/{yol}", headers=h).content)).active
        assert "resmi beyan yerine geçmez" in str(ws.cell(ws.max_row, 1).value), yol
