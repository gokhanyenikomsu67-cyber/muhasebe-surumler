"""Toplu fatura yüklemesinden sonra okuma hatalarının teşhisi ve "Okuma Hata Raporu".

Her taslak fatura aynı kurallarla denetlenir; her sorun bir kategori ve sade Türkçe nedenle yazılır, faturada
yazan ve programın bulduğu değerler yan yana konur. Rapor (ZIP) belgeleri, PDF metnini ve tablo dökümünü de
içerir: bir kez gönderildiğinde hangi firmanın düzeninde neyin neden yanlış okunduğu programda düzeltilebilir.
Okunamayan dosyalar (fatura açılmadı) okuma-hatalari klasörüne kaydedilir.
"""
import io
import json
import re
import uuid
import zipfile
from collections import Counter
from datetime import datetime

from . import kural
from .db import VERI_KLASORU

HATA_KLASORU = VERI_KLASORU / "okuma-hatalari"
KAYIT = HATA_KLASORU / "kayit.jsonl"

KATEGORILER = {
    "dosya": "Dosya okunamadı",
    "kalem_okunamadi": "Kalemler okunamadı",
    "kalem_adi": "Mal/hizmet adı eksik",
    "kalem_tutari": "Miktar x fiyat - iskonto tutarı tutmuyor",
    "kdv": "KDV farkı",
    "toplam": "Ödenecek tutar farkı",
    "tevkifat": "Tevkifat farkı",
    "cari": "Firma (cari) sorunu",
    "fatura_no": "Fatura no eksik",
    "tarih": "Tarih eksik",
    "mukerrer": "Mükerrer fatura",
    "resim": "Resimden okuma belirsiz",
    "kur": "Dövizli faturada kur yok",
    "diger": "Diğer uyarı",
}


def _p(x) -> float:
    return round(float(x or 0), 2)


def bilgi_notu_mu(u: str, f) -> bool:
    """Sorun değil, bilgi olan okuma notları: döviz cinsi (kur girilmişse), 1 TL'ye kadar yuvarlama farkı, fatura
    tutarıyla kayıt, karekodla doğrulama. 1 TL'yi aşan yuvarlama farkı uyarı kalır."""
    if re.search(r"fatura tutarıyla kaydedildi|karekodundan doğrulandı", u):
        return True
    if re.match(r"Fatura [A-Z]{3} cinsinden", u) and "kuru girin" not in u:
        return bool(float(getattr(f, "kur", None) or 0))
    m = re.search(r"Faturada ([\d.,]+) TL yuvarlama farkı var", u)
    if m:
        return (kural.para_coz(m.group(1)) or 0) <= 1
    return False


def fatura_teshis(f, sirket) -> list[dict]:
    """Bir faturanın sorunları: [{kategori, neden, faturada, programda}]. Boş liste = sorunsuz."""
    from .onarim import cari_sorunu
    o = f.ai_okuma or {}
    s = []

    def ekle(kategori, neden, faturada=None, programda=None):
        s.append({"kategori": kategori, "neden": neden, "faturada": faturada, "programda": programda})

    if not f.kalemler or o.get("kalemler_okunamadi") or any(
            re.search(r"ayrıştırılamadı|matrahı$|KDV'siz / istisna|^Hizmet bedeli \(KDV %|^Diğer fatura kalemleri \(taksit",
                      k.aciklama or "") for k in f.kalemler):
        ekle("kalem_okunamadi", "Kalem tablosu okunamadı; kalemler tek satıra toplandı ya da hiç yok.",
             o.get("ara_toplam"), f"{len(f.kalemler)} satır")
    for i, k in enumerate(f.kalemler, 1):
        ad = (k.aciklama or "").strip()
        if len(re.sub(r"[\W\d_]", "", ad)) < 2:
            ekle("kalem_adi", f"{i}. kalemin mal/hizmet adı okunamadı.", None, ad or "(boş)")
        brut = round(float(k.miktar or 0) * float(k.birim_fiyat or 0), 2)
        if abs(brut - float(k.iskonto_tutari or 0) - float(k.tutar or 0)) > max(0.05, 0.002 * abs(float(k.tutar or 0))):
            ekle("kalem_tutari", f"{i}. kalem ({ad[:40]}): {float(k.miktar):g} x {float(k.birim_fiyat):g} - iskonto "
                                 f"{float(k.iskonto_tutari or 0):g} = {brut - float(k.iskonto_tutari or 0):.2f}, tutar {float(k.tutar):.2f}.",
                 None, _p(k.tutar))
    if o.get("ara_toplam") is not None and f.kalemler and abs(_p(f.ara_toplam) + _p(f.yuvarlama) - o["ara_toplam"]) > 0.05:
        ekle("toplam", "Kalemlerin toplamı faturadaki mal/hizmet toplamını tutmuyor.", o["ara_toplam"], _p(f.ara_toplam))
    if o.get("kdv_toplam") is not None and abs(_p(f.kdv_toplam) - o["kdv_toplam"]) > 0.05:
        ekle("kdv", "Kalemlerden hesaplanan KDV faturadaki KDV'yi tutmuyor.", o["kdv_toplam"], _p(f.kdv_toplam))
    if o.get("genel_toplam") is not None and abs(_p(f.genel_toplam) - o["genel_toplam"]) > 0.05:
        ekle("toplam", "Hesaplanan ödenecek tutar faturadakini tutmuyor.", o["genel_toplam"], _p(f.genel_toplam))
    if not float(f.genel_toplam or 0):
        ekle("toplam", "Tutar okunamadı.", o.get("genel_toplam"), 0)
    if o.get("tevkifat") is not None and abs(_p(f.tevkifat) - float(o["tevkifat"] or 0)) > 0.05:
        ekle("tevkifat", "Kalemlerin tevkifatı faturadaki tevkifatı tutmuyor.", o["tevkifat"], _p(f.tevkifat))
    karsi = (o.get("satici" if f.yon == "gelen" else "alici") or {})
    if not f.cari_id:
        ekle("cari", "Firma okunamadı.", karsi.get("unvan"), None)
    elif f.cari and (neden := cari_sorunu(f.cari, sirket)):
        ekle("cari", f"Cari: {neden}.", karsi.get("unvan"), f.cari.unvan)
    if (f.para_birimi or "TRY") not in ("TRY", "TL") and not float(getattr(f, "kur", None) or 0):
        ekle("kur", f"Fatura {f.para_birimi} cinsinden; kur girilmemiş.", None, f.para_birimi)
    if not f.fatura_no:
        ekle("fatura_no", "Fatura numarası okunamadı.")
    if not f.tarih:
        ekle("tarih", "Fatura tarihi okunamadı.", o.get("tarih"))
    bilinen = ("Hesaplanan KDV (", "Hesaplanan toplam (", "Kalemlerin toplamı (", "Ara toplam + KDV", "Faturadaki firma okunamadı",
               "Kalemler ve tutarlar okunamadı", "Firmanın adı faturada yazmıyor", "Faturanın PDF'inde okunacak yazı yok")
    for u in o.get("uyarilar") or []:
        if not u or re.search(r"okundu\.?$|yeni cari açıldı", u) or u.startswith(bilinen) or bilgi_notu_mu(u, f):
            continue
        ekle("mukerrer" if "daha önce kaydedilmiş" in u else "resim" if u.startswith("Resim") else "tevkifat" if "evkifat" in u else
             "kalem_okunamadi" if "ayrıştırılamadı" in u else "diger", u)
    # Aynı kategori ve nedende bir kez
    tekil = {}
    for x in s:
        tekil.setdefault((x["kategori"], x["neden"]), x)
    return list(tekil.values())


def _tl(x) -> str:
    return f"{float(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# Genel Durum'daki onaylı fatura listesi için kullanıcının anlayacağı tek kısa cümle
_SADE = {
    "kalem_okunamadi": lambda x: "Kalemleri faturadan okunamadı; faturayı açıp kalemleri girin.",
    "kalem_adi": lambda x: "Bir kalemin mal/hizmet adı boş.",
    "kalem_tutari": lambda x: "Bir kalemde miktar x fiyat, tutarı tutmuyor.",
    "kdv": lambda x: f"KDV faturadakinden farklı (faturada {_tl(x['faturada'])}, programda {_tl(x['programda'])})."
    if x["faturada"] is not None else "KDV faturadakinden farklı.",
    "toplam": lambda x: f"Toplam faturadakinden farklı (faturada {_tl(x['faturada'])}, programda {_tl(x['programda'])})."
    if x["faturada"] is not None and x["programda"] is not None else "Toplam faturadakinden farklı.",
    "tevkifat": lambda x: "Tevkifat faturadakinden farklı.",
    "cari": lambda x: "Firma bilgisi eksik ya da hatalı.",
    "fatura_no": lambda x: "Fatura numarası boş.",
    "tarih": lambda x: "Fatura tarihi boş.",
    "kur": lambda x: f"{x['programda']} faturası; kur girilmemiş, raporlara döviz tutarıyla giriyor. Faturayı açıp kuru girin.",
}


def onayli_sorunlari(db, sirket) -> list[dict]:
    """Onaylı faturalarda kalan sorunlar (o anki kayıtlara göre): her fatura bir satır, sorun kısa ve sade."""
    from sqlalchemy import select

    from .models import Fatura
    sonuc = []
    for f in db.scalars(select(Fatura).where(Fatura.sirket_id == sirket.id, Fatura.durum == "onayli")
                        .order_by(Fatura.tarih, Fatura.id)).all():
        sorunlar = [x for x in fatura_teshis(f, sirket) if x["kategori"] in _SADE]
        if not sorunlar:
            continue
        # Kalemler okunamadıysa KDV/toplam farkları da ondandır; tek cümle yeter
        ana = next((x for x in sorunlar if x["kategori"] == "kalem_okunamadi"), sorunlar[0])
        cumleler = list(dict.fromkeys(_SADE[x["kategori"]](x) for x in ([ana] if ana["kategori"] == "kalem_okunamadi" else sorunlar)))
        sonuc.append({"id": f.id, "fatura_no": f.fatura_no or "", "tarih": str(f.tarih or ""), "cari": f.cari.unvan if f.cari else "",
                      "genel_toplam": float(f.genel_toplam or 0), "para_birimi": f.para_birimi or "TRY",
                      "sorun": " ".join(cumleler[:2])})
    return sonuc


def okunamayan_kaydet(sirket_id: int, dosya_adi: str | None, icerik: bytes, neden: str, yon: str) -> None:
    """Hiç fatura çıkmayan dosya: rapora girsin diye kopyası ve nedeni saklanır."""
    try:
        HATA_KLASORU.mkdir(parents=True, exist_ok=True)
        ad = f"{uuid.uuid4().hex[:8]}_{re.sub(r'[^\w.-]', '_', dosya_adi or 'dosya')}"[:120]
        (HATA_KLASORU / ad).write_bytes(icerik)
        with KAYIT.open("a", encoding="utf-8") as k:
            k.write(json.dumps({"sirket_id": sirket_id, "zaman": datetime.now().isoformat(timespec="seconds"), "dosya": dosya_adi,
                                "kopya": ad, "neden": neden, "yon": yon}, ensure_ascii=False) + "\n")
    except OSError:
        pass


def _kayitlar() -> list[dict]:
    try:
        return [json.loads(x) for x in KAYIT.read_text(encoding="utf-8").splitlines() if x.strip()]
    except (OSError, ValueError):
        return []


def okunamayanlar(sirket_id: int) -> list[dict]:
    # Aynı dosya tekrar denendiyse son kaydı geçerli; sonradan okunduysa listeden çıkar
    son = {}
    for x in _kayitlar():
        if x.get("sirket_id") == sirket_id:
            son[x.get("dosya")] = x
    return [x for x in son.values() if not x.get("cozuldu")]


def okundu(sirket_id: int, dosya_adi: str | None) -> None:
    """Daha önce okunamayan dosya yeniden yüklenip okunduysa rapordan düşer."""
    if any(x.get("dosya") == dosya_adi for x in okunamayanlar(sirket_id)):
        try:
            with KAYIT.open("a", encoding="utf-8") as k:
                k.write(json.dumps({"sirket_id": sirket_id, "dosya": dosya_adi, "cozuldu": True}, ensure_ascii=False) + "\n")
        except OSError:
            pass


def okunamayanlari_temizle(sirket_id: int) -> int:
    kalan, silinen = [], 0
    for x in _kayitlar():
        if x.get("sirket_id") == sirket_id:
            silinen += not x.get("cozuldu")
            (HATA_KLASORU / x.get("kopya", "-")).unlink(missing_ok=True)
        else:
            kalan.append(x)
    HATA_KLASORU.mkdir(parents=True, exist_ok=True)
    KAYIT.write_text("".join(json.dumps(x, ensure_ascii=False) + "\n" for x in kalan), encoding="utf-8")
    return silinen


def ozet(faturalar, sirket) -> dict:
    # sorunlar: fatura id -> [{kategori, neden}]; ekranda sorunlu faturalar üstte ve süzülerek gösterilir
    sonuc = {"toplam": len(faturalar), "sorunsuz": 0, "hatali": 0, "kategoriler": Counter(), "okunamayan_dosya": 0,
             "sorunlar": {}}
    for f in faturalar:
        sorunlar = fatura_teshis(f, sirket)
        if not sorunlar:
            sonuc["sorunsuz"] += 1
            continue
        sonuc["hatali"] += 1
        sonuc["sorunlar"][f.id] = [{"kategori": x["kategori"], "neden": x["neden"]} for x in sorunlar]
        for kat in dict.fromkeys(x["kategori"] for x in sorunlar):
            sonuc["kategoriler"][kat] += 1
    sonuc["okunamayan_dosya"] = len(okunamayanlar(sirket.id))
    sonuc["kategoriler"] = [{"kategori": k, "ad": KATEGORILER[k], "sayi": n} for k, n in sonuc["kategoriler"].most_common()]
    return sonuc


def _belge_dokumu(z: zipfile.ZipFile, klasor: str, veri: bytes, uzanti: str) -> None:
    """PDF ise metin, tablo ve kelime konumları; XML ise kendisi zaten belgelerde."""
    from .tablo import kelimeler, tablo_dokumu
    if uzanti != "pdf":
        return
    try:
        z.writestr(f"{klasor}/metin.txt", kural.pdf_metni(veri))
    except kural.KuralHatasi as e:
        z.writestr(f"{klasor}/metin.txt", f"(metin çıkarılamadı: {e})")
    z.writestr(f"{klasor}/tablolar.json", json.dumps(tablo_dokumu(veri), ensure_ascii=False, indent=1))
    z.writestr(f"{klasor}/kelimeler.json", json.dumps(kelimeler(veri), ensure_ascii=False))


def rapor_zip(faturalar, sirket, belge_klasoru) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws_ozet = wb.active
    ws_ozet.title = "Özet"
    ws_hata = wb.create_sheet("Hatalı Faturalar")
    ws_kalem = wb.create_sheet("Kalemler")
    ws_dosya = wb.create_sheet("Okunamayan Dosyalar")
    ws_hata.append(["Klasör", "Dosya", "Yön", "Fatura No", "Tarih", "Firma", "VKN", "Kategori", "Neden", "Faturada", "Programda",
                    "Okuma Yöntemi", "Tablo Başlıkları"])
    ws_kalem.append(["Klasör", "Fatura No", "Sıra No", "Mal Hizmet Adı", "Miktar", "Birim", "Birim Fiyat", "İskonto Oranı",
                     "İskonto Tutarı", "KDV Oranı", "KDV Tutarı", "Tevkifat", "Mal Hizmet Tutarı"])
    ws_dosya.append(["Klasör", "Dosya", "Yön", "Zaman", "Neden"])
    sayac, firmalar, toplam_hatali = Counter(), Counter(), 0
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as z:
        for f in faturalar:
            sorunlar = fatura_teshis(f, sirket)
            if not sorunlar:
                continue
            toplam_hatali += 1
            o = f.ai_okuma or {}
            karsi = o.get("satici" if f.yon == "gelen" else "alici") or {}
            firma = f.cari.unvan if f.cari else karsi.get("unvan")
            vkn = f.cari.vkn if f.cari else karsi.get("vkn")
            klasor = f"faturalar/{toplam_hatali:03d}_{re.sub(r'[^\w.-]', '_', f.fatura_no or f'fatura-{f.id}')}"
            basliklar = ""
            if (o.get("resimden") or {}).get("metin"):
                z.writestr(f"{klasor}/resim_metni.txt", o["resimden"]["metin"])
            for kaynak in dict.fromkeys(x for x in (f.kaynak_dosya, f.xml_dosya) if x):
                yol = belge_klasoru / kaynak
                if not yol.exists():
                    continue
                uzanti = kaynak.rsplit(".", 1)[-1].lower()
                veri = yol.read_bytes()
                z.writestr(f"{klasor}/belge.{uzanti}", veri)
                _belge_dokumu(z, klasor, veri, uzanti)
                if uzanti == "pdf":
                    from .tablo import tablo_dokumu
                    basliklar = " | ".join(" / ".join(" ".join((h or "").split()) for h in t["basliklar"])
                                           for t in tablo_dokumu(veri) if t["kalem_tablosu"]) or "(kalem tablosu bulunamadı)"
            z.writestr(f"{klasor}/okunan.json", json.dumps(o, ensure_ascii=False, indent=1, default=str))
            for x in sorunlar:
                sayac[x["kategori"]] += 1
                ws_hata.append([klasor, o.get("dosya_adi"), "Alış" if f.yon == "gelen" else "Satış", f.fatura_no,
                                f.tarih, firma, vkn, KATEGORILER[x["kategori"]], x["neden"], x["faturada"], x["programda"],
                                o.get("okuma_yontemi"), basliklar])
            firmalar[firma or "(firma yok)"] += 1
            for i, kl in enumerate(f.kalemler, 1):
                ws_kalem.append([klasor, f.fatura_no, i, kl.aciklama, float(kl.miktar), kl.birim, float(kl.birim_fiyat),
                                 float(kl.iskonto_orani or 0), float(kl.iskonto_tutari or 0), float(kl.kdv_orani), round(kl.kdv, 2),
                                 f"{float(kl.tevkifat_orani):g}/10 {kl.tevkifat_kodu or ''}".strip() if float(kl.tevkifat_orani or 0) else None,
                                 float(kl.tutar)])
        dosyalar = okunamayanlar(sirket.id)
        for i, x in enumerate(dosyalar, 1):
            klasor = f"okunamayan/{i:03d}_{re.sub(r'[^\w.-]', '_', x.get('dosya') or 'dosya')}"[:150]
            yol = HATA_KLASORU / x.get("kopya", "-")
            if yol.exists():
                uzanti = (x.get("dosya") or "").rsplit(".", 1)[-1].lower()
                veri = yol.read_bytes()
                z.writestr(f"{klasor}/belge.{uzanti}", veri)
                _belge_dokumu(z, klasor, veri, uzanti)
            ws_dosya.append([klasor, x.get("dosya"), "Alış" if x.get("yon") == "gelen" else "Satış", x.get("zaman"), x.get("neden")])
        sayac["dosya"] += len(dosyalar)
        ws_ozet.append(["Okuma Hata Raporu", datetime.now().strftime("%d.%m.%Y %H:%M")])
        ws_ozet.append(["Onay bekleyen fatura", len(faturalar)])
        ws_ozet.append(["Sorunsuz", len(faturalar) - toplam_hatali])
        ws_ozet.append(["Sorunlu", toplam_hatali])
        ws_ozet.append(["Okunamayan dosya", len(dosyalar)])
        ws_ozet.append([])
        ws_ozet.append(["Kategori", "Sorun sayısı"])
        for kat, n in sayac.most_common():
            if n:
                ws_ozet.append([KATEGORILER[kat], n])
        ws_ozet.append([])
        ws_ozet.append(["Firma", "Sorunlu fatura"])
        for firma, n in firmalar.most_common():
            ws_ozet.append([firma, n])
        kalin, dolgu = Font(bold=True), PatternFill("solid", fgColor="DDE6F0")
        for ws in (ws_hata, ws_kalem, ws_dosya):
            for h in ws[1]:
                h.font, h.fill = kalin, dolgu
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for sutun in ws.columns:
                ws.column_dimensions[sutun[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) + 2 for c in sutun[:200])))
        for hucre in ws_hata["E"][1:]:
            hucre.number_format = "DD.MM.YYYY"
        for r in ws_ozet.iter_rows():
            if r and r[0].value in ("Kategori", "Firma", "Okuma Hata Raporu"):
                for h in r:
                    h.font = kalin
        ws_ozet.column_dimensions["A"].width = 48
        x = io.BytesIO()
        wb.save(x)
        z.writestr("okuma-hata-raporu.xlsx", x.getvalue())
        z.writestr("BENI-OKU.txt", "Bu dosyayı olduğu gibi Claude'a gönderin. okuma-hata-raporu.xlsx: her sorun bir satır "
                   "(kategori, neden, faturada yazan ve programın bulduğu). faturalar/ ve okunamayan/ klasörlerinde her belge, "
                   "PDF metni (metin.txt), tablo dökümü (tablolar.json) ve kelime konumları (kelimeler.json) var.\n")
    return tampon.getvalue()
