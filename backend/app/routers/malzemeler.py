"""Malzeme kartları, benzer adların toplu eşleştirilmesi ve Malzeme Bazlı Alış/Satış raporu."""
import io
from collections import defaultdict
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import kural, malzeme
from ..auth import aktif_kullanici, yazabilir
from ..db import get_db
from ..hesap import firma_adi, isaret, satis_tarafi, tl_kur
from ..models import Kullanici, Malzeme
from .temel import sirket_kaydi

router = APIRouter(prefix="/api")


def _kart_cikis(m: Malzeme) -> dict:
    return {"id": m.id, "ad": m.ad, "eski_adlar": kural.kayit_adlari(m)[1:]}


@router.get("/malzemeler")
def liste(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    dizin, _ = malzeme.guncelle(db, k.sirket_id)
    return sorted((_kart_cikis(m) for m in dizin.kartlar), key=lambda x: kural.sade(x["ad"]))


@router.get("/malzemeler/bekleyen")
def bekleyen(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Bir karta benzeyen, henüz karara bağlanmamış kalem adları (en iyi öneri ilk sırada)."""
    return malzeme.guncelle(db, k.sirket_id)[1]


class Karar(BaseModel):
    ad: str
    malzeme_id: int | None = None  # boş = ayrı malzeme (yeni kart)


class Kararlar(BaseModel):
    kararlar: list[Karar]


@router.post("/malzemeler/esle")
def esle(istek: Kararlar, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Benzer adlar için kullanıcının kararı: seçilen kartın eski adlarına eklenir ya da ayrı kart açılır."""
    eslenen = ayri = 0
    for x in istek.kararlar:
        ad = " ".join(x.ad.split())
        if not ad:
            continue
        if x.malzeme_id:
            malzeme.ad_ekle(sirket_kaydi(db, Malzeme, x.malzeme_id, k), ad)
            eslenen += 1
        else:
            db.add(Malzeme(sirket_id=k.sirket_id, ad=ad[:255]))
            ayri += 1
        db.flush()
    db.commit()
    return {"eslenen": eslenen, "ayri": ayri}


class KartIstegi(BaseModel):
    ad: str


@router.put("/malzemeler/{malzeme_id}")
def ad_degistir(malzeme_id: int, istek: KartIstegi, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Kartın adını değiştirir; eski ad kartta kalır (o yazılışla gelen kalemler yine bu karta)."""
    m = sirket_kaydi(db, Malzeme, malzeme_id, k)
    yeni = " ".join(istek.ad.split())
    if not yeni:
        raise HTTPException(422, "Malzeme adı boş olamaz.")
    diger = malzeme.Dizin([x for x in db.scalars(select(Malzeme).where(Malzeme.sirket_id == k.sirket_id)) if x.id != m.id]).bul(yeni)
    if diger:
        raise HTTPException(409, f"Bu ad '{diger.ad}' malzemesinde var; iki kartı birleştirin.")
    kural.ad_degistir(m, yeni)
    db.commit()
    return _kart_cikis(m)


class Birlestir(BaseModel):
    idler: list[int]  # bu kartlar hedefe katılır


@router.post("/malzemeler/{malzeme_id}/birlestir")
def birlestir(malzeme_id: int, istek: Birlestir, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    hedef = sirket_kaydi(db, Malzeme, malzeme_id, k)
    for i in istek.idler:
        if i == hedef.id:
            continue
        m = sirket_kaydi(db, Malzeme, i, k)
        for ad in kural.kayit_adlari(m):
            malzeme.ad_ekle(hedef, ad)
        db.delete(m)
    db.commit()
    return _kart_cikis(hedef)


class AdCikar(BaseModel):
    ad: str


@router.post("/malzemeler/{malzeme_id}/ad-cikar")
def ad_cikar(malzeme_id: int, istek: AdCikar, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Yanlış eşleştirilen yazılışı karttan çıkarır; o kalemler yeniden sorulur ya da ayrı kart olur."""
    m = sirket_kaydi(db, Malzeme, malzeme_id, k)
    a = malzeme.anahtar(istek.ad)
    m.eski_adlar = "\n".join(x for x in kural.kayit_adlari(m)[1:] if not malzeme.ayni_mi(malzeme.anahtar(x), a)) or None
    db.commit()
    return _kart_cikis(m)


@router.delete("/malzemeler/{malzeme_id}")
def sil(malzeme_id: int, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Kartı siler; kalemleri yeniden eşleştirilir (benzer kart varsa sorulur, yoksa yeniden kart açılır)."""
    db.delete(sirket_kaydi(db, Malzeme, malzeme_id, k))
    db.commit()
    return {"silindi": True}


# --- Rapor ------------------------------------------------------------------

_BIRIM = {"ad": "adet", "adt": "adet", "kgm": "kg", "kilogram": "kg", "mt": "m", "metre": "m", "mtr": "m", "l": "lt",
          "ltr": "lt", "litre": "lt", "m2": "m²", "m3": "m³", "ton": "ton", "tn": "ton"}


def _birim(b: str | None) -> str:
    s = malzeme.anahtar(b).replace(" ", "")
    return _BIRIM.get(s, s) or "-"


def rapor_veri(db: Session, k: Kullanici, yon: str, cari_id: int | None, santiye_id: int | None,
               baslangic: date | None, bitis: date | None) -> tuple[list[dict], list[dict]]:
    """(özet, ayrıntı). Özet: malzeme + birim başına toplam miktar, tutar (TL, KDV hariç), ortalama birim fiyat.
    yon='gelen' alış tarafı (alış faturaları, tedarikçiye iade düşer), 'giden' satış tarafı."""
    dizin, _ = malzeme.guncelle(db, k.sirket_id)
    ayrinti = []
    for kalem in malzeme.onayli_kalemler(db, k.sirket_id):
        f = kalem.fatura
        if satis_tarafi(f) != (yon == "giden"):
            continue
        if cari_id and f.cari_id != cari_id:
            continue
        sid = kalem.santiye_id or f.santiye_id
        if santiye_id and sid != santiye_id:
            continue
        if (baslangic and f.tarih < baslangic) or (bitis and f.tarih > bitis):
            continue
        m = dizin.bul(kalem.aciklama)
        isa, kur = isaret(f), tl_kur(f)
        miktar = isa * float(kalem.miktar or 0)
        tutar = round(isa * float(kalem.tutar or 0) * kur, 2)
        ayrinti.append({
            "malzeme_id": m.id if m else None, "malzeme": m.ad if m else " ".join((kalem.aciklama or "").split()),
            "eslesme_bekliyor": m is None, "birim": _birim(kalem.birim), "tarih": f.tarih, "fatura_id": f.id,
            "fatura_no": f.fatura_no, "cari": firma_adi(f),
            "santiye": (kalem.santiye or f.santiye).ad if (kalem.santiye or f.santiye) else None,
            "kalem": kalem.aciklama, "miktar": miktar, "birim_fiyat": round(tutar / miktar, 4) if miktar else None,
            "tutar": tutar, "kdv": round(isa * kalem.kdv * kur, 2)})
    gruplar: dict = defaultdict(lambda: {"miktar": 0.0, "tutar": 0.0, "kdv": 0.0, "faturalar": set(), "cariler": set(), "adlar": set()})
    for x in ayrinti:
        g = gruplar[(x["malzeme_id"] or x["malzeme"], x["birim"])]
        g.update(malzeme_id=x["malzeme_id"], malzeme=x["malzeme"], birim=x["birim"], eslesme_bekliyor=x["eslesme_bekliyor"])
        g["miktar"] += x["miktar"]
        g["tutar"] += x["tutar"]
        g["kdv"] += x["kdv"]
        g["faturalar"].add(x["fatura_id"])
        g["cariler"].add(x["cari"])
        g["adlar"].add(" ".join((x["kalem"] or "").split()))
    ozet = []
    for g in gruplar.values():
        ozet.append({"malzeme_id": g["malzeme_id"], "malzeme": g["malzeme"], "birim": g["birim"],
                     "eslesme_bekliyor": g["eslesme_bekliyor"], "miktar": round(g["miktar"], 4),
                     "tutar": round(g["tutar"], 2), "kdv": round(g["kdv"], 2),
                     "ortalama_fiyat": round(g["tutar"] / g["miktar"], 4) if g["miktar"] else None,
                     "fatura_sayisi": len(g["faturalar"]), "cari_sayisi": len(g["cariler"] - {None}),
                     "yazilislar": sorted(g["adlar"])})
    ozet.sort(key=lambda x: (kural.sade(x["malzeme"]), x["birim"]))
    ayrinti.sort(key=lambda x: (kural.sade(x["malzeme"]), x["birim"], x["tarih"]))
    return ozet, ayrinti


@router.get("/raporlar/malzeme")
def rapor(yon: str = "gelen", cari_id: int | None = None, santiye_id: int | None = None, baslangic: date | None = None,
          bitis: date | None = None, malzeme_id: int | None = None, db: Session = Depends(get_db),
          k: Kullanici = Depends(aktif_kullanici)):
    ozet, ayrinti = rapor_veri(db, k, yon, cari_id, santiye_id, baslangic, bitis)
    if malzeme_id:
        ayrinti = [x for x in ayrinti if x["malzeme_id"] == malzeme_id]
    return {"ozet": ozet, "ayrinti": ayrinti}


@router.get("/excel/malzeme")
def excel(yon: str = "gelen", cari_id: int | None = None, santiye_id: int | None = None, baslangic: date | None = None,
          bitis: date | None = None, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    from .raporlar import BASLIK_DOLGU, BASLIK_STILI
    ozet, ayrinti = rapor_veri(db, k, yon, cari_id, santiye_id, baslangic, bitis)
    wb = Workbook()
    sayfalar = [
        ("Malzeme Toplamları", ["Malzeme", "Birim", "Toplam Miktar", "Toplam Tutar (KDV hariç, TL)", "KDV", "Ortalama Birim Fiyat",
                                "Fatura Sayısı", "Firma Sayısı", "Faturadaki Yazılışlar"],
         [[x["malzeme"], x["birim"], x["miktar"], x["tutar"], x["kdv"], x["ortalama_fiyat"], x["fatura_sayisi"],
           x["cari_sayisi"], " | ".join(x["yazilislar"])] for x in ozet]),
        ("Kalemler", ["Malzeme", "Tarih", "Fatura No", "Firma", "Şantiye", "Faturadaki Ad", "Miktar", "Birim",
                      "Birim Fiyat (TL)", "Tutar (KDV hariç, TL)", "KDV"],
         [[x["malzeme"], x["tarih"], x["fatura_no"], x["cari"], x["santiye"], x["kalem"], x["miktar"], x["birim"],
           x["birim_fiyat"], x["tutar"], x["kdv"]] for x in ayrinti]),
    ]
    for i, (ad, basliklar, satirlar) in enumerate(sayfalar):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = ad
        ws.append(basliklar)
        for h in ws[1]:
            h.font, h.fill = BASLIK_STILI, BASLIK_DOLGU
        for s in satirlar:
            ws.append(s)
        if i == 0:
            ws.append([])
            ws.append(["TOPLAM", None, None, round(sum(x["tutar"] for x in ozet), 2), round(sum(x["kdv"] for x in ozet), 2)])
            for h in ws[ws.max_row]:
                h.font = Font(bold=True)
        for j, b in enumerate(basliklar, start=1):
            uzunluk = max([len(b)] + [len(str(s[j - 1])) for s in satirlar if s[j - 1] is not None])
            ws.column_dimensions[ws.cell(1, j).column_letter].width = min(uzunluk + 2, 60)
        ws.freeze_panes = "A2"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    dosya = "malzeme-satislari" if yon == "giden" else "malzeme-alislari"
    return StreamingResponse(tampon, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{dosya}.xlsx"'})
