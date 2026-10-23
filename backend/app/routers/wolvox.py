"""Wolvox'tan bir kerelik geçmiş aktarımı: Wolvox'un "stok hareketleri / fatura kalemleri" Excel raporu (Stok Adı,
Tarihi, Fatura No, Ticari Unvanı, Birimi, Grubu, Miktarı, Tutarı, Tutarı (KDV Dahil) ve iade sütunları).

Aktarılan faturalar ayrı durumdadır (durum="aktarim"): cari açılmaz, cari bakiyesine, KDV ve mali raporlara girmez;
yalnız şantiye, depo ve malzeme raporlarında görünür. Wolvox'un "Grubu" sütunu (şantiye/depo grubu) önizlemede
kullanıcının seçtiği şantiyeye ya da taşeron deposuna eşlenir. Sistemde zaten olan fatura atlanır; aktarım tek
tuşla geri alınır (faturaları ve açtığı, başka yerde kullanılmayan şantiye/depolar silinir).
"""
import io
import json
import re
from collections import defaultdict
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session, joinedload

from .. import kural, malzeme, schemas
from ..auth import aktif_kullanici, yazabilir
from ..db import get_db
from ..hesap import firma_adi
from ..models import Cek, Depo, Fatura, FaturaKalem, Hareket, Kullanici, Santiye, TaseronGider, WolvoxAktarim
from ..yukleme import dosyayi_oku
from .faturalar import KAYIT_KILIDI, kalemleri_yaz
from .temel import sirket_kaydi

router = APIRouter(prefix="/api/wolvox")

# Sütun adı (sade yazılış) -> alan
SUTUNLAR = {
    "stok adi": "stok", "stok": "stok", "tarihi": "tarih", "tarih": "tarih", "fatura no": "no", "belge no": "no",
    "ticari unvani": "firma", "unvani": "firma", "cari adi": "firma", "birimi": "birim", "birim": "birim",
    "grubu": "grup", "grup": "grup", "miktari": "miktar", "tutari": "tutar", "tutari kdv dahil": "dahil",
    "miktari iade": "iade_miktar", "tutari iade": "iade_tutar", "tutari kdv dahil iade": "iade_dahil",
}
_ORANLAR = (0, 1, 8, 10, 18, 20)


def _sayi(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _tarih(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.match(r"\s*(\d{1,2})[./-](\d{1,2})[./-](\d{4})", str(v or ""))
    if m:
        try:
            return date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None
    return None


def _fatura_no(v) -> str | None:
    """Wolvox'un yazdığı biçimler: sayı, metin ya da formül gibi '="0001234"'."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    s = re.sub(r'^=?"(.*)"$', r"\1", s).strip()
    return s or None


def _kdv_orani(tutar: float, kdv: float) -> float:
    if not tutar:
        return 0
    oran = kdv / tutar * 100
    yakin = min(_ORANLAR, key=lambda o: abs(o - oran))
    return yakin if abs(yakin - oran) < 0.6 else round(oran, 2)


def oku(icerik: bytes) -> dict:
    """Excel'i faturalara böler: {"faturalar": [...], "atlanan_satir": [...]}. Her fatura:
    {yon, fatura_no, tarih, firma, iade, kalemler: [{aciklama, miktar, birim, tutar, kdv, grup}]}."""
    try:
        wb = load_workbook(io.BytesIO(icerik), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, "Excel dosyası açılamadı. Wolvox raporunu Excel (.xlsx) olarak kaydedip yükleyin.") from e
    for ws in wb.worksheets:
        satirlar = ws.iter_rows(values_only=True)
        sutun = None
        for i, r in enumerate(satirlar):
            adlar = {SUTUNLAR.get(kural.ad_anahtari(str(x or ""))): j for j, x in enumerate(r)}
            adlar.pop(None, None)
            if {"stok", "no", "tarih"} <= adlar.keys() and ({"tutar", "dahil"} & adlar.keys()):
                sutun = adlar
                break
            if i > 20:
                break
        if sutun:
            break
    else:
        raise HTTPException(400, "Wolvox raporu tanınmadı: 'Stok Adı', 'Tarihi', 'Fatura No' ve 'Tutarı' sütunları bulunamadı.")

    def al(r, alan):
        j = sutun.get(alan)
        return r[j] if j is not None and j < len(r) else None

    faturalar: dict[tuple, dict] = {}
    atlanan = []
    for n, r in enumerate(satirlar, start=2):
        if not any(x not in (None, "") for x in r):
            continue
        stok = " ".join(str(al(r, "stok") or "").split())
        no, tarih = _fatura_no(al(r, "no")), _tarih(al(r, "tarih"))
        firma = " ".join(str(al(r, "firma") or "").split()) or None
        grup = " ".join(str(al(r, "grup") or "").split()) or None
        birim = (str(al(r, "birim") or "").strip() or None)
        if not stok or not no or not tarih:
            if stok or no:
                atlanan.append({"satir": n, "stok": stok, "fatura_no": no,
                                "neden": "Fatura no yok" if not no else "Tarih okunamadı" if not tarih else "Stok adı yok"})
            continue
        for iade in (False, True):
            on = "iade_" if iade else ""
            tutar, dahil, miktar = _sayi(al(r, on + "tutar")), _sayi(al(r, on + "dahil")), _sayi(al(r, on + "miktar"))
            if iade and not (tutar or dahil or miktar):
                continue
            if not iade and not (tutar or dahil or miktar) and (_sayi(al(r, "iade_tutar")) or _sayi(al(r, "iade_dahil"))):
                continue  # yalnız iade satırı
            if not tutar and dahil:
                tutar = dahil
            kdv = round(dahil - tutar, 2) if dahil else 0.0
            yon = "giden" if iade else "gelen"  # tedarikçiye iade: bizim kestiğimiz iade faturası
            f = faturalar.setdefault((yon, no, tarih, kural.ad_anahtari(firma)), {
                "yon": yon, "fatura_no": no, "tarih": tarih, "firma": firma, "iade": iade, "kalemler": []})
            f["kalemler"].append({"aciklama": stok, "miktar": abs(miktar), "birim": birim, "tutar": abs(round(tutar, 2)),
                                  "kdv": abs(kdv), "grup": grup})
    for f in faturalar.values():
        f["toplam"] = round(sum(k["tutar"] + k["kdv"] for k in f["kalemler"]), 2)
    return {"faturalar": sorted(faturalar.values(), key=lambda f: (f["tarih"], f["fatura_no"])), "atlanan_satir": atlanan}


def _kayitli(db: Session, sirket_id: int) -> dict[tuple, list[tuple]]:
    """Sistemdeki faturalar (iptal hariç): (yön, fatura no) -> [(tarih, toplam, firma adı)]."""
    sonuc = defaultdict(list)
    for f in db.scalars(select(Fatura).where(Fatura.sirket_id == sirket_id, Fatura.durum != "iptal", Fatura.fatura_no.is_not(None))
                        .options(joinedload(Fatura.cari))):
        sonuc[(f.yon, f.fatura_no.strip())].append((f.tarih, float(f.genel_toplam or 0), kural.ad_anahtari(firma_adi(f))))
    return sonuc


def _sistemde_var(kayitli: dict, f: dict) -> bool:
    """Aynı fatura no ve aynı tarih ya da aynı toplam; iki firmanın adı da biliniyorsa ilk kelimeleri aynı olmalı
    (farklı firmalar aynı numarayı kullanabilir)."""
    firma = kural.ad_anahtari(f["firma"])
    return any((t == f["tarih"] or abs(top - f["toplam"]) < 1) and (not firma or not fa or firma.split()[0] == fa.split()[0])
               for t, top, fa in kayitli.get((f["yon"], f["fatura_no"]), []))


def _taseron_mu(ad: str) -> bool:
    return kural.ad_anahtari(ad).startswith(("tasaron", "taseron"))


class _Ad:
    def __init__(self, ad):
        self.ad = ad


def _oneri(grup: str | None, santiyeler: list, depolar: list, gruplar: list[str]) -> dict:
    """Grubun nereye aktarılacağı önerisi: kayıtlı depo/şantiye, taşeron grubu depo, öteki gruplar yeni şantiye."""
    if not grup:
        return {"tur": "yok", "ad": "", "santiye": ""}
    if d := next((d for d in depolar if kural.adi_bu_mu(d, grup)), None):
        return {"tur": "depo", "ad": d.ad, "santiye": d.santiye.ad if d.santiye else "", "kayitli": True}
    if s := next((s for s in santiyeler if kural.adi_bu_mu(s, grup)), None):
        return {"tur": "santiye", "ad": s.ad, "santiye": "", "kayitli": True}
    if _taseron_mu(grup):
        adaylar = santiyeler + [_Ad(g) for g in gruplar if not _taseron_mu(g)]
        icinde = kural.icindeki_kayit(grup, adaylar)
        return {"tur": "depo", "ad": grup, "santiye": icinde.ad if icinde else ""}
    return {"tur": "santiye", "ad": grup, "santiye": ""}


def _benzer_grup(grup: str, gruplar: list[str]) -> str | None:
    """Yazılışı çok yakın başka grup ("ÖRNEK DOĞALAR" / "ÖRNEK DOĞLAR"): kullanıcıya aynı şantiye mi diye gösterilir."""
    a = malzeme.anahtar(grup)
    en = None
    for g in gruplar:
        if kural.ad_anahtari(g) == kural.ad_anahtari(grup):
            continue  # Türkçe harf/boşluk farkı: zaten aynı şantiyeye gider
        o, guclu, _ = malzeme.karsilastir(a, malzeme.anahtar(g))
        if o >= 0.88 and (en is None or o > en[0]):
            en = (o, g)
    return en[1] if en else None


@router.post("/onizle")
def onizle(dosya: UploadFile = File(...), db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    veri = oku(dosyayi_oku(dosya))
    fs = veri["faturalar"]
    if not fs:
        raise HTTPException(400, "Dosyada aktarılacak fatura satırı bulunamadı.")
    kayitli = _kayitli(db, k.sirket_id)
    santiyeler = list(db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id)))
    depolar = list(db.scalars(select(Depo).where(Depo.sirket_id == k.sirket_id)))
    gruplar: dict = defaultdict(lambda: {"kalem": 0, "faturalar": set(), "toplam": 0.0})
    var = 0
    for i, f in enumerate(fs):
        if _sistemde_var(kayitli, f):
            var += 1
            continue
        for kl in f["kalemler"]:
            g = gruplar[kl["grup"] or ""]
            g["kalem"] += 1
            g["faturalar"].add(i)
            g["toplam"] += (-1 if f["iade"] else 1) * (kl["tutar"] + kl["kdv"])
    adlar = [g for g in gruplar if g]
    stoklar = {malzeme._bitisik(malzeme.anahtar(kl["aciklama"])) for f in fs for kl in f["kalemler"]}
    return {
        "dosya": dosya.filename, "fatura_sayisi": len(fs), "aktarilacak": len(fs) - var, "zaten_var": var,
        "iade_sayisi": sum(1 for f in fs if f["iade"]), "kalem_sayisi": sum(len(f["kalemler"]) for f in fs),
        "firma_sayisi": len({kural.ad_anahtari(f["firma"]) for f in fs}), "stok_adi_sayisi": len(stoklar),
        "ilk_tarih": fs[0]["tarih"], "son_tarih": fs[-1]["tarih"], "atlanan_satir": veri["atlanan_satir"][:50],
        "atlanan_satir_sayisi": len(veri["atlanan_satir"]),
        "gruplar": sorted(({"grup": g, "kalem_sayisi": x["kalem"], "fatura_sayisi": len(x["faturalar"]),
                            "toplam": round(x["toplam"], 2), "oneri": _oneri(g or None, santiyeler, depolar, adlar),
                            "benzer": _benzer_grup(g, adlar) if g else None}
                           for g, x in gruplar.items()), key=lambda x: (x["grup"] == "", kural.sade(x["grup"]))),
        "santiyeler": [s.ad for s in sorted(santiyeler, key=lambda s: kural.sade(s.ad))],
    }


class _Hedefler:
    """Eşleme kararlarını kayda çevirir: adıyla kayıtlı şantiye/depo bulunur, yoksa açılır (açılanlar not edilir)."""

    def __init__(self, db: Session, sirket_id: int):
        self.db, self.sirket_id = db, sirket_id
        self.santiyeler = list(db.scalars(select(Santiye).where(Santiye.sirket_id == sirket_id)))
        self.depolar = list(db.scalars(select(Depo).where(Depo.sirket_id == sirket_id)))
        self.acilan_santiyeler: list[int] = []
        self.acilan_depolar: list[int] = []

    def santiye(self, ad: str | None) -> Santiye | None:
        ad = " ".join((ad or "").split())
        if not ad:
            return None
        if s := next((s for s in self.santiyeler if kural.adi_bu_mu(s, ad)), None):
            return s
        s = Santiye(sirket_id=self.sirket_id, ad=ad[:255])
        self.db.add(s)
        self.db.flush()
        self.santiyeler.append(s)
        self.acilan_santiyeler.append(s.id)
        return s

    def depo(self, ad: str | None, santiye_adi: str | None) -> Depo | None:
        ad = " ".join((ad or "").split())
        if not ad:
            return None
        s = self.santiye(santiye_adi)
        ayni = [d for d in self.depolar if kural.adi_bu_mu(d, ad)]
        if d := next((d for d in ayni if s is None or d.santiye_id == s.id), None):
            return d
        d = Depo(sirket_id=self.sirket_id, ad=ad[:255], santiye_id=s.id if s else None)
        self.db.add(d)
        self.db.flush()
        self.depolar.append(d)
        self.acilan_depolar.append(d.id)
        return d

    def kalem_yeri(self, karar: dict | None) -> tuple[int | None, int | None]:
        """(şantiye id, depo id)"""
        karar = karar or {}
        if karar.get("tur") == "depo":
            d = self.depo(karar.get("ad"), karar.get("santiye"))
            return (d.santiye_id, d.id) if d else (None, None)
        if karar.get("tur") == "santiye":
            s = self.santiye(karar.get("ad"))
            return (s.id if s else None), None
        return None, None


@router.post("/aktar")
def aktar(dosya: UploadFile = File(...), eslesme: str = Form("{}"), db: Session = Depends(get_db),
          k: Kullanici = Depends(yazabilir)):
    """eslesme: {"grup adı": {"tur": "santiye" | "depo" | "yok", "ad": ..., "santiye": (depo için şantiye adı)}}.
    Eşlemesi gelmeyen grup önerilen yere aktarılır."""
    try:
        kararlar = json.loads(eslesme or "{}")
    except ValueError as e:
        raise HTTPException(400, "Eşleme okunamadı.") from e
    veri = oku(dosyayi_oku(dosya))
    with KAYIT_KILIDI:
        from .. import yedek
        try:
            yedek.yedek_al()
        except Exception as e:  # noqa: BLE001  yedeksiz toplu aktarım yapılmaz
            raise HTTPException(500, f"Aktarmadan önce yedek alınamadı, hiçbir şey aktarılmadı: {e}") from e
        kayitli = _kayitli(db, k.sirket_id)
        hedef = _Hedefler(db, k.sirket_id)
        gruplar = sorted({kl["grup"] for f in veri["faturalar"] for kl in f["kalemler"] if kl["grup"]})
        yerler: dict = {}
        for g in [None, *gruplar]:
            karar = kararlar.get(g or "") or _oneri(g, hedef.santiyeler, hedef.depolar, gruplar)
            yerler[g] = hedef.kalem_yeri(karar)
        a = WolvoxAktarim(sirket_id=k.sirket_id, dosya=(dosya.filename or "")[:255])
        db.add(a)
        db.flush()
        sayi = atlanan = 0
        toplam = 0.0
        for f in veri["faturalar"]:
            if _sistemde_var(kayitli, f):
                atlanan += 1
                continue
            yeni = Fatura(sirket_id=k.sirket_id, yon=f["yon"], durum="aktarim", fatura_no=f["fatura_no"][:50], tarih=f["tarih"],
                          fatura_tipi="IADE" if f["iade"] else None,
                          ai_okuma={"kaynak": "wolvox", "aktarim": a.id, "dosya_adi": dosya.filename,
                                    ("alici" if f["iade"] else "satici"): {"unvan": f["firma"]}})
            kalemleri_yaz(yeni, [schemas.KalemGiris(
                aciklama=kl["aciklama"], miktar=kl["miktar"], birim=kl["birim"],
                birim_fiyat=round(kl["tutar"] / kl["miktar"], 4) if kl["miktar"] else kl["tutar"],
                kdv_orani=_kdv_orani(kl["tutar"], kl["kdv"]), tutar=kl["tutar"], kdv_tutari=kl["kdv"],
                santiye_id=yerler[kl["grup"]][0], depo_id=yerler[kl["grup"]][1]) for kl in f["kalemler"]])
            db.add(yeni)
            sayi += 1
            toplam += -f["toplam"] if f["iade"] else f["toplam"]
        a.fatura_sayisi, a.atlanan, a.toplam = sayi, atlanan, round(toplam, 2)
        a.acilan_santiyeler, a.acilan_depolar = hedef.acilan_santiyeler, hedef.acilan_depolar
        db.commit()
        # Stok adları malzeme kartlarına bağlanır; yazılışı aynı olanlar tek kart, benzerleri sorulmak üzere bekler
        _, bekleyen = malzeme.guncelle(db, k.sirket_id)
    return {"id": a.id, "fatura_sayisi": sayi, "atlanan": atlanan, "toplam": a.toplam,
            "acilan_santiye": len(hedef.acilan_santiyeler), "acilan_depo": len(hedef.acilan_depolar),
            "benzer_malzeme": len(bekleyen)}


@router.get("/aktarimlar")
def aktarimlar(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    return [{"id": a.id, "tarih": a.tarih, "dosya": a.dosya, "fatura_sayisi": a.fatura_sayisi, "atlanan": a.atlanan,
             "toplam": float(a.toplam or 0), "acilan_santiye": len(a.acilan_santiyeler or []),
             "acilan_depo": len(a.acilan_depolar or [])}
            for a in db.scalars(select(WolvoxAktarim).where(WolvoxAktarim.sirket_id == k.sirket_id).order_by(WolvoxAktarim.id.desc()))]


@router.delete("/aktarimlar/{aktarim_id}")
def geri_al(aktarim_id: int, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Aktarımı geri alır: faturaları ve açtığı, başka kayıtta kullanılmayan şantiye/depolar silinir."""
    a = sirket_kaydi(db, WolvoxAktarim, aktarim_id, k)
    with KAYIT_KILIDI:
        from .. import yedek
        try:
            yedek.yedek_al()
        except Exception as e:  # noqa: BLE001
            raise HTTPException(500, f"Önce yedek alınamadı, hiçbir şey silinmedi: {e}") from e
        idler = [f.id for f in db.scalars(select(Fatura).where(Fatura.sirket_id == k.sirket_id, Fatura.durum == "aktarim"))
                 if (f.ai_okuma or {}).get("aktarim") == a.id]
        for i in range(0, len(idler), 500):
            parca = idler[i:i + 500]
            db.execute(delete(FaturaKalem).where(FaturaKalem.fatura_id.in_(parca)))
            db.execute(delete(Fatura).where(Fatura.id.in_(parca)))
        db.flush()
        silinen_depo = silinen_santiye = 0
        for did in a.acilan_depolar or []:
            if not any(db.scalar(select(m.id).where(m.depo_id == did).limit(1)) for m in (FaturaKalem, TaseronGider)) and \
                    (d := db.get(Depo, did)):
                db.delete(d)
                silinen_depo += 1
        db.flush()
        for sid in a.acilan_santiyeler or []:
            if not any(db.scalar(select(m.id).where(m.santiye_id == sid).limit(1))
                       for m in (Fatura, FaturaKalem, Hareket, Cek, Depo, TaseronGider)) and (s := db.get(Santiye, sid)):
                db.delete(s)
                silinen_santiye += 1
        db.delete(a)
        db.commit()
    return {"silinen_fatura": len(idler), "silinen_santiye": silinen_santiye, "silinen_depo": silinen_depo}
