"""Depolar (şantiye içindeki taşeron hesapları) ve taşeron sayfası.

Fatura kalemi bir depoya ayrılırsa şantiye raporunda o depo altında görünür. Depoya cari bağlıysa kalemin KDV dahil
tutarı o cariye borç yazılır; faturanın kendi carisi (satıcı) değişmez. Taşeron adına faturasız harcamalar (elden
harcama, personel ücreti/SGK) da taşeron giderleri olarak cariye borç yazılır.
"""
import io
import re
from calendar import monthrange
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .. import kural, schemas
from ..auth import aktif_kullanici, yazabilir
from ..db import get_db
from ..hesap import cari_hareketleri, depo_kalemleri, isaret, satis_tarafi, tl_kur
from ..models import Cari, Depo, Fatura, FaturaKalem, Hareket, Kullanici, Santiye, TaseronGider
from ..yukleme import dosyayi_oku
from .temel import sirket_kaydi

router = APIRouter(prefix="/api")


# --- Depolar -------------------------------------------------------------

def ayni_adli_depo(db: Session, sirket_id: int, ad: str | None, santiye_id: int | None, haric: int | None = None) -> Depo | None:
    """Aynı şantiyede aynı adlı depo (harf/Türkçe karakter farkı yok sayılır)."""
    anahtar = kural.ad_anahtari(ad)
    return next((d for d in db.scalars(select(Depo).where(Depo.sirket_id == sirket_id).order_by(Depo.id))
                 if d.id != haric and anahtar and kural.adi_bu_mu(d, ad) and d.santiye_id == santiye_id), None)


def _depo_cikis(d: Depo) -> dict:
    return {**schemas.DepoCikis.model_validate(d).model_dump(),
            "santiye_ad": d.santiye.ad if d.santiye else None, "cari_unvan": d.cari.unvan if d.cari else None}


def _depo_dogrula(db: Session, veri: schemas.DepoGiris, k: Kullanici):
    if not (veri.ad or "").strip():
        raise HTTPException(422, "Depo adını yazın.")
    veri.ad = veri.ad.strip()
    if veri.santiye_id:
        sirket_kaydi(db, Santiye, veri.santiye_id, k)
    if veri.cari_id:
        sirket_kaydi(db, Cari, veri.cari_id, k)


@router.get("/depolar")
def depolar(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    return [_depo_cikis(d) for d in sorted(db.scalars(select(Depo).where(Depo.sirket_id == k.sirket_id)),
                                           key=lambda d: kural.tr_sira(d.ad))]


@router.post("/depolar")
def depo_ekle(veri: schemas.DepoGiris, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    _depo_dogrula(db, veri, k)
    if (ayni := ayni_adli_depo(db, k.sirket_id, veri.ad, veri.santiye_id)) is not None:
        if veri.cari_id and not ayni.cari_id:
            ayni.cari_id = veri.cari_id
            db.commit()
        return _depo_cikis(ayni)  # aynı şantiyede aynı adla ikinci depo açılmaz; var olan seçilir
    d = Depo(sirket_id=k.sirket_id, **veri.model_dump())
    db.add(d)
    db.commit()
    return _depo_cikis(d)


@router.put("/depolar/{depo_id}")
def depo_guncelle(depo_id: int, veri: schemas.DepoGiris, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    d = sirket_kaydi(db, Depo, depo_id, k)
    _depo_dogrula(db, veri, k)
    if ayni_adli_depo(db, k.sirket_id, veri.ad, veri.santiye_id, haric=d.id) is not None:
        raise HTTPException(409, f"Bu şantiyede '{veri.ad}' adında bir depo zaten var.")
    kural.ad_degistir(d, veri.ad)  # depo kalemleri ve taşeron borcu aynı depoda kalır; eski ad saklanır
    for alan, deger in veri.model_dump(exclude={"ad"}).items():
        setattr(d, alan, deger)
    db.commit()
    return _depo_cikis(d)


class DepoDurumu(BaseModel):
    idler: list[int]
    durum: str  # aktif | pasif


@router.post("/depolar/durum")
def depo_durumu(istek: DepoDurumu, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Aktif / pasif: pasif depo seçim listelerinde görünmez, kayıtları durur."""
    if istek.durum not in ("aktif", "pasif"):
        raise HTTPException(422, "Durum aktif ya da pasif olmalı.")
    for i in istek.idler:
        sirket_kaydi(db, Depo, i, k).durum = istek.durum
    db.commit()
    return {"ok": True}


@router.delete("/depolar/{depo_id}")
def depo_sil(depo_id: int, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    d = sirket_kaydi(db, Depo, depo_id, k)
    kalem = db.scalar(select(func.count()).where(FaturaKalem.depo_id == d.id))
    gider = db.scalar(select(func.count()).where(TaseronGider.depo_id == d.id))
    if kalem or gider:
        neden = ", ".join(x for x in (f"{kalem} fatura kalemi" if kalem else "", f"{gider} gider kaydı" if gider else "") if x)
        raise HTTPException(409, f"Bu depoya ayrılmış {neden} var, silinemez. Önce kalemleri başka depoya alın.")
    db.delete(d)
    db.commit()
    return {"ok": True}


# --- Taşeron giderleri (elden harcama, personel/SGK) ---------------------------

class GiderGiris(BaseModel):
    cari_id: int
    tur: str = "elden"  # elden | personel
    tarih: date
    tutar: float | None = None
    aciklama: str | None = None
    santiye_id: int | None = None
    depo_id: int | None = None
    kasadan: bool = False  # elden harcama kasadan ödendiyse kasadan çıkış da yazılır
    donem: str | None = None
    personel: str | None = None
    tckn: str | None = None
    gun: int | None = None
    brut: float | None = None
    net: float | None = None
    sgk_isci: float | None = None
    vergi: float | None = None
    isveren: float | None = None


def _yuvarla(x) -> float | None:
    return round(float(x), 2) if x not in (None, "") else None


def personel_toplami(g) -> float:
    """Personel için taşerona yazılan tutar: brüt ücret (net + işçi SGK/işsizlik + vergiler + icra) + işveren payı.
    Brüt yoksa parçaların toplamı."""
    brut = _yuvarla(getattr(g, "brut", None))
    if brut is None:
        brut = sum(float(getattr(g, a, None) or 0) for a in ("net", "sgk_isci", "vergi"))
    return round(brut + float(getattr(g, "isveren", None) or 0), 2)


def _gider_cikis(g: TaseronGider) -> dict:
    return {a: getattr(g, a) for a in ("id", "cari_id", "tur", "tarih", "donem", "aciklama", "santiye_id", "depo_id", "personel",
                                       "tckn", "gun", "hareket_id")} | \
        {a: _yuvarla(getattr(g, a)) for a in ("brut", "net", "sgk_isci", "vergi", "isveren", "tutar")} | \
        {"santiye_ad": g.santiye.ad if g.santiye else None}


def _gider_yaz(db: Session, g: TaseronGider, veri: GiderGiris, k: Kullanici):
    if veri.tur not in ("elden", "personel"):
        raise HTTPException(422, "Geçersiz gider türü.")
    cari = sirket_kaydi(db, Cari, veri.cari_id, k)
    if veri.santiye_id:
        sirket_kaydi(db, Santiye, veri.santiye_id, k)
    if veri.depo_id:
        sirket_kaydi(db, Depo, veri.depo_id, k)
    for alan in ("cari_id", "tur", "tarih", "aciklama", "santiye_id", "depo_id", "donem", "personel", "tckn", "gun",
                 "brut", "net", "sgk_isci", "vergi", "isveren"):
        setattr(g, alan, getattr(veri, alan))
    g.tutar = personel_toplami(veri) if veri.tur == "personel" and not veri.tutar else _yuvarla(veri.tutar)
    if not g.tutar:
        raise HTTPException(422, "Tutarı yazın.")
    if veri.tur == "personel" and not (veri.personel or "").strip():
        raise HTTPException(422, "Personel adını yazın.")
    # Kasadan ödenen elden harcama kasadan düşer; cariye borcu bu kayıt yazar (kasa hareketi cariye bağlanmaz)
    h = db.get(Hareket, g.hareket_id) if g.hareket_id else None
    if veri.tur == "elden" and veri.kasadan:
        if h is None:
            h = Hareket(sirket_id=k.sirket_id, yon="cikis", tur="masraf", kaynak="elle", durum="onayli")
            db.add(h)
        h.tarih, h.tutar, h.santiye_id = veri.tarih, g.tutar, veri.santiye_id
        h.aciklama = f"{cari.unvan} adına elden harcama" + (f": {veri.aciklama}" if veri.aciklama else "")
        h.durum = "onayli"
        db.flush()
        g.hareket_id = h.id
    elif h is not None:
        h.durum, h.iptal_nedeni = "iptal", "Taşeron gideri kasadan ödenmedi olarak değiştirildi"
        g.hareket_id = None


@router.get("/taseron-giderleri")
def taseron_giderleri(cari_id: int | None = None, tur: str | None = None, db: Session = Depends(get_db),
                      k: Kullanici = Depends(aktif_kullanici)):
    q = select(TaseronGider).where(TaseronGider.sirket_id == k.sirket_id)
    if cari_id:
        q = q.where(TaseronGider.cari_id == cari_id)
    if tur:
        q = q.where(TaseronGider.tur == tur)
    return [_gider_cikis(g) for g in db.scalars(q.order_by(TaseronGider.tarih, TaseronGider.id))]


@router.post("/taseron-giderleri")
def gider_ekle(veri: GiderGiris, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    g = TaseronGider(sirket_id=k.sirket_id)
    _gider_yaz(db, g, veri, k)
    db.add(g)
    db.commit()
    return _gider_cikis(g)


@router.put("/taseron-giderleri/{gider_id}")
def gider_guncelle(gider_id: int, veri: GiderGiris, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    g = sirket_kaydi(db, TaseronGider, gider_id, k)
    _gider_yaz(db, g, veri, k)
    db.commit()
    return _gider_cikis(g)


class TopluGiderSil(BaseModel):
    idler: list[int]


@router.post("/taseron-giderleri/toplu-sil")
def gider_sil(istek: TopluGiderSil, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Seçilen giderleri siler; kasadan ödenmiş olanın kasa çıkışı iptal edilir (kasa bakiyesi geri gelir)."""
    n = 0
    for gid in istek.idler:
        g = sirket_kaydi(db, TaseronGider, gid, k)
        if g.hareket_id and (h := db.get(Hareket, g.hareket_id)):
            h.durum, h.iptal_nedeni = "iptal", "Taşeron gideri silindi"
        db.delete(g)
        n += 1
    db.commit()
    return {"silinen": n}


# --- Bordro Excel'i -------------------------------------------------------

# Bordro sütunları: başlık (harf/Türkçe karakter/nokta farkı yok sayılır) -> alan. Sıra önemli: ilk eşleşen alır.
BORDRO_SUTUNLARI = [
    ("sira", ("#", "sira", "sira no", "no")),
    ("personel", ("adi soyadi", "ad soyad", "adi soyad", "personel", "personel adi", "isim", "ad soyadi", "calisan")),
    ("tckn", ("tckn", "tc", "tc no", "tc kimlik", "tc kimlik no", "kimlik no")),
    ("gun", ("t gun", "gun", "prim gun", "prim gunu", "toplam gun", "calisma gunu")),
    ("brut", ("nor kazanc", "brut", "brut ucret", "brut kazanc", "toplam kazanc", "kazanc")),
    ("sgk", ("ssk isci", "sgk isci", "sgk isci payi", "ssk isci payi", "sgk", "sgk primi", "sgk tutari")),
    ("issizlik", ("iss p isci", "issizlik isci", "issizlik isci payi", "issizlik")),
    ("kalan_gv", ("kalan gv", "odenecek gv", "odenecek gelir vergisi")),
    ("gv", ("gel ver", "gelir vergisi", "gv")),
    ("damga", ("damga v", "damga vergisi", "kalan dv", "odenecek dv")),
    ("net", ("n odenen", "net odenen", "net ucret", "net", "net maas", "odenen")),
    ("icra", ("icra", "icra kesintisi")),
    ("isveren", ("isveren payi", "sgk isveren", "sgk isveren payi", "isveren sgk", "isveren primi")),
]
_AYLAR = {"ocak": 1, "subat": 2, "mart": 3, "nisan": 4, "mayis": 5, "haziran": 6, "temmuz": 7, "agustos": 8,
          "eylul": 9, "ekim": 10, "kasim": 11, "aralik": 12}


def _sayi(x) -> float | None:
    if x is None or x == "":
        return None
    if isinstance(x, (int, float)):
        return float(x)
    return kural.para_coz(str(x))


def _donem_bul(hucreler: list[str], dosya_adi: str | None) -> str | None:
    """Bordronun ayı: 'EYLÜL / 2026' gibi hücre, yoksa dosya adındaki 2026-09."""
    for metin in [*hucreler, dosya_adi or ""]:
        sade = kural.ad_anahtari(metin)
        for ad, no in _AYLAR.items():
            if m := re.search(rf"\b{ad}\b\D{{0,5}}(20\d\d)", sade):
                return f"{m.group(1)}-{no:02d}"
        if m := re.search(r"(20\d\d)[-_./ ](0[1-9]|1[0-2])\b", metin):
            return f"{m.group(1)}-{m.group(2)}"
    return None


def bordro_oku(veri: bytes, dosya_adi: str | None = None) -> dict:
    """Bordro Excel'inden personel satırları. Başlık satırı ad soyad ve net ödenen (ya da brüt) sütunlarından tanınır;
    TOPLAM satırı ve boş satırlar alınmaz."""
    try:
        wb = load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, "Excel dosyası açılamadı (.xlsx olmalı).") from e
    for ws in wb.worksheets:
        satirlar = [list(r) for r in ws.iter_rows(values_only=True)]
        for i, r in enumerate(satirlar[:40]):
            basliklar = [kural.ad_anahtari(str(c)) if c is not None else "" for c in r]
            if "#" in [str(c).strip() for c in r if c is not None]:
                basliklar = ["sira" if str(c).strip() == "#" else b for c, b in zip(r, basliklar)]
            sutun: dict[str, int] = {}
            for alan, adlar in BORDRO_SUTUNLARI:
                j = next((j for j, b in enumerate(basliklar) if b in adlar and j not in sutun.values()), None)
                if j is not None:
                    sutun[alan] = j
            if "personel" not in sutun or not ({"net", "brut"} & sutun.keys()):
                continue
            ust = [str(c) for rr in satirlar[:i] for c in rr if c is not None]
            donem = _donem_bul(ust + [ws.title], dosya_adi)
            kisiler = []
            for rr in satirlar[i + 1:]:
                al = lambda a: rr[sutun[a]] if a in sutun and sutun[a] < len(rr) else None  # noqa: E731
                ad = str(al("personel") or "").strip()
                if not ad or "toplam" in kural.ad_anahtari(ad):
                    continue
                if "sira" in sutun and _sayi(al("sira")) is None:
                    continue  # sıra numarası olmayan satır (toplam/açıklama)
                net, brut = _sayi(al("net")), _sayi(al("brut"))
                if net is None and brut is None:
                    continue
                vergi = (_sayi(al("kalan_gv")) if "kalan_gv" in sutun else _sayi(al("gv")) or 0) or 0
                vergi += _sayi(al("damga")) or 0
                sgk = (_sayi(al("sgk")) or 0) + (_sayi(al("issizlik")) or 0)
                tc = re.sub(r"\D", "", str(al("tckn") or ""))
                gun = _sayi(al("gun"))
                kisiler.append({"personel": ad, "tckn": tc if len(tc) == 11 else None, "gun": int(gun) if gun else None,
                                "brut": _yuvarla(brut), "net": _yuvarla(net), "sgk_isci": round(sgk, 2),
                                "vergi": round(vergi, 2), "icra": _yuvarla(_sayi(al("icra"))),
                                "isveren": _yuvarla(_sayi(al("isveren")))})
            if kisiler:
                return {"donem": donem, "kisiler": kisiler, "sayfa": ws.title}
    raise HTTPException(422, "Dosyada personel listesi bulunamadı: 'Adı Soyadı' ve 'Net Ödenen' (ya da 'Brüt') "
                             "başlıklı sütunlar olmalı.")


@router.post("/taseron-giderleri/bordro-onizle")
async def bordro_onizle(dosya: UploadFile = File(...), k: Kullanici = Depends(yazabilir)):
    return bordro_oku(dosyayi_oku(dosya), dosya.filename)


class BordroKaydi(BaseModel):
    cari_id: int
    donem: str
    santiye_id: int | None = None
    depo_id: int | None = None
    sgk_orani: float = 37.75  # SGK masrafı = bordrodaki normal kazanç (brüt) x oran (işçi + işveren payları)
    kisiler: list[dict]


def bordro_satiri(x: dict, sgk_orani: float) -> dict:
    """Taşerona yazılacak: net ödenen + vergiler + icra + SGK masrafı (brüt x oran). Oran yoksa brüt + işveren payı.
    İcra, brütten net, işçi SGK ve vergi düşülünce kalandır (bordroda ayrı sütun olmasa da)."""
    brut, net = _yuvarla(x.get("brut")), _yuvarla(x.get("net"))
    vergi, sgk_bordro = float(x.get("vergi") or 0), float(x.get("sgk_isci") or 0)
    if not (sgk_orani and brut is not None and net is not None):
        return {"sgk": _yuvarla(x.get("sgk_isci")), "tutar": personel_toplami(type("G", (), x))}
    icra = max(0.0, round(brut - net - sgk_bordro - vergi, 2))
    sgk = round(brut * sgk_orani / 100, 2)
    return {"sgk": sgk, "tutar": round(net + vergi + icra + sgk, 2)}


@router.post("/taseron-giderleri/bordro")
def bordro_kaydet(istek: BordroKaydi, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Seçilen bordro satırları taşeron giderine yazılır. Aynı dönem ve aynı kişi ikinci kez yüklenirse eskisi güncellenir."""
    if not re.fullmatch(r"20\d\d-(0[1-9]|1[0-2])", istek.donem or ""):
        raise HTTPException(422, "Bordro ayını seçin.")
    yil, ay = map(int, istek.donem.split("-"))
    tarih = date(yil, ay, monthrange(yil, ay)[1])
    eski = {(g.personel or "", g.tckn or ""): g for g in db.scalars(select(TaseronGider).where(
        TaseronGider.sirket_id == k.sirket_id, TaseronGider.cari_id == istek.cari_id, TaseronGider.tur == "personel",
        TaseronGider.donem == istek.donem))}
    yeni = guncel = 0
    for x in istek.kisiler:
        hesap = bordro_satiri(x, istek.sgk_orani)
        veri = GiderGiris(cari_id=istek.cari_id, tur="personel", tarih=tarih, donem=istek.donem,
                          santiye_id=istek.santiye_id, depo_id=istek.depo_id,
                          personel=(x.get("personel") or "").strip(), tckn=x.get("tckn"), gun=x.get("gun"), brut=_yuvarla(x.get("brut")),
                          net=_yuvarla(x.get("net")), sgk_isci=hesap["sgk"], vergi=_yuvarla(x.get("vergi")), tutar=hesap["tutar"],
                          aciklama=("İcra kesintisi " + f"{float(x['icra']):,.2f}") if x.get("icra") else None)
        g = eski.get((veri.personel, veri.tckn or ""))
        if g is None:
            g = TaseronGider(sirket_id=k.sirket_id)
            db.add(g)
            yeni += 1
        else:
            guncel += 1
        _gider_yaz(db, g, veri, k)
    db.commit()
    return {"eklenen": yeni, "guncellenen": guncel}


# --- Taşeron sayfası --------------------------------------------------------

@router.get("/taseron/{cari_id}")
def taseron_sayfasi(cari_id: int, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Taşeronun dört sekmesi (verilen paralar, adına harcanan fatura kalemleri, elden harcamalar, personel/SGK) ve özet.
    Bakiye cari ekstreyle aynıdır: pozitifse taşeron bize borçlu, negatifse biz ona borçluyuz."""
    c = sirket_kaydi(db, Cari, cari_id, k)
    satirlar = cari_hareketleri(db, k.sirket_id, c.id)
    paralar = []
    for h in db.scalars(select(Hareket).where(Hareket.sirket_id == k.sirket_id, Hareket.cari_id == c.id,
                                              Hareket.durum == "onayli").order_by(Hareket.tarih, Hareket.id)):
        kaynak = h.banka_hesap.ad if h.banka_hesap else ("Çek" if h.cek_id and h.tur not in ("cek_tahsil", "cek_odeme") else "Kasa")
        paralar.append({"id": h.id, "tarih": h.tarih, "aciklama": h.aciklama, "tur": h.tur, "kaynak": kaynak,
                        "verilen": float(h.tutar) if h.yon == "cikis" else 0.0,
                        "alinan": float(h.tutar) if h.yon == "giris" else 0.0})
    faturalari = []
    for f in db.scalars(select(Fatura).where(Fatura.sirket_id == k.sirket_id, Fatura.cari_id == c.id,
                                             Fatura.durum == "onayli").order_by(Fatura.tarih, Fatura.id)):
        tutar = round(float(f.genel_toplam) * tl_kur(f) * isaret(f), 2)
        faturalari.append({"id": f.id, "tarih": f.tarih, "fatura_no": f.fatura_no,
                           "tur": "Satış" if satis_tarafi(f) else "Alış", "tutar": tutar})
    kalemler = depo_kalemleri(db, k.sirket_id, cari_id=c.id)
    adlar = {s.id: s.ad for s in db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id))}
    for x in kalemler:
        x["santiye"] = adlar.get(x["santiye_id"])
    giderler = [_gider_cikis(g) for g in db.scalars(select(TaseronGider).where(
        TaseronGider.sirket_id == k.sirket_id, TaseronGider.cari_id == c.id).order_by(TaseronGider.tarih, TaseronGider.id))]
    bakiye = round(sum(s["borc"] - s["alacak"] for s in satirlar), 2)
    topla = lambda dizi, a: round(sum(x[a] or 0 for x in dizi), 2)  # noqa: E731
    elden = [g for g in giderler if g["tur"] == "elden"]
    personel = [g for g in giderler if g["tur"] == "personel"]
    return {
        "cari": {"id": c.id, "unvan": c.unvan, "vkn": c.vkn},
        "depolar": [_depo_cikis(d) for d in db.scalars(select(Depo).where(Depo.sirket_id == k.sirket_id, Depo.cari_id == c.id))],
        "paralar": paralar, "faturalari": faturalari, "kalemler": kalemler, "elden": elden, "personel": personel,
        "ozet": {
            "faturalari_alis": round(sum(x["tutar"] for x in faturalari if x["tur"] == "Alış"), 2),
            "faturalari_satis": round(sum(x["tutar"] for x in faturalari if x["tur"] == "Satış"), 2),
            "paralar": round(topla(paralar, "verilen") - topla(paralar, "alinan"), 2),
            "kalemler": topla(kalemler, "toplam"), "elden": topla(elden, "tutar"), "personel": topla(personel, "tutar"),
            "bakiye": bakiye,
        },
    }


@router.get("/taseronlar")
def taseronlar(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Aktif bir depoya bağlı ya da taşeron gideri olan cariler (taşeron sayfası seçimi için)."""
    idler = set(db.scalars(select(Depo.cari_id).where(Depo.sirket_id == k.sirket_id, Depo.cari_id.is_not(None),
                                                      or_(Depo.durum.is_(None), Depo.durum != "pasif")))) | \
        set(db.scalars(select(TaseronGider.cari_id).where(TaseronGider.sirket_id == k.sirket_id)))
    return [{"id": c.id, "unvan": c.unvan} for c in sorted(db.scalars(select(Cari).where(Cari.id.in_(idler))),
                                                                  key=lambda _x: kural.tr_sira(_x.unvan))]


# --- Depo bazlı harcama raporu -----------------------------------------------

def depo_rapor_veri(db: Session, k: Kullanici, santiye_id: int | None, depo_id: int | None,
                    baslangic: date | None, bitis: date | None) -> list[dict]:
    adlar = {s.id: s.ad for s in db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id))}
    satirlar = []
    for x in depo_kalemleri(db, k.sirket_id, depo_id=depo_id, aktarim=True):
        if santiye_id and x["santiye_id"] != santiye_id:
            continue
        if (baslangic and x["tarih"] and x["tarih"] < baslangic) or (bitis and x["tarih"] and x["tarih"] > bitis):
            continue
        satirlar.append({**x, "santiye": adlar.get(x["santiye_id"], "Şantiyesiz"),
                         "kaynak": "Wolvox aktarımı" if x["aktarim"] else "Fatura"})
    q = select(TaseronGider).where(TaseronGider.sirket_id == k.sirket_id)
    if depo_id:
        q = q.where(TaseronGider.depo_id == depo_id)
    if santiye_id:
        q = q.where(TaseronGider.santiye_id == santiye_id)
    depolar_ = {d.id: d.ad for d in db.scalars(select(Depo).where(Depo.sirket_id == k.sirket_id))}
    for g in db.scalars(q):
        if not g.depo_id and not depo_id:
            continue  # depoya ayrılmamış gider bu raporda yok (taşeron sayfasında)
        if (baslangic and g.tarih < baslangic) or (bitis and g.tarih > bitis):
            continue
        satirlar.append({"depo_id": g.depo_id, "depo": depolar_.get(g.depo_id), "santiye_id": g.santiye_id,
                         "santiye": adlar.get(g.santiye_id, "Şantiyesiz"), "fatura_id": None, "tarih": g.tarih,
                         "fatura_no": g.donem, "satici": None, "kaynak": "Elden" if g.tur == "elden" else "Personel/SGK",
                         "kalem": (g.personel if g.tur == "personel" else g.aciklama) or "-", "miktar": None, "birim": None,
                         "tutar": float(g.tutar), "kdv": 0.0, "toplam": float(g.tutar)})
    satirlar.sort(key=lambda r: (r["santiye"], r["depo"] or "", r["tarih"] or date.min))
    return satirlar


@router.get("/raporlar/depo-harcamalari")
def depo_harcamalari(santiye_id: int | None = None, depo_id: int | None = None, baslangic: date | None = None,
                     bitis: date | None = None, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    return depo_rapor_veri(db, k, santiye_id, depo_id, baslangic, bitis)


@router.get("/excel/depo-harcamalari")
def excel_depo_harcamalari(santiye_id: int | None = None, depo_id: int | None = None, baslangic: date | None = None,
                           bitis: date | None = None, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    from .raporlar import _excel
    satirlar, kalin = [], set()
    veri = depo_rapor_veri(db, k, santiye_id, depo_id, baslangic, bitis)
    anahtar = None
    toplam = [0.0, 0.0, 0.0]
    for x in veri + [None]:
        yeni = (x["santiye"], x["depo"]) if x else None
        if anahtar and yeni != anahtar:
            kalin.add(len(satirlar))
            satirlar.append([anahtar[0], anahtar[1], None, None, None, None, "TOPLAM", None, None, *[round(t, 2) for t in toplam]])
            toplam = [0.0, 0.0, 0.0]
        if x is None:
            break
        anahtar = yeni
        toplam = [toplam[0] + x["tutar"], toplam[1] + x["kdv"], toplam[2] + x["toplam"]]
        satirlar.append([x["santiye"], x["depo"], x["tarih"], x["kaynak"], x["fatura_no"], x["satici"], x["kalem"],
                         x["miktar"], x["birim"], x["tutar"], x["kdv"], x["toplam"]])
    return _excel("Depo Harcamaları", ["Şantiye", "Depo", "Tarih", "Kaynak", "Fatura No / Dönem", "Satıcı", "Mal Hizmet / Açıklama",
                                       "Miktar", "Birim", "Tutar (KDV hariç)", "KDV", "Toplam"], satirlar, "depo-harcamalari", kalin)


@router.get("/excel/taseron/{cari_id}")
def excel_taseron(cari_id: int, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Taşeron sayfasının dört sekmesi tek Excel dosyasında, her biri ayrı sayfa."""
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    v = taseron_sayfasi(cari_id, db, k)
    wb = Workbook()
    sayfalar = [
        ("Özet", ["Kalem", "Tutar"], [["Taşeronun faturaları (alış)", v["ozet"]["faturalari_alis"]],
                                      ["Verilen paralar (net)", v["ozet"]["paralar"]],
                                      ["Adına harcanan fatura kalemleri", v["ozet"]["kalemler"]],
                                      ["Elden harcamalar", v["ozet"]["elden"]], ["Personel ücreti / SGK", v["ozet"]["personel"]],
                                      ["Bakiye (+ taşeron borçlu, - biz borçluyuz)", v["ozet"]["bakiye"]]]),
        ("Verilen Paralar", ["Tarih", "Nereden", "Açıklama", "Verilen", "Alınan"],
         [[x["tarih"], x["kaynak"], x["aciklama"], x["verilen"], x["alinan"]] for x in v["paralar"]]),
        ("Fatura Çıkışları", ["Tarih", "Fatura No", "Satıcı", "Şantiye", "Depo", "Mal Hizmet", "Miktar", "Birim", "Tutar", "KDV", "Toplam"],
         [[x["tarih"], x["fatura_no"], x["satici"], x["santiye"], x["depo"], x["kalem"], x["miktar"], x["birim"], x["tutar"],
           x["kdv"], x["toplam"]] for x in v["kalemler"]]),
        ("Elden Harcamalar", ["Tarih", "Açıklama", "Şantiye", "Tutar"],
         [[x["tarih"], x["aciklama"], x["santiye_ad"], x["tutar"]] for x in v["elden"]]),
        ("Personel SGK", ["Dönem", "Personel", "TC", "Gün", "Normal kazanç", "Net", "SGK masrafı", "Vergi", "İşveren payı", "Toplam"],
         [[x["donem"], x["personel"], x["tckn"], x["gun"], x["brut"], x["net"], x["sgk_isci"], x["vergi"], x["isveren"], x["tutar"]]
          for x in v["personel"]]),
    ]
    for i, (ad, basliklar, satirlar) in enumerate(sayfalar):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = ad
        ws.append(basliklar)
        for h in ws[1]:
            h.font = Font(bold=True)
        for s in satirlar:
            ws.append(s)
        for j in range(1, len(basliklar) + 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = 18
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    dosya = re.sub(r"[^A-Za-z0-9]+", "-", kural.sade(v["cari"]["unvan"]))[:40].strip("-") or "taseron"
    return StreamingResponse(tampon, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="taseron-{dosya}.xlsx"'})
