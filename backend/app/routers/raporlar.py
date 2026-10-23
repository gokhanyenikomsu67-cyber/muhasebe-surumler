"""Raporlar ve Excel aktarımı."""
import io
import re
from collections import defaultdict
from datetime import date

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import kural
from ..auth import aktif_kullanici, yazabilir
from ..db import get_db
from ..hesap import (RAPOR_DURUMLARI, banka_bakiyeleri, cari_bakiyeleri, cari_hareketleri, firma_adi, isaret, kasa_bakiyesi,
                     satis_tarafi, tl_kur)
from ..models import ACIK_CEK_DURUMLARI, CEK_DURUMLARI, HAREKET_TURLERI, BankaHesabi, Cari, Cek, Fatura, Hareket, Kullanici, Santiye, Sirket
from ..yukleme import dosyayi_oku
from .temel import sirket_kaydi

router = APIRouter(prefix="/api")


def _onayli(db: Session, k: Kullanici):
    return db.scalars(select(Fatura).where(Fatura.sirket_id == k.sirket_id, Fatura.durum == "onayli")).all()


# --- Raporlar ------------------------------------------------------------

def santiye_paylari(f: Fatura) -> dict[int | None, dict]:
    """Faturanın hangi şantiyeye ne kadar düştüğü (TL karşılığıyla). Kalemin şantiyesi boşsa faturanın şantiyesi
    geçerlidir; tevkifat ve diğer vergiler kalem tutarı oranında paylaştırılır."""
    kur = tl_kur(f)
    if not f.kalemler:
        return {f.santiye_id: {"ara": round(float(f.ara_toplam) * kur, 2), "kdv": round(float(f.kdv_toplam) * kur, 2),
                               "genel": round(float(f.genel_toplam) * kur, 2), "kalemler": []}}
    ara_toplam = sum(float(kl.tutar) for kl in f.kalemler)
    ek = float(f.genel_toplam) - ara_toplam - sum(kl.kdv for kl in f.kalemler)
    paylar: dict[int | None, dict] = {}
    for kl in f.kalemler:
        p = paylar.setdefault(kl.santiye_id or f.santiye_id, {"ara": 0.0, "kdv": 0.0, "genel": 0.0, "kalemler": [], "detay": []})
        tutar, kdv = float(kl.tutar), kl.kdv
        oran = tutar / ara_toplam if ara_toplam else 1 / len(f.kalemler)
        p["ara"] += tutar
        p["kdv"] += kdv
        p["genel"] += tutar + kdv + ek * oran
        p["kalemler"].append(kl.aciklama)
        # Harcama dökümü için kalemin kendisi (TL karşılığıyla)
        p["detay"].append({"aciklama": kl.aciklama, "miktar": float(kl.miktar or 0), "birim": kl.birim,
                           "depo": kl.depo.ad if kl.depo else None,
                           "tutar": round(tutar * kur, 2), "kdv": round(kdv * kur, 2), "toplam": round((tutar + kdv + ek * oran) * kur, 2)})
    return {s: {**p, "ara": round(p["ara"] * kur, 2), "kdv": round(p["kdv"] * kur, 2), "genel": round(p["genel"] * kur, 2)}
            for s, p in paylar.items()}


def santiye_ozet_veri(db: Session, k: Kullanici) -> list[dict]:
    faturalar = [(f, santiye_paylari(f)) for f in db.scalars(select(Fatura).where(
        Fatura.sirket_id == k.sirket_id, Fatura.durum.in_(RAPOR_DURUMLARI)))]
    masraflar = db.scalars(select(Hareket).where(
        Hareket.sirket_id == k.sirket_id, Hareket.durum == "onayli", Hareket.santiye_id.is_not(None),
        Hareket.tur.in_(["masraf", "banka_masrafi"]))).all()
    sonuc = []
    for s in db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id).order_by(Santiye.ad)):
        gider = gelir = 0.0
        turler: dict[str, float] = defaultdict(float)
        for f, paylar in faturalar:
            if s.id not in paylar:
                continue
            tutar = paylar[s.id]["genel"] * isaret(f)  # iade faturası ilgili tarafı azaltır
            if not satis_tarafi(f):
                gider += tutar
                turler[f.gider_turu or "Belirtilmemiş"] += tutar
            else:
                gelir += tutar
        for h in masraflar:
            if h.santiye_id == s.id:
                gider += float(h.tutar)
                turler[h.gider_turu or "Belirtilmemiş"] += float(h.tutar)
        butce = float(s.butce) if s.butce is not None else None
        sonuc.append({
            "id": s.id, "ad": s.ad, "durum": s.durum, "butce": butce,
            "gider": round(gider, 2), "gelir": round(gelir, 2), "kar": round(gelir - gider, 2),
            "butce_kalan": round(butce - gider, 2) if butce is not None else None,
            "gider_turleri": {t: round(v, 2) for t, v in sorted(turler.items(), key=lambda x: -x[1])},
        })
    return sonuc


def santiye_fatura_veri(db: Session, k: Kullanici, yon: str | None, baslangic: date | None, bitis: date | None,
                        santiye_ids: str | None, durum: str | None = "onayli") -> list[dict]:
    """Şantiye bazlı fatura listesi: her satır bir faturanın bir şantiyeye düşen payıdır (kalem bazında).
    durum="onayli": onaylılarla birlikte Wolvox'tan aktarılan geçmiş faturalar."""
    secili = {int(x) for x in (santiye_ids or "").split(",") if x.strip().isdigit()}
    q = select(Fatura).where(Fatura.sirket_id == k.sirket_id)
    if durum:
        q = q.where(Fatura.durum.in_(RAPOR_DURUMLARI) if durum == "onayli" else Fatura.durum == durum)
    if yon:
        q = q.where(Fatura.yon == yon)
    if baslangic:
        q = q.where(Fatura.tarih >= baslangic)
    if bitis:
        q = q.where(Fatura.tarih <= bitis)
    adlar = {s.id: s.ad for s in db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id))}
    satirlar = []
    for f in db.scalars(q):
        for sid, p in santiye_paylari(f).items():
            if secili and sid not in secili:
                continue
            i = isaret(f)
            tur = ("Satış iadesi" if f.yon == "gelen" else "Alış iadesi") if i < 0 else "Alış" if f.yon == "gelen" else "Satış"
            satirlar.append({"santiye_id": sid, "santiye": adlar.get(sid, "Şantiyesiz"), "fatura_id": f.id, "tarih": f.tarih,
                             "fatura_no": f.fatura_no, "tur": tur,
                             "cari": firma_adi(f), "gider_turu": f.gider_turu, "aktarim": f.durum == "aktarim",
                             "kalemler": ", ".join(p["kalemler"]),
                             "kalem_detay": [{**d, "tutar": i * d["tutar"], "kdv": i * d["kdv"], "toplam": i * d["toplam"]}
                                             for d in p.get("detay", [])],
                             "ara_toplam": i * p["ara"], "kdv": i * p["kdv"],
                             "genel_toplam": i * p["genel"], "fatura_toplami": i * round(float(f.genel_toplam) * tl_kur(f), 2)})
    satirlar.sort(key=lambda r: (r["santiye"] == "Şantiyesiz", r["santiye"], r["tarih"] or date.min, r["fatura_id"]))
    return satirlar


@router.get("/raporlar/santiye-faturalari")
def santiye_faturalari(yon: str | None = None, baslangic: date | None = None, bitis: date | None = None,
                       santiye_ids: str | None = None, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    return santiye_fatura_veri(db, k, yon, baslangic, bitis, santiye_ids)


@router.get("/raporlar/santiye-ozet")
def santiye_ozet(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    return santiye_ozet_veri(db, k)


def cari_bakiye_veri(db: Session, k: Kullanici) -> list[dict]:
    toplam = cari_bakiyeleri(db, k.sirket_id)
    sonuc = []
    for c in db.scalars(select(Cari).where(Cari.sirket_id == k.sirket_id).order_by(Cari.unvan)):
        t = toplam.get(c.id, {"borc": 0.0, "alacak": 0.0})
        # Pozitif bakiye: cari bize borçlu. Negatif: biz cariye borçluyuz.
        sonuc.append({"id": c.id, "unvan": c.unvan, "vkn": c.vkn, "borc": round(t["borc"], 2),
                      "alacak": round(t["alacak"], 2), "bakiye": round(t["borc"] - t["alacak"], 2)})
    return sonuc


@router.get("/raporlar/cari-bakiye")
def cari_bakiye(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    return cari_bakiye_veri(db, k)


@router.get("/raporlar/cari-ekstre/{cari_id}")
def cari_ekstre(cari_id: int, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    c = sirket_kaydi(db, Cari, cari_id, k)
    satirlar = sorted(cari_hareketleri(db, k.sirket_id, c.id), key=lambda s: (s["tarih"] or date.min, s["sira"]))
    bakiye = 0.0
    for s in satirlar:
        bakiye += s["borc"] - s["alacak"]
        s["bakiye"] = round(bakiye, 2)
        del s["sira"], s["cari_id"]
    return {"cari": c.unvan, "satirlar": satirlar, "bakiye": round(bakiye, 2)}


@router.get("/raporlar/ozet")
def genel_ozet(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Gösterge paneli: banka bakiyeleri, yaklaşan çekler, onay bekleyenler."""
    bugun = date.today()
    cekler = db.scalars(select(Cek).where(Cek.sirket_id == k.sirket_id, Cek.durum.in_(ACIK_CEK_DURUMLARI))).all()
    yaklasan = lambda yon: [c for c in cekler if c.yon == yon and (c.vade - bugun).days <= 7]  # noqa: E731
    bakiyeler = cari_bakiye_veri(db, k)
    kart_idleri = set(db.scalars(select(BankaHesabi.id).where(BankaHesabi.sirket_id == k.sirket_id, BankaHesabi.tur == "kredi_karti")))
    banka = banka_bakiyeleri(db, k.sirket_id)
    return {
        "banka_toplam": round(sum(v for i, v in banka.items() if i not in kart_idleri), 2),
        "kart_borcu": round(-sum(v for i, v in banka.items() if i in kart_idleri), 2),
        "kasa_bakiye": kasa_bakiyesi(db, k.sirket_id),
        "portfoy_cek": round(sum(float(c.tutar) for c in cekler if c.yon == "alinan"), 2),
        "odenecek_cek": round(sum(float(c.tutar) for c in cekler if c.yon == "verilen"), 2),
        "yaklasan_alinan": len(yaklasan("alinan")),
        "yaklasan_verilen": len(yaklasan("verilen")),
        "yaklasan_verilen_tutar": round(sum(float(c.tutar) for c in yaklasan("verilen")), 2),
        "alacak_toplam": round(sum(b["bakiye"] for b in bakiyeler if b["bakiye"] > 0), 2),
        "borc_toplam": round(-sum(b["bakiye"] for b in bakiyeler if b["bakiye"] < 0), 2),
        "taslak_fatura": db.query(Fatura).filter(Fatura.sirket_id == k.sirket_id, Fatura.durum == "taslak").count(),
        "taslak_hareket": db.query(Hareket).filter(Hareket.sirket_id == k.sirket_id, Hareket.durum == "taslak").count(),
    }


# --- Excel ---------------------------------------------------------------

BASLIK_STILI = Font(bold=True, color="FFFFFF")
BASLIK_DOLGU = PatternFill("solid", fgColor="1F4E79")
ARA_TOPLAM_DOLGU = PatternFill("solid", fgColor="DDEBF7")


def _excel(sayfa_adi: str, basliklar: list[str], satirlar: list[list], dosya_adi: str,
           kalin: set[int] | None = None, dipnot: str | None = None) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = sayfa_adi[:31]
    ws.append(basliklar)
    for hucre in ws[1]:
        hucre.font, hucre.fill = BASLIK_STILI, BASLIK_DOLGU
    for i, s in enumerate(satirlar):
        ws.append(s)
        if kalin and i in kalin:  # ara toplam satırları
            for hucre in ws[ws.max_row]:
                hucre.font, hucre.fill = Font(bold=True), ARA_TOPLAM_DOLGU
    for i, b in enumerate(basliklar, start=1):
        uzunluk = max([len(str(b))] + [len(str(s[i - 1])) for s in satirlar if len(s) >= i and s[i - 1] is not None])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(uzunluk + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if dipnot:  # süzgecin dışında, tablonun altında
        ws.append([])
        ws.append([dipnot])
        ws.cell(ws.max_row, 1).font = Font(bold=True, color="B00020")
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return StreamingResponse(
        tampon, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{dosya_adi}.xlsx"'},
    )


CARI_SUTUNLARI = [("unvan", "Ünvan"), ("tur", "Tür"), ("vkn", "VKN/TCKN"), ("vergi_dairesi", "Vergi Dairesi"),
                  ("adres", "Adres"), ("telefon", "Telefon"), ("eposta", "E-posta"), ("iban", "IBAN")]


@router.get("/excel/cariler")
def excel_cariler(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    cariler = db.scalars(select(Cari).where(Cari.sirket_id == k.sirket_id).order_by(Cari.unvan)).all()
    return _excel("Cariler", [b for _, b in CARI_SUTUNLARI], [[getattr(c, a) for a, _ in CARI_SUTUNLARI] for c in cariler], "cariler")


@router.get("/excel/cari-sablon")
def excel_cari_sablon(k: Kullanici = Depends(aktif_kullanici)):
    return _excel("Cariler", [b for _, b in CARI_SUTUNLARI], [], "cari-yukleme-sablonu")


@router.post("/excel/cariler/yukle")
async def excel_cari_yukle(dosya: UploadFile = File(...), db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    icerik = dosyayi_oku(dosya)
    try:
        ws = load_workbook(io.BytesIO(icerik), read_only=True).active
    except Exception as e:
        raise HTTPException(400, "Excel dosyası okunamadı.") from e
    satirlar = ws.iter_rows(values_only=True)
    basliklar = [str(h or "").strip() for h in next(satirlar, [])]
    harita = {b: a for a, b in CARI_SUTUNLARI}
    if "Ünvan" not in basliklar:
        raise HTTPException(400, "İlk satırda 'Ünvan' sütunu bulunamadı. Şablonu kullanın.")
    from .temel import ayni_adli_cari
    rakam = lambda x: re.sub(r"\D", "", str(x or ""))  # noqa: E731
    kendi_vkn = rakam(db.get(Sirket, k.sirket_id).vkn)
    mevcut_vkn = {rakam(v) for v in db.scalars(select(Cari.vkn).where(Cari.sirket_id == k.sirket_id, Cari.vkn.is_not(None)))}
    eklenen = atlanan = 0
    for satir in satirlar:
        veri = {harita[b]: (str(v).strip() or None if v is not None else None) for b, v in zip(basliklar, satir) if b in harita}
        vkn = rakam(veri.get("vkn"))
        # Boş ünvan, kendi firmamız, aynı VKN'li ya da aynı adlı cari eklenmez (ikinci kez açılmaz)
        if not veri.get("unvan") or (vkn and (vkn in mevcut_vkn or vkn == kendi_vkn)) \
                or ayni_adli_cari(db, k.sirket_id, veri["unvan"], veri.get("vkn")) is not None:
            atlanan += 1
            continue
        veri["tur"] = veri.get("tur") if veri.get("tur") in ("musteri", "tedarikci", "her_ikisi") else "tedarikci"
        db.add(Cari(sirket_id=k.sirket_id, **veri))
        db.flush()
        if vkn:
            mevcut_vkn.add(vkn)
        eklenen += 1
    db.commit()
    return {"eklenen": eklenen, "atlanan": atlanan}


@router.get("/excel/faturalar")
def excel_faturalar(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    faturalar = db.scalars(select(Fatura).where(Fatura.sirket_id == k.sirket_id, Fatura.durum != "aktarim")
                           .order_by(Fatura.tarih)).all()
    basliklar = ["Tarih", "Fatura No", "Tür", "Durum", "Cari", "VKN", "Şantiye", "Gider Türü",
                 "Ara Toplam", "KDV", "Tevkifat", "Stopaj", "Diğer Vergiler", "Genel Toplam", "Para Birimi", "Fatura Tipi", "Kur",
                 "Genel Toplam (TL)"]
    satirlar = [[f.tarih, f.fatura_no, "Alış" if f.yon == "gelen" else "Satış", {"onayli": "Onaylı", "iptal": "İptal"}.get(f.durum, "Taslak"),
                 f.cari.unvan if f.cari else None, f.cari.vkn if f.cari else None, f.santiye.ad if f.santiye else None,
                 f.gider_turu, float(f.ara_toplam), float(f.kdv_toplam), float(f.tevkifat), float(f.stopaj or 0), float(f.diger_vergiler or 0), float(f.genel_toplam), f.para_birimi,
                 f.fatura_tipi, float(f.kur) if f.kur else None, round(float(f.genel_toplam) * tl_kur(f), 2)]
                for f in faturalar]
    return _excel("Faturalar", basliklar, satirlar, "faturalar")


@router.get("/excel/fatura-uyarilari")
def excel_fatura_uyarilari(yon: str | None = None, durum: str | None = "taslak", db: Session = Depends(get_db),
                           k: Kullanici = Depends(aktif_kullanici)):
    """Okunan faturalardaki uyarıların raporu: hangi dosyada ne sorun var."""
    q = select(Fatura).where(Fatura.sirket_id == k.sirket_id)
    if yon:
        q = q.where(Fatura.yon == yon)
    if durum:
        q = q.where(Fatura.durum == durum)
    satirlar = []
    for f in db.scalars(q.order_by(Fatura.tarih, Fatura.id)):
        o = f.ai_okuma or {}
        # Eski sürümler bilgi notlarını da uyarı olarak saklıyordu; onlar rapora girmez.
        uyarilar = [u for u in o.get("uyarilar") or [] if u and not re.search(r"okundu\.?$|yeni cari açıldı", u)]
        islenemedi = not f.kalemler or not float(f.genel_toplam) or not f.cari_id
        if not uyarilar and not islenemedi:
            continue
        satirlar.append(["İşlenemedi" if islenemedi else "Uyarı", f.tarih, f.fatura_no, f.cari.unvan if f.cari else None,
                         f.cari.vkn if f.cari else None, o.get("dosya_adi"), len(f.kalemler), float(f.ara_toplam),
                         float(f.kdv_toplam), float(f.genel_toplam), o.get("genel_toplam"), " | ".join(uyarilar)])
    satirlar.sort(key=lambda r: (r[0] != "İşlenemedi", r[1] or date.min))
    return _excel("Fatura Uyarıları", ["Durum", "Tarih", "Fatura No", "Cari", "VKN", "Dosya", "Kalem", "Ara Toplam", "KDV",
                                       "Genel Toplam", "Faturada Yazan Toplam", "Uyarılar"], satirlar, "fatura-uyari-raporu")


@router.get("/excel/santiye-faturalari")
def excel_santiye_faturalari(yon: str | None = None, baslangic: date | None = None, bitis: date | None = None,
                             santiye_ids: str | None = None, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Seçilen şantiyelerin (ve ayın) bütün harcamaları kalem kalem, şantiye şantiye gruplu ve ara toplamlı."""
    veri = santiye_fatura_veri(db, k, yon, baslangic, bitis, santiye_ids)
    # Her satır bir kalem (bütün harcamalar); kalemsiz faturada fatura tek satır
    basliklar = ["Şantiye", "Depo", "Tarih", "Fatura No", "Tür", "Cari", "Gider Türü", "Mal Hizmet Adı", "Miktar", "Birim",
                 "Tutar (KDV hariç)", "KDV", "Toplam"]
    satirlar, kalin = [], set()
    toplam = [0.0, 0.0, 0.0]

    def ara_toplam(ad, t):
        kalin.add(len(satirlar))
        satirlar.append([f"{ad} TOPLAMI", None, None, None, None, None, None, None, None, None, round(t[0], 2), round(t[1], 2), round(t[2], 2)])

    grup, t = None, [0.0, 0.0, 0.0]
    for r in veri:
        if grup is not None and r["santiye"] != grup:
            ara_toplam(grup, t)
            t = [0.0, 0.0, 0.0]
        grup = r["santiye"]
        kalemler = r["kalem_detay"] or [{"aciklama": r["kalemler"] or None, "miktar": None, "birim": None,
                                         "tutar": r["ara_toplam"], "kdv": r["kdv"], "toplam": r["genel_toplam"]}]
        for d in kalemler:
            satirlar.append([r["santiye"], d.get("depo"), r["tarih"], r["fatura_no"], r["tur"], r["cari"], r["gider_turu"], d["aciklama"],
                             d["miktar"], d["birim"], d["tutar"], d["kdv"], d["toplam"]])
        for i, a in enumerate(("ara_toplam", "kdv", "genel_toplam")):
            t[i] += r[a]
            toplam[i] += r[a]
    if grup is not None:
        ara_toplam(grup, t)
        ara_toplam("GENEL", toplam)
    ek = "-".join(str(x) for x in (baslangic, bitis) if x)
    return _excel("Şantiye Faturaları", basliklar, satirlar, f"santiye-faturalari{'-' + ek if ek else ''}", kalin)


@router.get("/excel/fatura-kalemleri")
def excel_fatura_kalemleri(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Wolvox'a aktarım için satır bazında onaylı faturalar (geçici sütun düzeni)."""
    basliklar = ["Fatura Tarihi", "Fatura No", "Fatura Türü", "Cari Ünvan", "Cari VKN", "Şantiye (Proje)", "Depo",
                 "Sıra No", "Mal Hizmet Adı", "Miktar", "Birim", "Birim Fiyat", "İskonto Oranı", "İskonto Tutarı",
                 "KDV Oranı", "KDV Tutarı", "Mal Hizmet Tutarı",
                 "Tevkifat Oranı", "Tevkifat Kodu", "Tevkifat Tutarı", "Stopaj Oranı", "Stopaj Tutarı"]
    satirlar = []
    for f in sorted(_onayli(db, k), key=lambda f: (f.tarih or date.min, f.id)):
        for sira, kl in enumerate(f.kalemler, 1):
            satirlar.append([f.tarih, f.fatura_no, "Alış" if f.yon == "gelen" else "Satış",
                             f.cari.unvan if f.cari else None, f.cari.vkn if f.cari else None,
                             (kl.santiye or f.santiye).ad if (kl.santiye or f.santiye) else None, kl.depo.ad if kl.depo else None,
                             sira, kl.aciklama, float(kl.miktar),
                             kl.birim, float(kl.birim_fiyat), float(kl.iskonto_orani or 0), float(kl.iskonto_tutari or 0),
                             float(kl.kdv_orani), round(kl.kdv, 2), float(kl.tutar),
                             f"{float(kl.tevkifat_orani):g}/10" if float(kl.tevkifat_orani or 0) else None, kl.tevkifat_kodu,
                             kural.tevkifat_tutari(kl.tutar, kl.kdv_orani, kl.tevkifat_orani, kl.kdv_tutari) or None,
                             f"%{float(kl.stopaj_orani):g}" if float(kl.stopaj_orani or 0) else None,
                             kural.stopaj_tutari(kl.tutar, kl.stopaj_orani) or None])
    return _excel("Fatura Kalemleri", basliklar, satirlar, "fatura-kalemleri")


@router.get("/excel/santiye-ozet")
def excel_santiye_ozet(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    veri = santiye_ozet_veri(db, k)
    return _excel("Şantiye Özeti", ["Şantiye", "Durum", "Bütçe", "Gider", "Gelir", "Kâr", "Bütçeden Kalan"],
                  [[s["ad"], s["durum"], s["butce"], s["gider"], s["gelir"], s["kar"], s["butce_kalan"]] for s in veri],
                  "santiye-ozet")


@router.get("/excel/cari-bakiye")
def excel_cari_bakiye(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    veri = cari_bakiye_veri(db, k)
    return _excel("Cari Bakiyeler", ["Ünvan", "VKN", "Borç", "Alacak", "Bakiye"],
                  [[c["unvan"], c["vkn"], c["borc"], c["alacak"], c["bakiye"]] for c in veri], "cari-bakiye")


def cek_raporu_sorgu(db: Session, k: Kullanici, yon: str | None = None, durum: str | None = None, cari_id: int | None = None,
                     ciro_cari_id: int | None = None, banka_hesap_id: int | None = None, vade_bas: date | None = None,
                     vade_bit: date | None = None, santiye_id: int | None = None) -> list[Cek]:
    """Çek/Senet Raporları: alınan ve verilen bütün çekler, süzgeçlerle. durum 'acik' = henüz kapanmamışlar."""
    q = select(Cek).where(Cek.sirket_id == k.sirket_id)
    if yon:
        q = q.where(Cek.yon == yon)
    if durum:
        q = q.where(Cek.durum.in_(ACIK_CEK_DURUMLARI if durum == "acik" else durum.split(",")))
    for sutun, deger in ((Cek.cari_id, cari_id), (Cek.ciro_cari_id, ciro_cari_id), (Cek.banka_hesap_id, banka_hesap_id),
                         (Cek.santiye_id, santiye_id)):
        if deger:
            q = q.where(sutun == deger)
    if vade_bas:
        q = q.where(Cek.vade >= vade_bas)
    if vade_bit:
        q = q.where(Cek.vade <= vade_bit)
    return list(db.scalars(q.order_by(Cek.vade, Cek.id)))


@router.get("/raporlar/cekler")
def cek_raporu(yon: str | None = None, durum: str | None = None, cari_id: int | None = None, ciro_cari_id: int | None = None,
               banka_hesap_id: int | None = None, vade_bas: date | None = None, vade_bit: date | None = None,
               santiye_id: int | None = None, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    from .cekler import cikis
    cekler = cek_raporu_sorgu(db, k, yon, durum, cari_id, ciro_cari_id, banka_hesap_id, vade_bas, vade_bit, santiye_id)
    satirlar = [cikis(c).model_dump() | {"santiye_ad": c.santiye.ad if c.santiye else None} for c in cekler]
    ozet = {y: {"adet": sum(1 for c in cekler if c.yon == y), "tutar": round(sum(float(c.tutar) for c in cekler if c.yon == y), 2),
                "acik": round(sum(float(c.tutar) for c in cekler if c.yon == y and c.durum in ACIK_CEK_DURUMLARI), 2)}
            for y in ("alinan", "verilen")}
    return {"satirlar": satirlar, "ozet": ozet}


@router.get("/excel/cekler")
def excel_cekler(yon: str | None = None, durum: str | None = None, cari_id: int | None = None, ciro_cari_id: int | None = None,
                 banka_hesap_id: int | None = None, vade_bas: date | None = None, vade_bit: date | None = None,
                 santiye_id: int | None = None, db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    cekler = cek_raporu_sorgu(db, k, yon, durum, cari_id, ciro_cari_id, banka_hesap_id, vade_bas, vade_bit, santiye_id)
    bugun = date.today()
    satirlar = [["Alınan" if c.yon == "alinan" else "Verilen", "Senet" if c.tur == "senet" else "Çek",
                 CEK_DURUMLARI[c.yon].get(c.durum, c.durum), c.alis_tarihi, c.vade,
                 (c.vade - c.alis_tarihi).days if c.alis_tarihi else None,
                 (c.vade - bugun).days if c.durum in ACIK_CEK_DURUMLARI else None, c.islem_tarihi,
                 c.cek_no, c.banka, c.sube, c.kesideci, c.kesideci_vkn, c.lehtar, c.cari.unvan if c.cari else None, c.ciro_cari.unvan if c.ciro_cari else None,
                 c.banka_hesap.ad if c.banka_hesap else None, c.santiye.ad if c.santiye else None, c.ozel_kod,
                 float(c.tutar), c.para_birimi] for c in cekler]
    kalin = set()
    for y, ad in (("alinan", "Alınan çekler toplamı"), ("verilen", "Verilen çekler toplamı")):
        if any(c.yon == y for c in cekler):
            kalin.add(len(satirlar))
            satirlar.append([ad] + [None] * 18 + [round(sum(float(c.tutar) for c in cekler if c.yon == y), 2), None])
    return _excel("Çekler", ["Yön", "Tür", "Durum", "Alış / Veriliş", "Vade", "Vade Günü", "Kalan Gün", "İşlem / Ödendiği Tarih",
                             "Çek No", "Banka", "Şube", "Çeki Yazan Firma", "Çeki Yazan VKN", "Lehtar (çekin yazıldığı firma)", "Kimden Alındı / Kime Verildi", "Ciro Edilen", "Banka Hesabı", "Şantiye", "Özel Kod",
                             "Tutar", "Para Birimi"], satirlar, "cekler", kalin=kalin)


@router.get("/excel/hareketler")
def excel_hareketler(db: Session = Depends(get_db), k: Kullanici = Depends(aktif_kullanici)):
    """Onaylı tahsilat/ödeme/masraf hareketleri (Wolvox'a aktarım için de kullanılır)."""
    hareketler = db.scalars(select(Hareket).where(Hareket.sirket_id == k.sirket_id, Hareket.durum == "onayli")
                            .order_by(Hareket.tarih, Hareket.id)).all()
    return _excel("Hareketler", ["Tarih", "Tür", "Giriş", "Çıkış", "Cari", "Cari VKN", "Banka Hesabı", "Şantiye",
                                 "Gider Türü", "Açıklama"],
                  [[h.tarih, HAREKET_TURLERI.get(h.tur, h.tur), float(h.tutar) if h.yon == "giris" else None,
                    float(h.tutar) if h.yon == "cikis" else None, h.cari.unvan if h.cari else None,
                    h.cari.vkn if h.cari else None, h.banka_hesap.ad if h.banka_hesap else None,
                    h.santiye.ad if h.santiye else None, h.gider_turu, h.aciklama] for h in hareketler], "hareketler")
