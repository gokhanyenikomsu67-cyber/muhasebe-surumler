"""Ana Depo: alış faturalarına girilen bütün malzemeler ana depoya giriş sayılır; kalemde (ya da faturada) şantiye
yazılıysa o malzeme o şantiyeye çıkmış sayılır, şantiyesi olmayan kalem ana depoda kalır. Stok sayımı değildir:
faturalardan kurulan giriş/çıkış görünümüdür. Adı "ANA DEPO" / "MERKEZ DEPO" gibi olan şantiye ana deponun kendisidir.

Ana depoda bekleyen kalemler buradan şantiyeye çıkarılır (kalemin şantiyesi yazılır), yanlış çıkış ana depoya geri alınır.
"""
import io
from collections import defaultdict
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import kural, malzeme
from ..auth import aktif_kullanici, yazabilir
from ..db import get_db
from ..hesap import firma_adi, isaret, satis_tarafi, tl_kur
from ..models import Depo, Fatura, FaturaKalem, Kullanici, Santiye
from .temel import sirket_kaydi

router = APIRouter(prefix="/api")

ANA_DEPO_ADI = "ANA DEPO"


def ana_depo_mu(s: Santiye) -> bool:
    """Şantiye ana deponun kendisi mi ("ANA DEPO", "MERKEZ DEPO")."""
    for ad in kural.kayit_adlari(s):
        kelimeler = set(kural.ad_anahtari(ad).split())
        if "depo" in kelimeler and kelimeler & {"ana", "merkez"}:
            return True
    return False


def _birim(b: str | None) -> str:
    from .malzemeler import _birim as birim
    return birim(b)


def satirlar(db: Session, k: Kullanici, baslangic: date | None = None, bitis: date | None = None) -> list[dict]:
    """Alış faturası kalemleri (onaylı ve Wolvox'tan aktarılan): her biri ana depoya giriş; şantiyesi varsa çıkış."""
    dizin, _ = malzeme.guncelle(db, k.sirket_id)
    santiyeler = {s.id: s for s in db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id))}
    ana = {i for i, s in santiyeler.items() if ana_depo_mu(s)}
    sonuc = []
    for kl in malzeme.onayli_kalemler(db, k.sirket_id):
        f = kl.fatura
        if satis_tarafi(f):
            continue  # satış faturası depoya giriş değildir (tedarikçiye iade aşağıda eksi giriş olur)
        if (baslangic and f.tarih and f.tarih < baslangic) or (bitis and f.tarih and f.tarih > bitis):
            continue
        sid = kl.santiye_id or f.santiye_id
        cikti = bool(sid) and sid not in ana
        m = dizin.bul(kl.aciklama)
        i, kur = isaret(f), tl_kur(f)
        miktar = i * float(kl.miktar or 0)
        tutar = round(i * float(kl.tutar or 0) * kur, 2)
        sonuc.append({
            "kalem_id": kl.id, "fatura_id": f.id, "tarih": f.tarih, "fatura_no": f.fatura_no, "firma": firma_adi(f),
            "wolvox": f.durum == "aktarim", "malzeme_id": m.id if m else None,
            "malzeme": m.ad if m else " ".join((kl.aciklama or "").split()), "kalem": kl.aciklama,
            "miktar": miktar, "birim": _birim(kl.birim), "birim_fiyat": round(tutar / miktar, 4) if miktar else None, "tutar": tutar,
            "santiye_id": sid if cikti else None, "santiye": santiyeler[sid].ad if cikti else None,
            "depo": kl.depo.ad if kl.depo else None, "yer": "santiye" if cikti else "ana"})
    sonuc.sort(key=lambda x: (x["tarih"] or date.min, x["fatura_id"], x["kalem_id"]), reverse=True)
    return sonuc


def ozetler(satir: list[dict]) -> tuple[list[dict], list[dict]]:
    """(malzeme özeti: giren / şantiyelere çıkan / ana depoda kalan, şantiye-malzeme çıkışları)."""
    mz: dict = defaultdict(lambda: {"giren_miktar": 0.0, "giren_tutar": 0.0, "cikan_miktar": 0.0, "cikan_tutar": 0.0,
                                    "santiyeler": set()})
    sm: dict = defaultdict(lambda: {"miktar": 0.0, "tutar": 0.0, "kalem_sayisi": 0, "son_tarih": None})
    for x in satir:
        g = mz[(x["malzeme_id"] or x["malzeme"], x["birim"])]
        g.update(malzeme_id=x["malzeme_id"], malzeme=x["malzeme"], birim=x["birim"])
        g["giren_miktar"] += x["miktar"]
        g["giren_tutar"] += x["tutar"]
        if x["yer"] == "santiye":
            g["cikan_miktar"] += x["miktar"]
            g["cikan_tutar"] += x["tutar"]
            g["santiyeler"].add(x["santiye"])
            s = sm[(x["santiye_id"], x["malzeme_id"] or x["malzeme"], x["birim"])]
            s.update(santiye_id=x["santiye_id"], santiye=x["santiye"], malzeme_id=x["malzeme_id"], malzeme=x["malzeme"], birim=x["birim"])
            s["miktar"] += x["miktar"]
            s["tutar"] += x["tutar"]
            s["kalem_sayisi"] += 1
            s["son_tarih"] = max(filter(None, [s["son_tarih"], x["tarih"]]), default=None)
    ozet = [{**{a: g[a] for a in ("malzeme_id", "malzeme", "birim")},
             "giren_miktar": round(g["giren_miktar"], 4), "giren_tutar": round(g["giren_tutar"], 2),
             "cikan_miktar": round(g["cikan_miktar"], 4), "cikan_tutar": round(g["cikan_tutar"], 2),
             "kalan_miktar": round(g["giren_miktar"] - g["cikan_miktar"], 4),
             "kalan_tutar": round(g["giren_tutar"] - g["cikan_tutar"], 2), "santiye_sayisi": len(g["santiyeler"])}
            for g in mz.values()]
    ozet.sort(key=lambda x: (kural.sade(x["malzeme"]), x["birim"]))
    cikis = [{**s, "miktar": round(s["miktar"], 4), "tutar": round(s["tutar"], 2)} for s in sm.values()]
    cikis.sort(key=lambda x: (kural.sade(x["santiye"]), kural.sade(x["malzeme"])))
    return ozet, cikis


@router.get("/ana-depo")
def ana_depo(baslangic: date | None = None, bitis: date | None = None, db: Session = Depends(get_db),
             k: Kullanici = Depends(aktif_kullanici)):
    s = satirlar(db, k, baslangic, bitis)
    ozet, cikis = ozetler(s)
    return {"ozet": ozet, "cikislar": cikis, "satirlar": s}


class Cikis(BaseModel):
    kalem_idler: list[int]
    santiye_id: int | None = None  # boş: ana depoya geri al
    depo_id: int | None = None  # şantiyedeki taşeron deposu (isteğe bağlı)


@router.post("/ana-depo/sevk")
def sevk(istek: Cikis, db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Seçilen kalemleri şantiyeye çıkarır (kalemin şantiyesi yazılır) ya da ana depoya geri alır. Kalem şantiyesi boş
    olup faturada şantiye yazılıysa kalem faturanın şantiyesine düşeceğinden, geri alınan kalem "ANA DEPO"ya yazılır."""
    d = sirket_kaydi(db, Depo, istek.depo_id, k) if istek.depo_id else None
    if istek.santiye_id:
        hedef = sirket_kaydi(db, Santiye, istek.santiye_id, k)
    elif d is not None:
        hedef = d.santiye
    else:
        hedef = next((s for s in db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id).order_by(Santiye.id))
                      if ana_depo_mu(s)), None)
        if hedef is None:
            hedef = Santiye(sirket_id=k.sirket_id, ad=ANA_DEPO_ADI)
            db.add(hedef)
            db.flush()
    if d is not None and hedef is not None and d.santiye_id not in (None, hedef.id):
        raise HTTPException(400, f"'{d.ad}' deposu başka bir şantiyede.")
    kalemler = db.scalars(select(FaturaKalem).join(Fatura).where(
        Fatura.sirket_id == k.sirket_id, FaturaKalem.id.in_(istek.kalem_idler))).all()
    for kl in kalemler:
        if kl.fatura.durum == "iptal":
            continue
        kl.santiye_id = hedef.id if hedef else None
        kl.depo_id = d.id if d else None
    db.commit()
    return {"tasinan": len(kalemler), "yer": hedef.ad if hedef else None}


@router.get("/excel/ana-depo")
def excel(baslangic: date | None = None, bitis: date | None = None, db: Session = Depends(get_db),
          k: Kullanici = Depends(aktif_kullanici)):
    from .raporlar import BASLIK_DOLGU, BASLIK_STILI
    s = satirlar(db, k, baslangic, bitis)
    ozet, cikis = ozetler(s)
    yer = lambda x: x["santiye"] or "Ana depoda"  # noqa: E731
    sayfalar = [
        ("Ana Depo Özeti", ["Malzeme", "Birim", "Giren Miktar", "Giren Tutar (TL)", "Şantiyelere Çıkan Miktar", "Çıkan Tutar (TL)",
                            "Ana Depoda Kalan Miktar", "Kalan Tutar (TL)"],
         [[x["malzeme"], x["birim"], x["giren_miktar"], x["giren_tutar"], x["cikan_miktar"], x["cikan_tutar"],
           x["kalan_miktar"], x["kalan_tutar"]] for x in ozet]),
        ("Şantiyelere Çıkışlar", ["Şantiye", "Malzeme", "Birim", "Miktar", "Tutar (TL)", "Kalem Sayısı", "Son Tarih"],
         [[x["santiye"], x["malzeme"], x["birim"], x["miktar"], x["tutar"], x["kalem_sayisi"], x["son_tarih"]] for x in cikis]),
        ("Bütün Hareketler", ["Tarih", "Fatura No", "Firma", "Malzeme", "Faturadaki Ad", "Miktar", "Birim", "Birim Fiyat (TL)",
                              "Tutar (KDV hariç, TL)", "Nerede", "Depo"],
         [[x["tarih"], x["fatura_no"], x["firma"], x["malzeme"], x["kalem"], x["miktar"], x["birim"], x["birim_fiyat"],
           x["tutar"], yer(x), x["depo"]] for x in s]),
    ]
    wb = Workbook()
    for i, (ad, basliklar, veri) in enumerate(sayfalar):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = ad
        ws.append(basliklar)
        for c in ws[1]:
            c.font, c.fill = BASLIK_STILI, BASLIK_DOLGU
        for r in veri:
            ws.append(r)
        if i == 0:
            ws.append([])
            ws.append(["TOPLAM", None, None, round(sum(x["giren_tutar"] for x in ozet), 2), None,
                       round(sum(x["cikan_tutar"] for x in ozet), 2), None, round(sum(x["kalan_tutar"] for x in ozet), 2)])
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
        for j, b in enumerate(basliklar, start=1):
            uzunluk = max([len(b)] + [len(str(r[j - 1])) for r in veri[:2000] if r[j - 1] is not None])
            ws.column_dimensions[ws.cell(1, j).column_letter].width = min(uzunluk + 2, 60)
        ws.freeze_panes = "A2"
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return StreamingResponse(tampon, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="ana-depo.xlsx"'})
