"""Excel şablonlarıyla toplu fatura ve çek girişi (yapay zekasız)."""
import io
import re
from collections import OrderedDict

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import kural, schemas
from ..auth import aktif_kullanici, yazabilir
from ..db import get_db
from ..models import GIDER_TURLERI, BankaHesabi, Cari, Fatura, Kullanici, Santiye, Sirket
from ..yukleme import dosyayi_oku
from .cekler import cek_olustur, cekten_ekstreyi_esle
from .faturalar import kalemleri_yaz
from .raporlar import _excel

router = APIRouter(prefix="/api/excel")

FATURA_SUTUNLARI = ["Fatura No", "Tarih", "Vade", "Tür (Alış/Satış)", "Cari Ünvan", "Cari VKN", "Şantiye", "Gider Türü",
                    "Kalem Açıklaması", "Miktar", "Birim", "Birim Fiyat", "KDV %", "Tevkifat", "Stopaj"]
# Şablon dışındaki listelerde (ör. Wolvox/muhasebe programından alınan "fatura listesi") geçen başlıklar.
FATURA_ES_ANLAMLILAR = {
    "Tarih": ["Tarihi", "Fatura Tarihi", "Belge Tarihi"],
    "Fatura No": ["Belge No", "Fatura Numarası", "Evrak No"],
    "Cari Ünvan": ["Ticari Unvanı", "Ticari Ünvanı", "Ünvan", "Unvan", "Cari", "Cari Adı", "Firma", "Firma Adı",
                   "Satıcı", "Tedarikçi"],
    "Cari VKN": ["VKN", "Vergi No", "Vergi Numarası", "VKN/TCKN", "TCKN"],
    "Şantiye": ["Grubu", "Grup", "Proje", "Şantiye Adı", "Masraf Merkezi"],
    "Kalem Açıklaması": ["Stok Adı", "Malzeme", "Malzeme Adı", "Ürün", "Ürün Adı", "Hizmet", "Açıklama", "Mal Hizmet"],
    "Birim": ["Birimi"],
    "Miktar": ["Miktarı", "Adet"],
    "Birim Fiyat": ["Fiyat", "Fiyatı", "Birim Fiyatı"],
    "KDV %": ["KDV Oranı", "KDV", "KDV (%)"],
    "Tutar (KDV Hariç)": ["Tutarı", "Tutar", "Tutar (KDV Hariç)", "KDV Hariç Tutar", "Matrah"],
    "Tutar (KDV Dahil)": ["Tutarı (KDV Dahil)", "KDV Dahil Tutar", "Tutar (KDV Dahil)", "Toplam Tutar"],
    "Tür (Alış/Satış)": ["Tür", "Fatura Türü", "Alış/Satış"],
}
KDV_ORANLARI = (0, 1, 8, 10, 18, 20)
CEK_SUTUNLARI = ["Yön (Alınan/Verilen)", "Tür (Çek/Senet)", "Çek No", "Banka", "Şube", "Hesap No", "Çeki Yazan Firma", "Tutar",
                 "Vade", "Alış/Veriliş Tarihi", "Cari Ünvan", "Cari VKN", "Şantiye", "Açıklama", "Lehtar (Çekin yazıldığı firma)"]
CEK_ES_ANLAMLILAR = {"Lehtar (Çekin yazıldığı firma)": ["Lehtar", "Çekin yazıldığı firma"], "Açıklama": ["Not", "Notlar", "Açıklama / Not"],
                     "Çeki Yazan Firma": ["Keşideci"], "Alış/Veriliş Tarihi": ["Alış Tarihi", "Veriliş Tarihi"]}


def _cek_santiye_aciklama(r: dict, santiyeler: list) -> tuple[str | None, str | None]:
    """Şantiye sütunundaki yazı şantiye adı değilse açıklamaya taşınır; şantiye boşsa açıklamada geçen var olan
    şantiye adı alınır."""
    metin = lambda a: str(r[a]).strip() if r.get(a) not in (None, "") else None  # noqa: E731
    santiye, ek = kural.santiye_ayikla(metin("Şantiye"), santiyeler)
    aciklama = " | ".join(x for x in dict.fromkeys([ek, metin("Açıklama")]) if x) or None
    # Yalnız ilk bölüme bakılır: "... | X firmasına ciro edildi" gibi sonraki notlardaki firma adı şantiye sanılmaz
    if not santiye and aciklama and (st := kural.icindeki_kayit(aciklama.split("|")[0], santiyeler)) is not None:
        santiye = st.ad
    return santiye, aciklama


# Akınsoft Wolvox "Çek/Senet Raporu": yalnız programın kullandığı sütunlar alınır, öteki sütunlar yok sayılır
AKINSOFT_CEK = {"tur": "Evrak Türü", "durum": "Evrak Durumu", "banka_durumu": "Banka Durumu", "vade": "Vadesi",
                "no": "Evrak No", "tutar": "KPB Tutarı", "kayit": "Kayıt Tarihi", "gl_unvan": "Gl.Ticari Unvanı",
                "gl_ad": "Gl.Adı Soyadı", "gd_unvan": "Gd.Ticari Unvanı", "gd_ad": "Gd.Adı Soyadı", "odeme": "Ödendiği Tarih",
                "banka": "Evrak bankası", "sube": "Banka Şubesi", "hesap": "IBAN/Hesap No", "takas": "Takas Banka",
                "borclu_unvan": "Asıl Borçlu Ünvanı", "borclu_ad": "Asıl Borçlu Adı", "borclu_soyad": "Asıl Borçlu Soyadı",
                "degisme": "Değiştirme Tarihi"}


def akinsoft_cekleri(veri: bytes) -> list[dict] | None:
    """Akınsoft çek raporunu tanırsa satırları döndürür (tanımazsa None): {yon, tur, cek_no, banka, sube, hesap_no, tutar,
    vade, alis_tarihi, cari, kesideci, durum, banka_durumu, takas, islem_tarihi}."""
    try:
        wb = load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        return None
    for ws in wb.worksheets:
        satirlar = ws.iter_rows(values_only=True)
        for i, baslik in enumerate(satirlar):
            adlar: dict[str, int] = {}
            for j, x in enumerate(baslik):
                adlar.setdefault(kural.ad_anahtari(str(x or "")), j)  # aynı adlı ikinci sütun (Vadesi saati) atlanır
            yer = {a: adlar.get(kural.ad_anahtari(b)) for a, b in AKINSOFT_CEK.items()}
            if all(yer[a] is not None for a in ("tur", "durum", "no", "tutar", "vade")):
                break
            if i > 10:
                break
        else:
            continue
        if not all(yer[a] is not None for a in ("tur", "durum", "no", "tutar", "vade")):
            continue
        sonuc = []
        for r in satirlar:
            al = lambda a: r[yer[a]] if yer[a] is not None and yer[a] < len(r) else None  # noqa: E731

            def yazi(*alanlar):
                return next((" ".join(str(al(a)).split()) for a in alanlar if str(al(a) or "").strip()), None)
            tur = kural.sade(yazi("tur") or "")
            if not tur:
                continue
            no = al("no")
            no = str(int(no)) if isinstance(no, float) and no.is_integer() else (str(no).strip() if no is not None else None)
            borclu = yazi("borclu_unvan") or " ".join(x for x in (yazi("borclu_ad"), yazi("borclu_soyad")) if x) or None
            yon = "verilen" if tur.startswith(("kendi", "firma")) or "borc" in tur else "alinan"
            sonuc.append({
                "yon": yon, "tur": "senet" if "senet" in tur else "cek", "cek_no": no, "banka": yazi("banka"), "sube": yazi("sube"),
                "hesap_no": yazi("hesap"), "tutar": kural.para_coz(al("tutar")), "vade": kural.tarih_coz(al("vade")),
                "alis_tarihi": kural.tarih_coz(al("kayit")),
                # Alınan çekte cari çeki aldığımız firma (Gl.), verilen çekte verdiğimiz firma (Gd.)
                "cari": yazi("gl_unvan", "gl_ad") if yon == "alinan" else yazi("gd_unvan", "gd_ad"),
                "ciro": yazi("gd_unvan", "gd_ad") if yon == "alinan" else None,
                "kesideci": borclu or (yazi("gl_unvan", "gl_ad") if yon == "alinan" else None),
                "durum": kural.sade(yazi("durum") or ""), "banka_durumu": kural.sade(yazi("banka_durumu") or ""),
                "takas": yazi("takas"),
                "islem_tarihi": kural.tarih_coz(al("odeme")) or kural.tarih_coz(al("degisme")) or kural.tarih_coz(al("kayit"))})
        return sonuc
    return None


class _BankaHesaplari:
    """Akınsoft'taki banka adına ("HALKBANK VADESİZ 10101594", "AKBANK VADESİZ - VADESİZ - TL") karşılık gelen hesabımız;
    yoksa banka adıyla açılır (o bankanın ekstresi yüklenince IBAN'ı eklenir)."""

    def __init__(self, db: Session, k: Kullanici):
        from .. import cek_okuma
        self.db, self.k, self.bul_banka = db, k, cek_okuma.banka_bul
        self.liste = [b for b in db.scalars(select(BankaHesabi).where(BankaHesabi.sirket_id == k.sirket_id)) if b.tur != "kredi_karti"]
        self.acilan: list[str] = []

    def bul(self, metin: str | None, hesap_no: str | None = None) -> BankaHesabi | None:
        banka = self.bul_banka(metin)
        if not banka:
            return None
        ayni = [b for b in self.liste if self.bul_banka(" ".join(x for x in (b.banka, kural.iban_banka(b.iban), b.ad) if x)) == banka]
        rakam = [x for x in (re.sub(r"\D", "", metin or ""), re.sub(r"\D", "", hesap_no or "")) if len(x) >= 5]
        for b in ayni:
            hesap = re.sub(r"\D", "", f"{b.iban or ''} {b.ad or ''}")
            if any(x in hesap for x in rakam):
                return b
        if ayni:
            return ayni[0]
        no = next((x for x in re.findall(r"\d{5,}", metin or "")), None)  # "HALKBANK VADESİZ 10101594" -> "Halkbank 10101594"
        b = BankaHesabi(sirket_id=self.k.sirket_id, ad=f"{banka} {no}" if no else banka, banka=banka)
        self.db.add(b)
        self.db.flush()
        self.liste.append(b)
        self.acilan.append(b.ad)
        return b


def _akinsoft_akibeti(x: dict, cariler: "_Cariler", hesaplar: _BankaHesaplari) -> tuple[dict | None, str | None]:
    """Evrak durumundan akıbet (Toplu Çek Girişi'ndeki seçimle aynı biçim) ve gerekiyorsa kullanıcıya not."""
    d, tarih = x["durum"], x["islem_tarihi"].isoformat() if x["islem_tarihi"] else None
    if x["yon"] == "verilen":
        if d.startswith("bankadan"):
            b = hesaplar.bul(x["banka"], x["hesap_no"])
            return ({"islem": "ode", "banka_hesap_id": b.id if b else None, "tarih": tarih},
                    None if b else "Ödendiği banka hesabını seçin.")
        if d.startswith("kasadan"):
            return {"islem": "ode_kasa", "tarih": tarih}, None
        if d.startswith("iade"):
            return {"islem": "iade", "tarih": tarih}, None
        return None, None  # ödenmedi
    if d.startswith("ciro"):
        c = cariler.bul(x["ciro"], None, "tedarikci") if x["ciro"] else None
        return {"islem": "ciro", "cari_id": c.id if c else None, "tarih": tarih}, None if c else "Ciro edilen firmayı seçin."
    tahsil = d.startswith("tahsil") or x["banka_durumu"].startswith("tahsil")
    if tahsil or d.startswith("takas"):
        b = hesaplar.bul(x["takas"]) if x["takas"] else None
        return ({"islem": "tahsil" if tahsil else "tahsile_ver", "banka_hesap_id": b.id if b else None, "tarih": tarih},
                None if b else "Akınsoft'ta tahsil edildiği hesap yazmıyor; banka hesabını seçin.")
    if d.startswith("teminat"):
        b = hesaplar.bul(x["takas"]) if x["takas"] else None
        return {"islem": "teminata_ver", "banka_hesap_id": b.id if b else None, "tarih": tarih}, None if b else "Banka hesabını seçin."
    if d.startswith("karsiliksiz") or d.startswith("protesto"):
        return {"islem": "karsiliksiz", "tarih": tarih}, None
    if d.startswith("iade"):
        return {"islem": "iade", "tarih": tarih}, None
    return None, None  # portföyde


def _satirlar(veri: bytes, beklenen: list[str], es_anlamlilar: dict | None = None) -> list[dict]:
    try:
        ws = load_workbook(io.BytesIO(veri), read_only=True, data_only=True).active
    except Exception as e:
        raise HTTPException(400, "Excel dosyası okunamadı. Şablonu kullanıp .xlsx olarak kaydedin.") from e
    harita = {kural.sade(b): b for b in beklenen}
    for asil, adlar in (es_anlamlilar or {}).items():
        for ad in adlar:
            harita.setdefault(kural.sade(ad), asil)
    tum = list(ws.iter_rows(values_only=True))
    # Başlık satırı ilk satır olmayabilir (üstte rapor başlığı vb.): tanınan en az iki sütunu olan ilk satır.
    bas = next((i for i, satir in enumerate(tum[:15])
                if sum(kural.sade(str(h or "")) in harita for h in satir) >= 2), None)
    if bas is None:
        raise HTTPException(400, "Başlık satırı tanınmadı. Şablonu indirip onu kullanın.")
    basliklar = [kural.sade(str(h or "")) for h in tum[bas]]
    sonuc = []
    for no, satir in enumerate(tum[bas + 1:], start=bas + 2):
        kayit = {harita[b]: v for b, v in zip(basliklar, satir) if b in harita and v not in (None, "")}
        if kayit:
            kayit["_satir"] = no
            sonuc.append(kayit)
    return sonuc


class _Cariler:
    """VKN veya ünvana göre cari bulur; yoksa yenisini açar."""

    def __init__(self, db: Session, k: Kullanici):
        self.db, self.k = db, k
        self.liste = db.scalars(select(Cari).where(Cari.sirket_id == k.sirket_id)).all()
        self.acilan = 0

    def bul(self, unvan, vkn, tur: str) -> Cari | None:
        vkn = str(vkn).strip().split(".")[0] if vkn not in (None, "") else None
        unvan = str(unvan).strip() if unvan else None
        sirket = self.db.get(Sirket, self.k.sirket_id)
        if vkn and vkn == (sirket.vkn or "").strip() or kural.kendi_unvani_mi(unvan, sirket.unvan):
            return None  # kendi firmamız cari olmaz
        anahtar = kural.unvan_cekirdegi(unvan) or kural.ad_anahtari(unvan)
        for c in self.liste:
            if vkn and c.vkn == vkn:
                return c
        for c in self.liste:  # aynı ad ("Ltd." / "Ltd. Şti." farkı yok sayılır); VKN'leri farklı iki firma ayrı kalır
            if anahtar and (kural.unvan_cekirdegi(c.unvan) or kural.ad_anahtari(c.unvan)) == anahtar and not (vkn and c.vkn and c.vkn != vkn):
                return c
        if not unvan:
            return None
        c = Cari(sirket_id=self.k.sirket_id, unvan=unvan, vkn=vkn, tur=tur)
        self.db.add(c)
        self.db.flush()
        self.liste.append(c)
        self.acilan += 1
        return c


class _Santiyeler:
    """Ada göre şantiye bulur; yoksa yenisini açar."""

    def __init__(self, db: Session, k: Kullanici):
        self.db, self.k = db, k
        self.liste = db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id)).all()
        self.acilan = 0

    def bul(self, ad) -> int | None:
        ad = str(ad).strip() if ad not in (None, "") else None
        if not ad:
            return None
        s = next((s for s in self.liste if kural.adi_bu_mu(s, ad)), None)
        if not s:
            s = Santiye(sirket_id=self.k.sirket_id, ad=ad)
            self.db.add(s)
            self.db.flush()
            self.liste.append(s)
            self.acilan += 1
        return s.id


def _kdv_orani(haric: float | None, dahil: float | None) -> float | None:
    """KDV hariç ve dahil tutardan oranı bulur; yuvarlama farkı için en yakın yasal orana oturtur."""
    if not haric or dahil is None:
        return None
    oran = (dahil / haric - 1) * 100
    yakin = min(KDV_ORANLARI, key=lambda o: abs(o - oran))
    return yakin if abs(yakin - oran) < 0.6 else round(oran, 2)


@router.get("/fatura-sablon")
def fatura_sablon(k: Kullanici = Depends(aktif_kullanici)):
    ornek = [["ABC2026000000001", "01.10.2026", "31.10.2026", "Alış", "Örnek Beton A.Ş.", "1234567890", "", "Beton",
              "C30 hazır beton", 10, "m3", 2500, 20, 0, 0],
             ["ABC2026000000001", "01.10.2026", "31.10.2026", "Alış", "Örnek Beton A.Ş.", "1234567890", "", "Beton",
              "Pompa", 1, "adet", 3000, 20, 0, 0]]
    return _excel("Faturalar", FATURA_SUTUNLARI, ornek, "fatura-yukleme-sablonu")


@router.post("/faturalar/yukle")
async def fatura_yukle(dosya: UploadFile = File(...), yon: str = Query("gelen", pattern="^(gelen|giden)$"),
                       db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Aynı fatura no + cari satırları tek faturada toplanır; faturalar taslak olarak eklenir.

    Programın şablonu da, şablon dışı fatura listeleri de (Tarihi, Fatura No, Ticari Unvanı, Stok Adı,
    Grubu, Miktarı, Tutarı, Tutarı (KDV Dahil) gibi başlıklarla) okunur. "Tür" sütunu yoksa `yon` kullanılır.
    Daha önce girilmiş faturalar (aynı yön, cari ve fatura no) atlanır.
    """
    satirlar = _satirlar(dosyayi_oku(dosya), FATURA_SUTUNLARI, FATURA_ES_ANLAMLILAR)
    cariler = _Cariler(db, k)
    santiyeler = _Santiyeler(db, k)
    gruplar: OrderedDict = OrderedDict()
    hatalar = []
    for r in satirlar:
        if not (r.get("Tarih") or r.get("Fatura No") or r.get("Cari Ünvan")):
            continue  # "TOPLAM" gibi alt toplam satırları
        try:
            tur = kural.sade(str(r.get("Tür (Alış/Satış)") or ""))
            satir_yon = ("giden" if tur.startswith("s") else "gelen") if tur else yon
            cari = cariler.bul(r.get("Cari Ünvan"), r.get("Cari VKN"), "tedarikci" if satir_yon == "gelen" else "musteri")
            tarih = kural.tarih_coz(r.get("Tarih"))
            miktar = kural.para_coz(r.get("Miktar")) or 1
            fiyat = kural.para_coz(r.get("Birim Fiyat"))
            haric = kural.para_coz(r.get("Tutar (KDV Hariç)"))
            dahil = kural.para_coz(r.get("Tutar (KDV Dahil)"))
            if haric is None and fiyat is not None:
                haric = round(miktar * fiyat, 2)
            if fiyat is None and haric is not None:
                fiyat = round(haric / miktar, 6) if miktar else haric
            if not tarih or fiyat is None:
                raise ValueError("tarih veya tutar okunamadı")
            kdv = kural.para_coz(r.get("KDV %"))
            if kdv is None:
                kdv = _kdv_orani(haric, dahil)
            anahtar = (satir_yon, str(r.get("Fatura No", f"satir{r['_satir']}")), cari.id if cari else None)
            if anahtar not in gruplar:
                gider = r.get("Gider Türü")
                if gider not in GIDER_TURLERI:
                    gider = kural.gider_turu_tahmin(str(r.get("Kalem Açıklaması") or ""))
                gruplar[anahtar] = {"f": schemas.FaturaGiris(
                    yon=satir_yon, fatura_no=str(r["Fatura No"]).strip() if r.get("Fatura No") else None, tarih=tarih,
                    vade=kural.tarih_coz(r.get("Vade")), cari_id=cari.id if cari else None,
                    santiye_id=santiyeler.bul(r.get("Şantiye")), gider_turu=gider,
                    tevkifat=kural.para_coz(r.get("Tevkifat")) or 0, stopaj=kural.para_coz(r.get("Stopaj")) or 0, kalemler=[]), "satir": r["_satir"]}
            gruplar[anahtar]["f"].kalemler.append(schemas.KalemGiris(
                aciklama=str(r.get("Kalem Açıklaması") or "-"), miktar=miktar, birim=r.get("Birim"),
                birim_fiyat=fiyat, kdv_orani=20 if kdv is None else kdv, tutar=haric))
        except ValueError as e:
            hatalar.append(f"{r['_satir']}. satır: {e}")
    if hatalar:
        db.rollback()
        raise HTTPException(400, "Dosyada hata var, hiçbir şey eklenmedi. " + "; ".join(hatalar[:10]))
    eklenen = atlanan = 0
    for g in gruplar.values():
        v = g["f"]
        if v.fatura_no and v.cari_id and db.scalar(select(Fatura.id).where(
                Fatura.sirket_id == k.sirket_id, Fatura.yon == v.yon, Fatura.cari_id == v.cari_id,
                Fatura.fatura_no == v.fatura_no)):
            atlanan += 1
            continue
        f = Fatura(sirket_id=k.sirket_id, durum="taslak", **v.model_dump(exclude={"kalemler"}))
        kalemleri_yaz(f, v.kalemler)
        db.add(f)
        eklenen += 1
    db.commit()
    return {"eklenen": eklenen, "atlanan": atlanan, "yeni_cari": cariler.acilan, "yeni_santiye": santiyeler.acilan}


@router.get("/cek-sablon")
def cek_sablon(k: Kullanici = Depends(aktif_kullanici)):
    ornek = [["Alınan", "Çek", "0012345", "Garanti BBVA", "Kartal", "6299999", "Ahmet Yılmaz", 4000, "22.12.2026",
              "18.10.2026", "Örnek Müşteri Ltd.", "", "", ""]]
    return _excel("Çekler", CEK_SUTUNLARI, ornek, "cek-yukleme-sablonu")


@router.post("/cekler/onizle")
async def cek_onizle(yon: str = "alinan", dosya: UploadFile = File(...), db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    """Çek şablonundaki satırları kaydetmeden toplu giriş ekranına getirir. Akınsoft çek raporu da olduğu gibi okunur."""
    from . import cekler as cek_r
    icerik = dosyayi_oku(dosya)
    if (akinsoft := akinsoft_cekleri(icerik)) is not None:
        return _akinsoft_onizle(akinsoft, yon, dosya.filename or "", db, k)
    satirlar, atlanan, toplam_satiri, excel_toplam, okunan_toplam = [], 0, 0, None, 0.0
    santiyeler = list(db.scalars(select(Santiye).where(Santiye.sirket_id == k.sirket_id)))
    hesaplar = [b for b in db.scalars(select(BankaHesabi).where(BankaHesabi.sirket_id == k.sirket_id)) if b.tur != "kredi_karti"]
    akibetler = []
    for r in _satirlar(icerik, CEK_SUTUNLARI, CEK_ES_ANLAMLILAR):
        # "TOPLAM / ARA TOPLAM / GENEL TOPLAM" satırı çek değildir; tutarı okunan çeklerle karşılaştırılır
        if not r.get("Çek No") and any("toplam" in kural.sade(str(v)) for a, v in r.items() if a != "_satir"):
            toplam_satiri += 1
            t = kural.para_coz(r.get("Tutar"))
            if t:
                excel_toplam = max(excel_toplam or 0.0, abs(t))
            continue
        okunan_toplam += abs(kural.para_coz(r.get("Tutar")) or 0)
        r_yon = "verilen" if kural.sade(str(r.get("Yön (Alınan/Verilen)", "alinan"))).startswith("v") else "alinan"
        if r_yon != yon:
            atlanan += 1
            continue
        metin = lambda a: str(r[a]).strip() if r.get(a) not in (None, "") else None  # noqa: E731
        vade, tutar = kural.tarih_coz(r.get("Vade")), kural.para_coz(r.get("Tutar"))
        alis = kural.tarih_coz(r.get("Alış/Veriliş Tarihi"))
        santiye, aciklama = _cek_santiye_aciklama(r, santiyeler)
        cari_unvan, akibet = metin("Cari Ünvan"), None
        notu = kural.cek_notu_coz(aciklama)
        if yon == "alinan" and notu["ciro"]:
            # "X firmasına ciro edildi": akıbet ciro; Cari sütununa ciro edilen firma yazılmışsa çeki aldığımız firma
            # (çeki yazan) caridir
            ciro = cek_r._cari_oner(db, k, notu["ciro"], None)
            if ciro is None:
                ciro = Cari(sirket_id=k.sirket_id, unvan=notu["ciro"], tur="tedarikci")
                db.add(ciro)
                db.flush()
            akibet = {"islem": "ciro", "cari_id": ciro.id}
            if cari_unvan and kural.ad_anahtari(cari_unvan) in (kural.ad_anahtari(ciro.unvan), kural.ad_anahtari(notu["ciro"])):
                cari_unvan = None
        elif yon == "alinan" and notu["banka"]:
            ayni = [b for b in hesaplar if (kural.iban_banka(b.iban) or b.banka) == notu["banka"]]
            if len(ayni) == 1:
                akibet = {"islem": "tahsile_ver", "banka_hesap_id": ayni[0].id}
        akibetler.append(akibet)
        satirlar.append({"tur": "senet" if kural.sade(str(r.get("Tür (Çek/Senet)", ""))).startswith("senet") else "cek",
                         "cek_no": metin("Çek No"), "banka": metin("Banka"), "sube": metin("Şube"), "hesap_no": metin("Hesap No"),
                         "kesideci": metin("Çeki Yazan Firma"), "tutar": abs(tutar) if tutar else None, "vade": vade.isoformat() if vade else None,
                         "alis_tarihi": alis.isoformat() if alis else None, "cari_unvan": cari_unvan,
                         "cari_vkn": metin("Cari VKN"), "santiye_ad": santiye, "aciklama": aciklama,
                         # Verilen çekte lehtar (çekin yazıldığı firma) caridir: Lehtar sütunu boşsa Cari Ünvan
                         "lehtar": metin("Lehtar (Çekin yazıldığı firma)") or (cari_unvan if yon == "verilen" else None),
                         "belirsiz": [a for a, v in (("tutar", tutar), ("vade", vade)) if not v]})
    hazir = cek_r._okunanlari_hazirla(db, k, satirlar, None, dosya.filename or "", cekten=False)
    for x, h, a in zip(satirlar, hazir, akibetler):
        if not x["cari_unvan"] and not x["cari_vkn"]:
            h["cari_id"], h["cari_unvan"] = None, None  # şablonda cari yazmıyorsa keşideci cari sayılmaz
        if a:
            h["akibet"] = a
        h["belirsiz"] = [b for b in (h.get("belirsiz") or []) if b != "lehtar"]  # Excel'de yazılan lehtar kesindir
    db.commit()  # ciro edilen firma için açılan cariler
    uyari = None
    if excel_toplam is not None and abs(excel_toplam - okunan_toplam) > 0.01:
        fmt = lambda x: f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")  # noqa: E731
        uyari = (f"Excel'deki toplam {fmt(excel_toplam)}, okunan çeklerin toplamı {fmt(okunan_toplam)}; "
                 f"{fmt(abs(excel_toplam - okunan_toplam))} fark var. Eksik ya da fazla satır olabilir, kontrol edin.")
    return {"satirlar": hazir, "atlanan": atlanan, "toplam_satiri": toplam_satiri, "uyari": uyari}


def _akinsoft_onizle(cekler: list[dict], yon: str, dosya_adi: str, db: Session, k: Kullanici) -> dict:
    """Akınsoft çek raporu: bu yöndeki çekler akıbetiyle (ciro, tahsil, ödendi...) toplu giriş ekranına gelir."""
    from . import cekler as cek_r
    cariler, hesaplar = _Cariler(db, k), _BankaHesaplari(db, k)
    secilen = [x for x in cekler if x["yon"] == yon]
    satirlar, akibetler, notlar = [], [], []
    for x in secilen:
        akibet, not_ = _akinsoft_akibeti(x, cariler, hesaplar)
        akibetler.append(akibet)
        notlar.append(not_)
        satirlar.append({"tur": x["tur"], "cek_no": x["cek_no"], "banka": x["banka"], "sube": x["sube"], "hesap_no": x["hesap_no"],
                         "kesideci": x["kesideci"], "tutar": abs(x["tutar"]) if x["tutar"] else None,
                         "vade": x["vade"].isoformat() if x["vade"] else None,
                         "alis_tarihi": x["alis_tarihi"].isoformat() if x["alis_tarihi"] else None,
                         "cari_unvan": x["cari"], "cari_vkn": None, "santiye_ad": None, "aciklama": None,
                         "lehtar": x["cari"] if yon == "verilen" else None,
                         "belirsiz": [a for a in ("tutar", "vade") if not x[a]]})
    hazir = cek_r._okunanlari_hazirla(db, k, satirlar, None, dosya_adi, cekten=False)
    for h, a, n in zip(hazir, akibetler, notlar):
        if a:
            h["akibet"] = a
        if n and not h.get("sorun"):
            h["sorun"] = n
        # Lehtar ve cari Akınsoft'ta yazılı, kesindir (caride yoksa satırda "Yeni cari açılacak" yazar, kırmızı olmaz)
        h["belirsiz"] = [b for b in (h.get("belirsiz") or []) if b not in ("lehtar", "cari_unvan")]
    db.commit()  # ciro edilen firmalar ve banka adıyla açılan hesaplar
    return {"satirlar": hazir, "atlanan": len(cekler) - len(secilen), "toplam_satiri": 0, "uyari": None, "akinsoft": True,
            "yeni_hesap": hesaplar.acilan}


@router.post("/cekler/yukle")
async def cek_yukle(dosya: UploadFile = File(...), db: Session = Depends(get_db), k: Kullanici = Depends(yazabilir)):
    satirlar = _satirlar(dosyayi_oku(dosya), CEK_SUTUNLARI, CEK_ES_ANLAMLILAR)
    cariler = _Cariler(db, k)
    santiyeler = _Santiyeler(db, k)
    hatalar, eklenen = [], 0
    for r in satirlar:
        try:
            yon = "verilen" if kural.sade(str(r.get("Yön (Alınan/Verilen)", "alinan"))).startswith("v") else "alinan"
            vade, tutar = kural.tarih_coz(r.get("Vade")), kural.para_coz(r.get("Tutar"))
            if not vade or not tutar:
                raise ValueError("vade veya tutar okunamadı")
            cari = cariler.bul(r.get("Cari Ünvan"), r.get("Cari VKN"), "musteri" if yon == "alinan" else "tedarikci")
            metin = lambda a: str(r[a]).strip() if r.get(a) not in (None, "") else None  # noqa: E731
            santiye, aciklama = _cek_santiye_aciklama(r, santiyeler.liste)
            c = cek_olustur(db, schemas.CekGiris(
                yon=yon, tur="senet" if kural.sade(str(r.get("Tür (Çek/Senet)", ""))).startswith("senet") else "cek",
                cek_no=metin("Çek No"), banka=metin("Banka"), sube=metin("Şube"), hesap_no=metin("Hesap No"),
                kesideci=metin("Çeki Yazan Firma"), tutar=abs(tutar), vade=vade,
                alis_tarihi=kural.tarih_coz(r.get("Alış/Veriliş Tarihi")), cari_id=cari.id if cari else None,
                santiye_id=santiyeler.bul(santiye), aciklama=aciklama,
                lehtar=metin("Lehtar (Çekin yazıldığı firma)") or (metin("Cari Ünvan") if yon == "verilen" else None)), k)
            cekten_ekstreyi_esle(db, c)
            eklenen += 1
        except (ValueError, HTTPException) as e:
            hatalar.append(f"{r['_satir']}. satır: {getattr(e, 'detail', e)}")
    if hatalar:
        db.rollback()
        raise HTTPException(400, "Dosyada hata var, hiçbir şey eklenmedi. " + "; ".join(hatalar[:10]))
    db.commit()
    return {"eklenen": eklenen, "yeni_cari": cariler.acilan, "yeni_santiye": santiyeler.acilan}
