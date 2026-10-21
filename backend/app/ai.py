"""PDF veya fotoğraftan yapay zeka ile fatura okuma (Claude API).

ANTHROPIC_API_KEY ortam değişkeni tanımlı olmalıdır.
"""
import base64
import json
import os

import anthropic

MODEL = os.environ.get("MUHASEBE_AI_MODEL", "claude-opus-5-5")

DESTEKLENEN_TURLER = {
    "application/pdf": "document",
    "image/jpeg": "image",
    "image/png": "image",
    "image/webp": "image",
    "image/gif": "image",
}

_metin = {"type": ["string", "null"]}
_sayi = {"type": ["number", "null"]}

FATURA_SEMASI = {
    "type": "object",
    "additionalProperties": False,
    "required": [
        "belge_fatura_mi", "fatura_no", "tarih", "vade", "satici", "alici", "kalemler",
        "ara_toplam", "kdv_toplam", "tevkifat", "genel_toplam", "para_birimi",
        "onerilen_santiye", "onerilen_gider_turu", "notlar",
    ],
    "properties": {
        "belge_fatura_mi": {"type": "boolean"},
        "fatura_no": _metin,
        "tarih": {"type": ["string", "null"], "description": "YYYY-AA-GG"},
        "vade": {"type": ["string", "null"], "description": "YYYY-AA-GG"},
        "satici": {
            "type": "object",
            "additionalProperties": False,
            "required": ["unvan", "vkn", "vergi_dairesi", "adres", "telefon", "iban"],
            "properties": {k: _metin for k in ["unvan", "vkn", "vergi_dairesi", "adres", "telefon", "iban"]},
        },
        "alici": {
            "type": "object",
            "additionalProperties": False,
            "required": ["unvan", "vkn"],
            "properties": {"unvan": _metin, "vkn": _metin},
        },
        "kalemler": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["aciklama", "miktar", "birim", "birim_fiyat", "kdv_orani", "tutar"],
                "properties": {
                    "aciklama": {"type": "string"},
                    "miktar": _sayi,
                    "birim": _metin,
                    "birim_fiyat": _sayi,
                    "kdv_orani": _sayi,
                    "tutar": _sayi,
                },
            },
        },
        "ara_toplam": _sayi,
        "kdv_toplam": _sayi,
        "tevkifat": _sayi,
        "genel_toplam": _sayi,
        "para_birimi": {"type": "string"},
        "onerilen_santiye": _metin,
        "onerilen_gider_turu": _metin,
        "notlar": {"type": "string"},
    },
}

SISTEM = """Türkiye'deki bir inşaat/taahhüt işletmesinin muhasebe asistanısın.
Sana bir fatura (e-Fatura, e-Arşiv, kağıt fatura fotoğrafı veya PDF) verilecek. Bilgileri eksiksiz ve
belgede yazdığı gibi çıkar.

Kurallar:
- Tarihleri YYYY-AA-GG biçiminde yaz. Tutarları nokta ondalık ayraçlı sayı olarak yaz (1.234,56 TL -> 1234.56).
- VKN 10 haneli, TCKN 11 hanelidir; boşluk ve tire olmadan yaz.
- Kalem 'tutar' alanı KDV hariç satır tutarıdır.
- Belgede olmayan veya okunamayan bilgi için null kullan; tahmin etme.
- 'tevkifat' alanına KDV tevkifatı tutarını yaz, yoksa 0.
- Belge fatura değilse belge_fatura_mi=false yap.
- onerilen_santiye: kalemlere, sevk adresine veya açıklamalara bakarak aşağıdaki şantiyelerden
  hangisine ait olduğunu seç. Emin değilsen null bırak. Listede olmayan bir ad yazma.
- onerilen_gider_turu: aşağıdaki gider türlerinden en uygun olanı seç.
- notlar: okunması zor yerler, tutarsızlıklar veya dikkat edilmesi gerekenler (Türkçe, kısa)."""


class OkumaHatasi(Exception):
    pass


def _belge_blogu(dosya: bytes, mime: str) -> dict:
    blok_turu = DESTEKLENEN_TURLER.get(mime)
    if not blok_turu:
        raise OkumaHatasi("Bu dosya türü desteklenmiyor. PDF, JPG, PNG veya WEBP yükleyin.")
    return {"type": blok_turu,
            "source": {"type": "base64", "media_type": mime, "data": base64.standard_b64encode(dosya).decode()}}


def anahtar_var() -> bool:
    return bool(os.environ.get("ANTHROPIC_API_KEY", "").strip())


def _cagir(sistem: str, icerik: list[dict], sema: dict, max_tokens: int = 16000) -> dict:
    """Claude'a yapılandırılmış çıktı isteği gönderir ve JSON sonucu döndürür."""
    if not anahtar_var():
        raise OkumaHatasi("Yapay zeka anahtarı tanımlı değil. Yapay zekasız okuma için e-Fatura XML'i, metinli PDF "
                          "veya Excel ekstre yükleyin.")
    client = anthropic.Anthropic()
    try:
        with client.beta.messages.stream(
            model=MODEL,
            max_tokens=max_tokens,
            system=sistem,
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
            output_config={"effort": "medium", "format": {"type": "json_schema", "schema": sema}},
            messages=[{"role": "user", "content": icerik}],
        ) as akis:
            yanit = akis.get_final_message()
    except anthropic.AuthenticationError as e:
        raise OkumaHatasi("Yapay zeka anahtarı geçersiz.") from e
    except anthropic.RateLimitError as e:
        raise OkumaHatasi("Yapay zeka şu an yoğun, biraz sonra tekrar deneyin.") from e
    except anthropic.APIStatusError as e:
        raise OkumaHatasi(f"Yapay zeka hatası: {e.message}") from e
    except anthropic.APIConnectionError as e:
        raise OkumaHatasi("Yapay zeka servisine bağlanılamadı.") from e

    if yanit.stop_reason == "refusal":
        raise OkumaHatasi("Yapay zeka bu belgeyi işlemedi.")
    if yanit.stop_reason == "max_tokens":
        raise OkumaHatasi("Belge çok uzun, okuma yarıda kaldı. Daha kısa bir dönem yükleyin.")
    metin = next((b.text for b in yanit.content if b.type == "text"), None)
    if not metin:
        raise OkumaHatasi("Yapay zekadan yanıt alınamadı.")
    return json.loads(metin)


def fatura_oku(dosya: bytes, mime: str, santiyeler: list[str], gider_turleri: list[str]) -> dict:
    istek = (
        "Şantiyeler: " + (", ".join(santiyeler) if santiyeler else "(henüz şantiye yok)") + "\n"
        "Gider türleri: " + ", ".join(gider_turleri) + "\n\n"
        "Bu faturayı oku."
    )
    return _cagir(SISTEM, [_belge_blogu(dosya, mime), {"type": "text", "text": istek}], FATURA_SEMASI)


# --- Çek ------------------------------------------------------------------

CEK_SEMASI = {
    "type": "object",
    "additionalProperties": False,
    "required": ["belge_cek_mi", "tur", "cek_no", "banka", "sube", "hesap_no", "kesideci", "tutar",
                 "para_birimi", "vade", "notlar"],
    "properties": {
        "belge_cek_mi": {"type": "boolean"},
        "tur": {"type": "string", "enum": ["cek", "senet"]},
        "cek_no": _metin, "banka": _metin, "sube": _metin, "hesap_no": _metin, "kesideci": _metin,
        "tutar": _sayi, "para_birimi": {"type": "string"},
        "vade": {"type": ["string", "null"], "description": "YYYY-AA-GG"},
        "notlar": {"type": "string"},
    },
}

CEK_SISTEM = """Türkiye'deki bir işletmenin muhasebe asistanısın. Sana bir çek veya senet görüntüsü verilecek.
Bilgileri belgede yazdığı gibi çıkar. Tarihleri YYYY-AA-GG, tutarı nokta ondalık ayraçlı sayı olarak yaz.
Rakamla ve yazıyla yazılan tutar farklıysa notlarda belirt. Okunamayan alan için null kullan; tahmin etme.
kesideci: çeki düzenleyen kişi/firma. notlar: dikkat edilmesi gerekenler (Türkçe, kısa)."""


def cek_oku(dosya: bytes, mime: str) -> dict:
    return _cagir(CEK_SISTEM, [_belge_blogu(dosya, mime), {"type": "text", "text": "Bu çeki oku."}], CEK_SEMASI, 4000)


# Toplu çek girişi (isteğe bağlı): sayfadaki bütün çekler. Tutar rakamla ve yazıyla ayrı istenir; program ikisi
# tutuyorsa emin sayar. Cari istenmez (çekten tahmin edilmez).
_YZ_CEK = {
    "type": "object",
    "additionalProperties": False,
    "required": ["cek_no", "banka", "sube", "hesap_no", "cekin_yazani", "cekin_yazani_vkn", "lehtar", "vade",
                 "tutar_rakamla", "tutar_yaziyla", "tutar_yaziyla_sayi", "para_birimi"],
    "properties": {
        "cek_no": _metin, "banka": _metin, "sube": _metin, "hesap_no": _metin,
        "cekin_yazani": {"type": ["string", "null"], "description": "Çeki yazan (düzenleyen) firma/kişi: çekin üstündeki basılı ad"},
        "cekin_yazani_vkn": _metin,
        "lehtar": {"type": ["string", "null"], "description": "'Bu çek karşılığında ... emrine' kısmındaki ad"},
        "vade": {"type": ["string", "null"], "description": "Elle yazılan keşide tarihi, YYYY-AA-GG"},
        "tutar_rakamla": {"type": ["number", "null"], "description": "Rakamla yazılan tutar"},
        "tutar_yaziyla": {"type": ["string", "null"], "description": "Yazıyla yazılan tutar, çekte yazdığı gibi"},
        "tutar_yaziyla_sayi": {"type": ["number", "null"], "description": "Yazıyla yazılan tutarın sayı karşılığı"},
        "para_birimi": {"type": "string"},
    },
}
CEKLER_SEMASI = {"type": "object", "additionalProperties": False, "required": ["cekler"],
                 "properties": {"cekler": {"type": "array", "items": _YZ_CEK}}}
CEKLER_SISTEM = """Türkiye'deki bir işletmenin muhasebe asistanısın. Sana taranmış bir sayfa verilecek; sayfada bir ya
da birden çok banka çeki olabilir. Her çek için ayrı kayıt çıkar, yukarıdan aşağı sırayla.
Kurallar:
- Yalnız çekte gördüğünü yaz. Okuyamadığın ya da emin olmadığın alanı null bırak; tahmin etme, tamamlama.
- Tarih ve tutar çoğu zaman el yazısıdır. Rakamla ve yazıyla yazılan tutarı birbirinden bağımsız oku: birini
  ötekine bakarak düzeltme.
- cek_no: çek seri no (yalnız rakamlar). hesap_no: hesap no ya da IBAN.
- cekin_yazani: çeki yazan firma (çekin üstünde basılı ad); cekin_yazani_vkn: onun VKN/TCKN'si.
- lehtar: "Bu çek karşılığında ... emrine" kısmında yazan ad."""


def cek_sayfasi_oku(resim_jpeg: bytes) -> list[dict]:
    """Taranmış sayfadaki çekler (yapay zekayla; anahtar gerekir)."""
    blok = {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                        "data": base64.standard_b64encode(resim_jpeg).decode()}}
    return _cagir(CEKLER_SISTEM, [blok, {"type": "text", "text": "Bu sayfadaki çekleri oku."}], CEKLER_SEMASI, 8000)["cekler"]


# --- Banka ekstresi -------------------------------------------------------

EKSTRE_SEMASI = {
    "type": "object",
    "additionalProperties": False,
    "required": ["satirlar", "notlar"],
    "properties": {
        "satirlar": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["tarih", "aciklama", "tutar", "bakiye", "onerilen_cari_id", "onerilen_tur", "onerilen_gider_turu"],
                "properties": {
                    "tarih": {"type": "string", "description": "YYYY-AA-GG"},
                    "aciklama": {"type": "string"},
                    "tutar": {"type": "number", "description": "Hesaba giren para pozitif, çıkan negatif"},
                    "bakiye": _sayi,
                    "onerilen_cari_id": {"type": ["integer", "null"]},
                    "onerilen_tur": {"type": "string", "enum": ["tahsilat", "odeme", "masraf", "banka_masrafi", "virman", "cek_tahsil", "cek_odeme", "diger"]},
                    "onerilen_gider_turu": _metin,
                },
            },
        },
        "notlar": {"type": "string"},
    },
}

EKSTRE_SISTEM = """Türkiye'deki bir inşaat işletmesinin muhasebe asistanısın. Sana bir banka hesap ekstresi verilecek.
Her hesap hareketini bir satır olarak çıkar; devir/toplam satırlarını atla.

Kurallar:
- tutar: hesaba giren para pozitif, hesaptan çıkan para negatif. Nokta ondalık ayraçlı sayı.
- tarih YYYY-AA-GG. aciklama: ekstredeki açıklamanın tamamı.
- onerilen_cari_id: açıklamadaki ünvan, VKN veya IBAN verilen cari listesindeki bir cariyle eşleşiyorsa o carinin
  id'si; emin değilsen null. Listede olmayan bir id yazma.
- onerilen_tur: cariden gelen para tahsilat, cariye giden para odeme; EFT/havale/komisyon/BSMV gibi banka
  kesintileri banka_masrafi; kendi hesapları arası virman; çek tahsilatı cek_tahsil, çek ödemesi cek_odeme;
  akaryakıt, kira, fatura ödemesi gibi carisiz giderler masraf.
- onerilen_gider_turu: masraf ise verilen gider türlerinden biri, değilse null.
- notlar: okunamayan yerler veya dikkat edilmesi gerekenler (Türkçe, kısa)."""

EKSTRE_TABLO_TURLERI = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/csv",
}


def tablo_metni(dosya: bytes, mime: str) -> str:
    """Excel/CSV ekstreyi yapay zekaya gönderilecek düz metne çevirir."""
    if mime == "text/csv":
        for kodlama in ("utf-8-sig", "cp1254", "latin-1"):
            try:
                return dosya.decode(kodlama)
            except UnicodeDecodeError:
                continue
    import io
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(dosya), read_only=True, data_only=True)
    except Exception as e:
        raise OkumaHatasi("Excel dosyası okunamadı. Dosyayı .xlsx olarak kaydedip tekrar deneyin.") from e
    satirlar = []
    for ws in wb.worksheets:
        for satir in ws.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() for v in satir):
                satirlar.append(" | ".join("" if v is None else str(v) for v in satir))
    return "\n".join(satirlar)


def ekstre_oku(dosya: bytes, mime: str, cariler: list[dict], gider_turleri: list[str]) -> dict:
    cari_listesi = "\n".join(
        f"{c['id']} | {c['unvan']} | {c.get('vkn') or ''} | {c.get('iban') or ''}" for c in cariler[:1000]
    ) or "(cari yok)"
    istek = f"Cariler (id | ünvan | VKN | IBAN):\n{cari_listesi}\n\nGider türleri: {', '.join(gider_turleri)}\n\nBu ekstreyi oku."
    if mime in EKSTRE_TABLO_TURLERI:
        icerik = [{"type": "text", "text": "Ekstre:\n" + tablo_metni(dosya, mime)}, {"type": "text", "text": istek}]
    else:
        icerik = [_belge_blogu(dosya, mime), {"type": "text", "text": istek}]
    return _cagir(EKSTRE_SISTEM, icerik, EKSTRE_SEMASI, 64000)


def tutarlilik_uyarilari(veri: dict) -> list[str]:
    """Okunan rakamları kendi içinde kontrol eder."""
    uyarilar = []
    kalem_toplami = sum((k.get("tutar") or 0) for k in veri.get("kalemler", []))
    ara = veri.get("ara_toplam")
    kalem_toplami += veri.get("yuvarlama") or 0  # satıcının kendi kuruş yuvarlaması (kural.pdf_fatura_oku)
    if ara is not None and veri.get("kalemler") and abs(kalem_toplami - ara) > 0.05:
        uyarilar.append(f"Kalemlerin toplamı ({kalem_toplami:.2f}) ara toplamla ({ara:.2f}) tutmuyor.")
    kdv, genel, tevk = veri.get("kdv_toplam"), veri.get("genel_toplam"), veri.get("tevkifat") or 0
    diger = veri.get("diger_vergiler") or 0
    if None not in (ara, kdv, genel) and abs(ara + kdv - tevk + diger - genel) > 0.05:
        uyarilar.append(
            f"Ara toplam + KDV - tevkifat + diğer vergiler ({ara + kdv - tevk + diger:.2f}) genel toplamla ({genel:.2f}) tutmuyor."
        )
    if not veri.get("belge_fatura_mi", True):
        uyarilar.append("Yüklenen belge fatura gibi görünmüyor.")
    return uyarilar
