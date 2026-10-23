"""Yapay zeka kullanmadan belge okuma.

- e-Fatura / e-Arşiv XML (UBL-TR): tüm bilgiler eksiksiz okunur.
- Metin içeren PDF (e-Fatura görüntüsü): başlık bilgileri ve toplamlar okunur,
  kalemler tek satır olarak girilir. Taranmış (resim) PDF ve fotoğraflar ocr.py ile metne çevrilip aynı kurallarla okunur.
- Banka ekstresi (Excel/CSV): sütun başlıkları tanınır, satırlar okunur.

Sonuçlar ai.py ile aynı biçimde döner; böylece onay ekranı ikisi için de aynıdır.
"""
import base64
import csv
import io
import re
from collections import defaultdict
import unicodedata
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree as ET


class KuralHatasi(Exception):
    pass


# --- Ortak yardımcılar ----------------------------------------------------

def sade(metin: str) -> str:
    """Büyük/küçük harf ve Türkçe karakter farkını kaldırır: 'İnşaat' -> 'insaat'."""
    metin = (metin or "").replace("İ", "i").replace("I", "ı").lower()
    metin = metin.translate(str.maketrans("ıçğöşü", "icgosu"))
    metin = unicodedata.normalize("NFKD", metin).encode("ascii", "ignore").decode()
    return " ".join(metin.split())


_TR_ALFABE = "0123456789abcçdefgğhıijklmnoöpqrsştuüvwxyz"
_TR_SIRA = {h: i for i, h in enumerate(_TR_ALFABE)}


def tr_sira(metin: str | None) -> tuple:
    """Türkçe alfabe sırası, büyük/küçük harf gözetmeden (A B C Ç ... I İ ... O Ö ... U Ü); listelerin sıralama anahtarı."""
    metin = " ".join((metin or "").replace("İ", "i").replace("I", "ı").lower().split())
    return tuple(_TR_SIRA.get(h, -1 if h == " " else 100 + ord(h)) for h in metin)


def ad_anahtari(ad: str | None) -> str:
    """Aynı kaydın adı mı: harf/Türkçe karakter, nokta ve boşluk farkı yok sayılır ('Ofis', 'OFİS ', 'Ltd. Şti.' = 'LTD ŞTİ')."""
    return " ".join(re.sub(r"[^a-z0-9]+", " ", sade(ad)).split())


def kayit_adlari(kayit) -> list[str]:
    """Şantiye/depo adı ve önceki adları: ad değişse de eski adla gelen belge aynı kayda eşlenir."""
    return [kayit.ad, *[x.strip() for x in (getattr(kayit, "eski_adlar", None) or "").splitlines() if x.strip()]]


def adi_bu_mu(kayit, ad: str | None) -> bool:
    anahtar = ad_anahtari(ad)
    return bool(anahtar) and any(ad_anahtari(x) == anahtar for x in kayit_adlari(kayit))


def icindeki_kayit(metin: str | None, kayitlar: list):
    """Serbest yazının içinde adı (ya da eski adı) geçen kayıt ("USTA BORDÜR ÖRNEK GES İÇİN" -> "ÖRNEK GES"
    şantiyesi); birden çok geçiyorsa en uzun ad. 4 harften kısa adlar aranmaz."""
    m = f" {ad_anahtari(metin)} "
    adaylar = [(len(a), k) for k in kayitlar for x in kayit_adlari(k) if len(a := ad_anahtari(x)) >= 4 and f" {a} " in m]
    return max(adaylar, key=lambda t: t[0])[1] if adaylar else None


_ACIKLAMA_KELIMESI = re.compile(r"\b(icin|verildi|verilmek|alindi|odeme|odendi|odencek|odenecek|ceki|cari|ciro|edildi|kalan|"
                                r"tutari|yazildi|isi|bankasinda|bankada|kontrol)\b")


def santiye_ayikla(metin: str | None, santiyeler: list) -> tuple[str | None, str | None]:
    """Excel'in Şantiye sütunundaki yazı gerçekten şantiye adı mı? (şantiye adı, açıklamaya taşınacak yazı) döner.
    Var olan şantiye adı (ya da yazının içinde geçen var olan şantiye) şantiyedir; uzun ya da açıklama gibi yazı
    ("RAMAZAN VİNÇ KET İÇİN VERİLDİ") şantiye açtırmaz, açıklamaya gider; kısa bilinmeyen ad yeni şantiye olur."""
    metin = (metin or "").strip()
    if not metin:
        return None, None
    if any(adi_bu_mu(st, metin) for st in santiyeler):
        return metin, None
    if (st := icindeki_kayit(metin, santiyeler)) is not None:
        return st.ad, metin
    if "|" in metin or len(ad_anahtari(metin).split()) >= 4 or _ACIKLAMA_KELIMESI.search(sade(metin)):
        return None, metin
    return metin, None


_CIRO = re.compile(r"([^,|]+?)\s+firmas[ıi]na\s+ciro\s+edildi", re.I)
_BANKA_ANAHTARLARI = [("halk", "Halkbank"), ("ziraat", "Ziraat Bankası"), ("vakif", "VakıfBank"), ("garanti", "Garanti BBVA"),
                      ("yapi kredi", "Yapı Kredi"), ("yapikredi", "Yapı Kredi"), ("is bank", "İş Bankası"), ("isbank", "İş Bankası"),
                      ("akbank", "Akbank"), ("teb", "TEB"), ("qnb", "QNB"), ("deniz", "DenizBank"), ("kuveyt", "Kuveyt Türk"),
                      ("albaraka", "Albaraka Türk"), ("sekerbank", "Şekerbank"), ("ing", "ING")]


def cek_notu_coz(metin: str | None) -> dict:
    """Çek listesindeki not: "X firmasına ciro edildi" -> ciro edilen firma; "HALKBANKASINDA | HALKBANK ..." ya da
    "takasa/tahsile verildi" -> çekin verildiği banka."""
    s = sade(metin or "")
    sonuc = {"ciro": None, "banka": None}
    if m := _CIRO.search(metin or ""):
        sonuc["ciro"] = " ".join(m.group(1).split()).strip(" .-")
    elif re.search(r"bank\w*(da|ta)\b|takasa|tahsile|teminata", s):
        for anahtar, ad in _BANKA_ANAHTARLARI:
            if re.search(rf"\b{anahtar}", s):
                sonuc["banka"] = ad
                break
    return sonuc


def ad_degistir(kayit, yeni: str) -> None:
    """Adı değiştirir; eski ad takma ad olarak saklanır (kayıtlar kimlikle bağlı, ad yalnız görüntü)."""
    yeni = (yeni or "").strip()
    if not yeni or yeni == kayit.ad:
        return
    eskiler = kayit_adlari(kayit)
    anahtar = ad_anahtari(yeni)
    kalan = []
    for x in eskiler:
        if ad_anahtari(x) != anahtar and all(ad_anahtari(x) != ad_anahtari(y) for y in kalan):
            kalan.append(x)
    kayit.ad = yeni
    kayit.eski_adlar = "\n".join(kalan) or None


# Ünvanın şirket türü ekleri: "Gelen Firma Ltd." ile "GELEN FİRMA LTD. ŞTİ." aynı firmadır.
_HUKUKI_EK = {"ltd", "sti", "as", "a", "s", "limited", "sirketi", "sirket", "anonim", "tic", "ticaret", "san", "sanayi",
              "ve", "ltdsti", "koll", "kom", "kolektif", "komandit"}
# Benzerlik aranırken faaliyet kelimeleri de yok sayılır: "Gülmez" ~ "Gulmez İnşaat".
_FAALIYET_EK = {"ins", "insaat", "muh", "muhendislik", "taah", "taahhut", "ith", "ihr", "ithalat", "ihracat", "paz",
                "pazarlama", "nak", "nakliyat", "nakliye", "elek", "elektrik", "enerji", "mak", "makina", "makine", "otomotiv",
                "gida", "tekstil", "yapi", "malz", "malzemeleri", "malzemesi", "hizmetleri", "hiz", "turizm", "madencilik",
                "petrol", "urunleri", "group", "grup", "dis", "danismanlik", "mimarlik", "hafriyat", "beton", "insaatcilik"}


def unvan_cekirdegi(ad: str | None, faaliyetsiz: bool = False) -> str:
    """Ünvanın ayırt edici kısmı: şirket türü ekleri (ve istenirse faaliyet kelimeleri) atılır."""
    atilan = _HUKUKI_EK | (_FAALIYET_EK if faaliyetsiz else set())
    return " ".join(k for k in ad_anahtari(ad).split() if k not in atilan)


def benzer_unvan(a: str | None, b: str | None) -> bool:
    """İki cari adı aynı firmayı gösteriyor olabilir mi: ekler, Türkçe karakter, nokta/boşluk farkı yok sayılır;
    biri ötekinin başını kapsıyorsa ("Gülmez" / "Gülmez İnşaat Ltd") ya da küçük yazım farkı varsa benzer sayılır."""
    import difflib
    x, y = unvan_cekirdegi(a, True), unvan_cekirdegi(b, True)
    if not x or not y:
        x, y = unvan_cekirdegi(a), unvan_cekirdegi(b)
    if not x or not y:
        return False
    if x == y or x.replace(" ", "") == y.replace(" ", ""):
        return True
    kx, ky = x.split(), y.split()
    kisa, uzun = (kx, ky) if len(kx) <= len(ky) else (ky, kx)
    if uzun[:len(kisa)] == kisa and len("".join(kisa)) >= 4:
        return True
    return len(x) >= 5 and len(y) >= 5 and difflib.SequenceMatcher(None, x, y).ratio() >= 0.88


def para_coz(deger) -> float | None:
    """'1.234,56', '-1,234.56', '1234,5 TL', 1234.5 gibi değerleri sayıya çevirir."""
    if deger is None or deger == "":
        return None
    if isinstance(deger, (int, float)):
        return float(deger)
    s = re.sub(r"[^\d,.\-+]", "", str(deger))
    if not re.search(r"\d", s):
        return None
    eksi = s.startswith("-") or str(deger).strip().endswith("-")
    s = s.strip("+-")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):  # 1.234,56
            s = s.replace(".", "").replace(",", ".")
        else:  # 1,234.56
            s = s.replace(",", "")
    elif "," in s:
        # Tek virgül: 3 hane sonrası binlik ayraç (1,234), diğerleri ondalık (12,5 / 998,9496)
        s = s.replace(",", "") if s.count(",") > 1 or len(s) - s.rfind(",") - 1 == 3 else s.replace(",", ".")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s) - s.rfind(".") - 1 == 3):
        s = s.replace(".", "")  # 1.234 (binlik ayraç)
    try:
        sayi = float(s)
    except ValueError:
        return None
    return -sayi if eksi else sayi


def tarih_coz(deger) -> date | None:
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    s = str(deger or "").strip()
    m = re.match(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", s)
    if m:
        g, a, y = map(int, m.groups())
    else:
        m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
        if not m:
            return None
        y, a, g = map(int, m.groups())
    try:
        return date(y, a, g)
    except ValueError:
        return None


GIDER_ANAHTARLARI = [
    ("Beton", ["beton", "pompa"]), ("Demir", ["demir", "nervurlu", "hasir celik", "insaat celigi"]),
    ("Çimento", ["cimento"]), ("Tuğla/Blok", ["tugla", "briket", "bims", "gazbeton", "ytong"]),
    ("Kum/Çakıl", ["kum", "cakil", "micir", "agrega", "stabilize"]), ("Kereste", ["kereste", "kalas", "plywood", "kontrplak"]),
    ("Elektrik malzemesi", ["kablo", "sigorta", "priz", "anahtar", "elektrik malz", "role", "pano", "hucre", "bara ",
                            "trafo", "armatur", "aydinlatma", "kesici", "ayirici", "bakir", "nayy", "nyy", "nya ", "nyaf",
                            "xlpe", "ttr "]),
    ("Tesisat malzemesi", ["boru", "vana", "tesisat", "fitting", "pprc", "pvc"]),
    ("Boya/Yalıtım", ["boya", "yalitim", "mantolama", "izolasyon", "strafor", "membran"]),
    ("Nakliye", ["nakliye", "tasima", "navlun", "kargo", "gonderici odemeli", "alici odemeli", "posta hizmet"]),
    ("Makine kiralama", ["kiralama", "vinc", "ekskavator", "kepce", "forklift", "iskele"]),
    ("Akaryakıt", ["motorin", "benzin", "akaryakit", "diesel", "lpg"]), ("Yemek", ["yemek", "lokanta", "catering"]),
    ("Konaklama", ["otel", "konaklama", "pansiyon", "oda ucreti", "roommate", "breakfast", "kahvalti", "accommodation"]),
    ("Elektrik/Su/Doğalgaz", ["dogalgaz", "su bedeli", "elektrik bedeli", "elektrik enerjisi", "enerji bedeli",
                              "elektrik tuketim", "su tuketim", "enerji tuketim", "aktif enerji", "dagitim bedeli"]), ("Kira", ["kira"]),
    ("İşçilik/Taşeron", ["iscilik", "taseron", "hakedis", "montaj"]), ("Danışmanlık", ["danismanlik", "musavirlik", "is guvenligi", "osgb", "is yeri hekimi", "isyeri hekimi"]),
    ("Haberleşme/İnternet", ["superbox", "modem", "gsm ", "paketcell", "dsl", "fiber", "alo paket", "internet tarife",
                             "hat ucreti", "telsiz kullanim"]),
    ("Vade farkı/Faiz", ["vade farki"]),
    ("Seyahat/Bilet", ["bilet", "ucus", "havayolu", "alan vergisi", "pnr", "otobus"]),
]


def gider_turu_tahmin(*metinler: str) -> str | None:
    tum = " " + sade(" ".join(m for m in metinler if m)) + " "
    puan = {}
    for tur, kelimeler in GIDER_ANAHTARLARI:
        n = sum(len(re.findall(r"(?<![a-z0-9])" + re.escape(k), tum)) for k in kelimeler)  # "okuma"da "kum" aranmaz
        if n:
            puan[tur] = n
    return max(puan, key=puan.get) if puan else None


def kalem_gider_turu(kalemler: list[dict], *ek_metinler: str) -> str | None:
    """Gider türünü kalem tutarlarıyla ağırlıklı tahmin eder (küçük bir montaj kalemi asıl malzemenin önüne geçmesin);
    kalemlerden çıkmazsa satıcı ünvanı gibi ek metinlere bakar."""
    puan: dict[str, float] = defaultdict(float)
    for k in kalemler:
        tur = gider_turu_tahmin(k.get("aciklama") or "")
        if tur:
            puan[tur] += abs(float(k.get("tutar") or 0)) or 1
    return max(puan, key=puan.get) if puan else gider_turu_tahmin(*ek_metinler)


def santiye_tahmin(santiyeler: list[str], *metinler: str) -> str | None:
    tum = sade(" ".join(m for m in metinler if m))
    bulunan = [s for s in santiyeler if s and sade(s) in tum]
    return bulunan[0] if len(bulunan) == 1 else None


# --- e-Fatura XML (UBL-TR) ------------------------------------------------

def _yerel(etiket: str) -> str:
    return etiket.rsplit("}", 1)[-1]


def _bul(el, *yol):
    """Ad alanından bağımsız olarak alt öğe bulur: _bul(el, 'AccountingSupplierParty', 'Party')."""
    for parca in yol:
        if el is None:
            return None
        el = next((c for c in el if _yerel(c.tag) == parca), None)
    return el


def _hepsi(el, ad):
    return [c for c in el if _yerel(c.tag) == ad] if el is not None else []


def _metin(el, *yol) -> str | None:
    e = _bul(el, *yol)
    return e.text.strip() if e is not None and e.text and e.text.strip() else None


def _sayi(el, *yol) -> float | None:
    return para_coz(_metin(el, *yol))


def _taraf(parti) -> dict:
    if parti is None:
        return {"unvan": None, "vkn": None, "vergi_dairesi": None, "adres": None, "telefon": None, "iban": None}
    vkn = None
    for pid in _hepsi(parti, "PartyIdentification"):
        kimlik = _bul(pid, "ID")
        if kimlik is not None and kimlik.get("schemeID", "").upper() in ("VKN", "TCKN") and kimlik.text:
            vkn = kimlik.text.strip()
            break
    unvan = _metin(parti, "PartyName", "Name")
    if not unvan:
        ad = " ".join(x for x in [_metin(parti, "Person", "FirstName"), _metin(parti, "Person", "FamilyName")] if x)
        unvan = ad or None
    adres_el = _bul(parti, "PostalAddress")
    adres = " ".join(x for x in [_metin(adres_el, "StreetName"), _metin(adres_el, "BuildingNumber"),
                                 _metin(adres_el, "CitySubdivisionName"), _metin(adres_el, "CityName")] if x) or None
    return {"unvan": unvan, "vkn": vkn, "vergi_dairesi": _metin(parti, "PartyTaxScheme", "TaxScheme", "Name"),
            "adres": adres, "telefon": _metin(parti, "Contact", "Telephone"), "iban": None}


KDV_KODU = "0015"  # UBL-TR TaxTypeCode: KDV. Diğer kodlar (ÖTV 0071, ÖİV 4080, konaklama 0059...) diğer vergidir.


def _vergi_kodu(alt) -> str | None:
    return _metin(alt, "TaxCategory", "TaxScheme", "TaxTypeCode")


def _vergiler(vergi_toplami_listesi) -> tuple[float, float, float | None]:
    """TaxTotal öğelerinden (KDV, diğer vergiler, KDV oranı) döner. Kodsuz alt toplam KDV sayılır."""
    kdv = diger = 0.0
    oran = None
    for t in vergi_toplami_listesi:
        altlar = _hepsi(t, "TaxSubtotal")
        if not altlar:
            kdv += _sayi(t, "TaxAmount") or 0
            continue
        for alt in altlar:
            kod = _vergi_kodu(alt)
            tutar = _sayi(alt, "TaxAmount") or 0
            if kod in STOPAJ_KODLARI:
                continue  # stopaj diğer vergi değildir: ödenecekten düşülür (_stopaj)
            if kod in (None, KDV_KODU):
                kdv += tutar
                if oran is None:
                    oran = _sayi(alt, "Percent")
            else:
                diger += tutar
    return round(kdv, 2), round(diger, 2), oran


STOPAJ_KODLARI = {"0003", "0011"}  # UBL-TR: gelir vergisi stopajı, kurumlar vergisi stopajı


def _stopaj(vergi_toplami_listesi) -> tuple[float, float | None]:
    """TaxTotal içindeki stopaj (0003/0011): (tutar, yüzde)."""
    tutar, oran = 0.0, None
    for t in vergi_toplami_listesi:
        for alt in _hepsi(t, "TaxSubtotal"):
            if _vergi_kodu(alt) in STOPAJ_KODLARI:
                tutar += _sayi(alt, "TaxAmount") or 0
                oran = oran or _sayi(alt, "Percent") or _sayi(alt, "TaxCategory", "Percent")
    return round(tutar, 2), oran


def _iskonto_orani(el) -> str:
    """Satır iskontosu açıklaması: 'İsk. %10' ya da tutar."""
    parcalar = []
    for ac in _hepsi(el, "AllowanceCharge"):
        if (_metin(ac, "ChargeIndicator") or "").lower() != "false":
            continue
        carpan = _sayi(ac, "MultiplierFactorNumeric")
        tutar = _sayi(ac, "Amount") or 0
        if not tutar:
            continue
        if carpan:
            yuzde = carpan * 100 if carpan <= 1 else carpan
            parcalar.append("İsk. %" + f"{yuzde:g}".replace(".", ","))
        else:
            parcalar.append("İsk. " + f"{tutar:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " TL")
    return ", ".join(parcalar)


def _xml_tevkifat(listesi) -> tuple[float | None, str | None]:
    """WithholdingTaxTotal: Percent (20 = 2/10) ve TaxTypeCode (624). (oran onda kaç, kod)."""
    for w in listesi:
        for alt in _hepsi(w, "TaxSubtotal"):
            yuzde = _sayi(alt, "Percent") or _sayi(alt, "TaxCategory", "Percent")
            kod = _vergi_kodu(alt)
            kod = kod if kod in TEVKIFAT_KODLARI else None
            pay = round(yuzde / 10, 2) if yuzde else (TEVKIFAT_KODLARI[kod][0] if kod else None)
            if pay:
                return pay, kod
    return None, None


def efatura_xml_oku(veri: bytes, santiyeler: list[str] | None = None) -> dict:
    try:
        kok = ET.fromstring(veri)
    except ET.ParseError as e:
        raise KuralHatasi("XML dosyası okunamadı.") from e
    if _yerel(kok.tag) != "Invoice":
        raise KuralHatasi("Bu XML bir e-Fatura değil.")

    kalemler = []
    satir_tevkifat = satir_diger = 0.0
    for satir in _hepsi(kok, "InvoiceLine"):
        miktar_el = _bul(satir, "InvoicedQuantity")
        miktar = (para_coz(miktar_el.text) if miktar_el is not None else None) or 1
        tutar = _sayi(satir, "LineExtensionAmount")
        fiyat = _sayi(satir, "Price", "PriceAmount")
        aciklama = _metin(satir, "Item", "Name") or _metin(satir, "Item", "Description") or "-"
        kod = _metin(satir, "Item", "SellersItemIdentification", "ID")
        if kod and kod not in aciklama:
            aciklama = f"{aciklama} ({kod})"
        satir_kdv, diger, oran = _vergiler(_hepsi(satir, "TaxTotal"))
        satir_diger += diger
        satir_tevkifat += sum(_sayi(w, "TaxAmount") or 0 for w in _hepsi(satir, "WithholdingTaxTotal"))
        tevk_pay, tevk_kod = _xml_tevkifat(_hepsi(satir, "WithholdingTaxTotal"))
        satir_stopaj, stopaj_yuzde = _stopaj(_hepsi(satir, "TaxTotal"))
        if satir_stopaj and not stopaj_yuzde and tutar:
            stopaj_yuzde = round(satir_stopaj / tutar * 100, 2)
        # Satır iskontosu: Price brüt, LineExtensionAmount net. Kayıtta tutar net olsun diye fiyat nete çevrilir.
        isk = _iskonto_orani(satir)
        if tutar is not None and (fiyat is None or abs(round(miktar * fiyat, 2) - tutar) > 0.01):
            if isk:
                aciklama = f"{aciklama} ({isk}, liste fiyatı " + f"{fiyat:g})".replace(".", ",") if fiyat else f"{aciklama} ({isk})"
            fiyat = round(tutar / miktar, 6) if miktar else tutar
        kalemler.append({
            "aciklama": aciklama,
            "miktar": miktar,
            "birim": BIRIMLER.get(miktar_el.get("unitCode", ""), miktar_el.get("unitCode")) if miktar_el is not None else None,
            "birim_fiyat": fiyat,
            "kdv_orani": oran if oran is not None else 0,
            "tutar": tutar,
            **({"tevkifat_orani": tevk_pay, "tevkifat_kodu": tevk_kod} if tevk_pay and oran else {}),
            **({"stopaj_orani": stopaj_yuzde} if satir_stopaj else {}),
            **({"kdv_tutari": satir_kdv} if tutar is not None and oran and abs(tutar * oran / 100 - satir_kdv) > 0.05 else {}),
        })
    toplam = _bul(kok, "LegalMonetaryTotal")
    tevkifat = sum(_sayi(w, "TaxAmount") or 0 for w in _hepsi(kok, "WithholdingTaxTotal")) or satir_tevkifat
    kdv, diger, _ = _vergiler(_hepsi(kok, "TaxTotal"))
    stopaj, stopaj_yuzde = _stopaj(_hepsi(kok, "TaxTotal"))
    stopaj = stopaji_dagit(kalemler, stopaj, stopaj_yuzde)
    if any("kdv_tutari" in k for k in kalemler) and abs(sum((k["tutar"] or 0) * k["kdv_orani"] / 100 for k in kalemler) - kdv) \
            <= abs(sum(kalem_kdv(k) for k in kalemler) - kdv):
        for k in kalemler:
            k.pop("kdv_tutari", None)
    if not diger:
        diger = satir_diger
    notlar = [n.text.strip() for n in _hepsi(kok, "Note") if n.text and n.text.strip()]
    teslim = _bul(kok, "Delivery", "DeliveryAddress")
    teslim_metni = " ".join((e.text or "") for e in teslim.iter()) if teslim is not None else ""
    vade = _metin(kok, "PaymentMeans", "PaymentDueDate") or _metin(kok, "PaymentTerms", "PaymentDueDate")
    satici = _taraf(_bul(kok, "AccountingSupplierParty", "Party"))
    iban = re.sub(r"\s", "", _metin(kok, "PaymentMeans", "PayeeFinancialAccount", "ID") or "").upper()
    if iban.startswith("TR") and len(iban) == 26:
        satici["iban"] = iban

    ara = _sayi(toplam, "TaxExclusiveAmount") if toplam is not None else None
    uyari = []
    kalem_toplami = round(sum(k["tutar"] or 0 for k in kalemler), 2)
    if ara is not None and kalemler and abs(kalem_toplami - ara) > 0.01:
        fark = round(kalem_toplami - ara, 2)
        if fark > 0:  # fatura altı iskonto
            kalemler += _iskonto_kalemleri(kalemler, fark) or []
        else:  # fatura altı ek bedel (nakliye vb.)
            nedeni = next((_metin(ac, "AllowanceChargeReason") for ac in _hepsi(kok, "AllowanceCharge")
                           if (_metin(ac, "ChargeIndicator") or "").lower() == "true"), None)
            oran = max(kalemler, key=lambda k: abs(k["tutar"] or 0))["kdv_orani"]
            kalemler.append({"aciklama": nedeni or "Fatura altı ek bedel", "miktar": 1, "birim": None,
                             "birim_fiyat": -fark, "kdv_orani": oran, "tutar": -fark})
    if tevkifat and not any(k.get("tevkifat_orani") for k in kalemler):
        # Tevkifat yalnız belge düzeyinde: oranı (Percent 20 = 2/10) ve kodu tutan satırlara dağıtılır.
        pay, kod = _xml_tevkifat(_hepsi(kok, "WithholdingTaxTotal"))
        if not _tevkifati_dagit(kalemler, round(tevkifat, 2), pay=pay, kod=kod):
            uyari.append("Faturada tevkifat var ama hangi kalemlere ait olduğu bulunamadı; kalemlerde tevkifat oranını seçin.")
    elif tevkifat:
        satir_toplam = round(sum(tevkifat_tutari(k["tutar"], k["kdv_orani"], k.get("tevkifat_orani"), k.get("kdv_tutari")) for k in kalemler), 2)
        if abs(satir_toplam - tevkifat) > 0.02 * len(kalemler):
            uyari.append(f"Satırlardan hesaplanan tevkifat ({_sayi_yazi(satir_toplam)} TL) faturadaki tevkifattan "
                         f"({_sayi_yazi(tevkifat)} TL) farklı; kalemlerdeki tevkifat oranlarını kontrol edin.")
    tip = (_metin(kok, "InvoiceTypeCode") or "").upper()
    if tip == "IADE":
        uyari.append("Bu bir İADE faturası; alış/satış yönünü kontrol edin.")
    para = _metin(kok, "DocumentCurrencyCode") or "TRY"
    kur = None
    if para.upper() not in ("TRY", "TL"):
        kur = _sayi(kok, "PricingExchangeRate", "CalculationRate")
        uyari.append(f"Fatura {para} cinsinden" + (f" (kur {kur:g})." if kur else "; kuru girin."))

    return {
        "belge_fatura_mi": True,
        "fatura_no": _metin(kok, "ID"),
        "ettn": _metin(kok, "UUID"),
        "fatura_tipi": tip or None,
        "tarih": _metin(kok, "IssueDate"),
        "vade": vade,
        "satici": satici,
        "alici": {k: v for k, v in _taraf(_bul(kok, "AccountingCustomerParty", "Party")).items() if k in ("unvan", "vkn")},
        "kalemler": [iskontoyu_ayir(k) for k in kalemler],
        "ara_toplam": ara,
        "kdv_toplam": kdv,
        "tevkifat": round(tevkifat, 2),
        "stopaj": stopaj,
        "diger_vergiler": round(diger, 2),
        "genel_toplam": _sayi(toplam, "PayableAmount") if toplam is not None else None,
        "para_birimi": para,
        "kur": kur,
        "onerilen_santiye": santiye_tahmin(santiyeler or [], teslim_metni, *notlar),
        "onerilen_gider_turu": kalem_gider_turu(kalemler, satici.get("unvan") or ""),
        "notlar": " ".join(uyari),
        "okuma_yontemi": "e-Fatura XML",
    }


def efatura_gorunumu(veri: bytes) -> tuple[bytes, str] | None:
    """e-Fatura XML'inin içindeki görüntü: gömülü PDF, yoksa gömülü XSLT şablonuyla üretilen HTML.

    XSLT için lxml gerekir; yoksa ya da şablon hatalıysa None döner (fatura yine okunur).
    """
    try:
        kok = ET.fromstring(veri)
    except ET.ParseError:
        return None
    xslt = None
    for ref in _hepsi(kok, "AdditionalDocumentReference"):
        nesne = _bul(ref, "Attachment", "EmbeddedDocumentBinaryObject")
        if nesne is None or not (nesne.text or "").strip():
            continue
        ad = (nesne.get("filename") or "").lower()
        tur = (nesne.get("mimeCode") or "").lower()
        try:
            icerik = base64.b64decode(nesne.text.strip(), validate=False)
        except ValueError:
            continue
        if tur == "application/pdf" or ad.endswith(".pdf") or icerik[:4] == b"%PDF":
            return icerik, "pdf"
        if ad.endswith((".xslt", ".xsl")) or b"xsl:stylesheet" in icerik[:2000]:
            xslt = xslt or icerik
    if not xslt:
        return None
    try:
        from lxml import etree
    except ImportError:
        return None
    try:
        guvenli = etree.XMLParser(resolve_entities=False, no_network=True)
        donustur = etree.XSLT(etree.fromstring(xslt, guvenli), access_control=etree.XSLTAccessControl.DENY_ALL)
        html = bytes(donustur(etree.fromstring(veri, guvenli)))
    except Exception:  # noqa: BLE001 - firmaların şablonları bozuk olabilir; görüntü şart değil
        return None
    return (html, "html") if html.strip() else None


BIRIMLER = {"C62": "adet", "NIU": "adet", "KGM": "kg", "TNE": "ton", "MTR": "m", "MTK": "m²", "MTQ": "m³",
            "LTR": "lt", "HUR": "saat", "DAY": "gün", "MON": "ay", "SET": "set", "PA": "paket", "BX": "kutu", "PR": "çift",
            "Çiti": "Çift"}  # "Çift"in bağlı harfi bazı PDF'lerde boş karakter çıkar, "ti" sanılır


def xml_dosyalari(veri: bytes, ad: str) -> list[bytes]:
    """Tek XML veya içinde XML'ler bulunan zip (GİB portalından indirilen) dosyasını açar."""
    return [x for x, _ in xml_ve_gorunumler(veri, ad)]


def xml_ve_gorunumler(veri: bytes, ad: str) -> list[tuple[bytes, tuple[bytes, str] | None]]:
    """Her e-Fatura XML'i ve (varsa) görüntüsü: zip'te aynı adlı PDF/HTML, yoksa XML'e gömülü görüntü.

    Zip içinde zip (portaldan toplu indirme) de açılır.
    """
    if not (ad.lower().endswith(".zip") or veri[:2] == b"PK"):
        return [(veri, efatura_gorunumu(veri))]
    try:
        dosyalar = _zip_ac(veri)
    except zipfile.BadZipFile as e:
        raise KuralHatasi("Zip dosyası açılamadı.") from e
    xmller = {n: v for n, v in dosyalar.items() if n.lower().endswith(".xml")}
    gorseller = {n: v for n, v in dosyalar.items() if n.lower().endswith((".pdf", ".html", ".htm"))}
    sonuc = []
    for n, x in xmller.items():
        kok = n.rsplit(".", 1)[0]
        eslesen = next((g for g in gorseller if g.rsplit(".", 1)[0] == kok), None)
        if eslesen is None and len(xmller) == 1 and len(gorseller) == 1:
            eslesen = next(iter(gorseller))
        gorunum = (gorseller[eslesen], "pdf" if eslesen.lower().endswith(".pdf") else "html") if eslesen else efatura_gorunumu(x)
        sonuc.append((x, gorunum))
    return sonuc


ZIP_ACILMIS_SINIR = 300_000_000  # açılınca şişen ZIP belleği doldurmasın


def _zip_ac(veri: bytes, derinlik: int = 0) -> dict[str, bytes]:
    dosyalar = {}
    with zipfile.ZipFile(io.BytesIO(veri)) as z:
        if sum(b.file_size for b in z.infolist()) > ZIP_ACILMIS_SINIR:
            raise KuralHatasi("ZIP dosyası açılınca çok büyük oluyor; faturaları birkaç ZIP'e bölüp yükleyin.")
        for bilgi in z.infolist():
            if bilgi.is_dir() or bilgi.file_size > 20_000_000:
                continue
            ad = bilgi.filename
            if ad.lower().endswith(".zip") and derinlik < 2:
                dosyalar.update({f"{ad}/{k}": v for k, v in _zip_ac(z.read(bilgi), derinlik + 1).items()})
            elif ad.lower().endswith((".xml", ".pdf", ".html", ".htm")):
                dosyalar[ad] = z.read(bilgi)
    return dosyalar


# --- PDF metni ------------------------------------------------------------

_BOLUNMUS_HARF = re.compile(r"(?<=[A-Za-zÇĞİÖŞÜçğıöşü]) [ÇĞİÖŞÜçğıöşü]{1,3} (?=[A-Za-zÇĞİÖŞÜçğıöşü])")


_METIN_ONBELLEK: dict[str, str] = {}


def pdf_metni(veri: bytes) -> str:
    """PDF'in yazısı. Aynı belge bir yüklemede birkaç kez okunur (yön, taraflar, fatura): son belgeler bellekte tutulur."""
    import hashlib
    anahtar = hashlib.sha1(veri).hexdigest()
    if anahtar not in _METIN_ONBELLEK:
        metin = _pdf_metni(veri)
        if len(_METIN_ONBELLEK) >= 16:
            _METIN_ONBELLEK.pop(next(iter(_METIN_ONBELLEK)))
        _METIN_ONBELLEK[anahtar] = metin
    return _METIN_ONBELLEK[anahtar]


def _pdf_metni(veri: bytes) -> str:
    from pypdf import PdfReader
    try:
        okuyucu = PdfReader(io.BytesIO(veri))
        metin = "\n".join((s.extract_text() or "") for s in okuyucu.pages[:10])
    except Exception as e:
        raise KuralHatasi("PDF açılamadı.") from e
    if len(_BOLUNMUS_HARF.findall(metin)) >= 3:
        # Bazı entegratörlerin PDF'inde Türkçe harfler ayrı yazı parçası: "ELEKTR İ K", "Y İĞİ THAN".
        # pdfplumber harfleri konumlarından birleştirir; o metin kullanılır.
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(veri)) as pdf:
                duz = "\n".join((s.extract_text() or "") for s in pdf.pages[:10])
            if len(duz.strip()) >= 30 and len(_BOLUNMUS_HARF.findall(duz)) < len(_BOLUNMUS_HARF.findall(metin)):
                metin = duz
        except Exception:  # noqa: BLE001
            pass
    # Bazı PDF'lerde 'ti' bağlı harfi boş karakter olarak çıkar ("Servis Ücre\x00" = "Servis Ücreti")
    metin = metin.replace("\x00", "ti").replace("ﬁ", "fi").replace("ﬂ", "fl")
    # Noktalı kısaltmalar tek yazıma: "K.D.V. Oranı", "Hesaplanan K.D.V. (%20)" -> KDV
    return re.sub(r"\bK\.\s?D\.\s?V\.?", "KDV", metin, flags=re.IGNORECASE)


def _ilk(desen: str, metin: str, bayrak=re.IGNORECASE) -> str | None:
    m = re.search(desen, metin, bayrak)
    return m.group(1).strip() if m else None


_VKN_SATIRI = re.compile(r"(?:\b(?:VKN|TCKN)|\bVergi\s*(?:Kimlik\s*)?(?:Numaras[ıi]|No)(?:\s*/\s*Dairesi)?|\bV\.\s?D\.|\bDairesi\s*/)\s*[:./]?\s*(?:(?:VKN|TCKN)\s*:?\s*)?"
                         r"(?P<no>\d{10,11}|\d{3} \d{3} \d{4})(?!\d)", re.IGNORECASE)  # "VKN: 0000000000TICARETSICILNO" bitişik yazılabilir
_ADRES = re.compile(r"\b(mah|mahallesi|mh|cad|cd|caddesi|sk|sok|sokak|sokagi|blv|bulvar|bulvari|no|kapi|blok|apt|site|"
                    r"kat|daire|osb|turkiye|bina|mudurluk|mudurlugu|sube|subesi|koy|koyu|mevkii|mevkisi|beldesi|kasabasi)\b|/")
# Taraf bloğunun başlığı ("DÜZENLEYEN", "SATICI BİLGİLERİ"): ünvanın parçası değil
_TARAF_BASLIGI = re.compile(r"^(duzenleyen|satici|alici|sayin|gonderen|gonderici|tedarikci|musteri|hizmet veren|fatura sahibi)"
                            r"(?: (bilgileri|firma|unvani))?$")
# "Tel" etiketi ama "GARANTİ TEL ÇİT" (tel örgü) ünvanın parçası: etiket ancak sonunda numara/iki nokta varsa.
_ETIKET = re.compile(r":|@|www\.|\.com|\.net|\.tr\b|\b(?:fax|faks|tel|telefon)\b\s*(?:no\b|[:(+\d]|$)|\be-?posta\b|"
                     r"\b(?:ticaret )?sicil no\b|\bmersis\b|\bvergi (?:no|dairesi|kimlik)\b")
_DUR = re.compile(r"\b(sayin|alici|alici bilgileri|e-?fatura|e-?arsiv|ettn|yalniz|not)\b|uretilmis|olusturulmus|hazirlanmis|e-?donusum|\d{2}[-./]\d{2}[-./]\d{4}|#")


_FAALIYET = re.compile(r"\b(her turlu|satisi|satis|taahhut|isleri|hizmetleri|nakliye|imalati|imalat|ticareti|malzemeleri|"
                       r"malzemesi|doseme|duzenleme|tamiri|bakimi|onarim|uygulama|uygulamalari|danismanlik)\b")


def faaliyeti_ayir(unvan: str | None) -> str | None:
    """Şahıs firmasının adından sonra yazılan faaliyet açıklaması ünvana katılmaz:
    'KEMAL GÜLMEZ Her Türlü İnşaat Malzemeleri Satışı- Kaldırım Parke...' -> 'KEMAL GÜLMEZ'.
    Büyük harfli ad en az iki kelime, devamı küçük harfli ve faaliyet anlatıyorsa (şirket eki yoksa) ayrılır."""
    if not unvan:
        return unvan
    kelimeler = unvan.split()
    for i in range(2, len(kelimeler) - 1):
        k = kelimeler[i]
        if re.search(r"[a-zçğıöşü]", k) and all(not re.search(r"[a-zçğıöşü]", x) for x in kelimeler[:i]):
            kalan = sade(" ".join(kelimeler[i:]))
            if _FAALIYET.search(kalan) and not SIRKET_EKLERI.search(kalan):
                return " ".join(kelimeler[:i]).rstrip(" -,.")
            break
    return unvan


def _unvan_bul(satirlar: list[str], vkn_i: int) -> str | None:
    unvan = faaliyeti_ayir(_unvan_satirlari(satirlar, vkn_i))
    if unvan:
        unvan = " ".join(unvan.split())
    if unvan and (m := re.fullmatch(r"(.+?) \1", unvan)):
        unvan = m.group(1)  # ad iki satırda aynen tekrar yazılmış ("YÜKSEL YALÇIN" / "YÜKSEL YALÇIN")
    elif unvan and (m := re.fullmatch(r"(.+ (\S+ \S+(?: \S+)?)) \2", unvan)):
        unvan = m.group(1)  # "... NUMARATAJ SAN. MEHMET SELÇUKİ MEHMET SELÇUKİ": sahibin adı iki kez
    return unvan


def _unvan_satirlari(satirlar: list[str], vkn_i: int) -> str | None:
    """VKN satırından geriye yürüyerek firma ünvanını bulur.

    e-Fatura görüntülerinde taraf bloğu "ünvan / adres / tel / e-posta / vergi dairesi / VKN" sırasıyladır:
    önce etiket ve adres satırları atlanır, sonra en fazla üç ünvan satırı toplanır.
    """
    i, adim = vkn_i - 1, 0
    while i >= 0 and adim < 10:
        sa = sade(satirlar[i])
        ust = sade(satirlar[i - 1]) if i > 0 else ""
        tek_kelime_sehir = len(sa.split()) == 1 and bool(_ADRES.search(ust))  # "KÜTAHYA" gibi adres devamı
        if not sa or _ETIKET.search(sa) or _ADRES.search(sa) or re.match(r"^\d", sa) or tek_kelime_sehir \
                or not re.search(r"[a-z]", sa):  # boş bırakılmış adres satırı: "-"
            i, adim = i - 1, adim + 1
            continue
        break

    def topla(i, en_fazla=3):
        toplanan = []
        while i >= 0 and len(toplanan) < en_fazla:
            ham = re.sub(r"\s*\(\d{1,3}\)$", "", satirlar[i].strip())  # "MEHMET BAYAR(00)": şube/sıra kodu
            sa = sade(ham)
            sa_ = re.sub(r"(ortakligi)\s*-\s*\d{1,2}$", r"\1", sa)  # "... İŞ ORTAKLIĞI -3": ortaklığın sıra numarası
            etiketli = re.match(r"^(?:satici|alici|duzenleyen|unvan|unvani|firma|firma adi|musteri)\b[^:]{0,15}:\s*(\S.*)$", sa)
            if etiketli and not _ADRES.search(etiketli.group(1)):
                toplanan.insert(0, ham[ham.index(":") + 1:].strip())
                break
            if re.search(r"\borganize sanayi bolgesi\b", sa) and not re.search(r"\b(mh|mah|cad|cd|sk|no)\b", sa):
                toplanan.insert(0, ham)  # "KÜTAHYA 30 AĞUSTOS ORGANİZE SANAYİ BÖLGESİ": rakam ve "OSB" adın parçası
                break
            if (not re.search(r"[a-z]", sa) or _ETIKET.search(sa) or _ADRES.search(sa) or _DUR.search(sa) or _TARAF_BASLIGI.match(sa)
                    or re.search(r"(?<![a-z0-9])\d+(?![a-z0-9])", sa_)
                    or re.fullmatch(r"[0-9a-f-]{8,}", sa)  # ETTN'nin alt satıra taşan parçası
                    or re.match(r"^(senaryo|fatura|ozellestirme|duzenleme|siparis|irsaliye)", sa)):
                break
            toplanan.insert(0, ham)
            i -= 1
        # Satır sonunda tek kalan ek harfi: "LİMİTED ŞİRKET⏎İ" / "ŞİRKET İ" -> "ŞİRKETİ"
        return re.sub(r"(?<=[A-ZÇĞİÖŞÜ]{3}) ([İI])(?= |$)", r"\1", " ".join(toplanan)) or None

    unvan = topla(i)
    if not unvan or not SIRKET_EKLERI.search(sade(unvan)):
        # "ANKARA Öveçler" gibi bir adres parçası alındıysa yukarıda şirket eki taşıyan ünvan bloğunu ara
        # (Nova Grup: ünvan üç satır, altında iki satır adres; Kalkersan: beş satır ünvan, "Adres:" ile başlayan üç satır;
        # Tek-Gökay: ünvanın altında numaralı adres satırı "GEBECELER KSB FATİH M 4").
        for j in range(i - 1, max(i - 8, -1), -1):
            if _VKN_SATIRI.search(satirlar[j]) or _DUR.search(sade(satirlar[j])):
                break  # öteki tarafın bloğuna geçildi (şahıs ünvanında üstteki alıcı firmayı alma)
            if SIRKET_EKLERI.search(sade(satirlar[j])) and not _ETIKET.search(sade(satirlar[j])):
                return topla(j, 5) or unvan
    return unvan


def pdf_taraflari(metin: str) -> tuple[dict, dict]:
    """(satıcı, alıcı) — 'SAYIN' / 'ALICI' işaretinden sonraki ilk VKN alıcıya aittir."""
    satirlar = metin.splitlines()
    vknler = [(i, m.group("no").replace(" ", "")) for i, sat in enumerate(satirlar) for m in [_VKN_SATIRI.search(sat)] if m]
    isaret = next((i for i, sat in enumerate(satirlar) if re.match(r"^(sayin|alici)\b", sade(sat))), None)
    alici_i = next((v for v in vknler if isaret is not None and v[0] > isaret), None)
    if alici_i is None and len(vknler) > 1:
        alici_i = vknler[1]
    # Alıcının VKN'si blokta iki kez yazılabilir ("VKN: 0000000000" adres altında ve vergi dairesi altında)
    satici_i = next((v for v in vknler if not alici_i or v[1] != alici_i[1]), None)

    def taraf(v):
        return {"unvan": _unvan_bul(satirlar, v[0]) if v else None, "vkn": v[1] if v else None}
    satici, alici = taraf(satici_i), taraf(alici_i)
    if isaret is not None and alici_i is not None and alici_i[0] > isaret:
        # "SAYIN" bloğunda ad ilk satırdadır. VKN'den geriye yürüyerek ad bulunamadıysa ya da şirket eki taşımayan bir
        # kurum adı ("ÖRNEK KÖYÜ MUHTARLIĞI", "... ORGANİZE SANAYİ BÖLGESİ") atlandıysa ad oradan alınır.
        ileri = _sayin_sonrasi_ad(satirlar, isaret, alici_i[0])
        if ileri and (not alici["unvan"] or kotu_unvan(alici["unvan"])
                      or (not SIRKET_EKLERI.search(sade(alici["unvan"])) and _KAMU.search(sade(ileri))
                          and not _KAMU.search(sade(alici["unvan"])))):
            alici["unvan"] = ileri
    return satici, alici


def pdf_yonu(veri: bytes, sirket_vkn: str | None, karekod_yonu=None) -> str | None:
    """PDF faturanın yönü belgeden: satıcı VKN'si bizimse satış (giden), alıcınınki bizimse alış (gelen); bilinmiyorsa None.
    Kullanıcının seçtiği tür ne olursa olsun belge esastır (satış faturası alışa düşmez). Yazıda yalnız bizim VKN'miz
    bulunduysa (Turkcell gibi satıcının VKN'si yazı olarak basılmamış) karar GİB karekodundan verilir (karekod_yonu())."""
    vkn = re.sub(r"\D", "", sirket_vkn or "")
    if not vkn:
        return None
    try:
        satici, alici = pdf_taraflari(pdf_metni(veri))
    except KuralHatasi:
        return None
    # Satış görünüyorsa da karekoda bakılır: "SAYIN" yazmayan düzende ilk VKN (bizimki) satıcı sanılabilir
    if karekod_yonu and (not (satici["vkn"] and alici["vkn"]) or satici["vkn"] == vkn):
        if (y := karekod_yonu()):
            return y
    if satici["vkn"] == vkn and alici["vkn"] and alici["vkn"] != vkn and not satici["unvan"] and alici["unvan"]:
        # Bizim VKN'miz adımız yazılmadan en üstte (müşteri bilgisi), karşı firma ünvanıyla: alış (TTNET faturası).
        # Kendi kestiğimiz faturada ünvanımız VKN'mizin yanında yazar.
        return "gelen"
    if satici["vkn"] == vkn and alici["vkn"] != vkn:
        return "giden"
    if alici["vkn"] == vkn and satici["vkn"] != vkn:
        return "gelen"
    return None


def _sayin_sonrasi_ad(satirlar: list[str], isaret: int, son: int) -> str | None:
    """'SAYIN' satırından sonraki ad: ilk dolu satır (etiketten sonra yazılmışsa o kısım) ve şirket ekiyle süren devamı.
    Adres, etiket ve rakamlı satırlar ad değildir; kurum adındaki köy/mahalle adı adres sayılmaz."""
    ilk = re.sub(r"^\s*(?:sayin|alici|SAYIN|ALICI|Sayın|SAYIN)\b[^:\n]{0,15}:?\s*", "", satirlar[isaret], flags=re.IGNORECASE).strip()
    adaylar = ([ilk] if ilk else []) + [x.strip() for x in satirlar[isaret + 1:son]]
    ad: list[str] = []
    for x in adaylar:
        sa = sade(x)
        if not re.search(r"[a-z]{2}", sa) or _ETIKET.search(sa) or _DUR.search(sa):
            if ad or not sa or not re.search(r"[a-z]{2}", sa):
                break
            continue
        kurum = bool(_KAMU.search(sa))
        adres = bool(_ADRES.search(_KAMU_BIRIMI.sub(" ", sa) if kurum else sa)) or bool(re.search(r"(?<![a-z0-9])\d+(?![a-z0-9])", sa)) \
            and not kurum and not SIRKET_EKLERI.search(sa)
        if adres:
            break
        if ad and not (SIRKET_EKLERI.search(sa) or _KAMU.search(sa) or re.fullmatch(r"(sirketi?|sti\.?)", sa)):
            break  # ikinci satır yalnız ünvanın devamıysa ("ŞİRKETİ") alınır
        ad.append(x)
        if len(ad) == 3:
            break
    return " ".join(ad) or None


_PB = r"(?:TL|TRY|₺)"  # para birimi yazımı


def _kodu_sona_al(aciklama: str) -> str:
    """'HG.SR.01.01.005 Demontaj Bedeli' -> 'Demontaj Bedeli (HG.SR.01.01.005)': malzeme kodu açıklamanın önüne geçmesin."""
    parca = aciklama.split(None, 1)
    if len(parca) == 2 and parca[0].endswith("-"):  # "CON2057516- 1 205/75R16 ..." satır kaymasıyla bölünmüş kod
        devam = parca[1].split(None, 1)
        if len(devam) == 2 and (re.fullmatch(r"\d{1,3}", devam[0])  # "NET- MU.VYU.00004 NETELSAN ..."
                                or re.fullmatch(r"[A-Za-z0-9.\-_/]*\d[A-Za-z0-9.\-_/]*", devam[0]) and re.search(r"[.\-_/]", devam[0])):
            parca = [parca[0] + devam[0], devam[1]]
    if len(parca) == 2:
        kod = parca[0].rstrip("-")
        ascii_kod = re.fullmatch(r"[A-Za-zÇĞİÖŞÜ0-9.\-_/]{4,}", kod) is not None
        harfli = re.search(r"\d", kod) and (re.search(r"[.\-_/]", kod) or len(re.findall(r"\d", kod)) >= 2) \
            and re.search(r"[A-Za-z]", kod)
        if ascii_kod and (harfli or re.fullmatch(r"\d{6,}", kod)) and len(parca[1]) > 3:  # "1010-Kamyonet PM" kalır  # "01010020008 VADE FARKI" gibi rakamlı kodlar da
            return f"{parca[1]} ({parca[0]})"
    return aciklama


# e-Fatura birim kodları ("1 C62" = 1 adet) ve KDV istisna sebebi ("13/b, Deniz ve hava taşıma ... hizmetler")
_BIRIM_KODU = "(?:" + "|".join(sorted(BIRIMLER, key=len, reverse=True)) + r"|[Mm][23²³])(?![A-Za-z0-9])"
# Satırdaki ek vergi: "Tevkifat (624 - YÜK TAŞIMA)(%20,00)=720 TL", "KDV TEVKİFAT (%20,00)=1.000,00 TL",
# "KONAKLAMA VERGİSİ (%1,00)=225,00 TL". Birden çok olabilir.
_TEVKIFAT_SATIRI = (r"(?:(?:[A-Za-zÇĞİÖŞÜçğıöşü][^%=]{0,80}?)?\(?\s*%\s*[\d.,]+\s*\)?\s*=\s*[\d.,]+\s*(?:TL|TRY|₺)\s+)*")
_MUAFIYET = r"(?:(?:\d{3}|\d{1,2}/[a-z])(?=,?\s+[A-Za-zÇĞİÖŞÜçğıöşü])[^%]{0,250}?\s)?"
_TABLO_BASLIGI = re.compile(r"\b(s[ıi]ra|s\.\s?n)\b|(?<!\S)#(?=\s)|(?<!\S)No\s+(?=(?:hizmet|mal|[üu]r[üu]n|a[çc][ıi]klama)\b)",
                            re.IGNORECASE)


# KDV tevkifatı (KDVGUT I/C-2.1.3): kod -> (oran onda kaç, açıklama). 650 ve 8xx oranı serbest/tam tevkifattır.
TEVKIFAT_KODLARI = {
    "601": (4, "Yapım işleri ve mühendislik-mimarlık, etüt-proje"),
    "602": (9, "Etüt, plan-proje, danışmanlık, denetim"),
    "603": (7, "Makine, teçhizat, demirbaş ve taşıt tadil, bakım, onarım"),
    "604": (5, "Yemek servis hizmeti"),
    "605": (5, "Organizasyon hizmeti"),
    "606": (9, "İşgücü temin hizmetleri"),
    "607": (9, "Özel güvenlik hizmeti"),
    "608": (9, "Yapı denetim hizmetleri"),
    "609": (7, "Fason tekstil, konfeksiyon, çanta ve ayakkabı"),
    "610": (9, "Turistik mağazalara müşteri bulma/götürme"),
    "611": (9, "Spor kulüplerinin yayın, reklam ve isim hakkı"),
    "612": (9, "Temizlik hizmeti"),
    "613": (9, "Çevre ve bahçe bakım hizmetleri"),
    "614": (5, "Servis taşımacılığı"),
    "615": (7, "Baskı ve basım hizmetleri"),
    "616": (5, "Diğer hizmetler"),
    "617": (7, "Hurda metalden elde edilen külçe teslimi"),
    "618": (7, "Bakır, çinko, demir-çelik, alüminyum, kurşun külçe teslimi"),
    "619": (7, "Bakır, çinko ve alüminyum ürünleri teslimi"),
    "620": (7, "İstisnadan vazgeçenlerin hurda ve atık teslimi"),
    "621": (9, "Hurda ve atıktan elde edilen hammadde teslimi"),
    "622": (9, "Pamuk, tiftik, yün, yapağı, ham post ve deri teslimi"),
    "623": (5, "Ağaç ve orman ürünleri teslimi"),
    "624": (2, "Yük taşımacılığı hizmeti"),
    "625": (3, "Ticari reklam hizmetleri"),
    "626": (2, "Diğer teslimler"),
    "627": (5, "Demir-çelik ürünleri teslimi"),
    "650": (None, "Diğerleri (oran faturada yazar)"),
    "801": (10, "Tam tevkifat"),
}
# Kısmi tevkifat alt sınırı (KDV dahil): 1.3.2024'ten beri o yılın fatura düzenleme sınırı (KDVGUT, 50 Seri No'lu Tebliğ).
# Bu tutarı aşmayan teslim/hizmette kısmi tevkifat yapılmaz. Yeni yılın tutarı açıklanınca buraya eklenir.
TEVKIFAT_SINIRI = {2024: 6900, 2025: 9900, 2026: 12000}


def tevkifat_siniri(yil: int) -> float | None:
    if yil < 2024:
        return None
    return TEVKIFAT_SINIRI.get(yil) or TEVKIFAT_SINIRI[max(TEVKIFAT_SINIRI)]


_TEVK = r"tevk[iİıI]fat"


def tevkifat_tutari(tutar: float, kdv_orani: float, pay: float, kdv_tutari: float | None = None) -> float:
    """Kalem tevkifatı = KDV x oran: 23.160,72 TL %20 KDV, 2/10 -> 926,43. Faturada yazan KDV tutarı varsa o esas."""
    kdv = float(kdv_tutari) if kdv_tutari is not None else float(tutar or 0) * float(kdv_orani or 0) / 100
    return round(kdv * float(pay or 0) / 10, 2)


def stopaj_tutari(tutar: float, oran: float | None) -> float:
    """Kalemin gelir/kurumlar vergisi stopajı: KDV hariç tutar x oran / 100."""
    return round(float(tutar or 0) * float(oran or 0) / 100, 2)


def stopaji_dagit(kalemler: list[dict], stopaj: float, oran: float | None) -> float:
    """Faturadaki stopaj kalemlere oranıyla yazılır. Satırda oran okunduysa o geçerlidir; yalnız faturanın altında
    yazıyorsa ve oran x kalemler tutuyorsa her kaleme o oran yazılır. Fatura stopajı (satırlardan ya da yazılandan) döner."""
    satir = round(sum(stopaj_tutari(k["tutar"], k.get("stopaj_orani")) for k in kalemler), 2)
    if satir:
        return stopaj or satir
    if stopaj and oran:
        if abs(round(sum(stopaj_tutari(k["tutar"], oran) for k in kalemler), 2) - stopaj) <= 0.05 + 0.01 * len(kalemler):
            for k in kalemler:
                k["stopaj_orani"] = oran
    return stopaj


def kalem_kdv(k: dict) -> float:
    return k["kdv_tutari"] if k.get("kdv_tutari") is not None else (k.get("tutar") or 0) * (k.get("kdv_orani") or 0) / 100


def kalemler_tutarli(o: dict) -> bool:
    """Kalemler okunmuş ve toplamları faturadaki mal/hizmet toplamını ve KDV'yi tutuyor mu?"""
    kalemler = o.get("kalemler") or []
    if not kalemler or o.get("kalemler_okunamadi"):
        return False
    pay = max(0.05, 0.01 * len(kalemler))
    ara, kdv = o.get("ara_toplam"), o.get("kdv_toplam")
    if ara is not None and abs(sum(k.get("tutar") or 0 for k in kalemler) - ara) > pay:
        return False
    return kdv is None or abs(sum(kalem_kdv(k) for k in kalemler) - kdv) <= pay


def _vergisiz(k: dict) -> dict:
    """Vergiler dahil yazılmış satırın net hali: satırda KDV/diğer vergi tutarı yazıyorsa düşülür, yazmıyorsa
    KDV oranıyla ayrılır. Fiyat ve iskonto da vergiler dahilse aynı oranda netleşir."""
    dahil, diger = k["tutar"], k.get("_diger") or 0
    if k.get("_kdvt") is not None:
        kdvt = k["_kdvt"]
        net = round(dahil - kdvt - diger, 2)
    else:
        net = round((dahil - diger) / (1 + (k.get("kdv_orani") or 0) / 100), 2)
        kdvt = round(dahil - diger - net, 2)
    y = {**k, "tutar": net}
    if dahil and abs(k["miktar"] * k["birim_fiyat"] - (k.get("iskonto_tutari") or 0) - dahil) <= max(0.05, 0.001 * dahil):
        oran = net / dahil
        y["birim_fiyat"] = round(k["birim_fiyat"] * oran, 6)
        if k.get("iskonto_tutari"):
            y["iskonto_tutari"] = round(k["iskonto_tutari"] * oran, 2)
    y.pop("kdv_tutari", None)
    if abs(net * (k.get("kdv_orani") or 0) / 100 - kdvt) > 0.01:
        y["kdv_tutari"] = kdvt
    return y


def _kalem_kendi_tutarli(k: dict) -> bool:
    return abs(k["miktar"] * k["birim_fiyat"] - (k.get("iskonto_tutari") or 0) - k["tutar"]) <= max(0.05, 0.001 * abs(k["tutar"]))


def kalemleri_uzlastir(kalemler: list[dict], hedefler: list[float | None], genel: float | None = None,
                       kdv: float | None = None) -> list[dict] | None:
    """Okunan satırları faturanın mal/hizmet toplamıyla karşılaştırır (hiçbir satır uydurulmaz):
    1) toplam tutuyorsa satırlar olduğu gibi;
    2) satır tutarları vergiler dahil yazılmışsa (market, otel, restoran, bilet) net tutarlara çevrilmiş halleri;
    3) tek satırlı faturada satır tutarı ödenecek tutarsa satır faturanın mal/hizmet toplamı ve KDV'si olur;
    4) satırlar kendi içinde tutarlı ama toplam faturada yazandan kuruşlarla farklıysa (fiyat yuvarlanmış) satırlar
       alınır; fark faturada kırmızı uyarı olarak görünür.
    Hiçbiri olmazsa None (eski okuma yöntemleri denenir)."""
    hed = [h for h in hedefler if h is not None]
    pay = max(0.05, 0.01 * len(kalemler or []))
    temiz = lambda ks: [{a: v for a, v in k.items() if a != "_kdvt"} for k in ks]  # noqa: E731
    if kalemler and not hed and genel is not None and all(k.get("_kdvt") is not None for k in kalemler) \
            and abs(sum(k["tutar"] + k["_kdvt"] for k in kalemler) - genel) <= pay:
        # Mal/hizmet toplamı okunamadı (etiketi kesik basılmış) ama satırların tutarı ve KDV'si ödenecek tutarı veriyor
        hed = [round(sum(k["tutar"] for k in kalemler), 2)]
    if not kalemler or not hed:
        return None

    def satir_kdvli(ks):
        # Satırlarda yazan KDV'lerin toplamı faturanın KDV'sini tutuyor, oranla hesaplanan kuruşla tutmuyorsa
        # (satıcı satır KDV'sini kendisi yuvarlamış) her satırın KDV'si faturada yazan olur.
        if kdv is not None and all(k.get("_kdvt") is not None for k in ks) \
                and abs(sum(k["_kdvt"] for k in ks) - kdv) <= 0.01 \
                and abs(sum(k["tutar"] * (k.get("kdv_orani") or 0) / 100 for k in ks) - kdv) > 0.01:
            ks = [{**k, "kdv_tutari": k["_kdvt"]} if abs(k["tutar"] * (k.get("kdv_orani") or 0) / 100 - k["_kdvt"]) > 0.004
                  else k for k in ks]
        return temiz(ks)
    toplam = round(sum(k["tutar"] for k in kalemler), 2)
    if any(abs(toplam - h) <= pay for h in hed):
        return satir_kdvli(kalemler)
    net = [_vergisiz(k) for k in kalemler]
    # Eksi satır (market faturasında poşet/kampanya indirimi) toplamın küçük bir parçasıysa olağandır
    eksi = -sum(k["tutar"] for k in net if k["tutar"] < 0)
    if eksi <= 0.1 * sum(k["tutar"] for k in net if k["tutar"] > 0) and any(abs(sum(k["tutar"] for k in net) - h) <= pay for h in hed):
        return temiz(net)
    if len(kalemler) == 1 and genel is not None and kdv is not None and abs(toplam - genel) <= 0.05:
        h = next((h for h in hed if abs(h + kdv + (kalemler[0].get("_diger") or 0) - genel) <= 0.05), None)
        if h is not None:
            k = {**_vergisiz(kalemler[0]), "tutar": h}
            if k["birim_fiyat"] * k["miktar"] != h and _kalem_kendi_tutarli(kalemler[0]):
                k["birim_fiyat"] = round(h / k["miktar"], 6)
                k.pop("iskonto_tutari", None), k.pop("iskonto_orani", None)
            k.pop("kdv_tutari", None)
            if abs(h * (k.get("kdv_orani") or 0) / 100 - kdv) > 0.01:
                k["kdv_tutari"] = kdv
            return temiz([k])
    if all(_kalem_kendi_tutarli(k) for k in kalemler) and any(abs(toplam - h) <= max(1.0, 0.002 * h) for h in hed):
        return satir_kdvli(kalemler)
    return None


def _tevkifat_bilgisi(metin: str) -> tuple[float | None, str | None]:
    """Satırdaki tevkifat: 'KDV TEVKİFAT (%20,00)=926,43 TL', 'Tevkifat Sebebi: 624-2/10', '(TEVKİFATLI 2/10)',
    'Tevkifat (624 - YÜK TAŞIMA)(%20,00)=720 TL'. (oran onda kaç, kod)."""
    pay = kod = None
    m = re.search(_TEVK + r"[^%=]{0,80}?\(?\s*%\s*([\d.,]+)\s*\)?\s*=", metin, re.IGNORECASE)
    if m and (y := para_coz(m.group(1))):
        pay = round(y / 10, 2)
    m = re.search(r"(?<!\d)([68]\d\d)\s*-\s*(\d{1,2})\s*/\s*10\b", metin) or re.search(
        _TEVK + r"[^0-9%=]{0,40}?(?<!\d)([68]\d\d)(?!\d)(?:\s*-\s*(\d{1,2})\s*/\s*10\b)?", metin, re.IGNORECASE)
    if m:
        kod = m.group(1) if m.group(1) in TEVKIFAT_KODLARI else None
        if pay is None and m.group(2):
            pay = float(m.group(2))
    if pay is None and (m := re.search(_TEVK + r"[^/]{0,20}?(\d{1,2})\s*/\s*10\b", metin, re.IGNORECASE)):
        pay = float(m.group(1))
    if pay is None and kod:
        pay = TEVKIFAT_KODLARI[kod][0]
    return (pay if pay and 0 < pay <= 10 else None), kod


def _kodlari_tamamla(kalemler: list[dict], metin: str) -> None:
    """Kodu satırda yazmayan tevkifatlı kalemlere faturanın altındaki 'Tevkifat Sebebi: 624-...' kodunu yazar
    (birden çok kod varsa oranı tutan)."""
    kodlar = list(dict.fromkeys(k for k in re.findall(_TEVK + r"\s*sebebi\s*:?\s*([68]\d\d)(?!\d)", metin, re.IGNORECASE)
                                if k in TEVKIFAT_KODLARI))
    for k in kalemler:
        if k.get("tevkifat_orani") and not k.get("tevkifat_kodu"):
            uyan = [x for x in kodlar if TEVKIFAT_KODLARI[x][0] in (None, k["tevkifat_orani"])]
            if len(uyan) == 1 or len(kodlar) == 1:
                k["tevkifat_kodu"] = (uyan or kodlar)[0]


def _tevkifati_dagit(kalemler: list[dict], tevk: float, metin: str = "", pay: float | None = None,
                     kod: str | None = None) -> bool:
    """Tevkifat yalnız fatura toplamında yazılıysa: tutarı veren satır grubuna (hepsi, bir KDV oranı ya da bir alt küme)
    oranı ve kodu yazar. Bulamazsa False."""
    if not tevk or not kalemler or any(k.get("tevkifat_orani") for k in kalemler):
        return False
    if pay is None:
        pay, kod = _tevkifat_bilgisi(metin)
    paylar = [pay] if pay else [2, 3, 4, 5, 7, 9, 10]
    adaylar = [k for k in kalemler if k.get("kdv_orani")]
    gruplar = [adaylar] + [[k for k in adaylar if k["kdv_orani"] == o] for o in sorted({k["kdv_orani"] for k in adaylar})]
    from itertools import combinations
    if len(adaylar) <= 12:
        gruplar += [list(c) for n in range(len(adaylar) - 1, 0, -1) for c in combinations(adaylar, n)]
    elif len(adaylar) <= 80:  # uzun faturada tevkifat çoğu zaman bir iki satırındır (hizmet satırı)
        gruplar += [list(c) for n in (1, 2, 3) for c in combinations(adaylar, n)]
    for p in paylar:
        for g in gruplar:
            if g and abs(sum(tevkifat_tutari(k["tutar"], k["kdv_orani"], p, k.get("kdv_tutari")) for k in g) - tevk) <= 0.02 * len(g):
                for k in g:
                    k["tevkifat_orani"], k["tevkifat_kodu"] = p, kod
                return True
    return False


def _oranlari_dagit(kalemler: list[dict], kdv_oranlari: dict[float, float]) -> bool:
    """Satırlarda KDV oranı yazmıyor (hepsine aynı oran verildi) ama faturanın altında oran başına KDV yazıyor
    ("KDV (%20): 200,00", "KDV (%10): 1,18"): azınlık oranının KDV'sini veren satırlara o oran yazılır."""
    oranlar = {k.get("kdv_orani") for k in kalemler}
    if len(oranlar) != 1 or len(kalemler) < 2:
        return False
    if abs(sum(k["tutar"] * (k.get("kdv_orani") or 0) / 100 for k in kalemler) - sum(kdv_oranlari.values())) <= 0.05:
        return False
    from itertools import combinations
    ana = max(kdv_oranlari, key=lambda r: kdv_oranlari[r])
    bos = list(kalemler)
    for r, t in kdv_oranlari.items():
        if r == ana:
            continue
        hedef = t * 100 / r if r else None
        if not hedef:
            continue
        n_max = len(bos) - 1 if len(bos) <= 12 else 3
        g = next((list(c) for n in range(1, n_max + 1) for c in combinations(bos, n)
                  if abs(sum(k["tutar"] for k in c) * r / 100 - t) <= 0.01 * n + 0.01), None)
        if g is None:
            return False
        for k in g:
            k["kdv_orani"] = r
            bos.remove(k)
    for k in bos:
        k["kdv_orani"] = ana
    return True


def _miktar(metin: str) -> float | None:
    """Miktarda virgül hep ondalıktır: '301,778 KWH' 301.778'dir (para_coz bunu binlik sanar)."""
    if re.fullmatch(r"\d+,\d+", metin):
        return float(metin.replace(",", "."))
    return para_coz(metin)


def _tablo_metni(bolum: str) -> str:
    """PDF'ten bozuk çıkan yazımları düzeltir."""
    bolum = re.sub(r"(?:\s\|)+\s", " ", bolum)  # boş sütun ayraçları: "0,00 TL | | %20,00"
    bolum = re.sub(r"%(\d+),(\d) (\d)\b", r"%\1,\2\3", bolum)  # "%20,0 0" -> "%20,00"
    bolum = re.sub(r"(\d)T L\b", r"\1TL", bolum)  # "152.915,00T L"
    return re.sub(r"(?<=[A-Za-zÇĞİÖŞÜçğıöşü])\((?:NIU|C62|[A-Z]{3})\)", "", bolum)  # "Adet(NIU)"


# Adın devamı aranırken: tutar/vergi sözcüğü görülürse bu bir toplam ya da vergi satırıdır, devam sayılmaz;
# belge etiketi (ETTN, Senaryo...) görülürse orada durulur, öncesi adın devamıdır.
_AD_VERGI = re.compile(r"^(?:tevkifat|kdv|iskonto|indirim|toplam|ara|genel|odenecek|vergi|vergiler|vergisi|fonu|payi|tl|try)$")
_AD_ETIKET = re.compile(r"^(?:not|notlar|ettn|irsaliye|siparis|yalniz|mal|hizmet|sayfa|senaryo|fatura|tipi|ozellestirme|"
                        r"duzenleme|tarih|saat)$")


def _ad_devami(aciklama: str, kuyruk: str) -> str:
    """Mal/hizmet adı tutarlardan sonraki satıra taşmış olabilir ("1 ELEKTRİK 1Adet 625TL % 20 125,00TL 625,00TL⏎MALZEMESİ").
    Satırın ardından, sonraki kalemden önce yalnız sözcük varsa ada eklenir; etiket, tutar ya da iki nokta görülünce durulur.
    Ad büyük harfle yazılmışsa devamı da büyük harfle olmalı (alttaki "Senaryo ..." gibi başlık bilgisi ada katılmaz)."""
    devam = []
    buyuk = aciklama == aciklama.upper()
    for k in kuyruk.split():
        sk = sade(k).strip(".,()")
        if _AD_VERGI.match(sk):
            return aciklama  # "Roommate⏎Konaklama Vergisi ...": vergi satırı
        if (buyuk and k != k.upper()) or ":" in k or "%" in k or re.fullmatch(r"[-\d.,()=/]+", k) or _AD_ETIKET.match(sk) \
                or not re.search(r"[A-Za-zÇĞİÖŞÜçğıöşü]{2}", k):
            break
        devam.append(k)
    return f"{aciklama} {' '.join(devam)}" if devam and len(devam) <= 8 else aciklama


def _kalemleri_coz(tek: str, hedefler: list[float], oran: float | None = None) -> list[dict] | None:
    """Kalem tablosunu satır satır çözer; toplamı fatura toplamını tutmazsa None döner.

    Başlık gibi görünen birden çok yer varsa ("# Yalnız ..." gibi) sırayla denenir.
    """
    for bas in _TABLO_BASLIGI.finditer(tek):
        kalemler = _tablodan_kalemler(tek[bas.end():], hedefler, oran)
        if kalemler:
            return kalemler
    return None


def _satir_tutarli(m: re.Match) -> bool:
    g = m.groupdict()
    tut, fiy = para_coz(g.get("tut") or ""), _miktar(g["fiy"]) if g.get("fiy") else None
    mik = (_miktar(g["mik"]) if g.get("mik") else 1) or 1
    kdv, kdvt = para_coz(g.get("kdv") or ""), para_coz(g.get("kdvt") or "")
    if tut is None:
        return fiy is not None
    if g.get("brt"):  # net tutar + KDV tutarı = KDV'li net tutar
        brt = para_coz(g["brt"])
        return brt is not None and kdvt is not None and abs(tut + kdvt - brt) <= 0.05
    if fiy is not None and abs(mik * fiy - tut) <= max(0.05, 0.002 * tut):
        return True
    isk = para_coz(g.get("isk") or "") or 0
    return bool(kdv and kdvt is not None and abs((tut - isk) * kdv / 100 - kdvt) <= 0.05 + 0.001 * kdvt
                or kdv and kdvt is not None and abs(tut * kdv / 100 - kdvt) <= 0.05 + 0.001 * kdvt)


def _tablodan_kalemler(bolum: str, hedefler: list[float], oran: float | None) -> list[dict] | None:
    son = re.search(r"(mal|malzeme)\s*/?\s*hizmet\s+toplam\s+tutar", bolum, re.IGNORECASE)
    bolum, konum, kalemler = _tablo_metni(bolum[:son.start()] if son else bolum), 0, []
    for n in range(1, 500):
        # Birim "TL" olamaz ("142,5 TL" fiyattır); miktar "1.000,21" gibi binlikli olabilir.
        bas_n = rf"(?<![\d.,]){n}\s+(?P<ac>.{{1,400}}?)\s+(?:(?P<mik>\d{{1,3}}(?:\.\d{{3}})+(?:,\d+)?|\d+(?:[.,]\d+)?)\s*(?P<bir>(?!(?:TL|TRY)\b){_BIRIM_KODU}|(?!(?:TL|TRY)\b)[A-Za-zÇĞİÖŞÜçğıöşü]+\.?)?\s+)?(?P<fiy>\d[\d.,]*)\s*{_PB}\s+"
        # İskonto oranı ve tutarı, iskonto nedeni ya da ürün açıklaması ("İskonto -", "AS26-00553-14-K") olabilir.
        ara_sutunlar = rf"(?:%\s*(?P<iskor>[\d.,]+)\s+)?(?:(?:\d{{3}}\s+)?(?P<isk>[\d.,]+)\s*{_PB}\s+)?(?:[^\s%]+\s+){{0,3}}?"
        desenler = [
            # ... fiyat, [iskonto], [tevkifat kodu ve tutarı], %KDV, KDV tutarı, [diğer], tutar
            bas_n + ara_sutunlar + rf"%\s*(?P<kdv>[\d.,]+)\s+(?P<kdvt>[\d.,]+)\s*{_PB}\s+"
                    rf"(?:[\d.,]+\s*{_PB}\s+)?{_TEVKIFAT_SATIRI}{_MUAFIYET}(?P<tut>[\d.,]+)\s*{_PB}",
            # ... fiyat, tutar, %KDV, KDV tutarı (tutar KDV'den önce gelen düzen)
            bas_n + rf"(?:(?P<isk>[\d.,]+)\s*{_PB}\s+)?(?P<tut>[\d.,]+)\s*{_PB}\s+%\s*(?P<kdv>[\d.,]+)\s+(?P<kdvt>[\d.,]+)\s*{_PB}",
            # ... fiyat, %KDV, "KDV Tutarı" etiketi, KDV tutarı; tutar sütunu yok (muayene istasyonu)
            bas_n + ara_sutunlar + rf"%\s*(?P<kdv>[\d.,]+)\s+KDV\s+Tutar[ıi]\s*:?\s*(?P<kdvt>[\d.,]+)\s*{_PB}",
            # ... fiyat, %KDV, tutar; KDV tutarı sütunu yok
            bas_n + ara_sutunlar + rf"%\s*(?P<kdv>[\d.,]+)\s+(?P<tut>[\d.,]+)\s*{_PB}(?!\s*(?:%|[\d.,]+\s*{_PB}))",
            # miktar, birim, tutar; fiyat ve KDV sütunu yok (akaryakıt): fiyat = tutar / miktar
            rf"(?<![\d.,]){n}\s+(?P<ac>.{{1,400}}?)\s+(?P<mik>\d[\d.,]*)\s*(?P<bir>{_BIRIM_KODU}|[A-Za-zÇĞİÖŞÜçğıöşü]+\.?)\s+"
            rf"(?P<tut>[\d.,]+)\s*{_PB}(?!\s*(?:%|[\d.,]+\s*{_PB}))",
            # KDV oranı sütunu yok: fiyat, İskonto1 %, [İskonto2 %], net birim fiyat, net tutar, KDV tutarı, KDV'li net tutar
            bas_n + rf"(?:%\s*(?P<iskor>[\d.,]+)\s+)?(?:%\s*[\d.,]+\s+)?(?:[\d.,]+\s*(?:{_PB}\s+|\s))?"
                    rf"(?P<tut>[\d.,]+)\s*{_PB}\s+(?P<kdvt>[\d.,]+)\s*{_PB}\s+(?P<brt>[\d.,]+)\s*{_PB}",
        ]
        adaylar = [(m.start(), not _satir_tutarli(m), i, m) for i, d in enumerate(desenler)
                   for m in [re.compile(d, re.IGNORECASE).search(bolum, konum)] if m and (i < 5 or _satir_tutarli(m))]
        # Aynı yerde birden çok düzen uyarsa rakamları kendi içinde tutan (miktar x fiyat = tutar), o da yoksa ilk düzen
        m = min(adaylar, key=lambda x: x[:3])[3] if adaylar else None
        if not m:
            break
        g = m.groupdict()
        miktar = (_miktar(g["mik"]) if g.get("mik") else 1) or 1
        kdvt = para_coz(g["kdvt"]) if g.get("kdvt") else None
        fiyat = _miktar(g["fiy"]) if g.get("fiy") else None  # "45,455 TL" 45,455'tir
        tutar = para_coz(g["tut"]) if g.get("tut") else (round(miktar * fiyat, 2) if fiyat is not None else None)
        if tutar is None:
            return None
        if fiyat is None:
            fiyat = round(tutar / miktar, 6)
        kdv = para_coz(g["kdv"]) if g.get("kdv") else oran
        if not g.get("kdv") and kdvt is not None and tutar:  # oran sütunu yoksa KDV tutarından
            kdv = min((0, 1, 8, 10, 18, 20), key=lambda o: abs(tutar * o / 100 - kdvt))
            if abs(tutar * kdv / 100 - kdvt) > 0.05 + 0.001 * kdvt:
                kdv = round(kdvt / tutar * 100, 2)
        aciklama, bilgi_satiri = _kodu_sona_al(g["ac"].strip()), False
        isk = para_coz(g["isk"]) if g.get("isk") else None
        if isk and kdvt is not None and kdv and 0 < isk < tutar and abs((tutar - isk) * kdv / 100 - kdvt) <= 0.05:
            # Satır iskontosu: tutar sütunu brüt yazılmış, KDV net üzerinden. Kalem net tutarla işlenir.
            oran_yazi = f"%{para_coz(g['iskor']):g}".replace(".", ",") if g.get("iskor") else f"{isk:.2f} TL"
            aciklama = f"{aciklama} (İsk. {oran_yazi}, liste fiyatı {_sayi_yazi(fiyat)})"
            tutar = round(tutar - isk, 2)
            fiyat = round(tutar / miktar, 6)
        iskor = para_coz(g["iskor"]) if g.get("iskor") else None
        if iskor and 0 < iskor < 100 and not isk and abs(miktar * fiyat * (1 - iskor / 100) - tutar) <= max(0.05, 0.002 * tutar):
            # Tutar net, fiyat liste fiyatı: "16 Adet 1.600 TL %60,00 640,00 10.240,00 TL"
            aciklama = f"{aciklama} (İsk. %{_sayi_yazi(iskor)}, liste fiyatı {_sayi_yazi(fiyat)})"
            fiyat = round(tutar / miktar, 6)
        if kdvt is not None:
            net = round(miktar * fiyat, 2)
            if abs(tutar - net - kdvt) <= 0.02 and abs(tutar - net) > 0.02:  # tutar KDV dahil yazılmış (Trendyol)
                tutar = net
            if kdvt == 0 and tutar > 0 and kdv:  # oran yazılı ama KDV alınmamış (aracılık edilen tutar)
                kdv, bilgi_satiri = 0, True
        # Faturada yazan KDV tutarı oran x tutarı tutmuyorsa (Öztaş: %20, 73.108,33 TL, KDV 833,00) yazan tutar esastır.
        kdv_yazan = kdvt if (kdvt is not None and kdv and not bilgi_satiri
                             and abs(tutar * kdv / 100 - kdvt) > 0.05 + 0.001 * kdvt) else None
        # Satır tevkifatı satırın içinde ya da altında ("Tevkifat Sebebi: 624-2/10"), sonraki satırdan önce yazılır.
        sonraki = re.compile(rf"(?<![\d.,]){n + 1}\s+\S").search(bolum, m.end(), m.end() + 400)
        kuyruk = bolum[m.end():sonraki.start() if sonraki else m.end() + 200]
        pay, tkod = _tevkifat_bilgisi(m.group(0) + " " + kuyruk)
        aciklama = _ad_devami(aciklama, kuyruk)
        kalemler.append({"aciklama": aciklama, "miktar": miktar,
                         "birim": BIRIMLER.get(g.get("bir") or "", g.get("bir")),
                         "birim_fiyat": fiyat, "kdv_orani": kdv or 0, "tutar": tutar, "_bilgi": bilgi_satiri,
                         **({"tevkifat_orani": pay, "tevkifat_kodu": tkod} if pay and kdv else {}),
                         **({"kdv_tutari": kdv_yazan} if kdv_yazan is not None else {})})
        konum = m.end()

    def tutuyor(liste):
        toplam = round(sum(k["tutar"] for k in liste), 2)
        return bool(liste) and any(h is not None and abs(toplam - h) <= max(0.05, 0.01 * len(liste)) for h in hedefler)
    # KDV'siz bilgi satırı (muayene istasyonunda "Kartlı ödemeye konu toplam tutar") faturanın toplamına girmiyorsa çıkar.
    for liste in (kalemler, [k for k in kalemler if not k["_bilgi"]]):
        if tutuyor(liste):
            return [{a: v for a, v in k.items() if a != "_bilgi"} for k in liste]
    return None


_SATIR_KALEM = re.compile(
    rf"^(?P<ac>\S.*?)\s+(?:(?P<mik>\d+(?:[.,]\d+)?)\s*)?(?P<bir>(?!(?:TL|TRY)\b)[A-Za-zÇĞİÖŞÜçğıöşü]+\.?)\s+(?P<fiy>\d[\d.,]*)\s*{_PB}\s+"
    rf"(?:%\s*[\d.,]+\s+)*?(?:%\s*(?P<kdv>[\d.,]+)\s+)?(?P<tut>\d[\d.,]*)\s*{_PB}$", re.IGNORECASE)


_MATRAHA_GIREN_VERGI = re.compile(r"^(?P<ad>BTV|TRT PAYI|ENERJ[İI] FONU|BELED[İI]YE T[ÜU]KET[İI]M VERG[İI]S[İI])\s*:?\s*"
                                  r"(?P<tut>\d[\d.,]*)\s*(?:TL|TRY)$", re.IGNORECASE)


def _kalemleri_satirdan(metin: str, hedefler: list[float], oran: float) -> list[dict] | None:
    """Sıra numarası olmayan tablolar (Turkcell, ALMERAM): başlıktaki 'Miktar' satırından sonra her satır bir kalemdir.

    'Superbox Modem Bedeli 1 Adet 111,67 TL 111,67 TL' ya da '0101 VADE FARKI 1,00 Adet 1.016,73 TL %20,00 1.016,73 TL'.
    """
    satirlar = [" ".join(x.split()) for x in metin.splitlines()]
    bas = next((i for i, x in enumerate(satirlar) if re.search(r"\bmiktar\b", sade(x)) and re.search(r"tutar|fiyat", sade(x))),
               None)
    if bas is None:
        return None
    kalemler = []
    for x in satirlar[bas + 1:]:
        if re.match(r"^(ara |genel )?toplam\b", sade(x)):
            break
        m = _SATIR_KALEM.match(x)
        if not m:
            continue
        tutar, fiyat = para_coz(m.group("tut")), _miktar(m.group("fiy"))
        if tutar is None or fiyat is None:
            return None
        kalemler.append({"aciklama": _kodu_sona_al(m.group("ac").strip()), "miktar": _miktar(m.group("mik") or "1") or 1,  # miktarsız sabit bedel: "ENERJI AÇMA/KESME BEDELI KWH 140 TL"
                         "birim": m.group("bir"), "birim_fiyat": fiyat,
                         "kdv_orani": para_coz(m.group("kdv")) if m.group("kdv") else oran, "tutar": tutar})
    if not kalemler:
        return None
    # Elektrik faturası: KDV matrahına giren BTV, TRT payı, enerji fonu ayrı satırlarda yazılır.
    for x in satirlar[bas + 1:]:
        m = _MATRAHA_GIREN_VERGI.match(x)
        if m and (tutar := _matrah(m.group("tut"))):
            kalemler.append({"aciklama": m.group("ad").strip(), "miktar": 1, "birim": None, "birim_fiyat": tutar,
                             "kdv_orani": oran, "tutar": tutar})
    toplam = round(sum(k["tutar"] for k in kalemler), 2)
    if not any(h is not None and abs(toplam - h) <= max(0.05, 0.01 * len(kalemler)) for h in hedefler):
        return None
    return kalemler


def _sayi_yazi(x: float) -> str:
    return f"{x:,.6f}".rstrip("0").rstrip(".").replace(",", "X").replace(".", ",").replace("X", ".")


def _kalemleri_denetle(kalemler: list[dict]) -> list[dict]:
    """Her kalemde miktar x birim fiyat = tutar olmalı. Tutmuyorsa PDF'in olası yanlış okunuşları denenir:
    açıklamaya karışmış miktar ("MOTORİN 1.000,21 lt"), kayan ondalık; sonunda tutar küçükse satır iskontosu sayılır
    ve fiyat net fiyata çevrilir (liste fiyatı açıklamada kalır). Tutar faturadaki gibi kalır."""
    for k in kalemler:
        t, m, p = k.get("tutar") or 0, k.get("miktar") or 1, k.get("birim_fiyat") or 0
        tol = max(0.05, 0.002 * abs(t))
        if t <= 0 or abs(m * p - t) <= tol:
            continue
        son = re.search(r"\s(\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+(?:[.,]\d+)?)\s*([A-Za-zÇĞİÖŞÜçğıöşü]{1,6}\.?)?$", k["aciklama"])
        n = _miktar(son.group(1)) if son else None
        if n and p and abs(n * p - t) <= tol:  # miktar açıklamaya karışmış
            k.update(miktar=n, birim=son.group(2) or k.get("birim"), aciklama=k["aciklama"][:son.start()].strip())
            continue
        if n and not p and abs(n * m - t) <= tol:  # miktar açıklamada, fiyat miktar sanılmış
            k.update(miktar=n, birim_fiyat=m, aciklama=k["aciklama"][:son.start()].strip())
            continue
        kayan = next((p / b for b in (10, 100, 1000, 10000) if p and abs(m * p / b - t) <= tol), None)
        if kayan is not None:  # "45,455" binlik sanılmış
            k["birim_fiyat"] = round(kayan, 6)
            continue
        if p and t < m * p:  # satır iskontosu yazılmamış ama tutar net
            yuzde = (1 - t / (m * p)) * 100
            k["aciklama"] = f"{k['aciklama']} (İsk. %{_sayi_yazi(round(yuzde, 2))}, liste fiyatı {_sayi_yazi(p)})"
        k["birim_fiyat"] = round(t / m, 6)
    return kalemler


def _tarihli_satirlar(tek: str, hedefler: list[float], oran: float) -> list[dict] | None:
    """Sıra numarasız, her satırı tarihle başlayan taşıma tabloları (kargo/ambar):
    '7.08.2026 275817 GELİŞ ... 1 PALET ... 0,000 TL 6.000,00 TL'. Birim fiyat yazılmamışsa tutar fiyat olur."""
    bas = re.search(r"birim\s+fiyat[ıi]?\s+tutar[ıi]?\b", tek, re.IGNORECASE)
    if not bas:
        return None
    bolum = tek[bas.end():]
    son = re.search(r"(mal|malzeme)\s*/?\s*hizmet\s+toplam|\btoplam\b", bolum, re.IGNORECASE)
    bolum = bolum[:son.start()] if son else bolum
    kalemler = []
    for m in re.finditer(rf"(?<!\S)(?P<tar>\d{{1,2}}\.\d{{2}}\.\d{{4}})\s+(?P<ac>.{{3,300}}?)\s+"
                         rf"(?:(?P<fiy>[\d.,]+)\s*{_PB}\s+)?(?P<tut>[\d.,]+)\s*{_PB}", bolum):
        tutar, fiyat = para_coz(m.group("tut")), para_coz(m.group("fiy") or "")
        if tutar is None:
            return None
        kalemler.append({"aciklama": f"{m.group('tar')} {m.group('ac').strip()}", "miktar": 1, "birim": None,
                         "birim_fiyat": fiyat if fiyat else tutar, "kdv_orani": oran, "tutar": tutar})
    toplam = round(sum(k["tutar"] for k in kalemler), 2)
    if not kalemler or not any(h is not None and abs(toplam - h) <= 0.05 for h in hedefler):
        return None
    return kalemler


def _fis_kalemleri(metin: str, hedefler: list, genel: float | None, kdv: float | None) -> list[dict] | None:
    """Kasa fişi düzeninde e-Fatura (mağaza): ürün adı bir-iki satır, altında "% 20 11.999,00" (KDV oranı ve KDV
    dahil tutar). Adın başındaki ürün kodu ("40594645.1248994") ada katılmaz, "Seri No:" satırı atlanır."""
    satirlar = [x.strip() for x in metin.splitlines()]
    kalemler, bas = [], 0
    for i, x in enumerate(satirlar):
        if kalemler and re.match(r"(kdv )?toplam|ara toplam|kdv orani", sade(x)):
            break  # fişin alt bölümü: "KDV Toplam", "KDV Oranı KDV Tutarı / % 20 2.000,00"
        m = re.fullmatch(r"%\s*(\d{1,2})\s+(-?\d[\d.,]*)", x)
        if not m or para_coz(m.group(2)) is None:
            continue
        adlar = []
        for y in reversed(satirlar[bas:i]):
            if not y or len(adlar) == 3:
                break
            if re.match(r"seri\s*no\b", sade(y)):
                continue
            if ":" in y or re.fullmatch(r"[\d\s.,%-]+", y):
                break
            adlar.insert(0, y)
        bas = i + 1
        if not adlar:
            return None
        adlar[0] = re.sub(r"^\d[\d.]{5,}\s+", "", adlar[0])  # ürün kodu
        tutar = para_coz(m.group(2))
        kalemler.append({"aciklama": " ".join(adlar), "miktar": 1, "birim": None, "birim_fiyat": tutar,
                         "kdv_orani": float(m.group(1)), "tutar": tutar})
    return kalemleri_uzlastir(kalemler, hedefler, genel, kdv) if kalemler else None


def _matrah(deger: str) -> float | None:
    """'1.092.49' (TTNET) ve '4,231.30' (Turkcell) gibi yazımlar: virgül yoksa son noktadan sonraki iki hane kuruştur."""
    if "," not in deger and re.search(r"\.\d{2}$", deger):
        deger = deger[:-3] + "," + deger[-2:]
    return para_coz(deger)


_OZET_VERGI = re.compile(r"^(?P<ad>.*?)\s*%\s*(?P<oran>[\d.,]+)\s*\(\s*matrah\s*:?\s*(?P<mat>[\d.,]+)\s*(?:TL|TRY)?\s*\)"
                         r"\s*(?P<tut>-?[\d.,]*\d)\s*(?:TL|TRY)?$", re.IGNORECASE)


_ABONELIK_SATIR = re.compile(r"^(?P<ad>.*?[A-Za-zÇĞİÖŞÜçğıöşü].*?)\s+(?P<sayilar>-?[\d.]+,\d{2}(?:\s+-?[\d.]+,\d{2}){0,2})$")


def _abonelik_kalemleri(metin: str, kdv_orani: float, kdv_matrah: float, oiv_orani: float, oiv_matrah: float
                        ) -> list[dict] | None:
    """Telekom faturasının vergiler dahil hizmet tablosu ('İnternet 1.900,01 866,01 1.034,00' = tutar, indirim, toplam;
    'Cihaz/ Bağlantı Ücreti 75,00') ve altındaki taksit/devir/yuvarlama satırları faturadaki adlarıyla kalem olur.
    Vergiler dahil tutar KDV'ye (ÖİV'li hizmette KDV+ÖİV'ye) bölünerek net tutar bulunur; hangi hizmetin ÖİV'li olduğu
    faturanın ÖİV matrahını tutan satır grubundan anlaşılır. Tutmazsa None (özet satırlarına dönülür)."""
    satirlar = [" ".join(x.split()) for x in metin.splitlines()]
    bas = next((i for i, x in enumerate(satirlar) if re.match(r"^tutar indirim toplam$", sade(x))), None)
    if bas is None:
        return None
    hizmet, ekler, bitti = [], [], False
    for x in satirlar[bas + 1:]:
        sa = sade(x)
        if re.match(r"^(odenecek|vergiler dahil|devlete odenen)", sa):
            break
        if re.match(r"^(toplam fatura tutari|fatura toplami|fatura tutari)\b", sa):
            bitti = True
            continue
        m = _ABONELIK_SATIR.match(x)
        if not m:
            if hizmet and not bitti:
                return None  # tanınmayan satır: tablo bu düzende değil
            continue
        ad = re.sub(r"\s+toplam[ıi]$", "", m.group("ad").strip(), flags=re.IGNORECASE)
        sayi = [para_coz(y) for y in m.group("sayilar").split()]
        if None in sayi:
            return None
        if bitti or re.match(r"^(yuvarlama|onceki aydan devir|gelecek aya devir|taksit)", sade(ad)):
            ekler.append({"aciklama": ad, "miktar": 1, "birim": None, "birim_fiyat": sayi[-1], "kdv_orani": 0, "tutar": sayi[-1]})
        else:
            hizmet.append((ad, sayi))
    if not hizmet or not bitti:
        return None
    # Ara toplam satırı ("Paket Ücreti 2.030,01 921,01 1.109,00" = İnternet + Telefon) kalem değildir
    kalan, acik = [], []
    for ad, sayi in hizmet:
        if len(acik) >= 2 and abs(sum(s[-1] for _, s in acik) - sayi[-1]) <= 0.01:
            acik = []
            continue
        kalan.append((ad, sayi))
        acik.append((ad, sayi))
    k, o = kdv_orani / 100, oiv_orani / 100
    n = len(kalan)
    if n > 12:
        return None
    for oivli_bolen in ((1 + k + o), (1 + k) * (1 + o)):  # KDV yalnız nete ya da ÖİV dahil tutara
        for maske in range(1 << n):
            oivli = sum(kalan[i][1][-1] for i in range(n) if maske >> i & 1)
            oivsiz = sum(kalan[i][1][-1] for i in range(n) if not maske >> i & 1)
            net_oivli = oivli / oivli_bolen
            if o and abs(net_oivli - oiv_matrah) > 0.02 * n:
                continue
            kdv_tabani = oivsiz / (1 + k) + net_oivli * (1 + o if oivli_bolen != 1 + k + o else 1)
            if abs(kdv_tabani - kdv_matrah) > 0.02 * n:
                continue
            kalemler = []
            for i, (ad, sayi) in enumerate(kalan):
                bolen = oivli_bolen if maske >> i & 1 else 1 + k
                net = round(sayi[-1] / bolen, 2)
                brut = round(sayi[0] / bolen, 2) if len(sayi) == 3 else net
                kalemler.append({"aciklama": ad, "miktar": 1, "birim": None, "birim_fiyat": brut, "kdv_orani": kdv_orani,
                                 "tutar": net, **({"iskonto_tutari": round(brut - net, 2)} if brut != net else {})})
            return kalemler + ekler
        if not o:
            break
    return None


_ADLI_SATIR = re.compile(r"^(?P<ad>[A-Za-zÇĞİÖŞÜçğıöşü][^:]*?[A-Za-zÇĞİÖŞÜçğıöşü)])\s+(?P<sayilar>-?[\d.,]*\d(?:\s+-?[\d.,]*\d){0,2})$")


def _adli_ozet_kalemleri(metin: str, kdv_orani: float, kdv_matrah: float, kdv: float, diger: float,
                          sifir_kalemler: list[dict]) -> tuple[list[dict], float] | None:
    """Abonelik faturasında (Turkcell, TTNET) faturanın kendi yazdığı ücret satırları ("Tarife ve Paket Ücretleri 8,958.38",
    "Gecikme Bedeli 184.08", "İnternet 1.900,00 790,01 1.109,99") adlarıyla kalem olur; "Hizmet bedeli" gibi uydurma ad
    kullanılmaz. Satırlar "Fatura Tutarı" satırının hemen üstündedir. Satırların toplamı fatura tutarını veriyorsa
    vergiler dahildir: KDV'li net tutarlar faturadaki KDV matrahı satırlara oranla dağıtılarak bulunur. Vermiyorsa
    satırlar vergisizdir: toplamı KDV matrahını tutan satırlar KDV'li, kalanlar (aracılık edilen ödemeler) KDV'sizdir.
    Satırların KDV'si faturada yazan KDV'ye göre dağıtılır. Tutmazsa None. Dönüş: (kalemler, fatura tutarı)."""
    satirlar = [" ".join(x.split()) for x in metin.splitlines()]
    t_i = next((i for i, x in enumerate(satirlar)
                if re.match(r"^(toplam )?fatura tutari\s*:?\s*-?[\d.,]*\d\s*(tl|try)?$", sade(x))), None)
    if t_i is None:
        return None
    t_tutar = para_coz(re.findall(r"-?[\d.,]*\d", satirlar[t_i])[-1])
    ust = []
    for x in reversed(satirlar[:t_i]):
        if not re.search(r"\d", x):
            if ust and re.match(r"^tutar indirim toplam$", sade(x)):
                break
            continue  # ara başlık ("KULLANIMLARINIZLA İLGİLİ ÜCRETLER")
        m = _ADLI_SATIR.match(x)
        if not m or re.search(r"\b(odenecek|son odeme|vergi|matrah|kdv|oiv)\b", sade(m.group("ad"))):
            break
        sayilar = [para_coz(y) for y in m.group("sayilar").split()]
        if None in sayilar:
            return None
        ust.append((re.sub(r"\s+toplam[ıi]$", "", m.group("ad").strip(), flags=re.IGNORECASE), sayilar[-1]))
    ust.reverse()
    # Ara toplam satırı ("Paket Ücreti 1.975,00 826,01 1.148,99" = İnternet + Telefon) kalem değildir
    hizmet, acik = [], []
    for ad, tutar in ust:
        if len(acik) >= 2 and abs(sum(t for _, t in acik) - tutar) <= 0.01:
            acik = []
            continue
        hizmet.append((ad, tutar))
        acik.append((ad, tutar))
    ekler = []  # fatura tutarından sonra ödenecek tutara kadar: "Önceki Aydan Devir 0,12"
    for x in satirlar[t_i + 1:]:
        if re.match(r"^(odenecek|vergiler dahil|devlete odenen|yalniz)", sade(x)):
            break
        m = _ADLI_SATIR.match(x)
        if m and re.search(r"devir|yuvarlama|taksit", sade(m.group("ad"))):
            tutar = para_coz(m.group("sayilar").split()[-1])
            if tutar is not None:
                ad = re.sub(r"\s+toplam[ıi]$", "", m.group("ad").strip(), flags=re.IGNORECASE)
                ekler.append({"aciklama": ad, "miktar": 1, "birim": None, "birim_fiyat": tutar, "kdv_orani": 0, "tutar": tutar})
    n = len(hizmet)
    if not n or n > 14 or t_tutar is None:
        return None
    toplam = round(sum(t for _, t in hizmet), 2)
    sifir = round(sum(k["tutar"] for k in sifir_kalemler), 2)
    tol = max(1.0, 0.002 * abs(t_tutar))

    def kalem(ad, tutar, oran, kdv_tutari=None):
        return {"aciklama": ad, "miktar": 1, "birim": None, "birim_fiyat": tutar, "kdv_orani": oran, "tutar": tutar,
                **({"kdv_tutari": kdv_tutari} if kdv_tutari is not None else {})}

    def dagit(tutarlar, hedef):  # hedef toplamı satırlara oranla dağıtır; kuruş farkı son satıra
        top = sum(tutarlar)
        parcalar = [round(hedef * t / top, 2) for t in tutarlar]
        parcalar[-1] = round(hedef - sum(parcalar[:-1]), 2)
        return parcalar
    if abs(toplam - t_tutar) <= 0.05:  # satırlar vergiler dahil
        # Vergilendirilmemiş küçük kısım (hangi satıra ait olduğu yazmaz) faturanın toplamıyla kaydedilir, yuvarlama olur
        if abs(toplam - (kdv_matrah + kdv + diger + sifir)) > max(1.0, 0.005 * abs(t_tutar)) or toplam <= 0:
            return None
        netler, kdvler = dagit([t for _, t in hizmet], kdv_matrah), dagit([t for _, t in hizmet], kdv)
        kalemler = [kalem(ad, net, kdv_orani, kv) for (ad, _), net, kv in zip(hizmet, netler, kdvler)]
    else:  # satırlar vergisiz: KDV matrahını tutan alt küme KDV'li
        if abs(toplam + sifir + kdv + diger - t_tutar) > tol:
            return None
        # Matrahı tutan kümelerden en çok satırı KDV'li sayan seçilir (KDV'siz olan genelde tek satırdır: aracılık edilen ödeme)
        en_iyi = None
        for maske in range(1, 1 << n):
            fark = abs(sum(hizmet[i][1] for i in range(n) if maske >> i & 1) - kdv_matrah)
            if fark <= tol and (en_iyi is None or (bin(maske).count("1"), -fark) > (bin(en_iyi[1]).count("1"), -en_iyi[0])):
                en_iyi = (fark, maske)
        if en_iyi is None:
            # Matrahı tutan küme yok (KDV'si satırlara faturada dağıtılmamış): aracılık edilen ödemeler KDV'siz, diğerleri
            # KDV'li sayılır; satırların KDV'si faturada yazan KDV toplamından dağıtılır, toplamlar faturayı tutar.
            maske = sum(1 << i for i in range(n) if not re.search(r"aracilik|adina tahsil", sade(hizmet[i][0])))
            if not maske:
                return None
        else:
            maske = en_iyi[1]
        kdvli = [i for i in range(n) if maske >> i & 1]
        kvler = dict(zip(kdvli, dagit([hizmet[i][1] for i in kdvli], kdv))) if all(hizmet[i][1] for i in kdvli) else {}
        kalemler = [kalem(ad, t, kdv_orani if i in kvler or maske >> i & 1 else 0, kvler.get(i)) for i, (ad, t) in enumerate(hizmet)]
    return kalemler + sifir_kalemler + ekler, t_tutar


def _ozetten_kalemler(metin: str, genel: float) -> tuple[list[dict], float, float, float, float] | None:
    """Kalem tablosu olmayan abonelik faturaları (TTNET, Turkcell kurumsal): vergi özetindeki matrah satırlarından
    kalem çıkarır. 'KDV %20 (Matrah 1.092.49) 218,50' -> %20'lik hizmet kalemi; ÖİV/ÖTV diğer vergilere gider;
    genel toplamı tamamlayan fark (taksit, devir, aracılık edilen ödemeler) %0'lık ayrı kalem olur.
    Faturada hizmet tablosu varsa (_abonelik_kalemleri) kalemler oradan adlarıyla gelir; kuruş farkı yuvarlama olur.
    Dönüş: (kalemler, kdv, diğer vergiler, yuvarlama, fatura tutarı)."""
    kalemler, kdv, diger = [], 0.0, 0.0
    kdv_oranlari, oiv = [], []
    for satir in metin.splitlines():
        m = _OZET_VERGI.match(" ".join(satir.split()))
        if not m:
            continue
        ad, oran = sade(m.group("ad")), para_coz(m.group("oran")) or 0
        matrah, tutar = _matrah(m.group("mat")), _matrah(m.group("tut"))
        if matrah is None or tutar is None:
            continue
        if re.search(r"\bkdv\b|katma deger", ad):
            kalemler.append({"aciklama": f"Hizmet bedeli (KDV %{oran:g})", "miktar": 1, "birim": None,
                             "birim_fiyat": matrah, "kdv_orani": oran, "tutar": matrah})
            kdv += tutar
            kdv_oranlari.append((oran, matrah))
        elif re.search(r"\boiv\b|ozel iletisim|\botv\b|ozel tuketim", ad):
            diger += tutar
            if re.search(r"\boiv\b|ozel iletisim", ad):
                oiv.append((oran, matrah))
        elif oran == 0:
            kalemler.append({"aciklama": " ".join(m.group("ad").split()), "miktar": 1, "birim": None,
                             "birim_fiyat": tutar, "kdv_orani": 0, "tutar": tutar})
    if not any(k["kdv_orani"] for k in kalemler):
        return None
    kdv, diger = round(kdv, 2), round(diger, 2)
    if len(kdv_oranlari) == 1 and len(oiv) <= 1 and len(kalemler) == 1:
        tablo = _abonelik_kalemleri(metin, kdv_oranlari[0][0], kdv_oranlari[0][1], *(oiv[0] if oiv else (0, 0)))
        if tablo:
            yuv = round(genel - sum(k["tutar"] for k in tablo) - kdv - diger, 2)
            if abs(yuv) <= 0.02 * len(tablo):
                return tablo, kdv, diger, yuv, genel
    if len(kdv_oranlari) == 1:
        adli = _adli_ozet_kalemleri(metin, kdv_oranlari[0][0], kdv_oranlari[0][1], kdv, diger,
                                    [k for k in kalemler if not k["kdv_orani"]])
        if adli:
            tablo, t_tutar = adli
            for hedef in (genel, t_tutar):  # ödenecek tutar, önceki ödeme/alacak düşülmüşse fatura tutarı
                yuv = round(hedef - sum(k["tutar"] for k in tablo) - kdv - diger, 2)
                if abs(yuv) <= max(1.0, 0.005 * hedef):
                    return tablo, kdv, diger, yuv, hedef
    # Faturanın kendi ücret satırları bulunamadı: "Hizmet bedeli" gibi uydurma adla kalem açılmaz (kullanıcı kuralı);
    # vergiler okunur, kalemler boş kalır, fatura kırmızı görünür ve kalemler elle girilir.
    return [], kdv, diger, 0.0, genel


def _iskonto_kalemleri(kalemler: list[dict], iskonto: float) -> list[dict] | None:
    """Fatura altı iskonto: kalemler iskontosuz yazılmışsa her KDV oranı için o orandaki kalemlerin payına düşen
    eksi tutarlı bir 'İskonto' satırı. Kalemlerin miktar ve fiyatı faturadaki gibi kalır, matrah ve KDV tutar."""
    toplam = sum(k["tutar"] for k in kalemler)
    if not toplam or iskonto <= 0 or iskonto >= toplam:
        return None
    oranlar: dict[float, float] = defaultdict(float)
    for k in kalemler:
        oranlar[k["kdv_orani"]] += k["tutar"]
    satirlar, kalan = [], iskonto
    for i, (oran, tutar) in enumerate(sorted(oranlar.items(), reverse=True)):
        pay = kalan if i == len(oranlar) - 1 else round(iskonto * tutar / toplam, 2)
        kalan = round(kalan - pay, 2)
        etiket = f"İskonto (%{iskonto / toplam * 100:.2f})".replace(".", ",") + (f" - KDV %{oran:g} kalemleri" if len(oranlar) > 1 else "")
        satirlar.append({"aciklama": etiket, "miktar": 1, "birim": None, "birim_fiyat": -pay, "kdv_orani": oran, "tutar": -pay})
    return satirlar


_ISK_NOTU = re.compile(r"\s*\((İsk\. [^()]*?)(?:, liste fiyatı ([\d.,]+))?\)")


def _tr_sayi(x: str) -> float | None:
    """'1.600' / '142,85' / '4.761,9': nokta binlik, virgül ondalık (programın yazdığı biçim)."""
    try:
        return float(x.replace(".", "").replace(",", "."))
    except ValueError:
        return None


_SUS = re.compile(r"(?:(?<=\s)|^)[#*=~_|]{2,}(?=\s|$)")


def susu_temizle(k: dict) -> dict:
    """Kalem adındaki süs işaretleri ("## ULS RACK KABİNET ##", "** ÜRÜN **") atılır; adın kendisi aynen kalır."""
    ad = k.get("aciklama") or ""
    yeni = " ".join(_SUS.sub(" ", ad).split())
    return {**k, "aciklama": yeni} if yeni and yeni != ad else k


def iskontoyu_ayir(k: dict) -> dict:
    """Standart kalem: birim fiyat iskontosuz liste fiyatıdır, iskonto oranı ve tutarı ayrı sütundadır,
    tutar (Mal Hizmet Tutarı) = miktar x birim fiyat - iskonto. Okuyucuların açıklamaya yazdığı
    '(İsk. %60, liste fiyatı 1.600)' notu bu sütunlara taşınır; tutar değişmez."""
    k.setdefault("iskonto_orani", 0)
    k.setdefault("iskonto_tutari", 0)
    m = _ISK_NOTU.search(k.get("aciklama") or "")
    tutar, miktar = k.get("tutar"), k.get("miktar") or 1
    if not m or not tutar or tutar <= 0:
        return k
    oranlar = [y for y in (_tr_sayi(x) for x in re.findall(r"%\s*([\d.,]+)", m.group(1))) if y]
    liste = _tr_sayi(m.group(2)) if m.group(2) else None
    if liste is None and len(oranlar) == 1 and 0 < oranlar[0] < 100:
        liste = round(tutar / (miktar * (1 - oranlar[0] / 100)), 6)
    if liste is None and (tl := re.search(r"([\d.,]+)\s*TL", m.group(1))) and (isk := _tr_sayi(tl.group(1))):
        liste = round((tutar + isk) / miktar, 6)
    if not liste:
        return k
    brut = round(miktar * liste, 2)
    isk = round(brut - tutar, 2)
    if isk <= 0:
        return k
    oran = oranlar[0] if len(oranlar) == 1 and abs(brut * oranlar[0] / 100 - isk) <= max(0.05, 0.002 * brut) \
        else round(isk / brut * 100, 2)
    k.update(aciklama=(k["aciklama"][:m.start()] + k["aciklama"][m.end():]).strip(), birim_fiyat=liste,
             iskonto_orani=oran, iskonto_tutari=isk)
    return k


def _ek_tutarlar(ham: list[str], s: str, kalemler: list[dict], ara: float, oran: float) -> list[dict]:
    """Kalem tablosunun dışında, toplamlar bölümünde matraha eklenen tutarlar: "Kargo Ücreti 124,92 TL / Toplam Ek Tutar
    124,92 TL" (Avansas). Kalemlerin toplamı + ek tutar = KDV hariç toplam ise her ek satır faturadaki adıyla kalem olur;
    KDV oranı, o oranın matrahından kalemlerle karşılanmayan kısma göre bulunur."""
    toplam = round(sum(k["tutar"] for k in kalemler), 2)
    satirlar = [sade(x) for x in ham]
    ek = None
    for i, x in enumerate(satirlar):
        m = re.search(r"\btoplam ek tutari?\s*:?\s*([\d.,]*\d)", x)  # resimden okunanda solda başka sütun olabilir
        if m:
            ek, son = para_coz(m.group(1)), i
            break
    if not ek or abs(toplam + ek - ara) > 0.05:
        return []
    bas = next((i for i in range(son - 1, -1, -1) if re.search(r"hizmet toplam", satirlar[i])), son - 4)
    adlar = re.compile(r"(?:^|(?<=[\s,;]))([a-z][a-z .()/-]{2,40}?)\s*:?\s*([\d.,]*\d)\s*(?:tl|try)?$")
    def ad(i, m):  # faturada yazdığı gibi; sade metinle hizalanmıyorsa sade hali
        x = " ".join(ham[i].split())[m.start(1):m.end(1)].strip(" :")
        return x if sade(x) == m.group(1).strip() else m.group(1).strip()
    satir = [(ad(i, m), para_coz(m.group(2))) for i in range(bas + 1, son) for m in [adlar.search(satirlar[i])] if m]
    satir = [(a, t) for a, t in satir if t and not re.search(r"toplam|iskonto|kdv|matrah", a)]
    if abs(sum(t for _, t in satir) - ek) > 0.05:
        satir = [("ek tutar", ek)]
    # Oran başına matrahtan kalemlerin karşıladığı kısım düşülünce kalan ek tutara ait
    matrah = {para_coz(m.group(1)): para_coz(m.group(2)) for m in re.finditer(
        r"kdv matrahi\s*\(\s*%\s*([\d.,]+)\s*\)\s*:?\s*([\d.,]*\d)", s)}
    kalan = {o: round(t - sum(k["tutar"] for k in kalemler if abs((k.get("kdv_orani") or 0) - o) < 0.01), 2)
             for o, t in matrah.items() if o is not None and t is not None}
    sonuc = []
    for ad, t in satir:
        o = next((o for o, k in kalan.items() if abs(k - t) <= 0.05), None)
        if o is None:
            o = next((o for o, k in kalan.items() if k >= t - 0.05), oran)
        kalan[o] = round(kalan.get(o, 0) - t, 2)
        sonuc.append({"aciklama": ad if ad != "ek tutar" else "Ek tutar", "miktar": 1, "birim": None,
                      "birim_fiyat": t, "kdv_orani": o, "tutar": t})
    return sonuc


_TUTAR_SATIRI = re.compile(r"^\s*-?\d[\d.,]*[.,]\d{2}\s*(?:TL|TRY|₺)?\s*$", re.IGNORECASE)
_VERGI_ETIKETI = re.compile(r"kdv|vergi|matrah|tevkifat|hesaplanan|stopaj|stpj", re.IGNORECASE)


def _alt_alta_etiketler(metin: str) -> str:
    """Etiketler alt alta, tutarları ayrı bir sütunda yazılmış toplam bloğu: "Hesaplanan KDV %1 / Hesaplanan KDV %10 /
    0,20 TRY / 80,00 TRY". Etiket ile tutar aynı satıra getirilir (yalnız ilk etiket ilk tutarla eşleşmesin)."""
    satirlar = metin.splitlines()
    i = 0
    while i < len(satirlar):
        j = i
        while j < len(satirlar) and _VERGI_ETIKETI.search(satirlar[j]) and not re.search(r"\d[.,]\d{2}\b", satirlar[j]):
            j += 1
        n = j - i
        if n >= 2 and j + n <= len(satirlar) and all(_TUTAR_SATIRI.match(satirlar[j + x]) for x in range(n)) \
                and not (j + n < len(satirlar) and _TUTAR_SATIRI.match(satirlar[j + n])):
            satirlar[i:j + n] = [f"{satirlar[i + x].rstrip()} {satirlar[j + x].strip()}" for x in range(n)]
            i += n
        else:
            i = max(j, i + 1)
    return "\n".join(satirlar)


def pdf_fatura_oku(veri: bytes, santiyeler: list[str] | None = None, sirket_unvan: str | None = None,
                   sirket_vkn: str | None = None, yon: str = "gelen", metin: str | None = None) -> dict:
    """metin verilirse (resimden metin tanıma ile okunmuş) PDF metni ve çizgili tablo okunmaz, yalnız metin çözülür."""
    resimden = metin is not None
    if metin is None:
        metin = pdf_metni(veri)
    if len(metin.strip()) < 30:
        raise KuralHatasi("Bu PDF taranmış bir görüntü; yapay zekasız okunamaz. Faturanın XML dosyasını yükleyin "
                          "ya da elle girin.")
    metin = re.sub(r"(Vergi\s+Kimlik)\s*\n\s*(No\b)", r"\1 \2", metin, flags=re.IGNORECASE)  # satıra bölünmüş etiket
    # Etiket bir satırda, numara alt satırda: "Vergi No\n6291734765 (İlyasbey)"
    metin = re.sub(r"^(\s*(?:Vergi\s*(?:Kimlik\s*)?(?:No|Numaras[ıi])|VKN(?:\s*/\s*TCKN)?)\s*:?)\s*\n\s*(\d{10,11}\b)",
                   r"\1 \2", metin, flags=re.IGNORECASE | re.MULTILINE)
    metin = _alt_alta_etiketler(metin)
    tek = " ".join(metin.split())
    s = re.sub(r"\bk\.\s?d\.\s?v\.?", "kdv", sade(metin))  # "K.D.V. MATRAHI % 1"
    # Döviz faturası: tutarlar önce döviz cinsinden, sonra kurla TL karşılıkları yazılır. Fatura döviz cinsinden
    # okunur (e-Fatura XML'indeki gibi); TL karşılıkları toplamlara karışmasın diye kur satırından sonrası atlanır.
    para_m = re.search(r"(?:odenecek tutar|vergiler dahil toplam(?: tutar)?|fatura tutari)\s*:?\s*[\d.,]*\d\s*(usd|eur|gbp|chf)\b", s)
    para = para_m.group(1) if para_m else None
    kur = tl_genel = tl_mal = None
    # Toplamlar iki blok: önce "(TL)" etiketli TL karşılıkları, sonra "Ödenecek Tutar (USD) 803,88 USD"
    para_m2 = None if para else re.search(r"odenecek tutar\s*\((usd|eur|gbp|chf)\)\s*:?\s*[\d.,]*\d", s)
    if para_m2:
        para = para_m2.group(1).upper()
        kur_m = re.search(r"(?:doviz kuru|uygulanan kur|fatura kuru)\s*:?\s*([\d.,]*\d)", s)
        kur = para_coz(kur_m.group(1)) if kur_m else None
        tl_genel = para_coz(_ilk(r"odenecek tutar\s*\(\s*(?:tl|try)\s*\)\s*:?\s*([\d.,]*\d)", s) or "")
        tl_mal = para_coz(_ilk(r"mal (?:ve )?hizmet toplam(?: tutari)?\s*\(\s*(?:tl|try)\s*\)\s*:?\s*([\d.,]*\d)", s) or "")
        s = re.sub(r"[a-z ()%\d]*\(\s*(?:tl|try)\s*\)\s*:?\s*-?[\d.,]*\d\s*(?:tl|try)?", " ", s)
        s = re.sub(r"\s*\(" + para_m2.group(1) + r"\)", "", s)
    elif para:
        para = para.upper()
        # "Döviz Kuru: 55.284700 TL", "Fatura Kuru: 45,9911 TL", "Kur : 1 USD = 45,99110 TRL"
        kur_m = (re.search(r"(?:doviz kuru|uygulanan kur|fatura kuru|kur)\s*(?:\(\s*[a-z]{3}\s*\))?\s*:?\s*([\d.,]*\d)\s*(?:tl|try|trl)\b", s)
                 or re.search(r"\b1\s*" + para.lower() + r"\s*=\s*([\d.,]*\d)\s*(?:tl|try|trl)?\b", s))
        kur = para_coz(kur_m.group(1)) if kur_m else None
        # Faturanın kendi yazdığı TL karşılıkları: kur yazılmamışsa onlardan bulunur, yazılmışsa kurla denetlenir
        tl_genel = para_coz(_ilk(r"(?:odenecek tutar|vergiler dahil toplam(?: tutar)?|fatura tutari)\s*\(\s*(?:tl|try)\s*\)\s*:?\s*([\d.,]*\d)", s) or "")
        tl_mal = para_coz(_ilk(r"mal (?:ve )?hizmet toplam(?: tutari)?\s*\(\s*(?:tl|try)\s*\)\s*:?\s*([\d.,]*\d)", s) or "")
        # TL karşılıkları toplamlara karışmasın: döviz cinsinden ödenecek tutardan sonraki ilk TL bölümünden kesilir.
        # Başlıktaki "Döviz Kuru: 55,28 TL" satırı (toplamlardan önce) kesmez.
        kes = next((m for m in re.finditer(r"\(\s*(?:tl|try)\s*\)|doviz kuru|uygulanan kur|fatura kuru", s)
                    if m.start() > para_m.start()), None)
        if kes:
            s = s[:kes.start()]
    s_satirlar = [re.sub(r"\bk\.\s?d\.\s?v\.?", "kdv", sade(x)) for x in metin.splitlines()]
    # Etiketsiz numara: 16 haneli MERSİS numarası da bu biçime uyar ("MERSISNO: 0632037786100014"), o atlanır
    fatura_no = _ilk(r"fatura\s*no\s*:?\s*([A-Z0-9]{3}20\d{11})", tek) or next(
        (m.group(1) for m in re.finditer(r"\b([A-Z0-9]{3}20\d{11})\b", tek)
         if not re.search(r"mersis|sicil", tek[max(m.start() - 20, 0):m.start()], re.IGNORECASE)), None)
    _TARIH = r"(\d{1,2}\s*[-./]\s*\d{1,2}\s*[-./]\s*\d{4})"

    def tarih_bul(*etiketler):
        for e in etiketler:
            t = _ilk(e + r"\s*:?\s*" + _TARIH, tek)
            if t:
                return tarih_coz(re.sub(r"\s", "", t))
        return None
    # Fiş tipi faturada yalnız "Tarih: 29. 03. 2026" yazar (irsaliye/sipariş/vade tarihi değil)
    tarih = (tarih_bul(r"fatura\s*tarihi?", r"d[üu]zenle(?:n)?me\s*tarihi",
                       r"(?<!rsaliye )(?<!pari[şs] )(?<!vade )(?<!deme )(?<![A-Za-zÇĞİÖŞÜçğıöşü])tarih(?![A-Za-zçğıöşü])")
             or tarih_coz(re.sub(r"\s", "", _ilk(r"(\d{2}\s?[./-]\s?\d{2}\s?[./-]\s?\d{4})", tek) or "")))
    vade = tarih_bul(r"son\s*[öo]deme\s*tarihi", r"vade\s*tarihi", r"fatura\s*vadesi", r"[öo]deme\s*vadesi")

    def tutar(*etiketler):
        for e in etiketler:
            for m in re.finditer(e + r"\s*:?\s*(-?[\d.,]*\d)", s):  # etiketten sonra sayı gelmeyen eşleşmeleri atla
                if para_coz(m.group(1)) is not None:
                    return para_coz(m.group(1))
        return None

    def satir_tutari(desen):
        """Satır başındaki etiket: 'Toplam : 111,67 TL', 'KDV 203,35 TL', '%20.0 KDV : 22,33 TL'."""
        for x in s_satirlar:
            m = re.match(desen + r"\s*:?\s*(-?[\d.,]*\d)\s*(?:tl|try)?$", x)
            if m and para_coz(m.group(1)) is not None:
                return para_coz(m.group(1))
        return None

    mal_hizmet = tutar(r"(?:mal|malzeme|urun) ?/? ?hizmet toplam(?: tutari?)?", r"ara toplam") or satir_tutari(r"toplam") \
        or tutar(r"urun bedeli(?=\s*:)")  # "Ürün Bedeli: 545,45TL / ... / Net Tutar: 600,00TL" (mağaza faturası)
    iskonto = tutar(r"toplam iskonto(?:\s*\(\s*%[^)]*\))?", r"toplam indirim") or satir_tutari(r"indirim") or 0
    ara = tutar(r"kdv haric toplam(?: tutari?)?", r"vergi(?:ler)? haric toplam(?: tutari?)?", r"net mal degeri")
    if ara is None:  # oran başına matrah: "KDV Matrahı (%20) 191,04" / "(%10): 5.800" toplanır
        # "KDV Matrahı (%20) 191,04", "K.D.V. MATRAHI % 1 : 2150.07", "Kdv %20 Matrah 41.668,00"
        matrahlar = [para_coz(m.group(1) or m.group(2)) for m in re.finditer(
            r"kdv matrahi?\s*(?:\(\s*%\s*[\d.,]+\s*\)|%\s*[\d.,]+)\s*:?\s*(-?[\d.,]*\d)|"
            r"kdv\s*%\s*[\d.,]+\s*matrahi?\s*:?\s*(-?[\d.,]*\d)", s)]
        ara = round(sum(matrahlar), 2) if matrahlar and None not in matrahlar else tutar(r"kdv matrahi")
    if ara is None and mal_hizmet is not None:
        ara = round(mal_hizmet - iskonto, 2)
    kdv_list, kdv_gorulen, kdv_oranlari = [], set(), {}
    # "Hesaplanan (%20.00) 168,27" (KDV yazılmamış), "Kdv % 20 Tutar 8.333,60"
    # "Hesaplanan %20 KDV dahil(%20.00) 149,83": satır tutarları KDV dahil, KDV içinden hesaplanmış
    kdv_dahil = bool(re.search(r"hesaplanan\s*%\s*[\d.,]+\s*kdv dahil", s))
    for m in re.finditer(r"(?:(?:hesaplanan\s*%\s*[\d.,]+\s*)?kdv(?: gercek| tutari(?=\s?\()| dahil(?=\())?|katma deger vergisi|hesaplanan)\s?(?:\((?![^)]*tevkifat)[^)]*\d[^)]*\)|%\s*[\d.,]+)\s*(?:%\s*[\d.,]+\s*)?(?:tutari?\s*)?:?\s*([\d.,]*\d)(?:\s*(?:tl|try))?"
                         r"(?:\s+([\d.,]*\d)(?!\s*%))?", s):
        k = para_coz(m.group(1))
        if k is not None and m.group(2) and k in (ara, mal_hizmet):  # 'KDV (%20) Matrah: Hesaplanan KDV: 18.000 3.600'
            k = para_coz(m.group(2))
        # Aynı oranın aynı KDV'si faturada iki yerde yazılabilir (not kutusunda "KDV %1 TUTARI : 5.15" ve altta
        # "Hesaplanan KDV (%1) 5,15"): bir kez sayılır
        o = re.search(r"%\s*([\d.,]+)", m.group(0))
        if o and k and (o.group(1), k) in kdv_gorulen:
            continue
        if o:
            kdv_gorulen.add((o.group(1), k))
            if k is not None and (r := para_coz(o.group(1))) is not None:
                kdv_oranlari[r] = round(kdv_oranlari.get(r, 0) + k, 2)
        kdv_list.append(k)
    kdv = round(sum(k for k in kdv_list if k is not None), 2) if kdv_list else (
        satir_tutari(r"(?:%\s*[\d.,]+\s*)?kdv(?: ?\(%? ?[\d.,]+\))?")
        or para_coz(_ilk(r"(?:kdv toplami?|toplam kdv)\s*:\s*([\d.,]*\d)", s) or ""))  # "KDV Toplam" sütun başlığı değil
    # Oran başına ayrı tevkifat satırları toplanır: "Hesaplanan KDV Tevkifat(%50) 1.000,16" + "(%20) 33,36"
    tevkler: dict[str, float] = {}
    for m in re.finditer(r"kdv tevkifat(?:i| tutari)?\s?(\([^)]*\))?\s*:?\s*(-?[\d.,]*\d)", s):
        if para_coz(m.group(2)) is not None:
            tevkler.setdefault(m.group(1) or "", para_coz(m.group(2)))
    if not tevkler:  # "Tevkifat (%50 - Kodu : 627) 2.530,93TL"
        for m in re.finditer(r"(?<!kdv )tevkifat\s*(\(\s*%\s*[\d.,]+[^)]*\))\s*:?\s*(-?[\d.,]*\d)", s):
            if para_coz(m.group(2)) is not None:
                tevkler.setdefault(m.group(1), para_coz(m.group(2)))
    tevk = tutar(r"toplam tevkifat") or round(sum(tevkler.values()), 2) or tutar(r"tevkifat tutari") or 0
    if not tevk:  # tevkifat yalnız kalem satırında: "Tevkifat ( 619 - ... ) (% 70,00 )= 12.348,00 TL"
        tevk = round(sum(para_coz(m.group(1)) or 0 for m in re.finditer(
            r"tevkifat\s*\([^)]*\)\s*\(?\s*%\s*[\d.,]+\s*\)?\s*=\s*([\d.,]*\d)", s)), 2)
    # Gelir/kurumlar vergisi stopajı ödenecek tutardan düşülür: "Hesaplanan STPJ(%5) 96.790,96",
    # kalemde "GELİR VERGİSİ STOPAJI (%5,00)=96.790,96 TL"
    stopaj_orani = None
    stopaj = 0.0
    if m := re.search(r"(?:hesaplanan\s*)?(?:stpj|stopaj\w*|gelir vergisi kesintisi)\s*\(?\s*%\s*([\d.,]+)\s*\)?\s*:?\s*([\d.,]*\d)(?!\s*%)(?!\s*\))", s):
        stopaj_orani, stopaj = para_coz(m.group(1)), para_coz(m.group(2)) or 0.0
        if m2 := re.search(r"(?:hesaplanan\s*)?(?:stpj|stopaj\w*)\s*\(\s*%\s*[\d.,]+\s*\)\s*=\s*([\d.,]*\d)", s):
            stopaj = para_coz(m2.group(1)) or stopaj  # satırdaki "(%5,00)=96.790,96" biçimi
    odenecek = tutar(r"odenecek tutar", r"odenecek toplam")
    genel = odenecek if odenecek is not None else tutar(r"genel toplam", r"fatura toplami", r"fatura tutari(?=\s*:)")
    if genel is None and re.search(r"urun bedeli\s*:", s):
        genel = tutar(r"net tutar(?=\s*:)")
    if odenecek is None and genel is not None and iskonto and kdv is not None:
        # "Genel Toplam" iskonto öncesi brüt olabilir (Genel Toplam 71.600, Toplam İskonto 42.960, KDV 5.728):
        # ödenecek tutar sanılırsa cari borcu iskonto kadar fazla görünür.
        brut = mal_hizmet if mal_hizmet is not None else genel
        net = round(brut - iskonto, 2)
        beklenen = round(net + kdv - tevk - stopaj, 2)
        if abs(genel - brut) <= 0.05 and net > 0 and (ara is None or abs(ara - net) <= 0.05) and abs(genel - beklenen) > 0.05:
            sayilar = {para_coz(x) for x in re.findall(r"\d[\d.,]*\d", s)}
            if (any(abs(kdv - net * r / 100) <= 0.05 + 0.001 * kdv for r in (1, 8, 10, 18, 20))
                    or beklenen in sayilar or ara is not None):
                mal_hizmet, ara, genel = brut, net, beklenen
    if genel is None:  # 'Vergi Dahil Toplam' tevkifattan öncedir; ödenecek tutar = bu - tevkifat
        dahil = tutar(r"vergi(?:ler)? dahil toplam(?: tutari?)?")
        genel = round(dahil - tevk - stopaj, 2) if dahil is not None else None
    # KDV dışındaki vergiler (konaklama vergisi, ÖTV, ÖİV...) genel toplama eklenir
    diger = round(sum(para_coz(m.group(2)) or 0 for m in re.finditer(
        r"(konaklama vergisi|hesaplanan kv|ozel tuketim vergisi|otv|ozel iletisim vergisi|oil?v|eglence vergisi)\s*"
        r"(?:\([^)]*\)|%\s*[\d.,]+)?\s*:?\s*"
        r"([\d.,]*\d)(?!\s*%)", s)), 2)
    if diger and ara is not None and kdv is not None and genel is not None and abs(ara + kdv - tevk - stopaj + diger - genel) > 0.05:
        diger = 0  # yakalanan sayı vergi tutarı değilmiş (ör. oran)
    if not diger and ara is not None and genel is not None and re.search(r"konaklama vergisi|otv|ozel tuketim|oil?v|ozel iletisim", s):
        kdv_ = kdv if kdv is not None else 0
        fark = round(genel - (ara + kdv_ - tevk - stopaj), 2)
        diger = fark if fark > 0.01 else 0
    if ara is None and genel is not None and kdv is not None:
        ara = round(genel + tevk + stopaj - kdv - diger, 2)
    if kdv is None and genel is not None and ara is not None:
        kdv = round(genel + tevk + stopaj - ara - diger, 2)  # KDV satırı okunamadıysa farktan bul

    if kdv_dahil and kdv and ara is not None and genel is not None and abs(ara - genel) <= 0.05:
        ara = round(genel - kdv, 2)

    def tutuyor(x):  # bu tutar mal/hizmet toplamıysa KDV ve vergilerle ödenecek tutarı verir
        return x is not None and None not in (kdv, genel) and abs(x + kdv - tevk - stopaj + diger - genel) <= 0.05
    if genel is None and None not in (ara, kdv, mal_hizmet) and abs(ara + kdv - tevk - stopaj + diger - mal_hizmet) <= 0.05 \
            and abs(mal_hizmet - ara) > 0.05:
        genel, mal_hizmet = mal_hizmet, ara  # fiş: "Toplam: 12.000" ödenecek tutar, "Net Mal Değeri: 10.000" mal/hizmet
    net_mal = round(mal_hizmet - iskonto, 2) if mal_hizmet is not None else None
    if not tutuyor(ara) and tutuyor(net_mal):
        ara = net_mal  # KDV'siz kısmı olan fatura (bilet: hizmet bedeli): KDV matrahı mal/hizmet toplamından küçük
    # Markette "Ara Toplam" KDV dahildir; ödenecek tutarı yalnız matrahlar veriyorsa kalemler matraha göre denetlenir
    hedefler = [ara] if tutuyor(ara) and kdv and net_mal is not None and genel is not None and abs(net_mal - genel) <= 0.05 \
        else [mal_hizmet, ara]
    oran = min((0, 1, 8, 10, 18, 20), key=lambda o: abs(o - kdv / ara * 100)) if ara and kdv is not None else 20

    satici, alici = pdf_taraflari(metin)
    kendi = kendi_kelimeleri(sirket_unvan)

    def bizim_mi(t):
        if sirket_vkn and t["vkn"] == sirket_vkn:
            return True
        kelimeler = re.sub(r"[^a-z0-9 ]", " ", sade(t["unvan"] or "")).split()
        return bool(kendi) and kelimeler[:len(kendi)] == kendi
    # Alış faturasında biz alıcıyız, satış faturasında satıcı; ters okunduysa tarafları değiştir.
    if (yon == "gelen" and bizim_mi(satici) and not bizim_mi(alici)) or (
            yon == "giden" and bizim_mi(alici) and not bizim_mi(satici)):
        satici, alici = alici, satici
    karsi_taraf = satici if yon == "gelen" else alici
    if sirket_vkn and karsi_taraf["vkn"] == sirket_vkn:
        # kendi VKN'miz hiçbir zaman karşı tarafın (satıcının) VKN'si olmaz; o VKN'nin yanındaki ad da bizimdir
        karsi_taraf["vkn"] = karsi_taraf["unvan"] = None
    # Kendi adımızın satırları satıcı adı olamaz: "ARMA GRUP ENERJİ⏎MÜHENDİSLİK İNŞAAT LİMİTED⏎ŞİRKETİ" bölünmüş
    # yazıldığında alt satırlar da şirket eki taşır. Ad, şirket ekiyle biten satıra kadar sürer.
    tum_satirlar = metin.splitlines()
    bizim_blok = set()
    for i, sat in enumerate(tum_satirlar):
        if not bizim_mi({"unvan": sat, "vkn": None}):
            continue
        j = i
        while j < len(tum_satirlar) and j - i < 4:
            bizim_blok.add(j)
            if re.search(r"\b(sirketi?|sti\.?|a\.?\s?s\.?|anonim sirketi)\s*$", sade(tum_satirlar[j]).strip()) \
                    or (j > i and re.search(r"\d", tum_satirlar[j])):
                break
            j += 1
    if yon == "gelen" and not satici["vkn"]:
        # VKN etiketsiz yazılmış olabilir (fiş tipi faturada ünvanın altında tek başına "6100000006"); MERSİS
        # numarası "0" + VKN + 5 hanedir: 0610000000600256 -> 6100000006
        adaylar = [m.group(1) for m in re.finditer(r"^\s*(\d{10,11})\s*$", metin, re.MULTILINE)]
        mersis = re.search(r"mersis\s*(?:no|numarasi)?\s*:?\s*([\d\s-]{16,22})", s)
        if mersis and len(re.sub(r"\D", "", mersis.group(1))) == 16:
            adaylar.append(re.sub(r"\D", "", mersis.group(1))[1:11])
        satici["vkn"] = next((v for v in adaylar if v not in (sirket_vkn, alici["vkn"])), None)

    if not satici["unvan"]:  # ünvan bloğu bozuk çıktıysa: şirket eki taşıyan, bize ait olmayan ilk satır
        satici["unvan"] = next((sat.strip() for i, sat in enumerate(tum_satirlar) if SIRKET_EKLERI.search(sade(sat))
                                and i not in bizim_blok
                                and not bizim_mi({"unvan": sat, "vkn": None}) and len(sat.strip()) < 120
                                and not _ETIKET.search(sade(sat))
                                and not re.search(r"\btr\s?\d{2}(?:\s?\d{4}){3}|\biban\b", sade(sat))  # banka/IBAN satırı
                                and not set(re.sub(r"[^a-z0-9 ]", " ", sade(sat)).split()) <= GENEL_KELIMELER), None)
    if satici["unvan"] and re.search(r"(limited|anonim|ticaret|sanayi)$", sade(satici["unvan"])):
        # "MEDİA MARKT TURKEY TİCARET LİMİTED⏎ŞİRKETİ": ünvanın son kelimesi alt satırda kalmış
        satirlar_ = [x.strip() for x in metin.splitlines()]
        i = next((i for i, x in enumerate(satirlar_) if x == satici["unvan"].strip()), None)
        if i is not None and i + 1 < len(satirlar_) and re.fullmatch(r"(sirketi?|sti\.?)", sade(satirlar_[i + 1])):
            satici["unvan"] = f"{satici['unvan'].strip()} {satirlar_[i + 1]}"
    if not satici["unvan"]:  # ünvan hiç yazılmamışsa (Turkcell kurumsal fatura) web adresinden: www.turkcell.com.tr
        alan = _ilk(r"www\.([a-z0-9-]{3,})\.com", s, 0)
        satici["unvan"] = tr_buyuk(alan) if alan else None
    from .tablo import tablodan_kalemler  # başlıklara göre genel tablo okuma; olmazsa eski düzenler
    kalemler = ((None if resimden else tablodan_kalemler(veri, hedefler, oran, genel, kdv)) or _kalemleri_coz(tek, hedefler, oran)
                or _kalemleri_satirdan(metin, hedefler, oran) or _tarihli_satirlar(tek, hedefler, oran)
                or (None if resimden else tablodan_kalemler(veri, hedefler, oran, genel, kdv, yazidan=True))
                or _fis_kalemleri(metin, hedefler, genel, kdv))
    if kalemler and len(kdv_oranlari) > 1:
        _oranlari_dagit(kalemler, kdv_oranlari)
    if kalemler and ara is None and genel is not None:  # toplam satırları okunamadı; satırlar ödenecek tutarı verdi
        ara = round(sum(k["tutar"] for k in kalemler), 2)
        if kdv is None:
            kdv = round(genel + tevk + stopaj - ara - diger, 2)
    # Satırda yazan diğer vergiler (konaklama vergisi) faturanın altında ayrıca yazmıyorsa oradan alınır
    satir_diger = round(sum(k.pop("_diger", 0) or 0 for k in kalemler or []), 2)
    if satir_diger and not diger and None not in (ara, kdv, genel) and abs(ara + kdv - tevk - stopaj + satir_diger - genel) <= 0.05:
        diger = satir_diger
    ozet, notlar_ek = None, ""
    if not kalemler and genel is not None:
        ozet = _ozetten_kalemler(metin, genel)
        if ozet:
            kalemler, kdv, diger, ozet_yuv, fatura_tutari = ozet
            if kalemler and abs(fatura_tutari - genel) > 0.05:
                notlar_ek = (f" Ödenecek tutar ({_sayi_yazi(genel)} TL) fatura tutarından ({_sayi_yazi(fatura_tutari)} TL) "
                             "farklı; fatura, fatura tutarıyla kaydedildi.")
                genel = fatura_tutari
            ara, tevk = round(sum(k["tutar"] for k in kalemler) + ozet_yuv, 2) if kalemler else round(genel - kdv - diger, 2), 0
        if not kalemler:
            ozet = None
    cozuldu = bool(kalemler)
    if cozuldu and not ozet:
        kalemler = _kalemleri_denetle(kalemler)
    notlar = ("Resimden yapay zekasız okundu." if resimden else "PDF'ten yapay zekasız okundu.") + notlar_ek
    if para:
        if not kur and genel and tl_genel:
            kur = round(tl_genel / genel, 4)
        elif not kur and mal_hizmet and tl_mal:
            kur = round(tl_mal / mal_hizmet, 4)
        notlar += f" Fatura {para} cinsinden" + (f" (kur {_sayi_yazi(kur)} TL)." if kur else "; faturada kur yazmıyor.")
        if kur and genel and tl_genel and abs(genel * kur - tl_genel) > 1 + 0.001 * tl_genel:
            notlar += (f" Faturada yazan TL tutarı ({_sayi_yazi(tl_genel)} TL) kurla hesaplanandan "
                       f"({_sayi_yazi(round(genel * kur, 2))} TL) farklı; kuru kontrol edin.")
    if cozuldu and not ozet:
        satir_tevk = round(sum(tevkifat_tutari(k["tutar"], k["kdv_orani"], k.get("tevkifat_orani"), k.get("kdv_tutari")) for k in kalemler), 2)
        if satir_tevk and not tevk:
            tevk = satir_tevk
            if genel is not None and ara is not None and kdv is not None and abs(genel - (ara + kdv + diger)) <= 0.05:
                genel = round(genel - tevk, 2)  # okunan "genel" tevkifattan önceki toplammış
        elif satir_tevk and abs(satir_tevk - tevk) > 0.02 * len(kalemler):
            notlar += (f" Satırlardan hesaplanan tevkifat ({_sayi_yazi(satir_tevk)} TL) faturadaki tevkifattan "
                       f"({_sayi_yazi(tevk)} TL) farklı; kalemlerdeki tevkifat oranlarını kontrol edin.")
        elif tevk and not satir_tevk and not _tevkifati_dagit(kalemler, tevk, metin):
            notlar += " Faturada tevkifat var ama hangi kalemlere ait olduğu okunamadı; kalemlerde tevkifat oranını seçin."
        _kodlari_tamamla(kalemler, metin)
        stopaj = stopaji_dagit(kalemler, stopaj, stopaj_orani)
    if kalemler and ara is not None:
        kalemler += _ek_tutarlar(metin.splitlines(), s, kalemler, ara, oran)
    # Fiyatsız ve tutarsız satır, metni faturanın "Not:" satırında da geçiyorsa ("KÜTAHYA TOROSGAZ FİRMASI ALDI") kalem
    # değil nottur; bedelsiz verilen gerçek mal (0 TL'lik diafon paneli) kalem kalır.
    notlar_metni = " ".join(sade(m.group(1)) for m in re.finditer(r"(?im)^\s*(?:not|aciklama|açıklama)\s*:\s*(.+)$", metin))
    yazi_satirlari = [k for k in kalemler or [] if not k.get("tutar") and not k.get("birim_fiyat") and not k.get("iskonto_tutari")
                      and len(sade(k.get("aciklama") or "")) >= 8 and sade(k.get("aciklama") or "") in notlar_metni]
    if yazi_satirlari and len(yazi_satirlari) < len(kalemler):
        kalemler = [k for k in kalemler if k not in yazi_satirlari]
    else:
        yazi_satirlari = []
    yuvarlama_farki = ozet[3] if ozet else 0.0
    if yuvarlama_farki and abs(yuvarlama_farki) > 1:
        # Kuruş değil: faturada hiçbir satıra yazılmamış tutar. Uydurma satır açılmaz; fark açıkça söylenir (uyarı kalır)
        notlar += (f" Faturadaki satırların toplamı fatura tutarından {_sayi_yazi(abs(yuvarlama_farki))} TL "
                   f"{'az' if yuvarlama_farki > 0 else 'fazla'}; fark faturada bir satıra yazılmamış. Kayıt fatura tutarıyla yapıldı.")
    elif yuvarlama_farki:
        notlar += f" Faturada {_sayi_yazi(abs(yuvarlama_farki))} TL yuvarlama farkı var; kayıt faturadaki toplamla yapıldı."
    if not kalemler:
        kalemler = []
        # Kullanıcı kuralı: kalemler asla "KDV %20 matrahı" gibi geçici bir satıra toplanmaz; boş kalır, kırmızı görünür
        # ve onaylanamaz. Faturadaki toplamlar ai_okuma'da kalır, elle girilen kalemler onlarla karşılaştırılır.
        notlar += " Kalemler tek tek okunamadı; faturadaki kalemleri (Mal Hizmet Adı, Miktar, Birim Fiyat, KDV) girin."
    elif ara is not None and mal_hizmet is not None and abs(sum(k["tutar"] for k in kalemler) - ara) > 0.05:
        # Faturada iskonto yoksa ("Toplam İskonto 0,00") fark iskonto satırıyla kapatılmaz
        iskonto_yok = re.search(r"(?:toplam )?iskonto(?: tutari)?\s*:?\s*0[.,]00(?![\d.,]*[1-9])", s) and not iskonto
        fark = round(ara - sum(k["tutar"] for k in kalemler), 2)
        iskonto_satirlari = None if iskonto_yok else _iskonto_kalemleri(kalemler, -fark)
        if iskonto_satirlari:
            kalemler += iskonto_satirlari
        elif abs(fark) <= 1 and all(_kalem_kendi_tutarli(k) for k in kalemler):
            # Satırlar kendi içinde doğru, faturanın toplamı kuruşla farklı: satıcının kendi yuvarlaması. Kalem uydurulmaz;
            # fark faturanın "yuvarlama" alanına yazılır, kayıt faturadaki toplamla yapılır.
            yuvarlama_farki = fark
            notlar += (f" Faturada {_sayi_yazi(abs(fark))} TL yuvarlama farkı var (kalemlerin toplamı "
                       f"{_sayi_yazi(round(ara - fark, 2))}, faturada {_sayi_yazi(ara)} TL); kayıt faturadaki toplamla yapıldı.")
        elif iskonto_yok:
            notlar += (f" Kalemlerin toplamı faturadaki mal/hizmet toplamından {_sayi_yazi(abs(fark))} TL farklı; "
                       "kalem tutarlarını kontrol edin.")
        else:
            notlar += " Faturada iskonto var; kalem tutarlarını kontrol edin."
    # Ödenecek tutarı kuruşa tamamlayan "Yuvarlama 0,05 TL" (elektrik, su): faturada yazan satır, KDV'siz kalem olur.
    # Elektrik faturasında bu ayın yuvarlaması eklenir, "Önceki Yuvarlama" düşülür (309,49 + 0,06 - 0,05 = 309,50).
    yuv = round(sum((-1 if m.group(1) else 1) * (para_coz(m.group(2)) or 0)
                    for m in re.finditer(r"(onceki )?\byuvarlama\s*:?\s*(-?[\d.,]*\d)", s)), 2)
    if kalemler and yuv and None not in (genel, ara, kdv) and abs(ara + kdv - tevk - stopaj + diger + yuv - genel) <= 0.01 \
            and not any(sade(k["aciklama"]).startswith("yuvarlama") for k in kalemler):
        kalemler.append({"aciklama": "Yuvarlama", "miktar": 1, "birim": None, "birim_fiyat": yuv, "kdv_orani": 0, "tutar": yuv})
        ara = round(ara + yuv, 2)

    if cozuldu and not ozet and kdv is not None:
        # Satırda yazan KDV tutarları yalnız faturanın KDV toplamına daha yakın sonuç veriyorsa esas alınır
        # (akaryakıt faturasında satır KDV'si iskonto sonrası yazılır, iskonto ayrıca satır olarak düşülür).
        oranla = round(sum((k.get("tutar") or 0) * (k.get("kdv_orani") or 0) / 100 for k in kalemler), 2)
        yazanla = round(sum(kalem_kdv(k) for k in kalemler), 2)
        if abs(oranla - kdv) <= abs(yazanla - kdv):
            for k in kalemler:
                k.pop("kdv_tutari", None)
        satir_kdv = round(sum(kalem_kdv(k) for k in kalemler), 2)
        if abs(satir_kdv - kdv) > 0.05 + 0.01 * len(kalemler):
            notlar += (f" Kalemlerin KDV toplamı ({_sayi_yazi(satir_kdv)} TL) faturadaki hesaplanan KDV'den "
                       f"({_sayi_yazi(kdv)} TL) farklı; kalemlerdeki KDV tutarlarını kontrol edin.")

    return {
        "belge_fatura_mi": bool(fatura_no or genel),
        "fatura_no": fatura_no,
        "tarih": tarih.isoformat() if tarih else None,
        "vade": vade.isoformat() if vade and vade != tarih else None,
        "satici": {**satici, "vergi_dairesi": None, "adres": None, "telefon": None, "iban": None},
        "alici": alici,
        "kalemler": [iskontoyu_ayir(susu_temizle(k)) for k in kalemler],
        "ara_toplam": ara, "kdv_toplam": kdv, "tevkifat": tevk, "stopaj": stopaj, "diger_vergiler": diger, "genel_toplam": genel,
        "yuvarlama": yuvarlama_farki, "aciklama": "; ".join(k["aciklama"] for k in yazi_satirlari) or None,
        "para_birimi": para or "TRY", "kur": kur if para else None, "tl_toplam": tl_genel if para else None,
        "fatura_tipi": "IADE" if re.search(r"fatura\s*tipi\s*:?\s*iade\b|\biade faturasi\b", sade(metin)) else None,
        "onerilen_santiye": santiye_tahmin(santiyeler or [], metin),
        "onerilen_gider_turu": (kalem_gider_turu(kalemler, satici["unvan"] or "") if cozuldu and not ozet else None)
        or gider_turu_tahmin(metin),
        "notlar": notlar,
        "kalemler_okunamadi": not cozuldu,
        "okuma_yontemi": "Resimden metin tanıma" if resimden else "PDF metni",
        "ettn": ettn_bul(metin),
        "red_isareti": red_isareti(metin),
    }


# --- Reddedilen / iptal edilen faturalar ------------------------------------------------------

_ETTN = re.compile(r"\b([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})\b", re.IGNORECASE)
# Portalın ya da entegratörün belgeye bastığı damga/yazı: "REDDEDİLDİ", "Fatura iptal edilmiştir", tek başına "RED".
# "Kredi", "EDENRED" gibi kelimelerin içi sayılmaz: tam kelime ya da satırın tamamı aranır.
_RED_YAZISI = re.compile(r"\b(?:reddedildi|reddedilmistir|reddedilen fatura|red edildi|red edilmistir|"
                         r"iptal edildi|iptal edilmistir|iptal edilen fatura)\b")
_RED_BASKASI = re.compile(r"\b(?:fis|fisi|fisler|fisleri|irsaliye\w*|siparis\w*|makbuz\w*|istinaden\w*|siparisi)\b")
_RED_SATIRI = re.compile(r"^(?:red|ret|iptal|reddedildi|iptal edildi|reddedilmistir|iptal edilmistir)[.!]?$")


def ettn_bul(metin: str) -> str | None:
    """Faturanın ETTN'si (UUID): yazının içinde 'ETTN' etiketinden sonra ya da metindeki ilk UUID."""
    m = re.search(r"ETTN\W{0,5}" + _ETTN.pattern[2:-2], metin or "", re.IGNORECASE) or _ETTN.search(metin or "")
    return m.group(1).lower() if m else None


def red_isareti(metin: str) -> bool:
    """Faturanın üstünde reddedildiğini / iptal edildiğini söyleyen yazı ya da damga var mı."""
    for sa in (metin or "").splitlines():
        sa = sade(sa)
        if _RED_SATIRI.match(sa):
            return True
        # "Bu faturaya istinaden 1 adet yazarkasa fişi iptal edilmiştir": iptal edilen fatura değil, fiş
        if any(not _RED_BASKASI.search(sa[:m.start()]) for m in _RED_YAZISI.finditer(sa)):
            return True
    return bool(_RED_YAZISI.search(sade(" ".join((metin or "").split()))) and not _RED_BASKASI.search(sade(metin or "")))


_RED_KODLARI = {"red", "ret", "rejected", "reject", "reddedildi", "iptal"}


def red_yanitlari(veri: bytes) -> set[str] | None:
    """e-Fatura yanıt belgesi (ApplicationResponse): RED ise reddedilen faturaların ETTN/numaraları; yanıt belgesi değilse None."""
    try:
        kok = ET.fromstring(veri)
    except ET.ParseError:
        return None
    if _yerel(kok.tag) != "ApplicationResponse":
        return None
    reddedilen = set()
    for yanit in _hepsi(kok, "DocumentResponse"):
        kod = sade(_metin(yanit, "Response", "ResponseCode") or "")
        if kod in _RED_KODLARI:
            ref = _bul(yanit, "DocumentReference")
            for alan in ("UUID", "ID"):
                if d := _metin(ref, alan):
                    reddedilen.add(d.lower())
    return reddedilen


# --- Banka ekstresi -------------------------------------------------------

def tablo_satirlari(veri: bytes, mime: str) -> list[list]:
    if mime == "text/csv":
        metin = None
        for kodlama in ("utf-8-sig", "cp1254", "latin-1"):
            try:
                metin = veri.decode(kodlama)
                break
            except UnicodeDecodeError:
                continue
        ayrac = ";" if metin.count(";") > metin.count(",") else ","
        return [satir for satir in csv.reader(io.StringIO(metin), delimiter=ayrac)]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    except Exception as e:
        raise KuralHatasi("Excel dosyası okunamadı. Dosyayı .xlsx olarak kaydedip tekrar deneyin.") from e
    ws = wb.worksheets[0]
    ws.reset_dimensions()  # Akbank dosyası boyutunu "A1" yazıyor: yalnız ilk satır okunuyordu
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _sutunlar(baslik: list) -> dict | None:
    adlar = [sade(str(h or "")) for h in baslik]
    sutun = {}
    for i, a in enumerate(adlar):
        if not a:
            continue
        if "tarih" in a and "valor" not in a and "tarih" not in sutun:
            sutun["tarih"] = i
        elif a in ("b/a", "ba", "b / a", "borc/alacak", "borc / alacak", "yon", "islem yonu", "(b/a)", "+/-"):
            sutun.setdefault("yon", i)
        elif "bakiye" in a:
            if "ek hesap" in a or "kmh" in a:  # İş Bankası "Ek Hesap Bakiyesi" (KMH): asıl bakiyeyle toplanır
                sutun.setdefault("ek_hesap", i)
            else:
                sutun.setdefault("bakiye", i)
        elif a in ("borc", "borc tutari", "cikan", "giden", "cikis") or a.startswith("borc"):
            sutun.setdefault("borc", i)
        elif a in ("alacak", "alacak tutari", "giren", "gelen", "giris") or a.startswith("alacak"):
            sutun.setdefault("alacak", i)
        elif "tutar" in a or a == "miktar":
            sutun.setdefault("tutar", i)
        elif "aciklama" in a:
            # Asıl açıklama sütunu "İşlem"/"İşlem Tipi" gibi kod sütunlarından önce gelir (İş Bankası: "EF", "FAST")
            if "aciklama" in sutun and "aciklama" not in adlar[sutun["aciklama"]]:
                sutun["islem"] = sutun["aciklama"]
                sutun["aciklama"] = i
            else:
                sutun.setdefault("aciklama", i)
        elif "islem tipi" in a or "islem turu" in a:
            sutun.setdefault("islem_tipi", i)
        elif any(k in a for k in ("islem", "detay", "karsi taraf", "gonderen", "alici")):
            if "aciklama" in sutun:
                sutun.setdefault("islem", i)
            else:
                sutun["aciklama"] = i
    if "tarih" in sutun and ("tutar" in sutun or "borc" in sutun or "alacak" in sutun):
        return sutun
    return None


IBAN_DESENI = r"TR\s?\d{2}(?:\s?\d{4}){5}\s?\d{2}"


def ekstre_iban(veri: bytes, mime: str) -> str | None:
    """Ekstrenin başlık kısmındaki hesap IBAN'ını bulur (Halkbank vb. ilk satırlarda verir)."""
    if mime == "application/pdf":
        return None
    try:
        satirlar = tablo_satirlari(veri, mime)
    except KuralHatasi:
        return None
    for satir in satirlar[:15]:
        metin = " ".join(str(v) for v in satir if v is not None)
        if "iban" in sade(metin):
            m = re.search(IBAN_DESENI, metin)
            if m:
                return re.sub(r"\s", "", m.group())
    return None


def ekstre_tablo_oku(veri: bytes, mime: str) -> list[dict]:
    satirlar = tablo_satirlari(veri, mime)
    for bas_i, satir in enumerate(satirlar[:40]):
        sutun = _sutunlar(satir)
        if sutun:
            break
    else:
        raise KuralHatasi("Ekstrede 'Tarih' ve 'Tutar' (ya da 'Borç'/'Alacak') sütunları bulunamadı.")

    def hucre(satir, ad):
        i = sutun.get(ad)
        return satir[i] if i is not None and i < len(satir) else None

    sonuc = []
    for satir in satirlar[bas_i + 1:]:
        tarih = tarih_coz(hucre(satir, "tarih"))
        if not tarih:
            continue
        if "tutar" in sutun:
            tutar = para_coz(hucre(satir, "tutar"))
            yon = sade(str(hucre(satir, "yon") or ""))
            if tutar and yon:  # B/A sütunu: B = borç (para çıktı), A = alacak (para girdi)
                tutar = -abs(tutar) if yon[0] in "b-" else abs(tutar)
        else:
            tutar = (para_coz(hucre(satir, "alacak")) or 0) - abs(para_coz(hucre(satir, "borc")) or 0)
        if not tutar:
            continue
        aciklama = str(hucre(satir, "aciklama") or "").strip()
        # "İşlem Tipi" (Ek Hesap, Para Çekme, Kredi Kartı...) türü bulmaya yardım eder; kod değil ad ise açıklamaya eklenir
        tip = str(hucre(satir, "islem_tipi") or "").strip()
        if aciklama and len(tip) > 4 and sade(tip) not in sade(aciklama):
            aciklama = f"{aciklama} ({tip})"
        if not aciklama:  # açıklama sütunu tanınmadıysa diğer metin hücrelerini birleştir
            aciklama = " ".join(str(v) for j, v in enumerate(satir) if isinstance(v, str) and j not in sutun.values()).strip()
        bakiye = para_coz(hucre(satir, "bakiye"))
        if bakiye is not None and "ek_hesap" in sutun:
            bakiye = round(bakiye + (para_coz(hucre(satir, "ek_hesap")) or 0), 2)
        sonuc.append({"tarih": tarih.isoformat(), "aciklama": aciklama, "tutar": round(tutar, 2), "bakiye": bakiye})
    if not sonuc:
        raise KuralHatasi("Ekstrede okunabilen hareket satırı bulunamadı.")
    if "ek_hesap" in sutun:
        sonuc = _ek_hesap_ici_ayikla(sonuc)
    return sonuc


def _ek_hesap_ici_ayikla(satirlar: list[dict]) -> list[dict]:
    """Ek hesaplı (KMH) ekstrede bakiye = hesap bakiyesi + ek hesap bakiyesi. Ek hesap borcunun hesaptan kapatılması
    ("KMH-ANAPARA BORCU TAHSİLATI") bu toplamı değiştirmez: hesabın içinde iki bölme arası aktarımdır, hareket sayılmaz."""
    yeniden_eskiye = satirlar[0]["tarih"] > satirlar[-1]["tarih"]
    sirali = list(reversed(satirlar)) if yeniden_eskiye else list(satirlar)
    kalan, onceki = [], None
    for s in sirali:
        b = s["bakiye"]
        if onceki is not None and b is not None and abs(b - onceki) < 0.01:
            continue
        if onceki is None and re.search(r"\bkmh\b.*anapara|anapara.*\bkmh\b", sade(s["aciklama"])):
            continue
        kalan.append(s)
        if b is not None:
            onceki = b
    return list(reversed(kalan)) if yeniden_eskiye else kalan


GENEL_KELIMELER = {"ltd", "sti", "as", "a.s", "a.s.", "san", "tic", "ins", "insaat", "ve", "limited", "sirketi", "anonim",
                   "ticaret", "sanayi", "taahhut", "ith", "ihr", "pazarlama", "ltd.", "sti.", "san.", "tic.", "ins."}
MASRAF_KELIMELERI = ["masraf", "komisyon", "bsmv", "ucret", "aidat", "kesinti", "faiz", "kkdf", "eft ucr", "havale ucr",
                     "gider vergisi", "kmh", "gayrinakdi", "ekspertiz", "degerli kagit", "sigorta prim", "police", "prim tahsil",
                     "fonfz", "kred.mev.vergi", "kredi mevduat vergi", "arti para", "artipara", " kom."]
VIRMAN_KELIMELERI = ["hesaplar arasi", "hesaplarim arasi", "hesaplar aras", "virman", "kendi hesab", "dbs hesab",
                     "hesabima transfer", "hesaba aktar", "vadeli h"]
NAKIT_CEKIM = re.compile(r"nakit cek|para cek|atm.*cek|kasaya")
NAKIT_YATIRMA = re.compile(r"nakit yat|para yat|atm.*yat")
SIRKET_EKLERI = re.compile(r"\b(ltd|sti|limited|sirket|a\.?s\b|anonim|tic\b|ticaret|kooperatif|ortakligi)", re.I)


# Kamu kurumu ünvanı ("T.C. KARAYOLLARI 14. BÖLGE MÜDÜRLÜĞÜ", "... BELEDİYESİ"): "müdürlüğü/şubesi" orada adres değil ad.
_KAMU = re.compile(r"^t\.? ?c\.? |\b(bakanligi|baskanligi|belediyesi|valiligi|kaymakamligi|universitesi|genel mudurlugu|"
                   r"bolge mudurlugu|il mudurlugu|ilce mudurlugu|kurumu|idaresi|muhtarligi|organize sanayi bolgesi|"
                   r"sanayi bolgesi|kooperatifi|birligi|odasi|dernegi|vakfi|okulu|mudurlugu|cemiyeti|sendikasi)\b")
# Kurum adının içindeki yer adı adres sayılmaz: "ÖRNEK KÖYÜ MUHTARLIĞI", "YENİ MAHALLESİ MUHTARLIĞI"
_KAMU_BIRIMI = re.compile(r"\b(mudurluk|mudurlugu|sube|subesi|koy|koyu|mahallesi|beldesi|osb)\b")


def kotu_unvan(unvan: str | None) -> bool:
    """Cari adı olamayacak metin: etiket ("Ticaret Sicil No"), adres ("ANKARA Öveçler" değil ama "Mah./Cad./No"),
    yalnız rakam ya da boş."""
    sa = sade(unvan or "").strip()
    return (not re.search(r"[a-z]{2}", sa) or bool(_ETIKET.search(sa))
            or bool(re.match(r"^(ticaret|sicil|mersis|vergi|vkn|tckn|adres|tel|e-?posta|web|duzenleyen|sayin|gonderen)\b", sa))
            or (bool(_ADRES.search(_KAMU_BIRIMI.sub(" ", sa) if _KAMU.search(sa) else sa)) and not SIRKET_EKLERI.search(sa)))


def daha_iyi_unvan(eski: str | None, yeni: str | None) -> bool:
    """Okunan yeni ad mevcut cari adının yerine geçmeli mi: eskisi bozuksa, şirket eki taşımıyorken yenisi taşıyorsa
    ya da eskisi yeni adın sonuna faaliyet açıklaması eklenmiş haliyse."""
    if not yeni or kotu_unvan(yeni) or sade(yeni) == sade(eski or ""):
        return False
    if eski and faaliyeti_ayir(eski) != eski and sade(faaliyeti_ayir(eski)) == sade(yeni):
        return True
    return kotu_unvan(eski) or (not SIRKET_EKLERI.search(sade(eski or "")) and bool(SIRKET_EKLERI.search(sade(yeni))))


def kendi_unvani_mi(unvan: str | None, sirket_unvan: str | None) -> bool:
    """Okunan/seçilen ad firmamızın kendi adı mı ("ARMA GRUP ENERJİ MÜH.İNŞ.LTD.ŞTİ" = "ARMA GRUP ENERJİ MÜHENDİSLİK...")."""
    kelimeler = kendi_kelimeleri(sirket_unvan)
    if not unvan or not kelimeler:
        return False
    ad = re.sub(r"[^a-z0-9 ]", " ", sade(unvan)).split()
    if len(kelimeler) == 1:  # tek ayırt edici kelimeli ünvan: ad onunla başlamalı ve şirket eki taşımalı
        return bool(ad) and ad[0] == kelimeler[0] and bool(SIRKET_EKLERI.search(sade(unvan)))
    return all(k in ad for k in kelimeler)


def firma_adi_benzer_mi(metin: str | None, sirket_unvan: str | None, esik: float = 0.75) -> bool:
    """El yazısıyla yazılıp bozuk okunan ad firmamız mı ("ARMA CRDP ENEPSIMUH." ~ "ARMA GRUP"): ünvanın ayırt edici
    ilk iki kelimesi bitişik yazılıp metindeki her aynı uzunluktaki parçayla karşılaştırılır."""
    import difflib
    anahtar = "".join(kendi_kelimeleri(sirket_unvan))
    m = re.sub(r"[^a-z0-9]", "", sade(metin or ""))
    if len(anahtar) < 5 or len(m) < len(anahtar) - 1:
        return False
    return any(difflib.SequenceMatcher(None, m[i:i + len(anahtar)], anahtar).ratio() >= esik
               for i in range(0, max(len(m) - len(anahtar), 0) + 1))


def kendi_kelimeleri(sirket_unvan: str | None) -> list[str]:
    """Şirketin kendi ünvanındaki ayırt edici ilk iki kelime (ör. 'afsa', 'grup')."""
    return [k for k in re.sub(r"[^a-z0-9 ]", " ", sade(sirket_unvan or "")).split()
            if len(k) >= 3 and k not in GENEL_KELIMELER][:2]


def cari_eslestir(aciklama: str, cariler: list[dict]) -> int | None:
    rakamlar = re.sub(r"\D", "", aciklama)
    duz = re.sub(r"\s", "", aciklama.upper())
    for c in cariler:
        if c.get("vkn") and len(c["vkn"]) >= 10 and c["vkn"] in rakamlar:
            return c["id"]
        if c.get("iban") and re.sub(r"\s", "", c["iban"].upper()) in duz:
            return c["id"]
    metin_kelimeleri = re.sub(r"[^a-z0-9 ]", " ", sade(aciklama)).split()

    def gecer(k: str) -> bool:  # kısaltmalara izin ver: "muh" ~ "muhendislik"
        return any(t == k or (min(len(t), len(k)) >= 3 and (t.startswith(k) or k.startswith(t))) for t in metin_kelimeleri)
    adaylar = []
    for c in cariler:
        kelimeler = [k for k in re.sub(r"[^a-z0-9 ]", " ", sade(c["unvan"])).split() if len(k) >= 3 and k not in GENEL_KELIMELER]
        if not kelimeler:
            continue
        aranan = kelimeler[:2]
        if f" {aranan[0]} " in f" {' '.join(metin_kelimeleri)} " and all(gecer(k) for k in aranan[1:]):
            adaylar.append((len(aranan), c["id"]))
    if not adaylar:
        return None
    adaylar.sort(reverse=True)
    if len(adaylar) > 1 and adaylar[0][0] == adaylar[1][0]:
        return None  # birden fazla cari eşit derecede uyuyor; kullanıcı seçsin
    return adaylar[0][1]


def virman_mi(aciklama: str, kendi: list[str], tutar: float = 0) -> bool:
    """Şirketin kendi hesapları arasındaki transfer mi? (cariye değil virmana işlenir)"""
    s = sade(aciklama)
    if any(k in s for k in VIRMAN_KELIMELERI):
        return True
    if kendi:
        # Banka ekstrelerinde ilk alanlar genelde "açıklama, alıcı ünvanı, ..." (giden) ya da
        # "gönderen ünvanı, ..." (gelen) olur. Karşı taraf şirketin kendisiyse bu bir virmandır.
        # Gidende ilk alan kullanıcının serbest açıklamasıdır ("afsa grup ... cari ödeme"), ona bakılmaz.
        parcalar = s.split(",")
        alanlar = parcalar[1:2] if tutar < 0 and len(parcalar) > 2 else parcalar[:1] if tutar > 0 else parcalar[:2]
        ham = _karsi_alan(aciklama, tutar)
        if ham:
            # Gönderen/alıcı alanı okunduysa karar odur: açıklamada firmamızın adı geçmesi ("FİRMAMIZ İŞ BANKASI
            # GEÇİŞ*ÖRNEK ENERJİ LTD*...") başka firmadan gelen parayı virman yapmaz.
            return re.sub(r"[^a-z0-9 ]", " ", sade(ham)).split()[:len(kendi)] == kendi
        for alan in alanlar:
            kelimeler = re.sub(r"[^a-z0-9 ]", " ", alan).split()
            if kelimeler[:len(kendi)] == kendi:
                return True
    return False


def tur_tahmin(tutar: float, aciklama: str, cari_id: int | None, kendi: list[str] | None = None) -> str:
    s = sade(aciklama)
    # Takas, çek tahsili, "ÇEK : 1000021 ..." ya da "0001000019  10100011 00000000 10101594" (çek no + hesap) satırları
    if ("takas" in s or ("cek" in s and "tahsil" in s) or re.match(r"cek\s*(no)?\s*:", s)
            or re.match(r"0*\d{7,10}\s+\d{8}\s+0{8}\s+\d{8}\s*$", s)
            or (re.search(r"\bcek\b", s) and re.search(r"\btks\b|\bodeme\b|\byatan\b", s))):  # "ÇEK 206/0038*TKS ÖDEME"
        return "cek_tahsil" if tutar > 0 else "cek_odeme"
    if tutar < 0 and NAKIT_CEKIM.search(s) and not ("kart" in s and "borc" in s):  # "KRE.KART BORÇ ÖDEME (Para Çekme)" kart ödemesidir
        return "nakit_cekim"
    if tutar > 0 and NAKIT_YATIRMA.search(s):
        return "nakit_yatirma"
    if tutar > 0 and re.search(r"ters islem|\biade", s) and any(k in s for k in MASRAF_KELIMELERI):
        return "diger"  # geri alınan banka kesintisi ("TERS İŞLEM - EFT KOMİSYONU"): MUHTELİF BANKA İŞLEMLERİ'ne
    if tutar < 0 and any(k in s.split(",")[0] for k in MASRAF_KELIMELERI):
        return "banka_masrafi"  # "KREDİ PEŞİN KOMİSYON, ŞİRKET, kredi hesabı" gibi kesintiler virman değildir
    if not cari_id and virman_mi(aciklama, kendi or [], tutar):
        return "virman"
    if tutar < 0 and any(k in s for k in MASRAF_KELIMELERI):
        return "banka_masrafi"
    if tutar < 0 and ("kredi" in s and ("geri odeme" in s or "taksit" in s or "kart odeme" in s or "kredi hesab" in s)
                      or "anapara" in s):
        return "kredi_odeme"
    parcalar = s.split(",")
    # "açıklama, alıcı ünvanı, alıcı banka/IBAN, ..." biçimindeki EFT/havale satırı
    havale = len(parcalar) >= 3 and re.search(r"bank|tr\d{2}", parcalar[2]) is not None
    if tutar > 0:
        return "tahsilat" if cari_id or havale or SIRKET_EKLERI.search(s) or "eft" in s or "havale" in s else "diger"
    if (cari_id or havale or " dbs" in f" {s}" or SIRKET_EKLERI.search(s) or "eft" in s or "havale" in s
            or "odeme" in s):
        return "odeme"
    return "masraf"


_PARA = r"-?\d{1,3}(?:\.\d{3})*,\d{2}-?|-?\d+,\d{2}-?|-?\d{1,3}(?:,\d{3})*\.\d{2}"


def ekstre_pdf_oku(veri: bytes) -> list[dict]:
    """Metin içeren PDF ekstreden satırları okur.

    Her hareket satırı tarihle başlar ve sonunda tutar ile bakiye bulunur. Tutarın yönü,
    bakiyedeki değişimden (yoksa eksi işaretinden) anlaşılır.
    """
    metin = pdf_metni(veri)
    if len(metin.strip()) < 30:
        raise KuralHatasi("Bu PDF taranmış bir görüntü; yapay zekasız okunamaz. Bankanızdan Excel ekstre indirin.")
    sonuc, onceki_bakiye = [], None
    for satir in metin.splitlines():
        m = re.match(r"\s*(\d{2}[./]\d{2}[./]\d{4})(?:\s+\d{2}:\d{2}(?::\d{2})?)?\s+(.*)$", satir)
        if not m:
            continue
        tarih = tarih_coz(m.group(1))
        govde = m.group(2)
        tutarlar = list(re.finditer(_PARA, govde))
        if not tarih or not tutarlar:
            continue
        if len(tutarlar) >= 2:
            tutar, bakiye = para_coz(tutarlar[-2].group()), para_coz(tutarlar[-1].group())
            aciklama = govde[:tutarlar[-2].start()]
        else:
            tutar, bakiye = para_coz(tutarlar[-1].group()), None
            aciklama = govde[:tutarlar[-1].start()]
        if not tutar:
            continue
        if bakiye is not None and onceki_bakiye is not None:
            fark = round(bakiye - onceki_bakiye, 2)
            if abs(abs(fark) - abs(tutar)) < 0.01:
                tutar = fark
        aciklama = re.sub(r"\s*\d{2}[./]\d{2}[./]\d{4}\s*", " ", aciklama).strip()  # valör tarihi
        sonuc.append({"tarih": tarih.isoformat(), "aciklama": aciklama, "tutar": round(tutar, 2), "bakiye": bakiye})
        if bakiye is not None:
            onceki_bakiye = bakiye
    if not sonuc:
        raise KuralHatasi("PDF ekstrede hareket satırı bulunamadı. Bankanızdan Excel ekstre indirin.")
    return sonuc


BANKA_ADI = re.compile(r"\bbank(?:a|asi|\b)|bankasi|katilim|t\.?a\.?o\b|enpara|papara|ziraat|garanti|akbank|halkbank|vakif|finansbank|"
                       r"kuveyt|albaraka|denizbank")


def tr_buyuk(metin: str) -> str:
    return metin.replace("i", "İ").replace("ı", "I").upper()


def _iban_alani(alan: str) -> bool:
    return re.match(r"TR\d{2}[\d*\s]{6,}", alan.replace(" ", "")) is not None


def _karsi_alan(aciklama: str, tutar: float) -> str | None:
    """Açıklamada karşı tarafın adının yazdığı alan (süzülmemiş). Bilinen düzenler genel kurallarla tanınır:
    - DBS: "... REF NOLU TEDARİKÇİ DBS K" ya da "ref |fatura no |GÖNDEREN|TEDARİKÇİ DBS"
    - Yıldızlı (İş Bankası vb.): "KARŞI TARAF*IBAN ya da banka kodu*açıklama*referans"
    - Virgüllü (Halkbank vb.): "açıklama, KARŞI TARAF, [BANKA], IBAN, ..." (ad virgül içerebilir; banka/IBAN alanına kadar)
      ya da gelen "GÖNDEREN, GÖNDEREN 'DAN açıklama, ..."
    """
    m = re.search(r"REF\s*NOLU\s+(.+?)\s+DBS\b", aciklama, re.IGNORECASE)
    if m:
        return m.group(1)
    # Akbank: "7777/MBL-1234567-Örnek Enerji Ltd Şti-not", "7777/MBL-HAV.AD SOYAD-not" (kanal/işlem-[referans-]AD-not)
    m = re.match(r"\s*\d{3,4}/[A-ZÇĞİÖŞÜ]{3}-(?:\d{5,}-)?(?:HAV\.|EFT\.)?([^-]{4,}?)\s*-", aciklama)
    if m and not re.search(r"\d{3,}", m.group(1)):
        return m.group(1)
    # "EFT: GÖNDEREN ÜNVANI not", "FAST:GÖNDEREN not": önek sonrasındaki büyük harfli ad (şirket ekiyle bitiyorsa oraya kadar)
    m = re.match(r"\s*(?:EFT|FAST|HAVALE|HVL)\s*:\s*(.+)", aciklama, re.IGNORECASE)
    if m:
        if ad := _serbest_unvan(m.group(1)):
            return ad
        kel = []
        sozcukler = m.group(1).split()
        # Ad büyük harfle yazılmışsa büyük harfli kelimeler, "Ad Soyad" yazılmışsa baş harfi büyük kelimeler alınır
        buyuk = bool(sozcukler) and sozcukler[0] == tr_buyuk(sozcukler[0])
        for w in sozcukler:
            if not (re.search(r"[A-ZÇĞİÖŞÜ]", w) and (w == tr_buyuk(w) if buyuk else w[0] == tr_buyuk(w[0]))) or re.search(r"\d", w) \
                    or re.match(r"(cari|hesab|mahsub|odeme|fatura|avans|bedel|kira|maas)", sade(w)):
                break  # ad bitti, büyük harfle yazılmış not başladı ("CARİ HESABA MAHSUBEN")
            kel.append(w)
        if len(kel) >= 2:
            return " ".join(kel)
    if "|" in aciklama:
        return re.sub(r"\s+DBS\s*$", "", aciklama.split("|")[-1].strip(), flags=re.IGNORECASE)
    if re.sub(r"\*{3,}", "", aciklama).count("*") >= 2:  # maskeli IBAN'daki yıldızlar ayıraç değildir
        # Genelde ilk alan karşı taraftır; ilk alan gönderenin serbest açıklamasıysa ("FİRMAMIZ BANKA GEÇİŞ*ÖRNEK ENERJİ LTD*...")
        # banka adı taşımayan, harfli ilk alan alınır
        for alan in [a.strip() for a in aciklama.split("*")[:3]]:
            if len(re.sub(r"[^a-z]", "", sade(alan))) >= 4 and not BANKA_ADI.search(sade(alan)) and not re.match(r"^[\d\s/.-]+$", alan):
                return alan
    # Tutardaki ondalık virgülü ("112.913,34 USD") alan ayıracı sayılmaz
    parcalar = [p.strip() for p in re.sub(r"(?<=\d),(?=\d)", ".", aciklama).split(",")]
    if len(parcalar) < 2:
        return _serbest_unvan(aciklama)
    if tutar > 0:
        ad = re.split(r"\s*'\s*DAN\b|\s*/", parcalar[0], flags=re.IGNORECASE)[0]
        if ad and sade(parcalar[1]).startswith(sade(ad)[:10]):
            return ad
    i = next((j for j in range(2, len(parcalar)) if BANKA_ADI.search(sade(parcalar[j])) or _iban_alani(parcalar[j])), None)
    if i:
        return ", ".join(p for p in parcalar[1:i] if p)
    if tutar > 0 and BANKA_ADI.search(sade(" ".join(parcalar[1:]))):
        return re.split(r"\s*'\s*DAN\b|\s*/", parcalar[0], flags=re.IGNORECASE)[0]
    return None


_UNVAN_SONU = {"sirketi", "sti", "a.s", "a.s.", "as", "a.s.,", "ltd", "ltd.", "anonim", "limited", "kooperatifi", "ortakligi"}


def _serbest_unvan(aciklama: str) -> str | None:
    """Ayıraçsız açıklamada büyük harfle yazılmış, şirket ekiyle biten ünvan (Ziraat: "avans ÖRNEK ENERJİ İNŞ SAN VE T
    ÖRNEK ENERJİ İNŞAAT SANAYİ VE TİCARET LİMİTED ŞİRKETİ Ziraat Mobil Havale"; Yapı Kredi: "Banka/IBAN-ÜNVAN/not").
    Gönderen adı önce kesik, sonra tam yazılmışsa tam olanı alınır."""
    kelimeler = re.split(r"[\s/|]+|(?<=\d)-", aciklama)
    son = max((i for i, w in enumerate(kelimeler) if sade(w).strip(",") in _UNVAN_SONU
               or re.fullmatch(r"sirk\w*", sade(w))), default=None)  # banka adı kesmiş olabilir: "LİMİTED ŞİRK"
    if son is None:
        return None
    bas = son
    while bas > 0 and (w := kelimeler[bas - 1]) and re.search(r"[A-ZÇĞİÖŞÜ]", w) and w == tr_buyuk(w) \
            and not re.search(r"\d", w) and not BANKA_ADI.search(sade(w)):
        bas -= 1
    ad = kelimeler[bas:son + 1]
    if sum(1 for w in ad if sade(w).strip(",.") not in _UNVAN_SONU) < 1 or len(re.sub(r"[^a-z]", "", sade(" ".join(ad)))) < 6:
        return None
    # Kesik + tam yazım: baştaki 2 kelime ileride yeniden başlıyorsa oradan al
    for i in range(2, len(ad) - 2):
        if [sade(w) for w in ad[i:i + 2]] == [sade(w) for w in ad[:2]]:
            ad = ad[i:]
            break
    return " ".join(ad)


def ekstre_kalibi(aciklama: str, tutar: float) -> str:
    """Ekstre satırının tekrar eden kalıbı: karşı tarafın adı ya da rakamları atılmış açıklamanın başı. Kullanıcının bu
    satır için seçtiği cari hatırlanır, aynı kalıp sonraki ekstrelerde gelince o cariye işlenir."""
    ad = _karsi_alan(aciklama, tutar)
    if ad and ad_anahtari(ad):
        return "ekstre:ad:" + ad_anahtari(ad)[:120]
    return "ekstre:metin:" + " ".join(ad_anahtari(re.sub(r"\d+", " ", aciklama)).split()[:6])


def hesap_numaralari(aciklama: str) -> list[str]:
    """Açıklamada geçen hesap numaraları ("1234/0567890", "1000-10102222", "AB001234"): kendi hesaplarımızdan hangisine
    gidildiğini IBAN'ın sonuyla karşılaştırmak için."""
    return [m.lstrip("0") for m in re.findall(r"(?<![\d*])\d{7,12}(?![\d*])", aciklama) if len(m.lstrip("0")) >= 6]


def vadeli_hesap_no(aciklama: str) -> str | None:
    """Vadeli hesap açılışı/kapanışı satırındaki vadeli hesap numarası ("AB001234 NOLU VADELİ HSP AÇILIŞI")."""
    s = sade(aciklama)
    if not re.search(r"vadeli|hesabin kapatilip|hsp acilis", s):
        return None
    m = re.search(r"\b([A-Z]{2}\d{5,8})\b", aciklama.upper())
    return m.group(1) if m else None


def karsi_taraf(aciklama: str, tutar: float, kendi: list[str]) -> dict | None:
    """EFT/havale açıklamasından karşı tarafın ünvanını (ve açık yazılmışsa IBAN'ını) çıkarır; şirketin kendisi,
    banka adı, numara ya da çok kısa metin karşı taraf sayılmaz."""
    ad = _karsi_alan(aciklama, tutar)
    if not ad:
        return None
    ad = " ".join(re.sub(r"\s+DBS(\s+K)?\s*$", "", ad, flags=re.IGNORECASE).split()).strip(" ,.-")
    kelimeler = re.sub(r"[^a-z0-9 ]", " ", sade(ad)).split()
    if len(re.sub(r"[^a-z]", "", sade(ad))) < 4 or BANKA_ADI.search(sade(ad)) or re.search(r"\d{5,}", ad):
        return None
    if kendi and kelimeler[:len(kendi)] == kendi:
        return None
    if ad == ad.lower():
        ad = tr_buyuk(ad)
    # Karşı tarafın IBAN'ı adından hemen sonra yazandır; açıklamanın sonundaki IBAN çoğu zaman bizim hesabımızdır
    ham = _karsi_alan(aciklama, tutar) or ""
    m = re.search(r"(?:^|[,*|])\s*" + re.escape(ham), aciklama) if ham else None  # ad kendi alanında (açıklamanın içinde değil)
    sonrasi = aciklama[m.end():] if m else ""
    alan = re.split(r"[,*|]", sonrasi.lstrip(",*| "), maxsplit=2)
    iban = next((re.sub(r"\s", "", m.group()) for a in alan[:2] if (m := re.fullmatch(f"\\s*({IBAN_DESENI})\\s*", a))), None)
    return {"unvan": ad, "iban": iban}


BANKA_KODLARI = {
    "00010": "Ziraat Bankası", "00012": "Halkbank", "00015": "VakıfBank", "00032": "TEB", "00046": "Akbank",
    "00059": "Şekerbank", "00062": "Garanti BBVA", "00064": "İş Bankası", "00067": "Yapı Kredi", "00099": "ING",
    "00111": "QNB", "00123": "HSBC", "00134": "DenizBank", "00203": "Albaraka Türk", "00205": "Kuveyt Türk",
    "00206": "Türkiye Finans", "00209": "Ziraat Katılım", "00210": "Vakıf Katılım", "00146": "Odeabank",
}
MASKELI_IBAN = r"TR\d{2}\*{4,}\d{4}"


def iban_banka(iban: str | None) -> str | None:
    """IBAN'ın 5-9. hanelerindeki banka kodundan banka adı."""
    return BANKA_KODLARI.get(iban[4:9]) if iban and len(iban) >= 9 and "*" not in iban[:9] else None


def iban_uyar(a: str | None, b: str | None) -> bool:
    """İki IBAN aynı hesabı mı gösteriyor? Biri maskeliyse (TR41*****1553) baştaki 4 ve sondaki 4 hane karşılaştırılır."""
    if not a or not b:
        return False
    a, b = re.sub(r"[^A-Z0-9*]", "", a.upper()), re.sub(r"[^A-Z0-9*]", "", b.upper())  # boşluk, tire, nokta yok sayılır
    if "*" in a or "*" in b:
        return a[:4] == b[:4] and a[-4:] == b[-4:]
    return a == b


def karsi_hesap(aciklama: str, kendi_iban: str | None) -> dict | None:
    """Hesaplar arası transfer açıklamasındaki şirketin diğer hesabı (IBAN'ı maskeli ya da açık)."""
    for m in re.finditer(f"{MASKELI_IBAN}|{IBAN_DESENI}", aciklama.upper()):
        iban = re.sub(r"\s", "", m.group())
        if iban_uyar(iban, kendi_iban):
            continue
        # Maskeli IBAN'da banka kodu görünmez; "hesaplarım arası"/DBS transferi aynı bankanın içindedir.
        ayni_banka = "*" in iban and re.search(r"arasi|dbs", sade(aciklama)) is not None
        return {"iban": iban, "banka": iban_banka(iban), "ayni_banka": ayni_banka}
    return None


def ekstre_kural_oku(veri: bytes, mime: str, cariler: list[dict], sirket_unvan: str | None = None) -> dict:
    satirlar = []
    ham = ekstre_pdf_oku(veri) if mime == "application/pdf" else ekstre_tablo_oku(veri, mime)
    iban = ekstre_iban(veri, mime)
    kendi = kendi_kelimeleri(sirket_unvan)
    # Şirketin kendisi cari olarak açılmışsa eşleştirmeye katma (her satırda gönderen olarak geçer).
    cariler = [c for c in cariler if not (kendi and kendi_kelimeleri(c["unvan"]) == kendi)]
    # Halkbank EFT ücretini aynı açıklamayla ayrı satır yazar: aynı gün aynı açıklamalı büyük bir çıkışın
    # yanındaki küçük çıkış, ödeme değil EFT/havale masrafıdır.
    gruplar: dict[tuple, list[dict]] = {}
    for s in ham:
        if s["tutar"] < 0:
            gruplar.setdefault((s["tarih"], " ".join(s["aciklama"].split()).lower()), []).append(s)
    ucretler = {id(s) for g in gruplar.values() if len(g) > 1
                for s in g if abs(s["tutar"]) <= 100 and any(abs(d["tutar"]) > 100 for d in g)}
    for s in ham:
        if id(s) in ucretler:
            satirlar.append({**s, "onerilen_cari_id": None, "onerilen_tur": "banka_masrafi", "onerilen_cari": None,
                             "karsi_hesap": None, "onerilen_gider_turu": None})
            continue
        # Açıklamada karşı tarafın adı ayrı alanda yazıyorsa cari o addan (ve IBAN'ından) aranır; açıklamanın serbest
        # metnindeki başka firma adı ("örnek metal damga vergisi, KARAYOLLARI ...") cari sanılmaz. Ad bizim firmamızsa virmandır.
        ad = _karsi_alan(s["aciklama"], s["tutar"])
        adk = re.sub(r"[^a-z0-9 ]", " ", sade(ad or "")).split()
        if ad and kendi and adk[:len(kendi)] == kendi:
            cari_id = None
        elif ad and len(re.sub(r"[^a-z]", "", sade(ad))) >= 4:
            k = karsi_taraf(s["aciklama"], s["tutar"], kendi)
            cari_id = cari_eslestir(f"{ad} {(k or {}).get('iban') or ''}", cariler)
        else:
            cari_id = cari_eslestir(s["aciklama"], cariler)
        tur = tur_tahmin(s["tutar"], s["aciklama"], cari_id, kendi)
        karsi = karsi_taraf(s["aciklama"], s["tutar"], kendi) if not cari_id and tur in ("tahsilat", "odeme", "diger", "masraf") else None
        if karsi and tur in ("diger", "masraf"):  # açıklamada karşı tarafın adı açıkça yazıyorsa para ona gitmiş/ondan gelmiştir
            tur = "tahsilat" if s["tutar"] > 0 else "odeme"
        satirlar.append({**s, "onerilen_cari_id": cari_id, "onerilen_tur": tur, "onerilen_cari": karsi,
                         "karsi_hesap": karsi_hesap(s["aciklama"], iban) if tur == "virman" else None,
                         "onerilen_gider_turu": gider_turu_tahmin(s["aciklama"]) if tur == "masraf" else None})
    return {"satirlar": satirlar, "notlar": "", "iban": iban}
